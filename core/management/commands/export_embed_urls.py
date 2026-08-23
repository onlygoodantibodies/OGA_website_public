"""
Management Command: export_embed_urls
======================================

File location: core/management/commands/export_embed_urls.py

Usage:
    python manage.py export_embed_urls

Generates: embed_urls.xlsx in the current directory
    - "Instructions" sheet with embedding guide
    - "All Antibodies" sheet with every embed URL
    - One sheet per supplier with their antibodies only

Filterable by supplier column. Each row has the iframe code ready to copy-paste.
"""

from django.core.management.base import BaseCommand
from django.db.models import Q
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from core.models import Gene, Antibody, Description


BASE_URL = 'https://onlygoodantibodies.co.uk'


class Command(BaseCommand):
    help = 'Export all embed card URLs to an xlsx file, grouped by supplier'

    def handle(self, *args, **options):
        wb = Workbook()

        # ── Styles ──
        header_font = Font(bold=True, color='FFFFFF', size=11, name='Arial')
        header_fill = PatternFill('solid', fgColor='10428A')
        subheader_font = Font(bold=True, size=10, name='Arial')
        body_font = Font(size=10, name='Arial')
        link_font = Font(size=10, name='Arial', color='10428A', underline='single')
        wrap = Alignment(wrap_text=True, vertical='top')
        top_align = Alignment(vertical='top')
        thin_border = Border(
            bottom=Side(style='thin', color='E2E8F0'),
        )

        # ── Instructions sheet ──
        ws_info = wb.active
        ws_info.title = 'Instructions'
        ws_info.sheet_properties.tabColor = '10428A'

        instructions = [
            ('How to embed an OGA antibody card on your website', ''),
            ('', ''),
            ('1. Find your antibody', 'Use the "All Antibodies" sheet or your supplier-specific sheet to find the antibody.'),
            ('2. Copy the iframe code', 'The "Iframe Code" column contains a ready-to-paste HTML snippet.'),
            ('3. Paste into your page', 'Add the iframe code to any HTML page, product listing, or CMS editor.'),
            ('', ''),
            ('Iframe example:', '<iframe src="https://onlygoodantibodies.co.uk/embed/?rrid=AB_2222383" width="100%" height="350" frameborder="0" style="border-radius:10px; border:1px solid #e2e8f0;"></iframe>'),
            ('', ''),
            ('URL parameters:', ''),
            ('?rrid=AB_2222383', 'Look up by RRID'),
            ('?catalogue=17987-1-AP', 'Look up by catalogue number'),
            ('?application=WB', 'Show only one application (WB, IP, ICC-IF, FC)'),
            ('?compact=true', 'Smaller card variant'),
            ('', ''),
            ('Notes:', ''),
            ('- The card is fully responsive and adapts to any container width.', ''),
            ('- Images are served live from onlygoodantibodies.co.uk — always current.', ''),
            ('- No API key needed — embed cards are public.', ''),
            ('- The card only shows experiment images that exist (no empty placeholders).', ''),
            ('', ''),
            ('Questions?', 'Contact onlygoodantibodies@gmail.com'),
        ]

        ws_info.column_dimensions['A'].width = 40
        ws_info.column_dimensions['B'].width = 80

        for row_idx, (a, b) in enumerate(instructions, 1):
            ws_info.cell(row=row_idx, column=1, value=a).font = Font(
                bold=(row_idx == 1 or a.endswith(':')),
                size=12 if row_idx == 1 else 10,
                name='Arial',
                color='10428A' if row_idx == 1 else '333333',
            )
            ws_info.cell(row=row_idx, column=2, value=b).font = body_font
            ws_info.cell(row=row_idx, column=2).alignment = wrap

        # ── Gather data ──
        antibodies = Antibody.objects.select_related(
            'gene', 'description'
        ).prefetch_related('experiments').all().order_by(
            'description__supplier', 'gene__name', 'name'
        )

        rows = []
        for ab in antibodies:
            desc = getattr(ab, 'description', None)
            if not desc:
                continue

            # Build embed URLs
            rrid_url = f'{BASE_URL}/embed/?rrid={desc.rrid}' if desc.rrid else ''
            cat_url = f'{BASE_URL}/embed/?catalogue={ab.name}'

            # Prefer RRID URL for iframe
            primary_url = rrid_url if rrid_url else cat_url
            iframe = f'<iframe src="{primary_url}" width="100%" height="350" frameborder="0" style="border-radius:10px; border:1px solid #e2e8f0;"></iframe>' if primary_url else ''

            # Recommendations
            rec = []
            if desc.wb_app: rec.append('WB')
            if desc.ip_app: rec.append('IP')
            if desc.icc_if_app: rec.append('ICC-IF')
            if desc.fc_app: rec.append('FC')

            # Experiments available
            exps = [e.experiment_type for e in ab.experiments.all() if e.file_path]

            # Gene page URL
            gene_url = f'{BASE_URL}/antibodies/{ab.gene.name}/'

            rows.append({
                'supplier': desc.supplier or 'Unknown',
                'gene': ab.gene.name,
                'catalogue': ab.name,
                'rrid': desc.rrid or '',
                'host': desc.host or '',
                'clonality': desc.clonality or '',
                'clone_id': desc.clone_ID or '',
                'discontinued': 'Yes' if desc.discontinued else '',
                'recommended': ', '.join(rec),
                'experiments': ', '.join(exps),
                'embed_rrid': rrid_url,
                'embed_cat': cat_url,
                'iframe': iframe,
                'gene_url': gene_url,
            })

        # ── Column definitions ──
        columns = [
            ('Supplier', 20),
            ('Gene', 12),
            ('Catalogue #', 18),
            ('RRID', 16),
            ('Host', 10),
            ('Clonality', 20),
            ('Clone ID', 14),
            ('Discontinued', 12),
            ('Recommended', 18),
            ('Experiments', 18),
            ('Embed URL (RRID)', 50),
            ('Embed URL (Cat #)', 50),
            ('Iframe Code', 80),
            ('Gene Page', 45),
        ]

        def write_sheet(ws, data_rows):
            # Header row
            for col_idx, (name, width) in enumerate(columns, 1):
                cell = ws.cell(row=1, column=col_idx, value=name)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal='left', vertical='center')
                ws.column_dimensions[get_column_letter(col_idx)].width = width

            # Data rows
            for row_idx, row in enumerate(data_rows, 2):
                values = [
                    row['supplier'], row['gene'], row['catalogue'],
                    row['rrid'], row['host'], row['clonality'],
                    row['clone_id'], row['discontinued'], row['recommended'],
                    row['experiments'], row['embed_rrid'], row['embed_cat'],
                    row['iframe'], row['gene_url'],
                ]
                for col_idx, val in enumerate(values, 1):
                    cell = ws.cell(row=row_idx, column=col_idx, value=val)
                    cell.font = body_font
                    cell.alignment = top_align
                    cell.border = thin_border

            # Freeze header row
            ws.freeze_panes = 'A2'
            # Auto-filter
            ws.auto_filter.ref = f'A1:{get_column_letter(len(columns))}{len(data_rows) + 1}'

        # ── All Antibodies sheet ──
        ws_all = wb.create_sheet('All Antibodies')
        ws_all.sheet_properties.tabColor = '4A8EC9'
        write_sheet(ws_all, rows)

        # ── Per-supplier sheets ──
        suppliers = sorted(set(r['supplier'] for r in rows))
        for supplier in suppliers:
            # Sheet name max 31 chars, no special chars
            sheet_name = supplier[:31].replace('/', '-').replace('\\', '-')
            ws = wb.create_sheet(sheet_name)
            supplier_rows = [r for r in rows if r['supplier'] == supplier]
            write_sheet(ws, supplier_rows)

        # ── Save ──
        output_path = 'embed_urls.xlsx'
        wb.save(output_path)

        self.stdout.write(self.style.SUCCESS(
            f'\nExported {len(rows)} antibodies to {output_path}\n'
            f'Sheets: Instructions + All Antibodies + {len(suppliers)} supplier sheets\n'
            f'Suppliers: {", ".join(suppliers)}\n'
        ))
