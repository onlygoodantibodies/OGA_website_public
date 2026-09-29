"""One manufacturer's results, as the email and the spreadsheet that go to them.

``manage.py supplier_mailing`` sends each manufacturer on file a summary of
everything OGA has published about their antibodies. The first email is a
one-off summary. After that it sends an update when something new has been
added, **at most once a week**. This module builds what is sent. The command
decides whether and when to send it.

Nothing here decides a verdict. Every result is ``core/recommendations.py``'s
``describe`` and the review list is ``core/not_supportive.py::rows``, so the
spreadsheet, the portal's own *Not supportive* tab and a public gene page
cannot say three different things about one product. The workbook's *Please
review* sheet uses that tab's default list (nothing supportive, and no
*Limited support*), so its count is the same as the tab's.

Four rules shape it.

**Scoped exactly like the key.** A supplier's rows come from
``api_views._get_supplier_company_ids`` over its ``APIConsumer.supplier_filter``,
the same function the portal and the pre-release feed ask. A workbook with
another manufacturer's products in it is the one mistake here that cannot be
taken back, and ``tests_supplier_mailing`` pins it.

**Pre-release is named, never graded.** Figures awaiting review are the
supplier's own, and ``api_pipeline`` already shows them to that supplier. The
spreadsheet lists *which* are waiting (gene, catalogue number, application)
and says nothing about how they came out. A provisional verdict in a file that
gets forwarded is the thing ``PRE_RELEASE_NOTE`` exists to stop.

**The fingerprint is what was reported**, stored on ``SupplierMailing``, so an
update says exactly what is new and a week with nothing new sends nothing.

**Untested is not failed.** ``All results`` lists only applications with a
published figure. An application nobody ran is absent, as on every other
surface.
"""
from __future__ import annotations

import io
from dataclasses import dataclass, field
from datetime import timedelta

from django.utils import timezone

from . import not_supportive as NS
from . import recommendations as R

DB = 'pipeline_db'
BASE_URL = 'https://onlygoodantibodies.co.uk'

#: A gap longer than this between two release days ends a review. The releases
#: from one consortium meeting land over several days (25, 26 and 28 Sep 2026),
#: and the next meeting is weeks away.
REVIEW_GAP_DAYS = 7


@dataclass
class Digest:
    consumer: object
    company_ids: list
    results: list            # one dict per (antibody, tested application)
    review: list             # not_supportive.rows(), the portal tab's default list
    pending: list            # one dict per figure awaiting release
    fingerprint: dict
    changes: dict = field(default_factory=dict)
    since: object = None     # the previous mailing, or None

    @property
    def antibody_count(self):
        return len({r['antibody_id'] for r in self.results})

    @property
    def genes(self):
        return sorted({r['gene'] for r in self.results if r['gene']})

    @property
    def pending_genes(self):
        return sorted({p['gene'] for p in self.pending if p['gene']})

    def counts(self):
        """Per application, how many results sit on each rung."""
        table = {app: {v: 0 for v in (R.SUPPORTIVE, R.LIMITED_SUPPORT,
                                      R.NOT_SUPPORTIVE)} for app in R.APPLICATIONS}
        for r in self.results:
            if r['oga_support'] in table[r['application']]:
                table[r['application']][r['oga_support']] += 1
        return table

    @property
    def change_count(self):
        c = self.changes
        return len(c.get('new', [])) + len(c.get('changed', []))

    def has_news(self):
        c = self.changes
        return bool(c.get('new') or c.get('changed') or c.get('new_pending'))

    def is_empty(self):
        return not self.results and not self.pending


def build(consumer, previous=None, base_url=BASE_URL):
    from .api_views import _get_supplier_company_ids

    company_ids = _get_supplier_company_ids(consumer.supplier_filter or '') or []
    genes = consumer.get_gene_filter_list()
    results = _results(company_ids, genes, base_url) if company_ids else []
    review = (NS.rows(company_ids=company_ids, allowed_genes=genes,
                      base_url=base_url) if company_ids else [])
    pending = _pending(company_ids) if company_ids else []
    fingerprint = {
        'results': {f"{r['antibody_id']}:{r['application']}": r['oga_support']
                    for r in results},
        'pending': sorted(p['id'] for p in pending),
    }
    digest = Digest(consumer, company_ids, results, review, pending, fingerprint,
                    since=previous)
    digest.changes = diff(previous.fingerprint if previous else None, fingerprint)
    return digest


def diff(before, after):
    """What is new since ``before``. ``before=None`` means everything is new."""
    before = before or {}
    old = before.get('results', {})
    new = after.get('results', {})
    return {
        'new': sorted(k for k in new if k not in old),
        'changed': sorted(k for k in new if k in old and old[k] != new[k]),
        'gone': sorted(k for k in old if k not in new),
        'new_pending': sorted(set(after.get('pending', []))
                              - set(before.get('pending', []))),
    }


def _results(company_ids, allowed_genes, base_url):
    antibodies = list(NS.scope(company_ids, allowed_genes))
    if not antibodies:
        return []
    curated = R.curated_gene_ids({a.target_id for a in antibodies})
    axes = R.capability_axes([a.pk for a in antibodies])
    out = []
    for antibody in antibodies:
        figures = {img.application_type: img
                   for img in antibody.publication_images.all()
                   if img.application_type in R.APPLICATIONS}
        tested = set(figures)
        gene = antibody.target.gene_name or ''
        company = antibody.company
        for application in R.APPLICATIONS:
            if application not in tested:
                continue
            d = R.describe(antibody, application, tested,
                           antibody.target_id in curated,
                           axes.get((antibody.pk, application)))
            out.append({
                'antibody_id': antibody.pk,
                'gene': gene,
                'catalogue_number': antibody.catalogue_number or '',
                'rrid': antibody.rrid or '',
                'supplier': (company.display_name or company.name) if company else '',
                'clone_id': antibody.clone_id or '',
                'application': application,
                'result': d['words'],
                'sentence': d['sentence'],
                'oga_support': d['support'],
                'discontinued': bool(antibody.out_of_market),
                'gene_page_url': f'{base_url}/antibodies/{gene}/' if gene else '',
                'figure_url': NS.figure_url(figures[application], base_url),
            })
    return out


def _pending(company_ids):
    from pipeline.models import PendingPublicationImage
    from pipeline.services.review import PENDING

    qs = (PendingPublicationImage.objects.using(DB)
          .filter(status=PENDING, antibody__company_id__in=company_ids)
          .select_related('antibody__target')
          .order_by('antibody__target__gene_name', 'antibody__catalogue_number',
                    'application_type'))
    return [{
        'id': p.pk,
        'gene': (p.antibody.target.gene_name or '') if p.antibody.target else '',
        'catalogue_number': p.antibody.catalogue_number or '',
        'application': p.application_type,
        'queued': p.created_at,
    } for p in qs]


def latest_review(since=None):
    """``(first_day, last_day, genes)`` of the most recent review's releases.

    A review is the latest run of release days with no gap longer than
    ``REVIEW_GAP_DAYS`` between them. ``since`` overrides the start. Only
    genes that are **public now** are listed, so a gene withdrawn since is
    never claimed to be live.
    """
    from pipeline.models import PendingPublicationImage
    from pipeline.public import public_targets
    from pipeline.services.review import RELEASED

    released = list(PendingPublicationImage.objects.using(DB)
                    .filter(status=RELEASED, released_at__isnull=False)
                    .values_list('released_at', 'antibody__target__gene_name'))
    if not released:
        return None, None, []
    days = sorted({timezone.localdate(at) for at, _ in released}, reverse=True)
    last = days[0]
    first = last
    if since:
        first = since
    else:
        for day in days[1:]:
            if (first - day).days > REVIEW_GAP_DAYS:
                break
            first = day
    live = set(public_targets().using(DB).values_list('gene_name', flat=True))
    genes = sorted({g for at, g in released
                    if g and g in live and first <= timezone.localdate(at) <= last})
    return first, last, genes


# ─────────────────────────────────────────────────────────
# The workbook
# ─────────────────────────────────────────────────────────

_WORDS = {R.SUPPORTIVE: 'Supportive', R.LIMITED_SUPPORT: 'Limited support',
          R.NOT_SUPPORTIVE: 'Not supportive'}


def workbook(digest, review=None):
    """The attachment, as bytes. ``review`` is ``latest_review()``'s answer."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    bold = Font(bold=True)
    head_fill = PatternFill('solid', fgColor='EEF4FB')
    fills = {R.SUPPORTIVE: PatternFill('solid', fgColor='DCFCE7'),
             R.LIMITED_SUPPORT: PatternFill('solid', fgColor='FEF3C7'),
             R.NOT_SUPPORTIVE: PatternFill('solid', fgColor='FEE2E2')}
    wrap = Alignment(wrap_text=True, vertical='top')

    def table(ws, headers, rows, widths):
        ws.append(headers)
        for cell in ws[ws.max_row]:
            cell.font = bold
            cell.fill = head_fill
        for row in rows:
            ws.append(row)
        for i, w in enumerate(widths):
            ws.column_dimensions[chr(65 + i)].width = w
        ws.freeze_panes = ws.cell(row=ws.max_row - len(rows) + 1, column=1)

    wb = Workbook()
    name = digest.consumer.name
    today = timezone.localdate()
    new_keys = set(digest.changes.get('new', [])) if digest.since else set()
    changed_keys = set(digest.changes.get('changed', [])) if digest.since else set()

    # ── Summary
    ws = wb.active
    ws.title = 'Summary'
    ws.column_dimensions['A'].width = 34
    for col in 'BCDE':
        ws.column_dimensions[col].width = 18
    ws.append([f'{name}: antibody characterisation results from Only Good Antibodies'])
    ws['A1'].font = Font(bold=True, size=14)
    ws.append([f'Prepared {today:%d %B %Y}'])
    ws.append([])
    if digest.since:
        c = digest.changes
        ws.append([f'Since the last email ({digest.since.sent_at:%d %B %Y})'])
        ws[f'A{ws.max_row}'].font = bold
        ws.append(['New results', len(c.get('new', []))])
        ws.append(['Results that changed', len(c.get('changed', []))])
        ws.append(['New figures awaiting review', len(c.get('new_pending', []))])
        ws.append(['The "What\'s new" sheet lists them.'])
        ws.append([])
    ws.append(['All results so far'])
    ws[f'A{ws.max_row}'].font = bold
    ws.append(['Antibodies with published results', digest.antibody_count])
    ws.append(['Genes', len(digest.genes)])
    ws.append(['Antibodies to review (no application supportive)', len(digest.review)])
    ws.append(['Figures awaiting OGA review (pre-release)', len(digest.pending)])
    ws.append([])
    ws.append(['Application', 'Supportive', 'Limited support', 'Not supportive'])
    for cell in ws[ws.max_row]:
        cell.font = bold
        cell.fill = head_fill
    counts = digest.counts()
    for app in R.APPLICATIONS:
        ws.append([app, counts[app][R.SUPPORTIVE], counts[app][R.LIMITED_SUPPORT],
                   counts[app][R.NOT_SUPPORTIVE]])
    ws.append([])
    if review and review[2]:
        ws.append(['Genes from the latest review, now live:'])
        ws[f'A{ws.max_row}'].font = bold
        ws.append([', '.join(review[2])])
        ws.append([])
    ws.append(['What the results mean'])
    ws[f'A{ws.max_row}'].font = bold
    for line in (
            'Supportive: the characterisation data supports that application.',
            'Limited support: tested; not supported overall, and the antibody was '
            'still seen to do what the application is for.',
            'Not supportive: tested, and nothing on-target was seen.',
            'An application nobody has tested is not listed. It is not a failure.',
            R.SCOPE_NOTE):
        ws.append([line])
        ws[f'A{ws.max_row}'].alignment = Alignment(wrap_text=False)
    ws.append([])
    ws.append(['Data portal', f'{BASE_URL}/portal/'])
    ws.append(["Request your company's key", f'{BASE_URL}/data-access/key/'])

    # ── What's new (updates only)
    if digest.since:
        ws = wb.create_sheet("What's new")
        rows = []
        for r in digest.results:
            key = f"{r['antibody_id']}:{r['application']}"
            if key in new_keys or key in changed_keys:
                rows.append(['New' if key in new_keys else 'Changed', r['gene'],
                             r['catalogue_number'], r['application'], r['sentence'],
                             r['gene_page_url']])
        pending_new = set(digest.changes.get('new_pending', []))
        for p in digest.pending:
            if p['id'] in pending_new:
                rows.append(['Awaiting review', p['gene'], p['catalogue_number'],
                             p['application'], 'Not published yet', ''])
        table(ws, ['What', 'Gene', 'Catalogue number', 'Application', 'Result',
                   'Gene page'], rows, [16, 12, 20, 12, 44, 52])

    # ── Please review
    ws = wb.create_sheet('Please review')
    ws.append(['Antibodies where no application tested came back supportive. '
               'Please review these products. Each row links to the public gene '
               'page with the figures.'])
    ws.append([NS.LIST_NOTE])
    ws['A2'].alignment = Alignment(wrap_text=False)
    ws.append([])
    rows = []
    for row in digest.review:
        findings = '; '.join(f"{f['application']}: {f['sentence']}" for f in row['findings'])
        rows.append([row['gene'], row['catalogue_number'], row['rrid'], row['clone_id'],
                     row['applications_tested'], findings,
                     'yes' if row['discontinued'] else '', row['gene_page_url'],
                     row['product_link']])
    table(ws, ['Gene', 'Catalogue number', 'RRID', 'Clone', 'Applications tested',
               'Results', 'Marked discontinued', 'Gene page', 'Product page'],
          rows, [12, 20, 16, 14, 12, 60, 12, 52, 40])
    for row in ws.iter_rows(min_row=5, min_col=6, max_col=6):
        for cell in row:
            cell.alignment = wrap

    # ── All results
    ws = wb.create_sheet('All results')
    rows = []
    for r in digest.results:
        key = f"{r['antibody_id']}:{r['application']}"
        flag = 'New' if key in new_keys else 'Changed' if key in changed_keys else ''
        rows.append([r['gene'], r['catalogue_number'], r['rrid'], r['clone_id'],
                     r['application'], _WORDS.get(r['oga_support'], r['result']),
                     r['sentence'], r['oga_support'], flag,
                     'yes' if r['discontinued'] else '',
                     r['gene_page_url'], r['figure_url']])
    table(ws, ['Gene', 'Catalogue number', 'RRID', 'Clone', 'Application', 'Result',
               'In full', 'oga_support', 'Since last email', 'Marked discontinued',
               'Gene page', 'Figure'],
          rows, [12, 20, 16, 14, 11, 16, 40, 16, 14, 12, 52, 60])
    for row in ws.iter_rows(min_row=2, min_col=6, max_col=6):
        for cell in row:
            fill = fills.get(ws.cell(row=cell.row, column=8).value)
            if fill:
                cell.fill = fill
    ws.auto_filter.ref = ws.dimensions

    # ── Pre-release
    ws = wb.create_sheet('Awaiting review')
    ws.append(['Your figures that OGA has cropped from a finished report and that '
               'are waiting for the review meeting. They are NOT published and '
               'may change. The portal\'s Pre-release tab shows them. Please tell '
               'us about a wrong catalogue number, lot or RRID before release.'])
    ws.append([])
    table(ws, ['Gene', 'Catalogue number', 'Application', 'Queued on'],
          [[p['gene'], p['catalogue_number'], p['application'],
            timezone.localtime(p['queued']).strftime('%d %b %Y')] for p in digest.pending],
          [12, 20, 12, 14])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def attachment_name(digest):
    safe = ''.join(ch if ch.isalnum() else '_' for ch in digest.consumer.name).strip('_')
    return f"OGA_results_{safe}_{timezone.localdate():%Y-%m-%d}.xlsx"
