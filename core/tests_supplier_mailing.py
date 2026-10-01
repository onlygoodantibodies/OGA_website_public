"""Manufacturer results emails and the key-request door: what must never go wrong.

Three failures here would be silent and could not be taken back once sent:
a manufacturer's spreadsheet holding **another manufacturer's** products; a
key request that **mints a second key** for an organisation (or sends one to
a personal address); and an update sent **more than once a week, or with
nothing new**. Those are pinned, and the wording is left alone.
"""
from __future__ import annotations

import io
from datetime import timedelta
from unittest import mock

from django.core import mail
from django.core.cache import cache
from django.core.management import call_command
from django.test import SimpleTestCase, TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from core import key_requests
from core import supplier_digest as SD
from core import supplier_mail as SM
from core.management.commands.supplier_mailing import decide
from core.models import APIConsumer, KeyRequest, SupplierContact, SupplierMailing
from core.tests_not_supportive import _figure
from pipeline.models import Antibody, Company, Target


class _Fixture(TestCase):
    databases = {"pipeline_db", "academy_db"}

    @classmethod
    def setUpTestData(cls):
        abcam = Company.objects.create(name="abcam", display_name="Abcam")
        other = Company.objects.create(name="Proteintech")
        gene = Target.objects.create(gene_name="TRPA1")
        good = Antibody.objects.create(target=gene, company=abcam,
                                       catalogue_number="ab-good", wb_recommended=True)
        _figure(good, "WB")
        bad = Antibody.objects.create(target=gene, company=abcam, catalogue_number="ab-bad")
        _figure(bad, "WB")
        theirs = Antibody.objects.create(target=gene, company=other, catalogue_number="pt-bad")
        _figure(theirs, "WB")
        cls.abcam = APIConsumer.objects.create(
            name="Abcam", consumer_type="manufacturer", supplier_filter="Abcam",
            email_domains="abcam.com")
        SupplierContact.objects.create(consumer=cls.abcam, email="contact@abcam.com",
                                       name="Alex Example")

    def setUp(self):
        cache.clear()


class ASupplierIsSentOnlyItsOwnProductsTests(_Fixture):

    def test_the_workbook_holds_no_other_suppliers_catalogue_numbers(self):
        digest = SD.build(self.abcam)
        book = load_workbook(io.BytesIO(SD.workbook(digest)))
        cells = {str(c.value) for ws in book for row in ws.iter_rows() for c in row if c.value}
        self.assertIn("ab-bad", cells)
        self.assertNotIn("pt-bad", cells, "Another manufacturer's product was in the file.")

    def test_the_review_sheet_is_the_portal_tabs_list(self):
        digest = SD.build(self.abcam)
        self.assertEqual([r["catalogue_number"] for r in digest.review], ["ab-bad"])

    def test_a_scope_matching_no_supplier_reports_nothing_rather_than_everything(self):
        stray = APIConsumer.objects.create(name="Typo", consumer_type="manufacturer",
                                           supplier_filter="Abcamm")
        self.assertTrue(SD.build(stray).is_empty())


@override_settings(SUPPLIER_MAILING_CC=["owner@le.ac.uk", "carl@mcgill.ca"])
class TheEmailTests(_Fixture):

    def test_the_owner_and_carl_are_copied_and_the_key_is_not_in_it(self):
        digest = SD.build(self.abcam)
        to, cc, contacts = SM.recipients(self.abcam)
        msg = SM.build_message(digest, (None, None, []), b"x", "f.xlsx", to, cc, contacts, False)
        self.assertEqual(msg.to, ["contact@abcam.com"])
        self.assertEqual(msg.cc, ["owner@le.ac.uk", "carl@mcgill.ca"])
        self.assertNotIn(str(self.abcam.api_key), msg.body)
        self.assertNotIn(str(self.abcam.api_key), msg.alternatives[0][0])
        self.assertIn("/data-access/key/", msg.body)


class OnceAWeekAndOnlyWithNewsTests(_Fixture):

    def _previous(self, days_ago, fingerprint):
        return SupplierMailing(consumer=self.abcam, kind="launch", subject="s",
                               recipients="r", fingerprint=fingerprint,
                               sent_at=timezone.now() - timedelta(days=days_ago))

    def test_first_email_is_never_sent_by_the_cron(self):
        digest = SD.build(self.abcam)
        self.assertIsNone(decide(digest, None, timezone.now(), updates_only=True)[0])
        self.assertEqual(decide(digest, None, timezone.now(), updates_only=False)[0], "launch")

    def test_nothing_new_sends_nothing(self):
        digest = SD.build(self.abcam)
        prev = self._previous(10, digest.fingerprint)
        self.assertIsNone(decide(SD.build(self.abcam, prev), prev, timezone.now(), True)[0])

    def test_news_within_a_week_waits(self):
        prev = self._previous(3, {"results": {}, "pending": []})
        digest = SD.build(self.abcam, prev)
        self.assertTrue(digest.has_news())
        self.assertIsNone(decide(digest, prev, timezone.now(), True)[0])

    def test_news_after_a_week_is_an_update(self):
        prev = self._previous(7, {"results": {}, "pending": []})
        digest = SD.build(self.abcam, prev)
        self.assertEqual(decide(digest, prev, timezone.now(), True)[0], "update")

    @mock.patch("core.management.commands.supplier_mailing.configuration_gaps", return_value=[])
    @override_settings(EMAIL_BACKEND="django.core.mail.backends.locmem.EmailBackend")
    def test_a_send_is_recorded_and_the_next_run_sends_nothing(self, _gaps):
        call_command("supplier_mailing", "--send", stdout=io.StringIO())
        self.assertEqual(len(mail.outbox), 1)
        self.assertEqual(SupplierMailing.objects.filter(consumer=self.abcam).count(), 1)
        call_command("supplier_mailing", "--send", stdout=io.StringIO())
        self.assertEqual(len(mail.outbox), 1, "A second email went out with nothing new.")

    @mock.patch("core.management.commands.supplier_mailing.configuration_gaps", return_value=[])
    @override_settings(EMAIL_BACKEND="django.core.mail.backends.locmem.EmailBackend")
    def test_a_preview_goes_only_to_the_previewer_and_records_nothing(self, _gaps):
        call_command("supplier_mailing", "--preview-to", "me@le.ac.uk", stdout=io.StringIO())
        self.assertEqual(mail.outbox[0].to, ["me@le.ac.uk"])
        self.assertEqual(mail.outbox[0].cc, [])
        self.assertFalse(SupplierMailing.objects.exists())


class AHalfConfiguredServiceRefusesTests(_Fixture):
    """The cron's environment is typed by hand; a missing database must not read as "0 sent"."""

    @override_settings(EMAIL_BACKEND="django.core.mail.backends.locmem.EmailBackend",
                       EMAIL_HOST_PASSWORD="x", USE_R2=False)
    def test_a_missing_setting_refuses_by_name_and_sends_nothing(self):
        from django.core.management.base import CommandError
        with self.assertRaises(CommandError) as caught:
            call_command("supplier_mailing", "--send", "--updates-only", stdout=io.StringIO())
        self.assertIn("ACADEMY_DATABASE_URL", str(caught.exception))
        self.assertIn("USE_R2", str(caught.exception))
        self.assertEqual(mail.outbox, [])


@override_settings(EMAIL_BACKEND="django.core.mail.backends.locmem.EmailBackend")
class AKeyRequestSendsTheOneSharedKeyTests(_Fixture):

    def test_an_address_on_the_domain_is_emailed_the_existing_key(self):
        before = APIConsumer.objects.count()
        outcome, consumer = key_requests.handle("someone@uk.abcam.com")
        self.assertEqual(outcome, KeyRequest.SENT)
        self.assertEqual(consumer, self.abcam)
        self.assertEqual(APIConsumer.objects.count(), before, "A request minted a key.")
        self.assertIn(str(self.abcam.api_key), mail.outbox[0].body)
        self.assertEqual(mail.outbox[0].to, ["someone@uk.abcam.com"])

    def test_a_lookalike_domain_is_not_matched(self):
        outcome, _ = key_requests.handle("someone@notabcam.com")
        self.assertEqual(outcome, KeyRequest.UNMATCHED)
        self.assertTrue(all(str(self.abcam.api_key) not in m.body for m in mail.outbox))

    def test_a_personal_address_never_gets_a_key_even_if_listed(self):
        APIConsumer.objects.filter(pk=self.abcam.pk).update(email_domains="abcam.com,gmail.com")
        outcome, _ = key_requests.handle("someone@gmail.com")
        self.assertEqual(outcome, KeyRequest.REFUSED)
        self.assertEqual(mail.outbox, [])

    def test_two_organisations_claiming_one_domain_sends_nothing(self):
        APIConsumer.objects.create(name="Abcam (old)", consumer_type="manufacturer",
                                   supplier_filter="Abcam", email_domains="abcam.com")
        outcome, _ = key_requests.handle("someone@abcam.com")
        self.assertEqual(outcome, KeyRequest.UNMATCHED)
        self.assertTrue(all(str(self.abcam.api_key) not in m.body for m in mail.outbox))

    def test_the_page_never_shows_the_key(self):
        response = self.client.post(reverse("request_key"), {"email": "someone@abcam.com"})
        self.assertEqual(response.status_code, 200)
        self.assertNotContains(response, str(self.abcam.api_key))
        self.assertEqual(len(mail.outbox), 1)


class TheContactsLoaderMakesOneKeyPerOrganisationTests(_Fixture):

    def _load(self, csv_text, *args):
        path = self._tmp_csv(csv_text)
        out = io.StringIO()
        call_command("supplier_contacts", "--file", path, *args, stdout=out)
        return out.getvalue()

    def _tmp_csv(self, text):
        import tempfile
        fh = tempfile.NamedTemporaryFile("w", suffix=".csv", delete=False)
        fh.write(text)
        fh.close()
        return fh.name

    HEAD = "organisation,supplier_filter,email_domains,name,email,role\n"

    def test_an_existing_key_is_reused(self):
        self._load(self.HEAD + "Abcam Ltd,Abcam,abcam.com,X,x@abcam.com,to\n", "--apply")
        self.assertEqual(APIConsumer.objects.filter(consumer_type="manufacturer").count(), 1)
        self.assertTrue(self.abcam.contacts.filter(email="x@abcam.com").exists())

    def test_a_new_supplier_gets_exactly_one_scoped_key(self):
        text = (self.HEAD + "Proteintech,Proteintech,ptglab.com,D,d@ptglab.com,to\n"
                "Proteintech,Proteintech,ptglab.com,E,e@ptglab.com,cc\n")
        self._load(text, "--apply")
        self._load(text, "--apply")
        keys = APIConsumer.objects.filter(supplier_filter="Proteintech")
        self.assertEqual(keys.count(), 1)
        self.assertEqual(keys[0].contacts.count(), 2)

    def test_an_internal_test_key_on_the_same_supplier_is_not_theirs(self):
        APIConsumer.objects.create(name="Test", consumer_type="manufacturer",
                                   supplier_filter="Abcam", is_internal=True)
        out = self._load(self.HEAD + "Abcam,Abcam,abcam.com,X,x@abcam.com,to\n", "--apply")
        self.assertNotIn("refused", out)
        self.assertTrue(self.abcam.contacts.filter(email="x@abcam.com").exists())

    def test_a_switched_off_key_gives_way_to_the_live_one(self):
        APIConsumer.objects.create(name="Old Abcam", consumer_type="manufacturer",
                                   supplier_filter="Abcam", is_active=False)
        out = self._load(self.HEAD + "Abcam,Abcam,abcam.com,X,x@abcam.com,to\n", "--apply")
        self.assertNotIn("refused", out)
        self.assertTrue(self.abcam.contacts.filter(email="x@abcam.com").exists())

    def test_dry_run_writes_nothing(self):
        self._load(self.HEAD + "Proteintech,Proteintech,ptglab.com,D,d@ptglab.com,to\n")
        self.assertFalse(APIConsumer.objects.filter(supplier_filter="Proteintech").exists())

    def test_the_shipped_file_parses(self):
        from core.management.commands.supplier_contacts import DEFAULT_FILE, read
        orgs = read(DEFAULT_FILE)
        self.assertNotIn("OriGene", orgs)
        self.assertTrue(all(e["contacts"] for e in orgs.values()))


class ThePublicFormsAreNeverCachedTests(SimpleTestCase):
    """A cached form page reaches every visitor without its CSRF cookie, so every
    submission 403s. The Cloudflare rule's exclusions mirror this list."""

    def test_each_form_is_under_a_never_cached_prefix(self):
        from OGA_website.cache_headers import NEVER_CACHED_PREFIXES
        for name in ("request_key", "contact", "nominate_gene"):
            with self.subTest(name=name):
                self.assertTrue(reverse(name).startswith(NEVER_CACHED_PREFIXES), reverse(name))
