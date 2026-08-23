"""A download says the same thing as the board it came from.

The owner's review: *"Download excel includes fields that are not on the board
(e.g. lot) — they will probably need aligning — also the order being different
is confusing"*, and *"the downloaded sheet does not match the board. They need
to be the same data fields and order."*

Both were true. Cell lines: the export led with `c number` and carried `lot` and
`clone` which the board does not draw, while KO validation, its reason,
`received` and `thawed` were in no sheet at all. Antibodies: no comments, none
of the OGA recommendations, and a different order again.

What is pinned here is the alignment, and the three places the sheet and the
screen *deliberately* differ — because building the sheet from the board's own
row payload is the obvious implementation and it would have been a silent data
change.
"""
from __future__ import annotations

import io
from pathlib import Path

import openpyxl
from django.conf import settings
from django.test import TestCase

from pipeline.models import Antibody, CellLine, Company, Site, Target
from pipeline.services import antibody_board, board_columns, cell_line_board
from pipeline.tests_timeouts import DB, _member_client


def _sheet(content):
    return openpyxl.load_workbook(io.BytesIO(content)).active


class TheSheetFollowsTheBoardTests(TestCase):
    databases = {"default", DB, "academy_db"}

    def setUp(self):
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)

    def _headings(self, kind, url):
        ws = _sheet(self.client.get(url).content)
        return [c.value for c in ws[1]]

    def test_every_board_column_is_in_the_download(self):
        for kind in ("antibodies", "cell-lines"):
            with self.subTest(kind=kind):
                on_board = {c.key for c in board_columns.registry(kind)
                            if c.board != board_columns.OFF}
                in_sheet = {c.key for c in board_columns.registry(kind)}
                self.assertTrue(on_board <= in_sheet)

    def test_the_download_is_in_board_order(self):
        for kind, url in (("antibodies", "/pipeline/antibodies/export/"),
                          ("cell-lines", "/pipeline/cell-lines/export/")):
            with self.subTest(kind=kind):
                headings = self._headings(kind, url)
                self.assertEqual(headings, board_columns.sheet_headers(kind))
                # …and the board's own headings appear in the same relative
                # order, which is the half the review actually noticed.
                board_keys = [c.key for c in board_columns.board_columns(kind)]
                sheet_keys = [c.key for c in board_columns.registry(kind)]
                self.assertEqual(board_keys,
                                 [k for k in sheet_keys if k in set(board_keys)])

    def test_the_cell_line_sheet_leads_with_the_name_not_the_batch(self):
        headings = self._headings("cell-lines", "/pipeline/cell-lines/export/")
        self.assertEqual(headings[0], "name")

    def test_the_columns_the_board_grew_are_in_the_sheet_now(self):
        headings = self._headings("cell-lines", "/pipeline/cell-lines/export/")
        for want in ("ko validated", "ko validation notes", "received", "thawed"):
            self.assertTrue(any(h.startswith(want) for h in headings), want)
        headings = self._headings("antibodies", "/pipeline/antibodies/export/")
        for want in ("comments", "OGA recommends"):
            self.assertTrue(any(h.startswith(want) for h in headings), want)


class AReadOnlyColumnIsNamedNotDroppedTests(TestCase):
    """An OGA recommendation is a verdict — the public gene pages, the browser
    extension index and the MCP server all read it — so a stale spreadsheet must
    not re-assert it. Keeping it silently is the other way to lose data."""

    databases = {"default", DB, "academy_db"}

    def setUp(self):
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)

    def test_the_header_says_so(self):
        for kind, url in (("antibodies", "/pipeline/antibodies/export/"),
                          ("cell-lines", "/pipeline/cell-lines/export/")):
            with self.subTest(kind=kind):
                ws = _sheet(self.client.get(url).content)
                headings = [c.value for c in ws[1]]
                marked = [h for h in headings if h.endswith("(read-only)")]
                self.assertEqual(sorted(marked),
                                 sorted(board_columns.read_only(kind)))
                self.assertTrue(marked, f"{kind} marks nothing read-only")

    def test_the_verdict_is_one_of_them(self):
        self.assertIn("OGA recommends (read-only)",
                      board_columns.read_only("antibodies"))


class TheSheetAndTheScreenDifferOnPurposeTests(TestCase):
    """Three values the sheet renders differently from the board — each of them
    a silent data change if the export were built from `row_for()`."""

    databases = {"default", DB}

    def setUp(self):
        self.site = Site.objects.create(name="Leicester", short_code="LEI")
        self.company = Company.objects.create(name="Abcam")

    def test_a_target_with_no_symbol_still_exports_a_gene(self):
        """`gene_name or protein_name`. The importer resolves a target by gene,
        so a blank cell is a row that cannot come back."""
        target = Target.objects.create(gene_name="", protein_name="Tau protein")
        ab = Antibody.objects.create(target=target, company=self.company,
                                     catalogue_number="ab1", site=self.site)
        gene_col = next(c for c in board_columns.ANTIBODIES if c.key == "gene")
        self.assertEqual(gene_col.value(ab), "Tau protein")
        self.assertEqual(antibody_board.row_for(ab)["gene"], "")

    def test_na_is_written_to_the_sheet_and_blanked_on_the_board(self):
        """`NA` is what `bulk_cell_lines` reads to settle a blank genotype as a
        wild type, so it must survive a round trip — and it must not appear in
        the GENE cell of a row whose point is that it has no gene."""
        na = Target.objects.create(gene_name="NA", protein_name="Not applicable")
        line = CellLine.objects.create(name="HeLa", genotype="WT", target=na,
                                       site=self.site)
        gene_col = next(c for c in board_columns.CELL_LINES if c.key == "gene")
        self.assertEqual(gene_col.value(line), "NA")
        self.assertEqual(cell_line_board.row_for(line)["gene"], "")

    def test_the_concentration_keeps_its_stored_precision(self):
        target = Target.objects.create(gene_name="SNCA")
        ab = Antibody.objects.create(target=target, company=self.company,
                                     catalogue_number="ab2", site=self.site,
                                     concentration="1.50")
        col = next(c for c in board_columns.ANTIBODIES if c.key == "concentration")
        self.assertEqual(col.value(ab), "1.50")
        self.assertEqual(antibody_board.row_for(ab)["concentration"], 1.5)

    def test_the_reference_column_is_filled_in_for_a_new_row(self):
        """`ab_number` was null for everything created since the Access import,
        so this column was blank for exactly the rows somebody had just added.

        It is not any more: a record is given its own bench's next A-number when
        it is created (`services/lab_numbers.py`), and that is what the column
        shows. A row that has no number shows nothing — the record-id fallback
        that used to fill that cell is what the first field test caught, since
        the board said *not numbered* and the sheet said `4550` about one
        vial."""
        from pipeline.services import lab_numbers

        target = Target.objects.create(gene_name="SNCA")
        ab = Antibody.objects.create(target=target, company=self.company,
                                     catalogue_number="ab3", site=self.site)
        col = next(c for c in board_columns.ANTIBODIES if c.key == "ab_number")
        self.assertEqual(ab.ab_number, 1, "the site's first antibody is A-1")
        self.assertEqual(col.value(ab), "A-1")

        # A row from before numbers were issued: an empty cell, not the record
        # id. The fallback that filled it said `4550` about a vial the board
        # called *not numbered*, which is a value somebody writes on a tube.
        with lab_numbers.suspended():
            old = Antibody.objects.create(target=target, company=self.company,
                                          catalogue_number="ab4", site=self.site)
        self.assertIsNone(old.ab_number)
        self.assertEqual(col.value(old), "")


class NoBoardKeepsItsOwnColumnListTests(TestCase):
    """The whole point: one registry, not one list per surface."""

    def test_the_old_export_constants_are_gone(self):
        src = (Path(settings.BASE_DIR) / "pipeline/views/search.py").read_text()
        for name in ("_EXPORT_COLUMNS", "_CL_EXPORT_COLUMNS", "_EXPORT_APPS"):
            self.assertNotIn(name, src)

    def test_the_boards_render_their_headings_from_the_registry(self):
        for name in ("antibody_board.html", "cell_line_board.html"):
            html = (Path(settings.BASE_DIR) /
                    "pipeline/templates/pipeline" / name).read_text()
            with self.subTest(name=name):
                self.assertIn("{% for c in columns %}", html)
                self.assertIn("{{ c.th }}", html)
                # A hard-coded colspan is how a grid ends up one column out of
                # step with its own headings.
                self.assertNotIn('colspan="12"', html)
                self.assertNotIn("colspan: 12,", html)
                self.assertNotIn("colspan: 10,", html)

    def test_a_dict_never_reaches_a_spreadsheet_cell(self):
        """`row_for` returns dicts for `recommended` and `supplier_validated`.
        `ws.append` raises on those, so the sheet must render them itself —
        which is why a column carries a `cell()` rather than a key."""
        for kind in ("antibodies", "cell-lines"):
            for col in board_columns.registry(kind):
                with self.subTest(kind=kind, key=col.key):
                    self.assertIsNotNone(col.cell,
                                         f"{col.key} has no sheet renderer")


class EveryDrawnColumnHasHoverTextTests(TestCase):
    """An empty `title=""` is invisible — it reads as a column nobody thought
    worth explaining, rather than as a bug.

    This is how it happened: the tips were keyed on display names
    (`cellosaurus`, `supplier`, `site_status`) while the registry keys on the
    `row_for` key (`cellosaurus_id`, `company`, `site`). The moment the board's
    `<thead>` started rendering from the registry, three headers silently lost
    their hover text. Nothing about the page looked wrong.

    Pinned as a comparison of the two lists rather than by reading the rendered
    page, so the failure names the key that drifted.
    """

    def test_the_registry_and_the_tips_use_the_same_vocabulary(self):
        from pipeline.services import antibody_board, cell_line_board
        for kind, svc in (("antibodies", antibody_board),
                          ("cell-lines", cell_line_board)):
            for col in board_columns.board_columns(kind):
                with self.subTest(kind=kind, key=col.key):
                    self.assertTrue(
                        svc.COLUMN_TIPS.get(col.key, "").strip(),
                        f"{kind}: column '{col.key}' ({col.th}) has no tip — a "
                        f"tip keyed on something other than the registry key "
                        f"renders an invisible empty tooltip")

    def test_no_tip_is_left_pointing_at_a_column_that_is_gone(self):
        """The other direction: a tip nothing draws is dead text that reads as
        current."""
        from pipeline.services import antibody_board, cell_line_board
        for kind, svc in (("antibodies", antibody_board),
                          ("cell-lines", cell_line_board)):
            drawn = {c.key for c in board_columns.registry(kind)}
            for key in svc.COLUMN_TIPS:
                with self.subTest(kind=kind, key=key):
                    self.assertIn(key, drawn,
                                  f"{kind}: tip '{key}' describes no column")


class AColumnTheSheetCallsWritableIsOneTheParserReadsTests(TestCase):
    """`board_columns` declares, per column, what an upload does with it. Nothing
    checked that the declaration was true.

    `supplier claims` was declared ``WRITE`` and had no alias in
    ``cropper/metadata.py::HEADER_ALIASES`` at all. So it was kept out of
    ``read_only()``, no preview named it as ignored, and a scientist who
    downloaded the antibodies sheet, corrected what the supplier recommends and
    uploaded it back had the change dropped in silence — the worst of the three
    possible outcomes, because the app claimed the opposite.

    The existing tests could not catch it: one asserts the *header text* of the
    read-only columns matches ``read_only()``, which is true of a column wrongly
    declared writable, and the round-trip tests only ever exercised columns that
    happened to work. This asks the parser.
    """

    databases = {"default", DB}

    def test_every_writable_antibody_column_has_a_parser_alias(self):
        from pipeline.services.cropper import metadata as meta
        for col in board_columns.registry("antibodies"):
            if col.upload == board_columns.READ:
                continue
            with self.subTest(sheet=col.sheet):
                self.assertIsNotNone(
                    meta.HEADER_ALIASES.get(meta._norm_header(col.sheet)),
                    f"the antibodies sheet writes '{col.heading}' and declares it "
                    f"upload={col.upload}, but no HEADER_ALIASES entry matches it — "
                    f"an upload would drop that column without saying so. Either add "
                    f"the alias or declare the column READ so it is named read-only.")

    def test_every_writable_cell_line_column_has_a_parser_alias(self):
        from pipeline.services import bulk_cell_lines as bulkcl
        for col in board_columns.registry("cell-lines"):
            if col.upload == board_columns.READ:
                continue
            with self.subTest(sheet=col.sheet):
                self.assertIsNotNone(
                    bulkcl.HEADER_ALIASES.get(col.sheet.strip().lower()),
                    f"the cell-lines sheet writes '{col.heading}' and declares it "
                    f"upload={col.upload}, but no HEADER_ALIASES entry matches it.")


class TheSupplierColumnHasOneNameTests(TestCase):
    """What the supplier says an antibody is for was called three things: the
    blank template said `applications`, the board's download said `supplier
    claims`, and the whole-dataset sheet said `applications` again — for one
    field, drawn next to `OGA recommends`, which is the other verdict entirely.

    It is `supplier recommendations` everywhere now. The older spellings stay
    readable, because sheets downloaded before the rename are on people's disks
    and a rename must not turn them into silent no-ops.
    """

    databases = {"default", DB}

    NAME = "supplier recommendations"

    def test_every_sheet_uses_the_one_name(self):
        from pipeline.services import dataset
        from pipeline.views.imports import ANTIBODY_COLUMNS
        self.assertIn(self.NAME, ANTIBODY_COLUMNS)
        self.assertIn(self.NAME, board_columns.sheet_headers("antibodies"))
        self.assertIn(self.NAME, dataset.ANTIBODY_COLS)
        for where, cols in (("blank template", ANTIBODY_COLUMNS),
                            ("board export", board_columns.sheet_headers("antibodies")),
                            ("dataset sheet", dataset.ANTIBODY_COLS)):
            with self.subTest(where=where):
                self.assertNotIn("applications", cols)
                self.assertNotIn("supplier claims", cols)

    def test_the_old_spellings_still_import(self):
        from pipeline.services.cropper import metadata as meta
        for spelling in ("supplier recommendations", "supplier claims",
                         "applications", "supplier applications"):
            with self.subTest(spelling=spelling):
                self.assertEqual(
                    meta.HEADER_ALIASES.get(meta._norm_header(spelling)), "apps_raw",
                    f"'{spelling}' no longer reads as the supplier's applications — "
                    f"a sheet downloaded before the rename would lose that column")

    def test_the_column_actually_round_trips(self):
        """The bug itself: parse the header the download writes."""
        from pipeline.services import bulk_antibodies as bulk
        rows = bulk.parse("gene\tcatalogue\tcompany\t%s\n"
                          "SOD1\t12A8\tProteintech\tWB, IF" % self.NAME)
        self.assertEqual(rows[0]["supplier_apps"], ["WB", "IF"])


class TheConcentrationExampleAgreesWithItsHeaderTests(TestCase):
    """A sheet that contradicts itself on the one column where a misread costs a
    factor of a thousand.

    The heading says ``concentration (ug/mL)``; the example row said
    ``1.0 mg/mL``. Both halves are individually defensible — the header names the
    stored unit, and the example exists to teach that a unit written in the cell
    is converted — and together they say two different things about what the
    column holds. A reader who resolves that in the header's favour writes a bare
    ``1.0`` for a milligram stock, which is the original thousand-fold bug typed
    by hand.
    """

    HEADER = "concentration (ug/mL)"

    def _example(self):
        from pipeline.views.imports import ANTIBODY_COLUMNS, ANTIBODY_EXAMPLE
        return ANTIBODY_EXAMPLE[ANTIBODY_COLUMNS.index(self.HEADER)]

    def test_the_example_is_in_the_unit_the_heading_names(self):
        from pipeline.services import concentration
        value, err = concentration.parse(self._example())
        self.assertEqual(err, "")
        bare, _ = concentration.parse(self._example().split()[0])
        self.assertEqual(value, bare,
                         "the example converts, so the heading's unit and the "
                         "example's unit are not the same unit")

    def test_the_example_still_writes_the_unit_out(self):
        """Dropping to a bare number would remove the only place the sheet says
        a unit may be written in the cell at all."""
        self.assertIn("/", self._example())

    def test_the_board_and_the_blank_template_name_the_same_unit(self):
        from pipeline.views.imports import ANTIBODY_COLUMNS
        self.assertIn(self.HEADER, ANTIBODY_COLUMNS)
        self.assertIn(self.HEADER, board_columns.sheet_headers("antibodies"))
