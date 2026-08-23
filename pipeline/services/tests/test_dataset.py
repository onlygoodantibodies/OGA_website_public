"""Whole-dataset download/upload round-trip + the never-destroy guarantees.

Runs against the isolated SQLite pipeline_db from conftest. Proves:
  * download builds a multi-sheet workbook + JSON tree with the id key column
  * gap-fill (empty → value) applies without confirmation
  * overwriting a populated value needs apply_overwrites=True (diff surfaces it)
  * a BLANK cell never clears a populated field (both modes)
  * a row absent from the file is left completely untouched
  * an empty / unrecognised file is refused (nothing written)
  * report DOIs (the headline use case) gap-fill across genes
"""
from __future__ import annotations

import io

import pytest

DB = "pipeline_db"


# ── seeding ─────────────────────────────────────────────────────────────────

@pytest.fixture()
def seeded(_pipeline_db):
    from pipeline.models import Target, Company, Antibody, Report
    # clean slate each test
    for M in (Report, Antibody, Company, Target):
        M.objects.using(DB).all().delete()

    snca = Target.objects.using(DB).create(protein_name="Alpha-synuclein", gene_name="SNCA")
    mapt = Target.objects.using(DB).create(protein_name="Tau", gene_name="MAPT")
    abcam = Company.objects.using(DB).create(name="Abcam")

    ab1 = Antibody.objects.using(DB).create(
        target=snca, company=abcam, catalogue_number="ab138501",
        rrid="", host_species="Rabbit", clonality="unknown",
        supplier_url="")                       # blank url + rrid → gap-fill targets
    ab2 = Antibody.objects.using(DB).create(
        target=mapt, company=abcam, catalogue_number="ab32057",
        host_species="Mouse", supplier_url="https://existing.example/keep")  # populated
    rpt = Report.objects.using(DB).create(target=snca, zenodo_doi="", f1000_doi="")
    return {"snca": snca, "mapt": mapt, "abcam": abcam, "ab1": ab1, "ab2": ab2, "rpt": rpt}


# ── helpers to build an upload file in memory ───────────────────────────────

def _xlsx(sheets: dict):
    """sheets = {title: [header_list, *data_rows]} → an in-memory .xlsx upload."""
    import openpyxl
    from django.core.files.uploadedfile import SimpleUploadedFile
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title[:31])
        for r in rows:
            ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return SimpleUploadedFile(
        "dataset.xlsx", buf.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def _ab_sheet(rows):
    from pipeline.services import dataset
    return {dataset.SHEET_ANTIBODIES: [dataset.ANTIBODY_COLS, *rows]}


# ── downloads ───────────────────────────────────────────────────────────────

def test_download_workbook_has_all_sheets(seeded):
    from pipeline.services import dataset
    import openpyxl
    data = dataset.build_workbook(dataset.resolve_targets())
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert set(wb.sheetnames) == {"Legend", "Targets", "Antibodies", "Cell lines", "Reports"}
    ab = wb["Antibodies"]
    assert [c.value for c in ab[1]] == dataset.ANTIBODY_COLS
    assert ab[1][0].value == "id"                       # id is the key column
    # both seeded antibodies present
    ids = {r[0].value for r in ab.iter_rows(min_row=2)}
    assert seeded["ab1"].pk in ids and seeded["ab2"].pk in ids


def test_download_json_tree(seeded):
    from pipeline.services import dataset
    payload = dataset.build_json(dataset.resolve_targets(genes="SNCA"))
    assert [t["gene"] for t in payload["targets"]] == ["SNCA"]
    t = payload["targets"][0]
    assert t["antibodies"][0]["catalogue"] == "ab138501"
    assert "_legend" in payload and "id column" in payload["_legend"]


# ── upload: gap-fill vs overwrite ───────────────────────────────────────────

def test_gapfill_applies_without_confirmation(seeded):
    from pipeline.services import dataset
    ab1 = seeded["ab1"]
    # fill the blank product url + rrid
    f = _xlsx(_ab_sheet([[ab1.pk, "SNCA", "ab138501", "Abcam", "AB_2537855", "Rabbit",
                          "", "", "", "", "https://www.abcam.com/new", ""]]))
    parsed = dataset.parse_upload(f)
    plan = dataset.plan_upload(parsed)
    assert plan["ok"]
    fill_fields = {i["field"] for i in plan["antibodies"]["fills"]}
    assert {"product url", "RRID"} <= fill_fields
    assert plan["antibodies"]["overwrites"] == []

    res = dataset.apply_upload(parsed, apply_overwrites=False)
    ab1.refresh_from_db(using=DB)
    assert ab1.supplier_url == "https://www.abcam.com/new"
    assert ab1.rrid == "AB_2537855"
    assert res["antibodies_updated"] == 1


def test_overwrite_needs_confirmation(seeded):
    from pipeline.services import dataset
    ab2 = seeded["ab2"]      # supplier_url already populated
    f = _xlsx(_ab_sheet([[ab2.pk, "MAPT", "ab32057", "Abcam", "", "Mouse",
                          "", "", "", "", "https://www.abcam.com/CHANGED", ""]]))
    parsed = dataset.parse_upload(f)
    plan = dataset.plan_upload(parsed)
    ov = plan["antibodies"]["overwrites"]
    assert any(i["field"] == "product url" and i["old"] == "https://existing.example/keep"
               and i["new"] == "https://www.abcam.com/CHANGED" for i in ov)

    # commit WITHOUT confirmation → unchanged
    dataset.apply_upload(dataset.parse_upload(_ab_sheet_file(ab2)), apply_overwrites=False)
    ab2.refresh_from_db(using=DB)
    assert ab2.supplier_url == "https://existing.example/keep"

    # commit WITH confirmation → changed
    dataset.apply_upload(dataset.parse_upload(_ab_sheet_file(ab2)), apply_overwrites=True)
    ab2.refresh_from_db(using=DB)
    assert ab2.supplier_url == "https://www.abcam.com/CHANGED"


def _ab_sheet_file(ab2):
    from pipeline.services import dataset
    return _xlsx(_ab_sheet([[ab2.pk, "MAPT", "ab32057", "Abcam", "", "Mouse",
                             "", "", "", "", "https://www.abcam.com/CHANGED", ""]]))


# ── the never-destroy guarantees ────────────────────────────────────────────

def test_blank_cell_never_clears(seeded):
    from pipeline.services import dataset
    ab2 = seeded["ab2"]      # populated supplier_url + host
    # blank the url cell entirely, even with overwrite on
    f = _xlsx(_ab_sheet([[ab2.pk, "MAPT", "ab32057", "Abcam", "", "", "",
                          "", "", "", "", ""]]))
    dataset.apply_upload(dataset.parse_upload(f), apply_overwrites=True)
    ab2.refresh_from_db(using=DB)
    assert ab2.supplier_url == "https://existing.example/keep"   # untouched
    assert ab2.host_species == "Mouse"                           # untouched


def test_absent_rows_untouched(seeded):
    from pipeline.services import dataset
    ab1, ab2 = seeded["ab1"], seeded["ab2"]
    # file mentions only ab1
    f = _xlsx(_ab_sheet([[ab1.pk, "SNCA", "ab138501", "Abcam", "", "Rabbit",
                          "", "", "", "", "https://www.abcam.com/x", ""]]))
    dataset.apply_upload(dataset.parse_upload(f), apply_overwrites=True)
    ab2.refresh_from_db(using=DB)
    assert ab2.supplier_url == "https://existing.example/keep"   # ab2 never visited
    # and nothing was deleted
    from pipeline.models import Antibody
    assert Antibody.objects.using(DB).count() == 2


def test_empty_and_malformed_files_refused(seeded):
    from pipeline.services import dataset
    # a workbook with only a header, no data rows
    empty = _xlsx({dataset.SHEET_ANTIBODIES: [dataset.ANTIBODY_COLS]})
    plan = dataset.plan_upload(dataset.parse_upload(empty))
    assert not plan["ok"]
    # a workbook with an unrecognised sheet
    junk = _xlsx({"Nonsense": [["a", "b"], ["1", "2"]]})
    parsed = dataset.parse_upload(junk)
    assert parsed["error"]
    assert not dataset.plan_upload(parsed)["ok"]


# ── reports: the DOI gap-fill headline use case ─────────────────────────────

def test_report_doi_gapfill(seeded):
    from pipeline.services import dataset
    rpt = seeded["rpt"]
    f = _xlsx({dataset.SHEET_REPORTS: [
        dataset.REPORT_COLS,
        [rpt.pk, "SNCA", "published", "https://zenodo.org/10.5281/x",
         "https://f1000.example/1", "", ""]]})
    parsed = dataset.parse_upload(f)
    plan = dataset.plan_upload(parsed)
    fill_fields = {i["field"] for i in plan["reports"]["fills"]}
    assert {"Zenodo DOI", "F1000 DOI"} <= fill_fields          # blank DOIs → gap-fill
    # status was already "draft" (the model default) → changing it is an overwrite
    assert any(i["field"] == "status" and i["old"] == "draft" and i["new"] == "published"
               for i in plan["reports"]["overwrites"])

    # gap-fill only: DOIs land, status stays draft (overwrite not confirmed)
    dataset.apply_upload(parsed, apply_overwrites=False)
    rpt.refresh_from_db(using=DB)
    assert rpt.zenodo_doi == "https://zenodo.org/10.5281/x"
    assert rpt.f1000_doi == "https://f1000.example/1"
    assert rpt.status == "draft"

    # with confirmation, status flips to published
    dataset.apply_upload(dataset.parse_upload(f), apply_overwrites=True)
    rpt.refresh_from_db(using=DB)
    assert rpt.status == "published"


def test_report_create_for_gene_without_report(seeded):
    from pipeline.services import dataset
    from pipeline.models import Report
    # MAPT has no report yet; a row with no id but a DOI creates one
    f = _xlsx({dataset.SHEET_REPORTS: [
        dataset.REPORT_COLS,
        ["", "MAPT", "", "https://zenodo.org/mapt", "", "", ""]]})
    parsed = dataset.parse_upload(f)
    res = dataset.apply_upload(parsed, apply_overwrites=False)
    assert res["reports_created"] == 1
    r = Report.objects.using(DB).get(target=seeded["mapt"])
    assert r.zenodo_doi == "https://zenodo.org/mapt"


def test_views_end_to_end(seeded):
    """Drive the real HTTP views (auth decorator + export/preview/commit)."""
    import json
    from django.test import RequestFactory
    from django.contrib.auth.models import User
    from pipeline.models import Site, Member
    from pipeline.services import dataset
    from pipeline import views

    site = Site.objects.using(DB).create(name="Test Site")
    user = User.objects.using(DB).create(username="tester")
    # cross-DB FK: assign via _id, never the object (CLAUDE.md rule + the router)
    Member.objects.using(DB).create(user_id=user.pk, site=site, role="admin", is_active=True)

    rf = RequestFactory()

    # landing page renders
    req = rf.get("/pipeline/data/"); req.user = user
    resp = views.data_io(req)
    assert resp.status_code == 200
    assert b"Download the dataset" in resp.content

    # excel export
    req = rf.get("/pipeline/data/export/?format=xlsx"); req.user = user
    resp = views.dataset_export(req)
    assert resp.status_code == 200
    assert "spreadsheetml" in resp["Content-Type"]

    # json export
    req = rf.get("/pipeline/data/export/?format=json&genes=SNCA"); req.user = user
    resp = views.dataset_export(req)
    payload = json.loads(resp.content)
    assert [t["gene"] for t in payload["targets"]] == ["SNCA"]

    # preview an antibody gap-fill
    upload = _xlsx(_ab_sheet([[seeded["ab1"].pk, "SNCA", "ab138501", "Abcam", "", "Rabbit",
                               "", "", "", "", "https://www.abcam.com/view-e2e", ""]]))
    req = rf.post("/pipeline/data/upload/preview/", {"file": upload}); req.user = user
    plan = json.loads(views.dataset_upload_preview(req).content)
    assert plan["ok"] and plan["counts"]["fills"] >= 1

    # commit it
    upload2 = _xlsx(_ab_sheet([[seeded["ab1"].pk, "SNCA", "ab138501", "Abcam", "", "Rabbit",
                                "", "", "", "", "https://www.abcam.com/view-e2e", ""]]))
    req = rf.post("/pipeline/data/upload/commit/",
                  {"file": upload2, "apply_overwrites": "false"}); req.user = user
    res = json.loads(views.dataset_upload_commit(req).content)
    assert res["ok"] and res["antibodies_updated"] == 1
    seeded["ab1"].refresh_from_db(using=DB)
    assert seeded["ab1"].supplier_url == "https://www.abcam.com/view-e2e"


def _json_upload(payload):
    import json as _json
    from django.core.files.uploadedfile import SimpleUploadedFile
    return SimpleUploadedFile("dataset.json",
                              _json.dumps(payload).encode("utf-8"),
                              content_type="application/json")


def test_json_round_trip_gapfill(seeded):
    """Download JSON, fill gaps in the tree, upload the JSON back (the LLM flow)."""
    from pipeline.services import dataset
    # download the whole dataset as JSON
    payload = dataset.build_json(dataset.resolve_targets())
    # an LLM fills the blank product url on ab1 and a blank report DOI on SNCA
    for t in payload["targets"]:
        if t["gene"] == "SNCA":
            for ab in t["antibodies"]:
                if ab["catalogue"] == "ab138501":
                    ab["product url"] = "https://www.abcam.com/from-json"
            for r in t["reports"]:
                r["zenodo doi"] = "https://zenodo.org/from-json"

    parsed = dataset.parse_upload(_json_upload(payload))
    plan = dataset.plan_upload(parsed)
    assert plan["ok"]
    assert any(i["field"] == "product url" for i in plan["antibodies"]["fills"])
    assert any(i["field"] == "Zenodo DOI" for i in plan["reports"]["fills"])

    dataset.apply_upload(parsed, apply_overwrites=False)
    seeded["ab1"].refresh_from_db(using=DB)
    seeded["rpt"].refresh_from_db(using=DB)
    assert seeded["ab1"].supplier_url == "https://www.abcam.com/from-json"
    assert seeded["rpt"].zenodo_doi == "https://zenodo.org/from-json"


def test_json_blank_never_clears(seeded):
    from pipeline.services import dataset
    payload = dataset.build_json(dataset.resolve_targets(genes="MAPT"))
    # ab2 already has a populated supplier_url; a JSON with it blank must NOT clear
    for t in payload["targets"]:
        for ab in t["antibodies"]:
            ab["product url"] = ""
    dataset.apply_upload(dataset.parse_upload(_json_upload(payload)), apply_overwrites=True)
    seeded["ab2"].refresh_from_db(using=DB)
    assert seeded["ab2"].supplier_url == "https://existing.example/keep"


def test_json_empty_refused(seeded):
    from pipeline.services import dataset
    parsed = dataset.parse_upload(_json_upload({"targets": []}))
    assert not dataset.plan_upload(parsed)["ok"]


def test_new_antibody_created_by_natural_key(seeded):
    from pipeline.services import dataset
    from pipeline.models import Antibody
    # no id, new catalogue on an existing gene → create
    f = _xlsx(_ab_sheet([["", "SNCA", "ab-NEW-1", "Abcam", "AB_999", "Rabbit",
                          "", "", "", "", "https://x", "WB"]]))
    parsed = dataset.parse_upload(f)
    plan = dataset.plan_upload(parsed)
    assert any(c["catalogue"] == "ab-NEW-1" for c in plan["antibodies"]["creates"])
    res = dataset.apply_upload(parsed, apply_overwrites=False)
    assert res["antibodies_created"] == 1
    ab = Antibody.objects.using(DB).get(catalogue_number="ab-NEW-1")
    assert ab.supplier_url == "https://x" and ab.supplier_validated_wb is True


# ── field selector (Phase 1) ────────────────────────────────────────────────

def test_parse_selection_basic():
    from pipeline.services import dataset
    assert dataset.parse_selection("") is None
    assert dataset.parse_selection("   ") is None
    allsel = dataset.parse_selection("all")
    assert set(allsel.keys()) == set(dataset.FAMILY_ORDER)
    # unknown family + unknown column are dropped; keys need not be listed
    s = dataset.parse_selection("antibodies:antigen,bogus|nope:x")
    assert s == {"antibodies": ["antigen"]}


def test_selection_workbook_only_selected_family_and_extra_fields(seeded):
    from pipeline.services import dataset
    import openpyxl
    ab1 = seeded["ab1"]
    ab1.antigen = "synuclein peptide"
    ab1.in_kind_value = 100
    ab1.save(using=DB)

    sel = dataset.parse_selection("antibodies:antigen,in_kind_value")
    data = dataset.build_workbook(dataset.resolve_targets(), selection=sel)
    wb = openpyxl.load_workbook(io.BytesIO(data))
    # only the selected family sheet (+ Legend) is present
    assert set(wb.sheetnames) == {"Legend", "Antibodies"}
    hdr = [c.value for c in wb["Antibodies"][1]]
    assert hdr[:3] == ["id", "gene", "catalogue"]       # keys always lead
    assert "antigen" in hdr and "in_kind_value" in hdr
    assert "company" not in hdr                          # unselected default dropped
    row = next(r for r in wb["Antibodies"].iter_rows(min_row=2, values_only=True) if r[0] == ab1.pk)
    d = dict(zip(hdr, row))
    assert d["antigen"] == "synuclein peptide"
    assert str(d["in_kind_value"]).startswith("100")


def test_selection_json_nests_only_selected_children(seeded):
    from pipeline.services import dataset
    sel = dataset.parse_selection("targets:protein_type|antibodies:antigen")
    payload = dataset.build_json(dataset.resolve_targets(genes="SNCA"), selection=sel)
    t = payload["targets"][0]
    assert "protein_type" in t and "antibodies" in t
    assert "cell_lines" not in t and "reports" not in t
    assert set(payload["_columns"].keys()) == {"targets", "antibodies"}
    assert "antigen" in payload["_columns"]["antibodies"]
    # keys are always carried even if not explicitly listed
    assert "id" in payload["_columns"]["antibodies"] and "catalogue" in payload["_columns"]["antibodies"]


def test_view_export_honours_fields_param(seeded):
    from django.test import RequestFactory
    from django.contrib.auth.models import User
    from pipeline.models import Site, Member
    from pipeline import views
    import openpyxl

    site = Site.objects.using(DB).create(name="Sel Site", short_code="SEL")
    user = User.objects.using(DB).create(username="seltester")
    Member.objects.using(DB).create(user_id=user.pk, site=site, role="admin", is_active=True)

    rf = RequestFactory()
    req = rf.get("/pipeline/data/export/?format=xlsx&fields=antibodies:antigen"); req.user = user
    resp = views.dataset_export(req)
    assert resp.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    assert set(wb.sheetnames) == {"Legend", "Antibodies"}
    assert "antigen" in [c.value for c in wb["Antibodies"][1]]


def test_every_family_the_picker_offers_can_produce_a_sheet(seeded):
    """Ten families on the picker, four sheets in the workbook.

    The field test ticked what the panel showed as ten selected families and got
    a four-sheet file. The panel was the half that was wrong — its locked *key*
    boxes were drawn ticked whether or not the family was going anywhere — but
    the claim underneath is worth pinning at the server too: a family listed by
    `family_catalogue()` and ticked in the URL is a family that gets a sheet.

    Data-independent on purpose. Antibodies, cell lines and reports arrived as
    header-only sheets in that run, which is the correct behaviour for an empty
    scope and is not what "produces nothing" meant.
    """
    from pipeline.services import dataset
    import openpyxl

    offered = [f["key"] for f in dataset.family_catalogue()]
    assert offered == dataset.FAMILY_ORDER

    for key in offered:
        # One non-key field is what the panel requires before it sends a family.
        first = next(h for h, kind, _g in dataset.FAMILIES[key]["fields"]
                     if kind != "key")
        sel = dataset.parse_selection(f"{key}:{first}")
        data = dataset.build_workbook(dataset.resolve_targets(), selection=sel)
        names = openpyxl.load_workbook(io.BytesIO(data)).sheetnames
        assert dataset.FAMILIES[key]["title"][:31] in names, (
            f"{key} is offered by the picker and writes no sheet")


def test_selecting_everything_writes_every_sheet(seeded):
    """The `Everything` preset, end to end — ten families, ten sheets."""
    from pipeline.services import dataset
    import openpyxl

    data = dataset.build_workbook(dataset.resolve_targets(),
                                  selection=dataset.parse_selection("all"))
    names = set(openpyxl.load_workbook(io.BytesIO(data)).sheetnames)
    expected = {dataset.FAMILIES[k]["title"][:31] for k in dataset.FAMILY_ORDER}
    assert expected <= names
    assert len(names) == len(expected) + 1          # + Legend


def test_the_page_reports_every_counter_the_save_returns(seeded):
    """A count the writer returns and the page does not render is a save that
    says nothing.

    `apply_upload` returns eleven counters; `data_io.html` had eight
    hand-written lines, so editing an assignment, a reagent request, a cell
    culture event or an inventory location through the round trip wrote
    correctly and reported "Done — no changes needed". Adding a family to the
    importer without adding it to the receipt is a one-line omission that reads
    as success, which is the shape of nearly every defect in this repo's notes.
    """
    import pathlib
    import re

    from django.conf import settings
    from pipeline.services import dataset

    written = dataset.apply_upload({})               # nothing to write, full shape
    # `not isinstance(v, bool)` because `ok` and `overwrites_applied` are flags,
    # and in Python a bool *is* an int.
    counters = {k for k, v in written.items()
                if isinstance(v, int) and not isinstance(v, bool)}

    page = (pathlib.Path(settings.BASE_DIR)
            / "pipeline/templates/pipeline/data_io.html").read_text()
    block = re.search(r"const SAVED_COUNTS = \[(.*?)\];", page, re.S)
    assert block, "data_io.html no longer lists the counters it renders"
    rendered = set(re.findall(r'\["(\w+)"', block.group(1)))

    assert counters <= rendered, (
        f"apply_upload returns {sorted(counters - rendered)} and the receipt "
        "renders nothing for them")
