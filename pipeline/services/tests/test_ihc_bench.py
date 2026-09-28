"""IHC at the bench, and "the report needs this" on every bench sheet.

PLATFORM_ROADMAP #102 (owner, 26 Sep 2026). What is pinned here is what would
be *silently* wrong:

* an IHC bench sheet's readings landing on the wrong antibody, or its session
  conditions (the block under the table) being dropped or stored under a key
  nothing reads;
* an IHC sheet accepted into another procedure's session, or IHC's headings
  stealing the IF plate map's one distinctive heading so *its* refusal stops;
* the upload preview and the generated report disagreeing about what is
  missing — one declaration (`report_needs`) behind both, checked from the
  report's own printed brackets in both directions;
* a result table that the deletion, renumbering or merge lists do not know,
  which is data lost by cascade with nothing on screen.
"""
from __future__ import annotations

import io
import re
from datetime import date as _date

import pytest

DB = "pipeline_db"

# Gaps a draft prints that are not bench fields, so no preview can warn about
# them (`report_needs` module docstring).
NOT_BENCH = {"[WT cell line]", "[KO cell line]", "[XX]", "[UniProt ID]", "[X.X]"}


@pytest.fixture
def lab(_pipeline_db):
    from django.contrib.auth.models import User
    from pipeline.models import (Antibody, CellLine, Company, ExperimentSession,
                                 FcResult, IfResult, IhcResult, IpResult, Member,
                                 Sample, Site, Target, WbResult)
    for M in (WbResult, IpResult, IfResult, FcResult, IhcResult,
              ExperimentSession, Sample, Antibody, CellLine, Company, Target):
        M.objects.using(DB).all().delete()
    # A site of its own: the pytest modules share one database, and another
    # module creating "McGill" after this one would collide on the name.
    site, _ = Site.objects.using(DB).get_or_create(
        short_code="IHCT", defaults={"name": "IHC bench test site"})
    user, _ = User.objects.using(DB).get_or_create(username="ihcuser")
    member, _ = Member.objects.using(DB).get_or_create(
        user_id=user.pk, defaults={"site": site, "role": "experimenter",
                                   "is_active": True, "display_name": "IHC User"})
    target = Target.objects.using(DB).create(
        gene_name="TP53", protein_name="Cellular tumor antigen p53",
        theoretical_mass_kda=43.7)
    company = Company.objects.using(DB).create(name="Abcam")
    ab1 = Antibody.objects.using(DB).create(
        target=target, company=company, catalogue_number="ab-p53-1",
        lot_number="L1", site=site,
        supplier_recommended_dilutions={"IHC": "1/100"})
    ab2 = Antibody.objects.using(DB).create(
        target=target, company=company, catalogue_number="ab-p53-2",
        lot_number="L2", site=site)
    wt = CellLine.objects.using(DB).create(name="HAP1", genotype="WT", site=site)
    ko = CellLine.objects.using(DB).create(
        name="HAP1 TP53 KO", genotype="KO", target=target, site=site, parent_line=wt)
    return {"site": site, "member": member, "target": target, "ab1": ab1,
            "ab2": ab2, "wt": wt, "ko": ko}


def _session(lab, proc, **kw):
    from pipeline.models import ExperimentSession
    return ExperimentSession.objects.using(DB).create(
        procedure_type=proc, target=lab["target"], experimenter=lab["member"],
        date=_date(2026, 9, 26), site=lab["site"], cell_line_wt=lab["wt"],
        cell_line_ko=lab["ko"], status=ExperimentSession.SessionStatus.PLANNED, **kw)


def _planned_ihc(lab):
    """An IHC session planned the way the board plans one: a blank result row
    per antibody."""
    from pipeline.models import IhcResult
    s = _session(lab, "IHC")
    for ab in (lab["ab1"], lab["ab2"]):
        IhcResult.objects.using(DB).create(session=s, antibody=ab)
    return s


def _sheet(session):
    import openpyxl
    from pipeline.services.planning import generate_bench_sheet
    buf = io.BytesIO()
    generate_bench_sheet(session).save(buf)
    buf.seek(0)
    return openpyxl.load_workbook(buf)


def _upload(wb):
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    buf.name = "filled.xlsx"
    return buf


def _header(ws):
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if any(str(c or "").strip().lower() == "ab#" for c in row):
            return i, [("" if c is None else str(c)) for c in row]
    raise AssertionError("no header row")


def _row_of(ws, header_row, header, catalogue):
    cat = header.index("CatNumber") + 1
    for r in range(header_row + 1, ws.max_row + 1):
        if ws.cell(row=r, column=cat).value == catalogue:
            return r
    raise AssertionError(f"no row for {catalogue}")


def _condition_row(ws, label):
    for r in range(1, ws.max_row + 1):
        if ws.cell(row=r, column=1).value == label:
            return r
    raise AssertionError(f"no conditions-block row labelled {label!r}")


# ── The IHC bench sheet round trip ─────────────────────────────────────────

def test_an_ihc_bench_sheet_round_trips_onto_the_right_antibodies(lab):
    """Download → fill → preview → save → IhcResult rows on the right vials,
    and the conditions block stored under the keys the report reads."""
    from pipeline.models import IhcResult
    from pipeline.services import bench_results

    session = _planned_ihc(lab)
    wb = _sheet(session)
    ws = wb["IHC"]
    i, header = _header(ws)
    col = {h: header.index(h) + 1 for h in header if h}
    assert "Recommended dilution" in col
    r1 = _row_of(ws, i, header, "ab-p53-1")
    r2 = _row_of(ws, i, header, "ab-p53-2")
    assert ws.cell(row=r1, column=col["Recommended dilution"]).value == "1/100"

    ws.cell(row=r1, column=col["Used dilution"], value="1/100")
    ws.cell(row=r1, column=col["Dilution source"], value="supplier-recommended")
    ws.cell(row=r1, column=col["Specific staining"], value="Yes")
    ws.cell(row=r1, column=col["Staining location"], value="nuclear")
    ws.cell(row=r2, column=col["Specific staining"], value="No")
    ws.cell(row=r2, column=col["Secondary / detection"], value="OmniMap anti-rabbit HRP")

    # The block under the table: label in A, value in B, example in D.
    ab_row = _condition_row(ws, "Antigen retrieval")
    assert "95" in str(ws.cell(row=ab_row, column=4).value)   # the example…
    assert not ws.cell(row=ab_row, column=2).value             # …is not the value
    ws.cell(row=ab_row, column=2, value="CC1, 95 °C, 64 min")
    ws.cell(row=_condition_row(ws, "Section thickness (µm)"), column=2, value="4")

    parsed = bench_results.parse(_upload(wb), "IHC")
    assert parsed["sheet_procedure"] == "IHC"
    assert parsed["conditions"] == {"antigen_retrieval": "CC1, 95 °C, 64 min",
                                    "section_thickness_um": "4"}

    preview = bench_results.plan(session, parsed)
    assert preview["ok"] is True
    assert {i["antibody"] for i in preview["items"]} == {str(lab["ab1"]), str(lab["ab2"])}
    # Nothing on this sheet is "not recognised": the block is the sheet's own.
    assert preview["unknown_columns"] == []
    assert {c["key"] for c in preview["conditions_read"]} == {
        "antigen_retrieval", "section_thickness_um"}

    out = bench_results.apply(session, parsed)
    assert out["ok"] is True
    one = IhcResult.objects.using(DB).get(session=session, antibody=lab["ab1"])
    two = IhcResult.objects.using(DB).get(session=session, antibody=lab["ab2"])
    assert (one.primary_ab_dilution, one.dilution_source, one.specific_signal,
            one.staining_location) == ("1/100", "supplier-recommended", "Yes", "nuclear")
    assert (two.specific_signal, two.secondary_ab) == ("No", "OmniMap anti-rabbit HRP")
    assert IhcResult.objects.using(DB).filter(session=session).count() == 2
    session.refresh_from_db(using=DB)
    assert session.session_conditions["antigen_retrieval"] == "CC1, 95 °C, 64 min"
    assert session.session_conditions["section_thickness_um"] == "4"


def test_the_downloaded_block_carries_what_is_already_recorded(lab):
    """A round trip, not a blank form: a condition on file comes back in
    column B, so re-uploading an untouched sheet changes nothing."""
    session = _planned_ihc(lab)
    session.session_conditions = {"stainer": "BenchMark ULTRA"}
    session.save(using=DB)
    ws = _sheet(session)["IHC"]
    assert ws.cell(row=_condition_row(ws, "Stainer"), column=2).value == "BenchMark ULTRA"


# ── Which procedure a sheet is ─────────────────────────────────────────────

def test_an_ihc_sheet_is_refused_by_a_western_blot_session_and_the_reverse(lab):
    from pipeline.services import bench_results
    ihc = _planned_ihc(lab)
    wb_session = _session(lab, "WB")

    ihc_sheet = _sheet(ihc)
    ihc_sheet["IHC"].cell(row=2, column=1, value="an unstamped copy")
    parsed = bench_results.parse(_upload(ihc_sheet), "WB")
    assert parsed["sheet_procedure"] == "IHC"
    refused = bench_results.plan(wb_session, parsed)
    assert refused["ok"] is False and "Immunohistochemistry" in refused["error"]
    assert bench_results.apply(wb_session, parsed)["ok"] is False

    wb_sheet = _sheet(wb_session)
    ws = wb_sheet.worksheets[0]
    ws.cell(row=2, column=1, value="an unstamped copy")
    parsed = bench_results.parse(_upload(wb_sheet), "IHC")
    assert parsed["sheet_procedure"] == "WB"
    assert bench_results.plan(ihc, parsed)["ok"] is False


def test_every_procedure_keeps_a_heading_of_its_own(lab):
    """A heading two procedures share identifies neither. IHC's must not take
    "Specific signal", the IF plate map's only distinctive heading — that is
    what refuses a plate map fed to a western blot (run 11)."""
    from pipeline.services import bench_results
    from pipeline.services.session_board import PROCEDURES
    for proc in PROCEDURES:
        assert bench_results.DISTINCTIVE_HEADERS.get(proc), proc
    assert "specific signal" in bench_results.DISTINCTIVE_HEADERS["IF"]
    assert bench_results.sheet_procedure(["Ab#", "Specific signal"]) == "IF"
    # `Comments` is a column anybody adds to any sheet — never an identity.
    assert bench_results.sheet_procedure(["Ab#", "Used dilution", "Comments"]) is None


# ── The report's needs: one declaration, warned at the upload ──────────────

def test_the_preview_warns_for_each_missing_field_and_the_save_still_happens(lab):
    from pipeline.services import bench_results
    session = _planned_ihc(lab)
    wb = _sheet(session)
    ws = wb["IHC"]
    i, header = _header(ws)
    r1 = _row_of(ws, i, header, "ab-p53-1")
    ws.cell(row=r1, column=header.index("Specific staining") + 1, value="Yes")
    ws.cell(row=_condition_row(ws, "Fixation"), column=2, value="10% NBF, 30 min")

    parsed = bench_results.parse(_upload(wb), "IHC")
    preview = bench_results.plan(session, parsed)
    fields = {w["field"] for w in preview["report_needs"]}
    assert "Antigen retrieval" in fields and "Scanner" in fields
    assert "Fixation" not in fields                 # the sheet fills it
    assert not {"Tissue species", "Tissue organs", "Tissue fixation"} & fields
    retrieval = next(w for w in preview["report_needs"] if w["field"] == "Antigen retrieval")
    assert "[antigen retrieval]" in retrieval["message"]
    dilution = next(w for w in preview["report_needs"] if w["field"] == "Used dilution")
    assert set(dilution["antibodies"]) == {str(lab["ab1"]), str(lab["ab2"])}

    out = bench_results.apply(session, parsed)
    assert out["ok"] is True and out["created"] + out["updated"]
    assert out["report_needs"]                     # said again on the receipt


def test_a_tissue_run_is_asked_for_the_rest_of_its_tissue(lab):
    from pipeline.services import report_needs
    gaps = {w["gap"] for w in report_needs.missing(
        "IHC", {"tissue_species": "mouse"}, [])}
    assert {"[tissue organs]", "[tissue fixation]"} <= gaps
    assert "[tissue species]" not in gaps


def test_every_report_need_names_a_field_somebody_can_fill(lab):
    """A warning that names a field with no control is the refusal that sends
    you nowhere — so every need has a writer in the one registry the form, the
    sessions board and the bench sheet's label map read."""
    from pipeline.services import report_needs
    for proc, needs in report_needs.SESSION_NEEDS.items():
        for n in needs:
            assert report_needs.writer_key(proc, n), (proc, n.name)
    from pipeline.services.session_board import result_field_names
    for proc, needs in report_needs.RESULT_NEEDS.items():
        for n in needs:
            assert set(n.fields) & set(result_field_names(proc)), (proc, n.name)


def _texts(proc, target, session, cell_lines):
    """Every paragraph the draft prints for one procedure's session."""
    from pipeline.services import report_generator as rg
    by_type = {proc: [session]}
    if proc == "WB":
        return [rg._build_wb_methods(target, [session], cell_lines, []),
                rg._wb_legend(target, by_type, cell_lines, [])]
    if proc == "IP":
        return [rg._build_ip_methods(target, [session], cell_lines),
                rg._ip_legend(target, by_type, cell_lines)]
    if proc == "IF":
        return [rg._build_if_methods(target, [session], cell_lines),
                rg._if_legend(target, by_type, cell_lines)]
    if proc == "FC":
        return [rg._build_fc_methods(target, [session], cell_lines),
                rg._fc_legend(target, by_type, cell_lines)]
    return [rg._build_ihc_methods(target, [session], cell_lines),
            rg._ihc_legend(target, by_type, cell_lines)]


_A_READING = {"WB": {"signal": "clean band"}, "IP": {"enrichment": "yes"},
              "IF": {"specific_signal": "yes"}, "FC": {"histogram_shift": "yes"},
              "IHC": {"specific_signal": "yes"}}


@pytest.mark.parametrize("proc", ["WB", "IP", "IF", "FC", "IHC"])
def test_the_preview_and_the_report_name_the_same_gaps(lab, proc):
    """Generate the report's paragraphs from a session with a reading and no
    conditions, and read the brackets they print. They must be exactly the gaps
    `report_needs.missing` warns about — nothing the report prints unwarned,
    nothing warned that the report would not print."""
    from pipeline.services import report_needs
    from pipeline.services.session_board import RESULT_MODELS
    session = _session(lab, proc)
    RESULT_MODELS[proc].objects.using(DB).create(
        session=session, antibody=lab["ab1"], **_A_READING[proc])
    session.refresh_from_db(using=DB)
    cell_lines = [lab["wt"], lab["ko"]]

    printed = set()
    for text in _texts(proc, lab["target"], session, cell_lines):
        printed |= set(re.findall(r"\[[^\]]+\]", text))
    printed -= NOT_BENCH

    rows = list(RESULT_MODELS[proc].objects.using(DB).filter(session=session))
    warned = {w["gap"] for w in report_needs.missing(
        proc, report_needs.session_conditions(session), rows)}
    assert printed == warned


@pytest.mark.parametrize("proc", ["WB", "IP", "IF", "FC", "IHC"])
def test_a_fully_recorded_session_prints_no_gap_and_warns_about_nothing(lab, proc):
    from pipeline.services import report_needs
    from pipeline.services.session_board import RESULT_MODELS
    conditions = {n.keys[0]: f"value-{n.name}"
                  for n in report_needs.SESSION_NEEDS[proc]}
    session = _session(lab, proc, session_conditions=conditions)
    fields = dict(_A_READING[proc])
    for n in report_needs.RESULT_NEEDS[proc]:
        fields[n.fields[0]] = f"value-{n.name}"
    RESULT_MODELS[proc].objects.using(DB).create(
        session=session, antibody=lab["ab1"], **fields)
    session.refresh_from_db(using=DB)

    printed = set()
    for text in _texts(proc, lab["target"], session, [lab["wt"], lab["ko"]]):
        printed |= set(re.findall(r"\[[^\]]+\]", text))
    assert printed - NOT_BENCH == set()
    rows = list(RESULT_MODELS[proc].objects.using(DB).filter(session=session))
    assert report_needs.missing(proc, report_needs.session_conditions(session), rows) == []


def test_the_ihc_paragraph_prints_what_was_recorded_and_names_the_rest(lab):
    from pipeline.models import IhcResult
    from pipeline.services import report_generator as rg
    session = _session(lab, "IHC", session_conditions={
        "fixation": "10% NBF, 30 min", "stainer": "BenchMark ULTRA",
        "mosaic_ratio": "1:1"})
    IhcResult.objects.using(DB).create(
        session=session, antibody=lab["ab1"], specific_signal="yes",
        primary_ab_dilution="1/100", dilution_source="as recommended by the supplier",
        secondary_ab="OmniMap anti-rabbit HRP")
    methods = rg._build_ihc_methods(lab["target"], [session], [lab["wt"], lab["ko"]])
    assert "fixed with 10% NBF, 30 min" in methods
    assert "stained using the BenchMark ULTRA" in methods
    assert "OmniMap anti-rabbit HRP" in methods
    assert "1:1" in methods                          # optional, recorded
    assert "Antigen retrieval was performed with [antigen retrieval]" in methods
    assert "[chromogen]" in methods and "[scanner]" in methods
    assert "tissue" not in methods.lower()           # no tissue recorded
    assert "RIPA" not in methods and "Aperio" not in methods   # nothing invented
    legend = rg._ihc_legend(lab["target"], {"IHC": [session]}, [lab["wt"], lab["ko"]])
    assert "1/100 (as recommended by the supplier)" in legend


def test_a_draft_for_a_gene_with_only_ihc_readings_has_its_methods(lab, tmp_path):
    """End to end once: the antibody is in Table 2, the procedure is named, and
    the IHC methods heading is there — none of which existed before."""
    from docx import Document
    from pipeline.models import IhcResult
    from pipeline.services import report_generator as rg
    session = _session(lab, "IHC", session_conditions={"chromogen": "DAB, 8 min"})
    IhcResult.objects.using(DB).create(session=session, antibody=lab["ab1"],
                                       specific_signal="yes")
    path = rg.generate_report(lab["target"].pk, output_path=str(tmp_path / "d.docx"))
    doc = Document(path)
    text = "\n".join(p.text for p in doc.paragraphs)
    cells = " ".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    assert "Antibody screening by immunohistochemistry" in text
    assert "DAB, 8 min" in text
    assert "ab-p53-1" in cells
    # The procedure is named where the draft names what was done — the title
    # and abstract read `procedures_list`, which did not know IHC.
    assert "[no application has a recorded result yet]" not in text
    assert "Figure 5" in text


# ── The step-by-step form / payload door ───────────────────────────────────

def test_a_bench_payload_cannot_set_the_ihc_verdict(lab):
    """IHC's verdict is judged by eye (owner, 26 Sep 2026); the payload door
    records readings and refuses `selective` by name rather than writing the
    public flag."""
    from pipeline.services import sessions as sess
    payload = {"procedure_type": "IHC", "gene": "TP53", "date": "2026-09-26",
               "experimenter_id": lab["member"].pk, "site_id": lab["site"].pk,
               "results": [{"antibody": "ab-p53-1", "specific_signal": "yes",
                            "selective": True}]}
    out = sess.apply(payload, member=lab["member"])
    assert out["ok"] is False
    assert any("selective" in e for e in out["errors"])
    del payload["results"][0]["selective"]
    out = sess.apply(payload, member=lab["member"])
    assert out["ok"] is True
    lab["ab1"].refresh_from_db(using=DB)
    assert lab["ab1"].ihc_recommended is False


# ── Every list of "the result tables" knows IHC ────────────────────────────

def test_deletion_renumbering_and_merges_know_every_result_table(lab):
    """Each of these loses data silently if a result table is missing: an
    antibody with readings deletes unwarned, a printed A-number is moved, a
    merge deletes the loser's readings by cascade."""
    from pipeline import dedup_utils
    from pipeline.management.commands import fix_biotechne_brands
    from pipeline.services import duplicates, renumber
    from pipeline.services.session_board import RESULT_MODELS
    accessors = {m._meta.get_field("antibody").remote_field.related_name
                 for m in RESULT_MODELS.values()}
    assert "ihc_results" in accessors
    assert accessors <= set(renumber._PROCEDURES)
    assert accessors <= set(duplicates.CHILD_RELATIONS)
    assert accessors <= set(dedup_utils.CHILD_RELATIONS)
    assert accessors <= set(fix_biotechne_brands.CHILD_RELATIONS)

    from pipeline.models import IhcResult
    from pipeline.services import deletion
    s = _session(lab, "IHC")
    IhcResult.objects.using(DB).create(session=s, antibody=lab["ab2"], specific_signal="yes")
    assert deletion._result_count(antibody=lab["ab2"]) == 1


# ── Review fixes (26 Sep 2026) ─────────────────────────────────────────────

def test_the_ihc_draft_prints_the_recorded_embedding_and_nothing_from_ppp2r5d(lab):
    """The legend and the results paragraph said "paraffin embedded", and the
    methods "placed on charged slides" and "each core", whatever the bench
    recorded — the PPP2R5D run's values, printed as this run's."""
    from pipeline.models import IhcResult
    from pipeline.services import report_generator as rg
    session = _session(lab, "IHC", session_conditions={
        "embedding": "OCT, snap frozen", "fixation": "acetone"})
    IhcResult.objects.using(DB).create(session=session, antibody=lab["ab1"],
                                       specific_signal="yes")
    texts = _texts("IHC", lab["target"], session, [lab["wt"], lab["ko"]])
    joined = " ".join(texts).lower()
    assert "oct, snap frozen" in texts[1].lower()         # the legend says it
    for invented in ("paraffin", "charged", "each core"):
        assert invented not in joined, invented


@pytest.mark.parametrize("proc,row", [
    ("WB", {"gel": "4-20% TGX", "membrane": "PVDF", "ecl": "Clarity",
            "detection_system": "ChemiDoc", "secondary_ab": "HRP anti-rabbit"}),
    ("IP", {"bead_type": "protein G", "gel": "4-20% TGX"}),
    ("IF", {"fixative": "4% PFA", "blocking": "5% BSA",
            "permeabilisation": "0.1% Triton", "secondary_ab": "AF555 anti-rabbit"}),
])
def test_a_value_recorded_on_the_result_row_is_neither_warned_nor_a_gap(lab, proc, row):
    """The workbook's WB, IP and IF tabs offer these as per-row columns. A
    value typed there was warned blank and printed as `[membrane]`."""
    from pipeline.services import report_needs
    from pipeline.services.session_board import RESULT_MODELS
    session = _session(lab, proc)
    RESULT_MODELS[proc].objects.using(DB).create(
        session=session, antibody=lab["ab1"], **_A_READING[proc], **row)
    session.refresh_from_db(using=DB)
    printed = set()
    for text in _texts(proc, lab["target"], session, [lab["wt"], lab["ko"]]):
        printed |= set(re.findall(r"\[[^\]]+\]", text))
    rows = list(RESULT_MODELS[proc].objects.using(DB).filter(session=session))
    warned = {w["gap"] for w in report_needs.missing(
        proc, report_needs.session_conditions(session), rows)}
    covered = {f"[{n.gap}]" for n in report_needs.SESSION_NEEDS[proc]
               if set(n.result_fields) & set(row)}
    assert covered and not covered & printed and not covered & warned
    assert printed - NOT_BENCH == warned                  # still one list


def _corrupt_cell(wb):
    """An xlsx whose first sheet holds a numeric cell that is not a number —
    openpyxl raises a bare ValueError reading it."""
    import zipfile
    raw = io.BytesIO()
    wb.save(raw)
    src = zipfile.ZipFile(io.BytesIO(raw.getvalue()))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as z:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = data.replace(b"</sheetData>",
                                    b'<row r="90"><c r="A90" t="n"><v>1.2.3x</v></c></row>'
                                    b"</sheetData>", 1)
            z.writestr(item, data)
    out.seek(0)
    out.name = "damaged.xlsx"
    return out


def test_a_damaged_sheet_is_refused_in_our_words_not_the_librarys(lab):
    """Only `SheetRefusal` may reach the page; openpyxl's ValueError (`could
    not convert string to float`) went into the reply verbatim."""
    from django.test import RequestFactory
    from pipeline.services import bench_results, session_import
    from pipeline.views import session_bulk
    session = _planned_ihc(lab)

    with pytest.raises(session_import.SheetRefusal) as refused:
        session_import.parse_template(_corrupt_cell(_sheet(session)))
    assert "convert" not in str(refused.value) and "damaged" in str(refused.value)

    with pytest.raises(ValueError) as raised:
        bench_results.parse(_corrupt_cell(_sheet(session)), "IHC")
    assert not isinstance(raised.value, bench_results.SheetRefusal)

    f = _corrupt_cell(_sheet(session))
    from django.core.files.uploadedfile import SimpleUploadedFile
    upload = SimpleUploadedFile("damaged.xlsx", f.getvalue())
    request = RequestFactory().post("/x/", {"file": upload})
    reply = session_bulk.session_results_upload.__wrapped__.__wrapped__(request, session.pk)
    body = reply.content.decode()
    assert "convert" not in body and "1.2.3x" not in body


def test_the_conditions_block_fills_only_blanks(lab):
    """The block is printed pre-filled; an older print re-uploaded must not
    put its value back over a correction made on the board since."""
    from pipeline.services import bench_results
    session = _planned_ihc(lab)
    session.session_conditions = {"stainer": "Leica BOND"}
    session.save(using=DB)
    wb = _sheet(session)
    ws = wb["IHC"]
    ws.cell(row=_condition_row(ws, "Stainer"), column=2, value="BenchMark ULTRA")
    ws.cell(row=_condition_row(ws, "Scanner"), column=2, value="Aperio GT450")
    parsed = bench_results.parse(_upload(wb), "IHC")
    preview = bench_results.plan(session, parsed)
    read = {c["key"]: c for c in preview["conditions_read"]}
    assert read["stainer"]["held"] is True and read["scanner"]["held"] is False
    bench_results.apply(session, parsed)
    session.refresh_from_db(using=DB)
    assert session.session_conditions["stainer"] == "Leica BOND"
    assert session.session_conditions["scanner"] == "Aperio GT450"


def test_a_warning_says_whether_the_sheet_has_a_place_for_it(lab):
    """Only IHC's bench sheet has a conditions block; a WB sheet sending you
    to fill in the membrane "on the sheet" sent you nowhere."""
    from pipeline.models import WbResult
    from pipeline.services import bench_results
    wb_session = _session(lab, "WB")
    WbResult.objects.using(DB).create(session=wb_session, antibody=lab["ab1"])
    warnings = bench_results.report_warnings(wb_session, {}, [{"antibody": str(lab["ab1"])}])
    kinds = {(w["kind"], w["on_sheet"]) for w in warnings}
    assert ("session", False) in kinds and ("antibody", True) in kinds
    ihc = bench_results.report_warnings(_planned_ihc(lab), {}, [])
    assert ihc and all(w["on_sheet"] for w in ihc)


def test_both_importers_spell_a_condition_heading_the_same_way(lab):
    from pipeline.services import bench_results, session_import
    for header in ("Section thickness (µm)", "Mosaic ratio (WT:KO)", "Owner notes"):
        assert (session_import._condition_key(header, "IHC")
                == bench_results._condition_key(header, "IHC"))
    assert session_import._condition_key("Section thickness (µm)", "IHC") == "section_thickness_um"
    assert session_import._condition_key("Owner notes", "IHC") == "owner_notes"
