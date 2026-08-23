"""The workbook round trip: a gene tab creates, a session tab fills in.

Before this, the per-gene workbook could **only** create and a session's bench
sheet could **only** update, so a session planned through the step-by-step form
could not be filled in from a spreadsheet at all — and the gene page needed a
paragraph of prose explaining which of two doors you were standing in. The
`session_ref` column carries the answer now: `(new)` creates, `#480` fills 480 in.

Also pinned here: the wild type surviving the trip. `session_import` looked
`cell_line_wt` up in `target.cell_lines`, and a wild type has no gene, so the
column the workbook itself pre-fills was discarded on every row of every tab.
"""
from __future__ import annotations

import io
from datetime import date as _date

import pytest

DB = "pipeline_db"


@pytest.fixture
def gene(_pipeline_db):
    from pipeline.models import (Antibody, CellLine, Company, ExperimentSession,
                                 Member, ProtocolTemplate, Sample, Site, Target,
                                 WbResult)
    for M in (WbResult, ExperimentSession, Sample, ProtocolTemplate, Antibody,
              CellLine, Company, Target):
        M.objects.using(DB).all().delete()

    from django.contrib.auth.models import User
    site, _ = Site.objects.using(DB).get_or_create(
        short_code="LEI", defaults={"name": "Leicester"})
    # Cross-DB FK: the Member row lives in pipeline_db and points at the auth
    # user by id, never by object (CLAUDE.md).
    user, _ = User.objects.using(DB).get_or_create(username="benchuser")
    member, _ = Member.objects.using(DB).get_or_create(
        user_id=user.pk, defaults={"site": site, "role": "experimenter",
                                   "is_active": True, "display_name": "Bench User"})
    target = Target.objects.using(DB).create(gene_name="ELP3", protein_name="ELP3")
    company = Company.objects.using(DB).create(name="Proteintech")
    ab1 = Antibody.objects.using(DB).create(
        target=target, company=company, catalogue_number="A-ELP3-1")
    ab2 = Antibody.objects.using(DB).create(
        target=target, company=company, catalogue_number="A-ELP3-2")
    # The shape that matters: the WT carries NO gene, the KO points at it.
    wt = CellLine.objects.using(DB).create(name="HAP1", genotype="WT", site=site)
    ko = CellLine.objects.using(DB).create(
        name="HAP1 ELP3 KO", genotype="KO", target=target, site=site, parent_line=wt)
    ProtocolTemplate.objects.using(DB).create(
        name="Leicester WB Standard", site=site, procedure_type="WB", is_default=True,
        conditions={"lysis_buffer": "RIPA", "protein_loading_ug": "30"})
    return {"site": site, "member": member, "target": target,
            "wt": wt, "ko": ko, "ab1": ab1, "ab2": ab2}


def _load(data):
    import openpyxl
    return openpyxl.load_workbook(io.BytesIO(data))


def _tab(wb, name):
    ws = wb[name]
    rows = [[("" if c is None else str(c)) for c in r] for r in ws.iter_rows(values_only=True)]
    return rows[0], rows[1:]


def _to_file(wb):
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    buf.name = "filled.xlsx"
    return buf


def _set(header, row, col, value):
    row[header.index(col)] = value


# ── The wild type survives the round trip ─────────────────────────────────

def test_the_workbook_ships_the_wild_type_and_the_importer_keeps_it(gene):
    """F3: the app filled this column in and then threw it away.

    `target.cell_lines` cannot contain a wild type — it has no gene — so the
    lookup could only ever return None. Three sessions from one upload came out
    with a blank WT, and it surfaced two steps later as `[WT cell line]` in a
    generated Data Note.
    """
    from pipeline.models import ExperimentSession
    from pipeline.services import session_import, session_template

    wb = _load(session_template.build_gene_template("ELP3", site_id=gene["site"].pk))
    header, rows = _tab(wb, "WB")
    assert rows[0][header.index("cell_line_wt")].startswith("HAP1")

    ws = wb["WB"]
    for r in range(2, 2 + len(rows)):
        ws.cell(row=r, column=header.index("signal") + 1, value="single band")
    res = session_import.apply_import(
        session_import.parse_template(_to_file(wb)), uploader=gene["member"])

    assert res["ok"] and res["sessions_created"] == 1
    session = ExperimentSession.objects.using(DB).get(procedure_type="WB")
    assert session.cell_line_wt_id == gene["wt"].pk
    assert session.cell_line_ko_id == gene["ko"].pk


def test_the_preview_names_the_wild_type_it_will_store(gene):
    """A preview is only true if it is compared against the write."""
    from pipeline.services import session_import, session_template
    wb = _load(session_template.build_gene_template("ELP3", site_id=gene["site"].pk))
    header, rows = _tab(wb, "WB")
    ws = wb["WB"]
    ws.cell(row=2, column=header.index("signal") + 1, value="single band")
    plan = session_import.plan_import(
        session_import.parse_template(_to_file(wb)), uploader=gene["member"])
    s = plan["sessions"][0]
    assert s["cell_line_wt"].startswith("HAP1")
    assert s["cell_line_wt_error"] is None


# ── A session tab fills that session in ───────────────────────────────────

def _planned_session(gene, proc="WB"):
    from pipeline.models import ExperimentSession
    return ExperimentSession.objects.using(DB).create(
        procedure_type=proc, target=gene["target"], experimenter=gene["member"],
        date=_date(2026, 8, 1), site=gene["site"],
        status=ExperimentSession.SessionStatus.PLANNED,
        session_conditions={"lysis_buffer": "RIPA"})


def test_a_session_workbook_is_stamped_with_its_number(gene):
    from pipeline.services import session_template
    session = _planned_session(gene)
    wb = _load(session_template.build_session_template(session))
    assert wb.sheetnames == ["How to use", "WB"]
    header, rows = _tab(wb, "WB")
    assert rows, "a planned session falls back to the gene's vials as a picking list"
    for row in rows:
        assert row[header.index("session_ref")].endswith(f"#{session.pk}")


def test_uploading_it_fills_that_session_in_rather_than_creating_a_second(gene):
    """The whole point. This used to make a second, Complete session beside the
    Planned one, with no explanation on either screen."""
    from pipeline.models import ExperimentSession, WbResult
    from pipeline.services import session_import, session_template

    session = _planned_session(gene)
    wb = _load(session_template.build_session_template(session))
    header, rows = _tab(wb, "WB")
    ws = wb["WB"]
    for r in range(2, 2 + len(rows)):
        ws.cell(row=r, column=header.index("signal") + 1, value="clean band")
        ws.cell(row=r, column=header.index("rating") + 1, value="5")

    plan = session_import.plan_import(
        session_import.parse_template(_to_file(wb)), uploader=gene["member"])
    assert plan["counts"]["sessions"] == 0        # nothing created
    assert plan["counts"]["updated"] == 1
    assert plan["sessions"][0]["mode"] == "update"
    assert plan["sessions"][0]["session_id"] == session.pk
    assert "Complete" in plan["sessions"][0]["status_change"]

    res = session_import.apply_import(
        session_import.parse_template(_to_file(wb)), uploader=gene["member"])
    assert res["sessions_created"] == 0 and res["sessions_updated"] == 1
    assert ExperimentSession.objects.using(DB).count() == 1

    session.refresh_from_db()
    assert session.status == ExperimentSession.SessionStatus.COMPLETE
    assert WbResult.objects.using(DB).filter(session=session).count() == 2
    assert all(r.signal == "clean band"
               for r in WbResult.objects.using(DB).filter(session=session))


def test_a_second_upload_updates_the_same_rows_instead_of_doubling_them(gene):
    from pipeline.models import WbResult
    from pipeline.services import session_import, session_template

    session = _planned_session(gene)

    def _filled(signal):
        wb = _load(session_template.build_session_template(session))
        header, rows = _tab(wb, "WB")
        ws = wb["WB"]
        for r in range(2, 2 + len(rows)):
            ws.cell(row=r, column=header.index("signal") + 1, value=signal)
        return session_import.parse_template(_to_file(wb))

    session_import.apply_import(_filled("first pass"), uploader=gene["member"])
    res = session_import.apply_import(_filled("corrected"), uploader=gene["member"])

    assert res["results_created"] == 0 and res["results_updated"] == 2
    rows = WbResult.objects.using(DB).filter(session=session)
    assert rows.count() == 2
    assert {r.signal for r in rows} == {"corrected"}


def test_a_blank_cell_does_not_clear_what_is_recorded(gene):
    """Fill-only-blank, the rule every other write path holds."""
    from pipeline.models import WbResult
    from pipeline.services import session_import, session_template

    session = _planned_session(gene)
    wb = _load(session_template.build_session_template(session))
    header, rows = _tab(wb, "WB")
    ws = wb["WB"]
    for r in range(2, 2 + len(rows)):
        ws.cell(row=r, column=header.index("signal") + 1, value="a band")
        ws.cell(row=r, column=header.index("rating") + 1, value="4")
    session_import.apply_import(session_import.parse_template(_to_file(wb)),
                                uploader=gene["member"])

    # Download again, clear `rating`, put it back.
    wb2 = _load(session_template.build_session_template(session))
    h2, r2 = _tab(wb2, "WB")
    assert r2[0][h2.index("rating")] == "4"      # pre-filled from what is recorded
    ws2 = wb2["WB"]
    for r in range(2, 2 + len(r2)):
        ws2.cell(row=r, column=h2.index("rating") + 1, value="")
        ws2.cell(row=r, column=h2.index("signal") + 1, value="a band")
    session_import.apply_import(session_import.parse_template(_to_file(wb2)),
                                uploader=gene["member"])

    assert all(r.rating == "4" for r in WbResult.objects.using(DB).filter(session=session))


# ── A ref that names nothing is refused, never turned into a create ───────

def test_an_unknown_session_number_is_refused(gene):
    from pipeline.models import ExperimentSession
    from pipeline.services import session_import, session_template

    session = _planned_session(gene)
    wb = _load(session_template.build_session_template(session))
    header, rows = _tab(wb, "WB")
    ws = wb["WB"]
    for r in range(2, 2 + len(rows)):
        ws.cell(row=r, column=header.index("session_ref") + 1, value="WB · ELP3 #99999")
        ws.cell(row=r, column=header.index("signal") + 1, value="a band")

    plan = session_import.plan_import(
        session_import.parse_template(_to_file(wb)), uploader=gene["member"])
    assert plan["counts"]["sessions"] == 0 and plan["counts"]["updated"] == 0
    assert any("#99999 is not in the database" in b["reason"] for b in plan["blocked"])

    before = ExperimentSession.objects.using(DB).count()
    res = session_import.apply_import(
        session_import.parse_template(_to_file(wb)), uploader=gene["member"])
    assert res["sessions_created"] == 0 and res["sessions_updated"] == 0
    assert ExperimentSession.objects.using(DB).count() == before


def test_a_ref_naming_another_procedures_session_is_refused(gene):
    from pipeline.services import session_import, session_template

    ip = _planned_session(gene, proc="IP")
    wb = _load(session_template.build_session_template(_planned_session(gene, "WB")))
    header, rows = _tab(wb, "WB")
    ws = wb["WB"]
    for r in range(2, 2 + len(rows)):
        ws.cell(row=r, column=header.index("session_ref") + 1, value=f"WB · ELP3 #{ip.pk}")
        ws.cell(row=r, column=header.index("signal") + 1, value="a band")

    plan = session_import.plan_import(
        session_import.parse_template(_to_file(wb)), uploader=gene["member"])
    assert plan["counts"]["updated"] == 0
    assert any("Immunoprecipitation session" in b["reason"] for b in plan["blocked"])


def test_an_untouched_session_tab_is_still_not_an_experiment(gene):
    """`_has_result` gates the fill-in path too: a picking list nobody wrote on
    must not become a Complete session with blank results."""
    from pipeline.models import ExperimentSession
    from pipeline.services import session_import, session_template

    session = _planned_session(gene)
    wb = _load(session_template.build_session_template(session))
    res = session_import.apply_import(
        session_import.parse_template(_to_file(wb)), uploader=gene["member"])
    assert res["sessions_updated"] == 0 and res["results_created"] == 0
    session.refresh_from_db()
    assert session.status == ExperimentSession.SessionStatus.PLANNED
