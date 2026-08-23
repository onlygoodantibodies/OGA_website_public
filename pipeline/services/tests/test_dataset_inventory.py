"""Flat Inventory family in the whole-dataset download/upload.

InventoryLocation is polymorphic (antibody | cell line | vial), so it can't nest
under a single gene — it exports as a FLAT top-level array in JSON and its own
`Inventory` sheet in Excel. These tests prove: it exports one row per location
across all three parent types with a derived gene, edits by id (gap-fill /
overwrite-confirm / blank-never-clears / no-id blocked), and that the polymorphic
parent links can never be re-pointed via upload."""
from __future__ import annotations

import io
import json as _json

import pytest

DB = "pipeline_db"


@pytest.fixture()
def seeded(_pipeline_db):
    from pipeline.models import (Target, Company, Antibody, CellLine, CellLineVial,
                                 InventoryLocation, Site)
    for M in (InventoryLocation, CellLineVial, Antibody, CellLine, Company, Target):
        M.objects.using(DB).all().delete()
    site, _ = Site.objects.using(DB).get_or_create(short_code="INV", defaults={"name": "Inv Site"})
    tgt = Target.objects.using(DB).create(protein_name="Alpha-synuclein", gene_name="SNCA")
    comp = Company.objects.using(DB).create(name="Abcam")
    ab = Antibody.objects.using(DB).create(target=tgt, company=comp, catalogue_number="ab138501")
    cl = CellLine.objects.using(DB).create(target=tgt, name="SNCA KO", genotype="KO")
    vial = CellLineVial.objects.using(DB).create(cell_line=cl, c_number=42)
    loc_ab = InventoryLocation.objects.using(DB).create(
        antibody=ab, site=site, storage_type="-20", building="", box="A1")
    loc_cl = InventoryLocation.objects.using(DB).create(
        cell_line=cl, site=site, storage_type="ln2", box="")
    loc_vial = InventoryLocation.objects.using(DB).create(
        vial=vial, site=site, storage_type="-80", box="")
    return {"site": site, "tgt": tgt, "ab": ab, "cl": cl, "vial": vial,
            "loc_ab": loc_ab, "loc_cl": loc_cl, "loc_vial": loc_vial}


def _xlsx(sheets):
    import openpyxl
    from django.core.files.uploadedfile import SimpleUploadedFile
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title[:31])
        for r in rows:
            ws.append(r)
    buf = io.BytesIO(); wb.save(buf)
    return SimpleUploadedFile("d.xlsx", buf.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def _json_upload(payload):
    from django.core.files.uploadedfile import SimpleUploadedFile
    return SimpleUploadedFile("d.json", _json.dumps(payload).encode(),
                              content_type="application/json")


# ── export ───────────────────────────────────────────────────────────────────

def test_inventory_exports_flat_top_level_with_derived_gene(seeded):
    from pipeline.services import dataset
    payload = dataset.build_json(dataset.resolve_targets(genes="SNCA"),
                                 selection=dataset.parse_selection("all"))
    # flat: a top-level array, NOT nested under the gene
    assert "inventory" in payload
    for t in payload["targets"]:
        assert "inventory" not in t
    inv = {row["id"]: row for row in payload["inventory"]}
    assert set(inv) == {seeded["loc_ab"].pk, seeded["loc_cl"].pk, seeded["loc_vial"].pk}
    # each row knows what it hangs off and derives the gene from that parent
    assert inv[seeded["loc_ab"].pk]["attached_to"] == "antibody"
    assert inv[seeded["loc_cl"].pk]["attached_to"] == "cell_line"
    assert inv[seeded["loc_vial"].pk]["attached_to"] == "vial"
    for row in inv.values():
        assert row["gene"] == "SNCA"
    assert inv[seeded["loc_ab"].pk]["parent"] == "ab138501"
    assert inv[seeded["loc_vial"].pk]["parent"] == "C42"


def test_inventory_gets_its_own_excel_sheet(seeded):
    import openpyxl
    from pipeline.services import dataset
    xlsx = dataset.build_workbook(dataset.resolve_targets(genes="SNCA"),
                                  selection=dataset.parse_selection("all"))
    wb = openpyxl.load_workbook(io.BytesIO(xlsx))
    assert "Inventory" in wb.sheetnames
    # re-uploading the unedited export is a clean no-op
    from django.core.files.uploadedfile import SimpleUploadedFile
    up = SimpleUploadedFile("d.xlsx", xlsx,
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    plan = dataset.plan_upload(dataset.parse_upload(up))
    assert plan["ok"]
    assert plan["inventory"]["fills"] == [] and plan["inventory"]["overwrites"] == []


# ── write-back ───────────────────────────────────────────────────────────────

def test_inventory_gapfill_and_overwrite_by_id(seeded):
    from pipeline.services import dataset
    loc = seeded["loc_ab"]                        # building='', box='A1', storage_type='-20'
    sheet = {"Inventory": [["id", "gene", "building", "box", "storage_type"],
                           [loc.pk, "SNCA", "Genetics", "B7", "-80"]]}
    parsed = dataset.parse_upload(_xlsx(sheet))
    plan = dataset.plan_upload(parsed)
    assert any(i["field"] == "building" for i in plan["inventory"]["fills"])   # '' → value
    ov = {i["field"] for i in plan["inventory"]["overwrites"]}
    assert {"box", "storage_type"} <= ov                                       # populated → confirm

    res = dataset.apply_upload(parsed, apply_overwrites=False)
    assert res["inventory_updated"] == 1
    loc.refresh_from_db(using=DB)
    assert loc.building == "Genetics"             # gap-fill applied
    assert loc.box == "A1" and loc.storage_type == "-20"   # overwrites held back

    dataset.apply_upload(dataset.parse_upload(_xlsx(sheet)), apply_overwrites=True)
    loc.refresh_from_db(using=DB)
    assert loc.box == "B7" and loc.storage_type == "-80"


def test_inventory_parent_links_never_repointed(seeded):
    """Even with antibody/cell_line/vial columns in the file, the polymorphic
    parent is never changed — those columns aren't editable."""
    from pipeline.services import dataset
    loc = seeded["loc_ab"]
    other_cl = seeded["cl"]
    sheet = {"Inventory": [["id", "attached_to", "cell_line id", "vial id", "building"],
                           [loc.pk, "cell_line", other_cl.pk, "999", "moved?"]]}
    dataset.apply_upload(dataset.parse_upload(_xlsx(sheet)), apply_overwrites=True)
    loc.refresh_from_db(using=DB)
    assert loc.antibody_id == seeded["ab"].pk     # still the antibody it was
    assert loc.cell_line_id is None and loc.vial_id is None
    assert loc.building == "moved?"               # the editable field did apply


def test_inventory_blank_never_clears(seeded):
    from pipeline.services import dataset
    loc = seeded["loc_ab"]                         # box='A1'
    sheet = {"Inventory": [["id", "gene", "box"], [loc.pk, "SNCA", ""]]}
    dataset.apply_upload(dataset.parse_upload(_xlsx(sheet)), apply_overwrites=True)
    loc.refresh_from_db(using=DB)
    assert loc.box == "A1"


def test_inventory_no_id_is_blocked_not_created(seeded):
    from pipeline.services import dataset
    from pipeline.models import InventoryLocation
    before = InventoryLocation.objects.using(DB).count()
    sheet = {"Inventory": [["id", "gene", "box"], ["", "SNCA", "Z9"]]}
    parsed = dataset.parse_upload(_xlsx(sheet))
    plan = dataset.plan_upload(parsed)
    assert plan["inventory"]["blocked"]
    dataset.apply_upload(parsed, apply_overwrites=True)
    assert InventoryLocation.objects.using(DB).count() == before


def test_inventory_json_roundtrip(seeded):
    from pipeline.services import dataset
    loc = seeded["loc_cl"]
    payload = dataset.build_json(dataset.resolve_targets(genes="SNCA"),
                                 selection=dataset.parse_selection("all"))
    for row in payload["inventory"]:
        if row["id"] == loc.pk:
            row["building"] = "Cryo room"
            row["shelf"] = "3"
    dataset.apply_upload(dataset.parse_upload(_json_upload(payload)), apply_overwrites=False)
    loc.refresh_from_db(using=DB)
    assert loc.building == "Cryo room" and loc.shelf == "3"
