"""Replay the seventh field test's own workbook against the fixed code.

Not a substitute for the unit tests — a sanity check that the real artefact,
byte for byte as `hsv6` uploaded it, now lands where the report says it should
have. `cddc2247-session_template_ELP3_filled.xlsx` is the file that created
sessions 481/482/483 with a blank wild type.
"""
from __future__ import annotations

import io
import pathlib
import shutil
from datetime import date as _date

import pytest

DB = "pipeline_db"

# The real file, vendored — 12 KB, and it is the strongest test here precisely
# because nobody wrote it to be one. `hsv6` downloaded it, filled in the IP, IF
# and FC tabs at a bench, added a column of their own, and uploaded it; three
# sessions were created with a blank wild type and the placeholder reached a
# generated Data Note two steps later.
UPLOAD = pathlib.Path(__file__).parent / "fixtures" / "run7_ELP3_workbook_filled.xlsx"


@pytest.fixture
def elp3(_pipeline_db):
    from django.contrib.auth.models import User
    from pipeline.models import (Antibody, CellLine, Company, ExperimentSession,
                                 Member, Sample, Site, Target)
    for M in (ExperimentSession, Sample, Antibody, CellLine, Company, Target):
        M.objects.using(DB).all().delete()
    site, _ = Site.objects.using(DB).get_or_create(
        short_code="LEI", defaults={"name": "Leicester"})
    # A second site with its own HAP1 parental — the shape that made a bare name
    # ambiguous in the real database.
    mcg, _ = Site.objects.using(DB).get_or_create(
        short_code="MCG", defaults={"name": "McGill"})
    user, _ = User.objects.using(DB).get_or_create(username="hsv6")
    member, _ = Member.objects.using(DB).get_or_create(
        user_id=user.pk, defaults={"site": site, "role": "experimenter",
                                   "is_active": True, "display_name": "Harvinder Virk"})
    t = Target.objects.using(DB).create(gene_name="ELP3", protein_name="Elongator complex protein 3")
    for name in ("abcam", "Proteintech", "Cell Signaling Technology"):
        Company.objects.using(DB).get_or_create(name=name)
    for cat, comp in (("A-ELP3-R7A", "abcam"), ("A-ELP3-R7B", "Proteintech"),
                      ("A-ELP3-R7C", "Cell Signaling Technology")):
        Antibody.objects.using(DB).create(
            target=t, company=Company.objects.using(DB).get(name=comp),
            catalogue_number=cat)
    wt = CellLine.objects.using(DB).create(name="HAP1", genotype="WT", site=site)
    CellLine.objects.using(DB).create(name="HAP1", genotype="WT", site=mcg)
    CellLine.objects.using(DB).create(
        name="HAP1 ELP3 KO", genotype="KO", target=t, site=site, parent_line=wt)
    return {"site": site, "member": member, "target": t, "wt": wt}


def _upload():
    buf = io.BytesIO()
    with open(UPLOAD, "rb") as fh:
        shutil.copyfileobj(fh, buf)
    buf.seek(0)
    buf.name = "session_template_ELP3_filled.xlsx"
    return buf


def test_the_real_run7_workbook_now_records_its_wild_type(elp3):
    from pipeline.models import ExperimentSession
    from pipeline.services import session_import

    res = session_import.apply_import(
        session_import.parse_template(_upload()), uploader=elp3["member"])
    assert res["ok"], res
    # IP, IF and FC were filled in; WB was left blank and must stay uncreated.
    assert res["sessions_created"] == 3, res
    made = ExperimentSession.objects.using(DB).all()
    assert {s.procedure_type for s in made} == {"IP", "IF", "FC"}
    for s in made:
        assert s.cell_line_wt_id == elp3["wt"].pk, f"{s.procedure_type} lost its WT"
        assert s.cell_line_ko is not None
    print("\n  sessions:", sorted(f"{s.procedure_type}:{s.cell_line_wt.name}" for s in made))


def test_the_invented_column_is_named_and_every_value_kept(elp3):
    """`hsv6` wrote a different page reference on each of the three IP rows.

    Run 7 stored none of them. The fix after it stored the first and *said* it
    was discarding two, which is the run-11 finding one importer over: naming a
    loss is not preventing one. All three survive into the one condition now —
    asserted against the real file, which is the only place the three values
    exist without somebody having written them to make a point.
    """
    from pipeline.services import session_import
    plan = session_import.plan_import(
        session_import.parse_template(_upload()), uploader=elp3["member"])
    ip = next(s for s in plan["sessions"] if s["procedure"] == "IP")
    assert ip["unknown_columns"], "Owner notes should still be reported"
    u = ip["unknown_columns"][0]
    assert u["key"] == "owner_notes"
    assert u["value"] == ("HV bench book p.31 | HV bench book p.32 "
                          "| HV bench book p.33")
    assert u["values"] == 3 and u["rows"] == 3
    # …and the WT it will store is named, not the cell as typed.
    assert ip["cell_line_wt"].startswith("HAP1")
    assert ip["cell_line_wt_error"] is None
    print("\n  IP preview WT:", ip["cell_line_wt"], "| unknown:", u)


def test_the_report_no_longer_prints_the_placeholder(elp3):
    """F1 end to end: the Method section names the wild type the tables do."""
    import docx
    from pipeline.services import report_generator, session_import
    session_import.apply_import(session_import.parse_template(_upload()),
                                uploader=elp3["member"])
    path = report_generator.generate_report(
        target_pk=elp3["target"].pk,
        output_path="/tmp/claude-0/-home-user-OGA-website/"
                    "34eaeee6-308d-5fc4-a2e8-ab8ef0e6517f/scratchpad/ELP3_check.docx")
    doc = docx.Document(path)
    offenders = [p.text for p in doc.paragraphs if "[WT cell line]" in p.text]
    assert not offenders, offenders
    table1 = doc.tables[0]
    rows = [[c.text for c in r.cells] for r in table1.rows[1:]]
    print("\n  Table 1:", rows)
    assert any(r[-1] == "WT" and r[-2] == "HAP1" for r in rows)
    # Nothing from another gene or another site.
    assert all(r[-1] in ("WT", "ELP3 KO") for r in rows), rows


def test_the_method_leaves_a_named_gap_rather_than_inventing_a_number(elp3):
    """A draft that invents a number is worse than one with a hole in it.

    `RUN7_STMN2_report.docx` stated *"RIPA buffer"* and *"40 µg of protein"* for
    a session whose conditions were empty and whose `protein_loading_ug` was
    null — fluent, specific, and describing an experiment nobody had recorded.
    The hole gets filled in; the number gets published.
    """
    import docx
    from pipeline.models import ExperimentSession, Member, Site, Target, WbResult, Antibody
    from pipeline.services import report_generator

    # A WB session with nothing recorded about how it was run.
    session = ExperimentSession.objects.using(DB).create(
        procedure_type="WB", target=elp3["target"], experimenter=elp3["member"],
        date=_date(2026, 8, 1), site=elp3["site"], session_conditions={},
        status=ExperimentSession.SessionStatus.COMPLETE)
    WbResult.objects.using(DB).create(
        session=session, antibody=Antibody.objects.using(DB).first(),
        signal="single band", rating="5")

    path = report_generator.generate_report(
        target_pk=elp3["target"].pk,
        output_path="/tmp/claude-0/-home-user-OGA-website/"
                    "34eaeee6-308d-5fc4-a2e8-ab8ef0e6517f/scratchpad/ELP3_gaps.docx")
    text = "\n".join(p.text for p in docx.Document(path).paragraphs)

    # Two paragraphs open "For western blot experiments" — the Results summary
    # and the Method. Both stated details nobody had recorded, so both are checked.
    paras = [p for p in text.split("\n") if "western blot experiments" in p]
    assert len(paras) >= 2, paras
    method = next(p for p in paras if "lysates were prepared" in p)
    summary = next(p for p in paras if "SDS-PAGE" in p)
    print("\n  Method:  ", method[:220])
    print("  Results: ", summary[:220])

    # The gaps are named, so whoever edits the draft knows what to supply…
    for named in ("[lysis buffer]", "[protein loading]", "[gel chemistry]",
                  "[membrane]", "[blocking buffer]"):
        assert named in method, f"{named} missing from the Method paragraph"
    assert "[membrane]" in summary

    # …and nothing plausible was invented in their place, in either.
    for invented in ("RIPA", "40 \u00b5g", "4-20% Tris-Glycine", "nitrocellulose",
                     "5% milk in TBST"):
        assert invented not in method, f"Method still invents {invented!r}"
        assert invented not in summary, f"Results still invents {invented!r}"
