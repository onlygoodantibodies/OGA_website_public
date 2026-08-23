"""Per-gene session template: tabs per application, pre-filled from the DB +
the site default protocol. Download-only — proves shape, not writes."""
from __future__ import annotations

import io

import pytest

DB = "pipeline_db"


@pytest.fixture()
def seeded(_pipeline_db):
    from pipeline.models import (Target, Company, Antibody, CellLine, Site,
                                 ProtocolTemplate, Sample)
    # Sample before CellLine/Target; leave Sites alone (Samples elsewhere PROTECT them).
    for M in (Sample, ProtocolTemplate, Antibody, CellLine, Company, Target):
        M.objects.using(DB).all().delete()

    site, _ = Site.objects.using(DB).get_or_create(short_code="LEI", defaults={"name": "Leicester"})
    ProtocolTemplate.objects.using(DB).filter(site=site).delete()
    tgt = Target.objects.using(DB).create(protein_name="SOD1", gene_name="SOD1")
    company = Company.objects.using(DB).create(name="Proteintech")
    Antibody.objects.using(DB).create(target=tgt, company=company, catalogue_number="12A8")
    Antibody.objects.using(DB).create(target=tgt, company=company, catalogue_number="sc-101523")
    CellLine.objects.using(DB).create(target=tgt, name="SW620 WT", genotype="WT")
    CellLine.objects.using(DB).create(target=tgt, name="SW620 SOD1-KO", genotype="KO")
    ProtocolTemplate.objects.using(DB).create(
        name="Leicester WB Standard", site=site, procedure_type="WB", is_default=True,
        conditions={"lysis_buffer": "RIPA", "protein_loading_ug": "20"})
    return {"site": site, "target": tgt}


def _load(data):
    import openpyxl
    return openpyxl.load_workbook(io.BytesIO(data))


def test_template_has_tab_per_application(seeded):
    from pipeline.services import session_template
    wb = _load(session_template.build_gene_template("SOD1", site_id=seeded["site"].pk))
    assert {"WB", "IP", "IF", "FC"} <= set(wb.sheetnames)


def test_template_case_insensitive_and_unknown(seeded):
    from pipeline.services import session_template
    # lower-case gene still resolves
    assert session_template.build_gene_template("sod1", site_id=seeded["site"].pk)
    with pytest.raises(ValueError):
        session_template.build_gene_template("NOTAGENE")


def test_wb_tab_prefilled_and_shared_session(seeded):
    from pipeline.services import session_template
    wb = _load(session_template.build_gene_template("SOD1", site_id=seeded["site"].pk))
    ws = wb["WB"]
    hdr = [c.value for c in ws[1]]
    assert hdr[:6] == ["session_ref", "gene", "antibody", "company", "cell_line_wt", "cell_line_ko"]
    # Protocol conditions appear as pre-filled columns, under the `cond:` prefix
    # that keeps them from colliding with a result field of the same name. This
    # test asserted the bare names and had been failing since the prefix landed.
    assert "cond:lysis_buffer" in hdr and "cond:protein_loading_ug" in hdr
    # result columns present, to be completed
    assert "signal" in hdr and "rating" in hdr

    rows = [[c.value for c in r] for r in ws.iter_rows(min_row=2)]
    assert len(rows) == 2                                   # one row per antibody
    d0 = dict(zip(hdr, rows[0]))
    d1 = dict(zip(hdr, rows[1]))
    assert d0["session_ref"] == d1["session_ref"]          # ONE shared session
    assert d0["session_ref"].endswith("(new)")             # …and this file CREATES
    assert d0["cell_line_wt"] == "SW620 WT" and d0["cell_line_ko"] == "SW620 SOD1-KO"
    assert d0["cond:lysis_buffer"] == "RIPA"               # from site default protocol
    assert {d0["antibody"], d1["antibody"]} == {"12A8", "sc-101523"}
    assert d0["signal"] in (None, "")                      # result blank


def test_tab_without_protocol_still_builds(seeded):
    from pipeline.services import session_template
    # IF has no template seeded → tab still exists with context + result columns
    wb = _load(session_template.build_gene_template("SOD1", site_id=seeded["site"].pk))
    hdr = [c.value for c in wb["IF"][1]]
    assert "fixative" in hdr and "permeabilisation" in hdr
    assert hdr[0] == "session_ref"


def test_view_downloads_xlsx(seeded):
    from django.test import RequestFactory
    from django.contrib.auth.models import User
    from pipeline.models import Member
    from pipeline import views

    user = User.objects.using(DB).create(username="sessuser")
    Member.objects.using(DB).create(user_id=user.pk, site=seeded["site"], role="admin", is_active=True)
    rf = RequestFactory()
    req = rf.get("/pipeline/session/template/?gene=SOD1"); req.user = user
    resp = views.session_template_export(req)
    assert resp.status_code == 200
    assert "spreadsheetml" in resp["Content-Type"]
    assert "session_template_SOD1.xlsx" in resp["Content-Disposition"]

    # unknown gene → 404 text, nothing written
    req = rf.get("/pipeline/session/template/?gene=ZZZ"); req.user = user
    assert views.session_template_export(req).status_code == 404
