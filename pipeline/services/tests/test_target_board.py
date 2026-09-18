"""The target board and Carl's-workbook round-trip.

The cases here are the ones the real file forced: headers on row 5, free-text
dates, ``-`` as a placeholder, one gene on two rows under different projects, and
a gene written two different ways. Plus the two things the spreadsheet cannot do
— cross-site duplicate detection and per-application completion.
"""
from __future__ import annotations

import datetime as dt
import io

import pytest

DB = "pipeline_db"

# Carl's layout: four rows of preamble, headers on row 5, data from row 6.
CARL_HEADERS = [
    "Granting agencies", "Project", "Funding", "Date of nomination", "Comments",
    "Which YcharOS site", "protein name", "gene name", "alternative protein name",
    "Uniprot ID", "Theoretical Molecular Mass (kDa)", "Essential gene (depmap)?",
    "Antibodies requested?", "Status", "Conclusion", "Zenodo DOI",
    "Date added on Zenodo", "To prioritize for F1000", "On F1000",
    "Published on F1000",
]


def _carl_workbook(rows, *, header_row=5):
    """Build a workbook shaped like the real one."""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Complete target list"
    ws.cell(1, 26, "legend")
    ws.cell(4, 7, "Target information")
    for i, h in enumerate(CARL_HEADERS, start=1):
        ws.cell(header_row, i, h)
    for r, row in enumerate(rows, start=header_row + 1):
        for i, v in enumerate(row, start=1):
            ws.cell(r, i, v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _upload(data, name="carl.xlsx"):
    from django.core.files.uploadedfile import SimpleUploadedFile
    return SimpleUploadedFile(name, data)


@pytest.fixture()
def clean(_pipeline_db):
    from pipeline.models import (GrantingAgency, Project, Report, Target,
                                 TargetClassification, TargetNomination)
    for M in (TargetClassification, TargetNomination, Report, Target):
        M.objects.using(DB).all().delete()
    Project.objects.using(DB).all().delete()
    GrantingAgency.objects.using(DB).all().delete()
    return True


# ---------------------------------------------------------------------------
# Parsing his actual file shape
# ---------------------------------------------------------------------------

def test_finds_headers_on_row_five(clean):
    from pipeline.services import target_list_io as tio
    data = _carl_workbook([["NIH", "AMP-AD", "Funding available", "2020", "",
                            "McGill", "Alpha-synuclein", "SNCA", "", "P37840",
                            14.46, "NO", "YES", "report uploaded on Zenodo",
                            "High quality antibodies identified", "", "", "", "", ""]])
    parsed = tio.parse_workbook(_upload(data))
    assert parsed["ok"]
    assert parsed["header_row"] == 5
    assert len(parsed["rows"]) == 1
    assert parsed["columns_missing"] == []


def test_finds_headers_on_row_one_too(clean):
    """A cleaned-up export has headers on row 1 — both must round-trip."""
    from pipeline.services import target_list_io as tio
    data = _carl_workbook([["NIH", "", "", "", "", "", "", "SNCA", "", "", "", "",
                            "", "", "", "", "", "", "", ""]], header_row=1)
    parsed = tio.parse_workbook(_upload(data))
    assert parsed["ok"] and parsed["header_row"] == 1


def test_refuses_a_file_that_is_not_the_target_list(clean):
    import openpyxl
    from pipeline.services import target_list_io as tio
    wb = openpyxl.Workbook()
    wb.active.append(["something", "else", "entirely"])
    buf = io.BytesIO()
    wb.save(buf)
    parsed = tio.parse_workbook(_upload(buf.getvalue()))
    assert not parsed["ok"]
    assert "gene name" in parsed["error"]


@pytest.mark.parametrize("raw,expected,kept", [
    ("2020", dt.date(2020, 1, 1), "2020"),
    ("2021-05", dt.date(2021, 5, 1), "2021-05"),
    ("april 2024", dt.date(2024, 4, 1), "april 2024"),
    ("03/22/2023", dt.date(2023, 3, 22), ""),
    ("2021-12-06", dt.date(2021, 12, 6), ""),
    ("-", None, "-"),
    ("", None, ""),
])
def test_date_parsing_keeps_what_was_written(raw, expected, kept):
    """A year-only entry must not silently become 1 January in the export."""
    from pipeline.services import target_list_io as tio
    assert tio.parse_date(raw) == (expected, kept)


def test_funding_words_and_blank(clean):
    from pipeline.services import target_list_io as tio
    assert tio.parse_funding("Funding available") is True
    assert tio.parse_funding("Funding not available") is False
    assert tio.parse_funding("") is None          # blank means "leave it alone"
    assert tio.parse_funding("-") is None


# ---------------------------------------------------------------------------
# The duplicate-gene case: 2 rows, 1 target, 2 nominations
# ---------------------------------------------------------------------------

def test_one_gene_on_two_rows_becomes_two_nominations(clean):
    from pipeline.models import Target, TargetNomination
    from pipeline.services import target_list_io as tio
    data = _carl_workbook([
        ["NIH", "AMP-AD", "Funding available", "2020", "", "McGill", "", "SYNGAP1",
         "", "", "", "", "", "", "", "", "", "", "", ""],
        ["", "Autism-related", "", "", "", "", "", "SYNGAP1",
         "", "", "", "", "", "", "", "", "", "", "", ""],
    ])
    tio.apply(tio.parse_workbook(_upload(data)))
    assert Target.objects.using(DB).filter(gene_name__iexact="SYNGAP1").count() == 1
    target = Target.objects.using(DB).get(gene_name__iexact="SYNGAP1")
    noms = TargetNomination.objects.using(DB).filter(target=target)
    assert noms.count() == 2
    assert {n.project.name for n in noms} == {"AMP-AD", "Autism-related"}
    assert [n.funded for n in noms.order_by("project__name")] == [True, False]


def test_case_variant_gene_is_the_same_target(clean):
    """The real sheet has both ``Rab3A`` and ``RAB3A``."""
    from pipeline.models import Target, TargetNomination
    from pipeline.services import target_list_io as tio
    data = _carl_workbook([
        ["", "Rab'ome", "Funding available", "", "", "McGill", "", "Rab3A",
         "", "", "", "", "", "", "", "", "", "", "", ""],
        ["NIH", "TREAT-AD_A", "Funding available", "", "", "McGill", "", "RAB3A",
         "", "", "", "", "", "", "", "", "", "", "", ""],
    ])
    tio.apply(tio.parse_workbook(_upload(data)))
    assert Target.objects.using(DB).filter(gene_name__iexact="RAB3A").count() == 1
    assert TargetNomination.objects.using(DB).count() == 2


def test_reupload_is_idempotent(clean):
    from pipeline.models import Target, TargetNomination
    from pipeline.services import target_list_io as tio
    data = _carl_workbook([
        ["NIH", "AMP-AD", "Funding available", "2020", "", "McGill", "", "SNCA",
         "", "", "", "", "", "", "", "", "", "", "", ""]])
    tio.apply(tio.parse_workbook(_upload(data)))
    second = tio.apply(tio.parse_workbook(_upload(data)))
    assert second["nominations_created"] == 0
    assert Target.objects.using(DB).count() == 1
    assert TargetNomination.objects.using(DB).count() == 1


# ---------------------------------------------------------------------------
# Write rules
# ---------------------------------------------------------------------------

def test_blank_cell_never_clears_and_populated_value_is_not_overwritten(clean):
    from pipeline.models import Target
    from pipeline.services import target_list_io as tio
    first = _carl_workbook([
        ["", "", "", "", "", "McGill", "Alpha-synuclein", "SNCA", "", "", "", "NO",
         "", "", "", "", "", "", "", ""]])
    tio.apply(tio.parse_workbook(_upload(first)))

    # Same gene, blank protein name and a *different* essentiality call.
    second = _carl_workbook([
        ["", "", "", "", "", "McGill", "", "SNCA", "", "", "", "YES",
         "", "", "", "", "", "", "", ""]])
    tio.apply(tio.parse_workbook(_upload(second)))
    t = Target.objects.using(DB).get(gene_name__iexact="SNCA")
    assert t.protein_name == "Alpha-synuclein"   # blank did not clear it
    assert t.essential_gene == "NO"              # populated value not overwritten


def test_overwrite_is_opt_in(clean):
    from pipeline.models import Target
    from pipeline.services import target_list_io as tio
    tio.apply(tio.parse_workbook(_upload(_carl_workbook([
        ["", "", "", "", "", "McGill", "", "SNCA", "", "", "", "NO",
         "", "", "", "", "", "", "", ""]]))))
    tio.apply(tio.parse_workbook(_upload(_carl_workbook([
        ["", "", "", "", "", "McGill", "", "SNCA", "", "", "", "YES",
         "", "", "", "", "", "", "", ""]]))), apply_overwrites=True)
    assert Target.objects.using(DB).get(gene_name__iexact="SNCA").essential_gene == "YES"


def test_disagreement_is_reported_before_anything_is_written(clean):
    from pipeline.services import target_list_io as tio
    tio.apply(tio.parse_workbook(_upload(_carl_workbook([
        ["", "", "", "", "", "McGill", "AT-rich protein 1B", "ARID1B", "", "", "",
         "", "", "", "", "", "", "", "", ""]]))))
    plan = tio.plan(tio.parse_workbook(_upload(_carl_workbook([
        ["", "", "", "", "", "McGill", "T-rich protein 1B", "ARID1B", "", "", "",
         "", "", "", "", "", "", "", "", ""]]))))
    assert plan["summary"]["conflicts"] == 1
    conflict = plan["items"][0]["conflicts"][0]
    assert conflict["field"] == "protein name"
    assert conflict["in_database"] == "AT-rich protein 1B"


def test_f1000_column_splits_links_from_prose(clean):
    from pipeline.models import Report
    from pipeline.services import target_list_io as tio
    tio.apply(tio.parse_workbook(_upload(_carl_workbook([
        ["", "", "", "", "", "McGill", "", "SNCA", "", "", "", "", "", "", "", "",
         "", "", "https://doi.org/10.12688/f1000research.1", ""],
        ["", "", "", "", "", "McGill", "", "MAPT", "", "", "", "", "", "", "", "",
         "", "", "coming", ""],
    ]))))
    snca = Report.objects.using(DB).get(target__gene_name__iexact="SNCA")
    mapt = Report.objects.using(DB).get(target__gene_name__iexact="MAPT")
    assert snca.f1000_doi.endswith("f1000research.1") and snca.f1000_note == ""
    assert mapt.f1000_note == "coming" and mapt.f1000_doi == ""


# ---------------------------------------------------------------------------
# Completion, cross-site duplicates, family routing
# ---------------------------------------------------------------------------

def test_completed_means_zenodo_report_or_f1000_published(clean):
    from pipeline.models import Report, Target
    from pipeline.services import target_board as board
    zen = Target.objects.using(DB).create(protein_name="A", gene_name="AAA1")
    f10 = Target.objects.using(DB).create(protein_name="B", gene_name="BBB1")
    nei = Target.objects.using(DB).create(protein_name="C", gene_name="CCC1")
    Report.objects.using(DB).create(target=zen, zenodo_doi="https://doi.org/10.5281/zenodo.1")
    Report.objects.using(DB).create(target=f10, f1000_date=dt.date(2025, 1, 1))
    Report.objects.using(DB).create(target=nei, f1000_priority="YES")  # queued only
    done = {r["gene"]: r["completed"] for r in board.board_rows()}
    assert done["AAA1"] is True
    assert done["BBB1"] is True
    assert done["CCC1"] is False


def test_same_gene_at_two_sites_is_flagged(clean):
    from pipeline.models import Site, Target, TargetNomination
    from pipeline.services import target_board as board
    mcgill, _ = Site.objects.using(DB).get_or_create(
        short_code="MCG", defaults={"name": "McGill"})
    leicester, _ = Site.objects.using(DB).get_or_create(
        short_code="LEI", defaults={"name": "Leicester"})
    t = Target.objects.using(DB).create(protein_name="Midkine", gene_name="MDK")
    TargetNomination.objects.using(DB).create(target=t, site=mcgill, funded=True)
    TargetNomination.objects.using(DB).create(target=t, site=leicester, funded=True)

    dupes = board.duplicate_targets()
    assert len(dupes) == 1
    assert dupes[0]["gene"] == "MDK"
    assert sorted(dupes[0]["sites"]) == ["Leicester", "McGill"]

    check = board.nomination_check("MDK")
    assert check["exists"]
    assert any("already on the list" in w for w in check["warnings"])


def test_family_routing_points_at_the_site_with_the_expertise(clean):
    """Carl's Rab case: a Rab nominated elsewhere should be routed to whoever
    has already done most of the family."""
    from pipeline.models import Report, Site, Target, TargetNomination
    from pipeline.services import target_board as board
    mcgill, _ = Site.objects.using(DB).get_or_create(
        short_code="MCG", defaults={"name": "McGill"})
    for n in range(1, 6):
        t = Target.objects.using(DB).create(protein_name=f"Rab{n}", gene_name=f"RAB{n}A")
        TargetNomination.objects.using(DB).create(target=t, site=mcgill, funded=True)
        Report.objects.using(DB).create(
            target=t, zenodo_doi=f"https://doi.org/10.5281/zenodo.{n}")

    check = board.nomination_check("RAB44")
    assert check["exists"] is False              # genuinely new gene…
    assert check["family"] == "RAB"
    assert check["family_expertise"]["McGill"]["completed"] == 5
    assert any("McGill has 5 RAB targets" in n for n in check["notes"])


def test_gene_family_heuristic():
    from pipeline.services.protein_class import gene_family
    assert gene_family("RAB11A") == "RAB"
    assert gene_family("SLC17A7") == "SLC"
    assert gene_family("TRIM33") == "TRIM"
    assert gene_family("TP53") == ""     # two-letter prefix is too noisy
    assert gene_family("SNCA") == ""     # no digit, no family


def test_classify_works_with_no_network():
    """UniProt is routinely unreachable in dev; the family label must survive."""
    from pipeline.services import protein_class
    labels = protein_class.classify("RAB11A", accession="", terms={"found": False})
    assert [d["label"] for d in labels] == ["RAB family"]


def test_classify_maps_keywords_to_classes():
    from pipeline.services import protein_class
    labels = protein_class.classify("ADRB2", terms={
        "found": True,
        "keywords": ["G-protein coupled receptor", "Cell membrane", "Glycoprotein"],
        "go_terms": ["nucleus"], "locations": [],
    })
    got = {d["label"]: d["source"] for d in labels}
    assert got["GPCR"] == "uniprot"
    assert got["Membrane"] == "uniprot"
    assert got["Nuclear"] == "go"     # GO only fills what keywords missed


# ---------------------------------------------------------------------------
# Export
# ---------------------------------------------------------------------------

def test_both_export_layouts_reimport(clean):
    from pipeline.services import target_list_io as tio
    tio.apply(tio.parse_workbook(_upload(_carl_workbook([
        ["NIH", "AMP-AD", "Funding available", "2020", "note", "McGill",
         "Alpha-synuclein", "SNCA", "", "P37840", 14.46, "NO", "YES", "done",
         "High quality antibodies identified",
         "https://doi.org/10.5281/zenodo.1", "2024-03-14", "YES", "", ""]]))))

    for layout, header_row in (("carl", 5), ("board", 1)):
        data = tio.export_bytes(layout=layout)
        back = tio.parse_workbook(_upload(data, f"{layout}.xlsx"))
        assert back["ok"], layout
        assert back["header_row"] == header_row
        assert len(back["rows"]) == 1
        assert back["columns_missing"] == []


def test_carl_layout_keeps_his_column_order(clean):
    import openpyxl
    from pipeline.services import target_list_io as tio
    wb = openpyxl.load_workbook(io.BytesIO(tio.export_bytes(layout="carl")))
    ws = wb["Complete target list"]
    assert [ws.cell(5, i).value for i in range(1, 21)] == CARL_HEADERS
    assert ws.cell(4, 7).value == "Target information"


def test_board_layout_adds_the_computed_columns(clean):
    import openpyxl
    from pipeline.services import target_list_io as tio
    wb = openpyxl.load_workbook(io.BytesIO(tio.export_bytes(layout="board")))
    ws = wb["Complete target list"]
    headers = [c.value for c in ws[1]]
    for extra in ("Sites pursuing", "Protein classes", "Completed",
                  "Applications done"):
        assert extra in headers
    assert "Read me" in wb.sheetnames


def test_year_only_date_exports_as_the_year_not_a_made_up_day(clean):
    import openpyxl
    from pipeline.services import target_list_io as tio
    tio.apply(tio.parse_workbook(_upload(_carl_workbook([
        ["NIH", "", "", "2020", "", "McGill", "", "SNCA", "", "", "", "", "", "",
         "", "", "", "", "", ""]]))))
    wb = openpyxl.load_workbook(io.BytesIO(tio.export_bytes(layout="carl")))
    ws = wb["Complete target list"]
    assert ws.cell(6, 4).value == "2020"


def test_no_nomination_is_not_the_same_as_not_funded(clean):
    """Before a target list is imported, every target has no nomination. The
    board must not render that as "not funded" — it is a claim we cannot make."""
    from pipeline.models import Target, TargetNomination
    from pipeline.services import target_board as board

    unknown = Target.objects.using(DB).create(protein_name="A", gene_name="UNKNOWN1")
    known = Target.objects.using(DB).create(protein_name="B", gene_name="KNOWN1")
    TargetNomination.objects.using(DB).create(target=known, funded=False)

    rows = {r["gene"]: r for r in board.board_rows()}
    assert rows["UNKNOWN1"]["funded"] is False
    assert rows["UNKNOWN1"]["has_nomination"] is False   # renders as "—"
    assert rows["KNOWN1"]["funded"] is False
    assert rows["KNOWN1"]["has_nomination"] is True      # renders as "not funded"


def test_watchlist_counts_nominations_not_bare_targets(clean):
    from pipeline.models import Target, TargetNomination
    from pipeline.services import target_board as board

    Target.objects.using(DB).create(protein_name="A", gene_name="BARE1")
    Target.objects.using(DB).create(protein_name="B", gene_name="BARE2")
    wanted = Target.objects.using(DB).create(protein_name="C", gene_name="WISHED1")
    TargetNomination.objects.using(DB).create(target=wanted, funded=False)

    totals = board.portfolio()["totals"]
    assert totals["unfunded_watchlist"] == 1     # only the nominated one
    assert totals["no_nomination"] == 2


def test_carl_layout_carries_tips_without_breaking_the_parser(clean):
    """The tips sit in the four rows above the headers — the same rows the old
    workbook used for its colour legend — so they must not be read as data."""
    import openpyxl
    from pipeline.services import target_list_io as tio
    tio.apply(tio.parse_workbook(_upload(_carl_workbook([
        ["NIH", "AMP-AD", "Funding available", "2020", "", "McGill", "", "SNCA",
         "", "", "", "", "", "", "", "", "", "", "", ""]]))))

    data = tio.export_bytes(layout="carl")
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb["Complete target list"]
    assert "Read me" in wb.sheetnames
    assert "upload" in (ws.cell(1, 1).value or "").lower() or \
           "target board" in (ws.cell(1, 1).value or "").lower()
    assert ws.cell(2, 1).value and "preview" in ws.cell(2, 1).value.lower()
    # Headers still on row 5, data still on row 6 — the tips changed nothing.
    assert ws.cell(5, 8).value == "gene name"
    assert ws.cell(6, 8).value == "SNCA"

    back = tio.parse_workbook(_upload(data, "carl.xlsx"))
    assert back["ok"] and back["header_row"] == 5
    assert [r["gene"] for r in back["rows"]] == ["SNCA"]


def test_headers_carry_hover_notes(clean):
    import openpyxl
    from pipeline.services import target_list_io as tio
    wb = openpyxl.load_workbook(io.BytesIO(tio.export_bytes(layout="carl")))
    ws = wb["Complete target list"]
    gene_note = ws.cell(5, 8).comment
    assert gene_note is not None
    assert "must be filled in" in gene_note.text


def test_overview_drilldown_by_stage_still_narrows_the_list(clean):
    """Overview links to the board as ?status=in_progress. The board doesn't show
    Target.status, but it must still honour it — ignoring it would turn
    "45 in progress" into a link showing everything."""
    from pipeline.models import Target
    from pipeline.services import target_board as board

    Target.objects.using(DB).create(protein_name="A", gene_name="INPROG1",
                                    status=Target.Status.IN_PROGRESS)
    Target.objects.using(DB).create(protein_name="B", gene_name="INPROG2",
                                    status=Target.Status.IN_PROGRESS)
    Target.objects.using(DB).create(protein_name="C", gene_name="ONHOLD1",
                                    status=Target.Status.ON_HOLD)

    assert len(board.board_rows()) == 3
    narrowed = [r["gene"] for r in board.board_rows(status="in_progress")]
    assert sorted(narrowed) == ["INPROG1", "INPROG2"]


def test_overview_drilldown_by_gene_still_works(clean):
    """The other Overview links are ?q=<gene>."""
    from pipeline.models import Target
    from pipeline.services import target_board as board
    Target.objects.using(DB).create(protein_name="A", gene_name="SNCA")
    Target.objects.using(DB).create(protein_name="B", gene_name="MAPT")
    assert [r["gene"] for r in board.board_rows(q="SNCA")] == ["SNCA"]


def test_page_tooltips_come_from_the_same_source_as_the_workbook(clean):
    """The hover text on the board and the note in the downloaded file must not
    drift — they are the same dict."""
    from pipeline.services import target_list_io as tio
    from pipeline.views.target_board import _HEADER_TIP_KEYS
    for column, key in _HEADER_TIP_KEYS.items():
        assert key in tio.COLUMN_TIPS, f"no tip for board column '{column}'"
        assert tio.COLUMN_TIPS[key].strip(), f"empty tip for '{column}'"


def test_import_makes_no_network_calls(clean, monkeypatch):
    """A file of this size once made one UniProt call per new gene, inside the
    request — minutes of latency and a gateway timeout. The file already carries
    the identity fields, so the import must not reach the network at all;
    enrich_targets_from_uniprot fills any gaps afterwards."""
    import requests
    from pipeline.services import target_list_io as tio

    def explode(*a, **k):
        raise AssertionError("import attempted a network call")

    monkeypatch.setattr(requests, "get", explode)
    monkeypatch.setattr(requests, "post", explode)

    data = _carl_workbook([
        ["NIH", "AMP-AD", "Funding available", "2020", "", "McGill",
         "Alpha-synuclein", "SNCA", "", "P37840", 14.46, "NO", "", "", "",
         "", "", "", "", ""]])
    out = tio.apply(tio.parse_workbook(_upload(data)))
    assert out["targets_created"] == ["SNCA"]

    from pipeline.models import Target
    t = Target.objects.using(DB).get(gene_name="SNCA")
    # Identity came from the file, not from a lookup.
    assert t.protein_name == "Alpha-synuclein"
    assert t.uniprot_id == "P37840"
    assert float(t.theoretical_mass_kda) == 14.46


def test_preview_cost_does_not_grow_with_row_count(clean):
    """Preview must be a fixed number of queries, not two per row."""
    from django.db import connections, reset_queries
    from django.test.utils import override_settings
    from pipeline.services import target_list_io as tio

    rows = [["NIH", "P", "Funding available", "2020", "", "McGill", "", f"GENE{i}",
             "", "", "", "", "", "", "", "", "", "", "", ""] for i in range(60)]
    parsed = tio.parse_workbook(_upload(_carl_workbook(rows)))
    with override_settings(DEBUG=True):
        conn = connections[DB]
        reset_queries()
        tio.plan(parsed, default_site="McGill")
        assert len(conn.queries) < 15, f"{len(conn.queries)} queries for 60 rows"


def test_a_bogus_sheet_dimension_does_not_blow_up_the_parse(clean):
    """Excel records a dimension covering every row it has ever formatted. Carl's
    file claims 1,048,539 rows and holds 373. Believing the declaration meant
    materialising fifty million cells, which is what hung the upload."""
    import openpyxl
    from pipeline.services import target_list_io as tio

    data = _carl_workbook([
        ["NIH", "AMP-AD", "Funding available", "2020", "", "McGill", "", "SNCA",
         "", "", "", "", "", "", "", "", "", "", "", ""]])
    # Re-declare the dimension the way Excel does.
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb["Complete target list"]
    ws.print_area = None
    buf = io.BytesIO()
    wb.save(buf)
    raw = buf.getvalue().replace(
        b'<dimension ref="A1:T6"/>', b'<dimension ref="A1:AV1048539"/>')

    parsed = tio.parse_workbook(_upload(raw))
    assert parsed["ok"]
    assert [r["gene"] for r in parsed["rows"]] == ["SNCA"]


def test_parse_stops_after_a_run_of_blank_rows(clean):
    """The real file has 20 blank rows before the padding starts; a short gap
    must not truncate the data."""
    from pipeline.services import target_list_io as tio
    rows = [["", "", "", "", "", "", "", "GENE1", "", "", "", "", "", "", "", "", "", "", "", ""]]
    rows += [[""] * 20 for _ in range(10)]          # a gap
    rows += [["", "", "", "", "", "", "", "GENE2", "", "", "", "", "", "", "", "", "", "", "", ""]]
    parsed = tio.parse_workbook(_upload(_carl_workbook(rows)))
    assert [r["gene"] for r in parsed["rows"]] == ["GENE1", "GENE2"]


def test_offline_backfill_makes_no_network_calls(clean, monkeypatch):
    """--offline must mean offline. Passing terms=None let classify() fetch them
    itself, so the flag looked like it worked while still calling UniProt."""
    import requests
    from django.core.management import call_command
    from pipeline.models import Target, TargetClassification

    monkeypatch.setattr(requests, "get", lambda *a, **k: (_ for _ in ()).throw(
        AssertionError("--offline attempted a network call")))
    Target.objects.using(DB).create(gene_name="RAB11A", protein_name="Rab-11A",
                                    uniprot_id="P62491")
    call_command("backfill_protein_classes", offline=True, sleep=0, verbosity=0)
    labels = list(TargetClassification.objects.using(DB)
                  .values_list("label", flat=True))
    assert labels == ["RAB family"]      # the offline-derivable one


# ── the gene check on the spreadsheet door ───────────────────────────────────
#
# The paste box and the single-gene Add have confirmed a symbol against UniProt
# since the day `bulk_targets` was written; the *upload* went straight to a
# write, and the panel above it said "the same check, the same preview". So a
# typo on row 40 of a 590-row workbook became a permanent target with a
# nomination hanging off it, and nothing on any screen contradicted it. These
# pin the third door, which is the one nobody was watching.

def _uniprot(monkeypatch, table):
    """Stand in for the network. `table` maps GENE → lookup dict."""
    from pipeline.services import bulk_targets

    def fake(gene):
        return table.get(gene.upper(), {"found": False, "error": "no such gene"})
    monkeypatch.setattr(bulk_targets.uniprot, "lookup_gene", fake)


def _found(gene):
    return {"found": True, "gene_name": gene, "protein_name": f"{gene} protein",
            "uniprot_id": f"P{abs(hash(gene)) % 99999:05d}", "mass_kda": 42.0,
            "gene_synonyms": []}


def _sheet(*genes):
    return _carl_workbook([
        ["NIH", "AMP-AD", "Funding available", "2020", "", "McGill", "", g,
         "", "", "", "", "", "", "", "", "", "", "", ""] for g in genes])


def test_preview_refuses_a_gene_uniprot_does_not_know(clean, monkeypatch):
    from pipeline.services import target_list_io as tio
    _uniprot(monkeypatch, {"SNCA": _found("SNCA")})

    out = tio.plan(tio.parse_workbook(_upload(_sheet("SNCA", "ZZZZZZ"))),
                   default_site="McGill")

    assert out["confirmed_genes"] == ["SNCA"]
    assert out["summary"]["new_targets"] == 1, "the bad symbol was previewed as a creation"
    assert out["summary"]["genes_unknown"] == 1
    bad = [i for i in out["items"] if i["gene"] == "ZZZZZZ"][0]
    assert bad["target_action"] == "refused"
    # Nothing on the row is written, not just the target — a nomination and a
    # report both hang off one.
    assert bad["nomination_action"] == "none"
    assert bad["report_action"] == "none"
    assert "check the spelling" in bad["gene_note"].lower()


def test_an_unreachable_uniprot_is_unchecked_not_a_bad_spelling(clean, monkeypatch):
    """`found=False` is two answers, and only one of them is the reader's fault."""
    from pipeline.services import target_list_io as tio
    _uniprot(monkeypatch, {"TRPA1": {"found": False, "unavailable": True,
                                     "error": "proxy refused"}})

    out = tio.plan(tio.parse_workbook(_upload(_sheet("TRPA1"))), default_site="McGill")

    assert out["confirmed_genes"] == []
    assert out["summary"]["genes_unchecked"] == 1
    assert out["summary"]["genes_unknown"] == 0
    note = out["items"][0]["gene_note"].lower()
    assert "could not be reached" in note
    assert "spelling" not in note, "an outage was reported as a typo"


def test_a_gene_already_on_the_list_costs_no_lookup(clean, monkeypatch):
    """The latency answer for a real workbook: it is nearly all updates."""
    from pipeline.models import Target
    from pipeline.services import target_list_io as tio
    Target.objects.using(DB).create(gene_name="SNCA", protein_name="Alpha-synuclein")

    def explode(gene):
        raise AssertionError(f"looked up {gene}, which is already on the list")
    from pipeline.services import bulk_targets
    monkeypatch.setattr(bulk_targets.uniprot, "lookup_gene", explode)

    out = tio.plan(tio.parse_workbook(_upload(_sheet("SNCA"))), default_site="McGill")
    assert out["summary"]["existing_targets"] == 1
    assert out["summary"]["genes_unchecked"] == 0


def test_apply_will_not_create_a_gene_the_preview_did_not_confirm(clean, monkeypatch):
    from pipeline.models import Target
    from pipeline.services import target_list_io as tio
    _uniprot(monkeypatch, {"SNCA": _found("SNCA")})

    parsed = tio.parse_workbook(_upload(_sheet("SNCA", "ZZZZZZ")))
    out = tio.apply(parsed, default_site="McGill", confirmed_genes={"SNCA"})

    assert out["targets_created"] == ["SNCA"]
    assert out["unconfirmed"] == ["ZZZZZZ"]
    assert not Target.objects.using(DB).filter(gene_name="ZZZZZZ").exists()
    # And nothing hung off the refused row either.
    from pipeline.models import TargetNomination
    assert TargetNomination.objects.using(DB).count() == 1


def test_apply_with_no_confirmed_list_is_the_old_ungated_upsert(clean, monkeypatch):
    """`None` means "no check was run" — what a management command asks for.

    The *view* always passes a list, so the door a person presses is gated; this
    keeps the service usable from a shell without a network.
    """
    from pipeline.models import Target
    from pipeline.services import target_list_io as tio

    out = tio.apply(tio.parse_workbook(_upload(_sheet("SNCA"))), default_site="McGill")
    assert out["targets_created"] == ["SNCA"]
    assert out["unconfirmed"] == []
    assert Target.objects.using(DB).filter(gene_name="SNCA").exists()


def test_the_budget_bounds_a_workbook_rather_than_the_row_count(clean, monkeypatch):
    """A file with more new genes than the deadline allows must not hang.

    The rest come back `unchecked`, which creates nothing and asks to be
    previewed again — the cost of a slow UniProt is a second press, never a
    gateway timeout part way through a write.
    """
    import time as _time

    from pipeline.services import bulk_targets
    from pipeline.services import target_list_io as tio

    # A lookup that does **not** return instantly. `as_completed` yields futures
    # that are already finished before it consults its timeout, so an instant
    # mock proves nothing about the deadline — it just races it.
    def slow(gene):
        _time.sleep(0.2)
        return _found(gene)
    monkeypatch.setattr(bulk_targets.uniprot, "lookup_gene", slow)

    parsed = tio.parse_workbook(_upload(_sheet(*[f"GENE{i}" for i in range(40)])))
    started = _time.monotonic()
    out = tio.plan(parsed, default_site="McGill", budget_seconds=0.0)
    # The wall clock, not the recorded results: the failure this guards against
    # is `shutdown(wait=True)` sitting there for the slowest call while the
    # budget decides only what got written down.
    assert _time.monotonic() - started < 0.2, "plan waited for the abandoned lookups"

    s = out["summary"]
    assert out["confirmed_genes"] == []
    assert s["genes_unchecked"] == 40
    assert s["new_targets"] == 0
    assert all(i["target_action"] == "refused" for i in out["items"])
