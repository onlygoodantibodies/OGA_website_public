"""The printed bench sheet keeps its upload, and stops losing things quietly.

Three run-7 findings live here:

* an invented column was swallowed **in silence** — three different values, three
  identical "3 field(s) filled in" lines, and `session_conditions == {}`
  afterwards. The workbook already named such a column and said what it would
  keep; this surface said nothing, and dropped everything rather than all-but-one;
* the concentration column was headed `mg/mL` over µg/mL values — a thousandfold
  mislabel on the one artefact somebody dilutes from at a bench;
* nothing on the sheet said which session it belonged to, so the wrong file in
  the wrong box produced a completely plausible record.
"""
from __future__ import annotations

import io
from datetime import date as _date

import pytest

DB = "pipeline_db"


@pytest.fixture
def bench(_pipeline_db):
    from django.contrib.auth.models import User
    from pipeline.models import (Antibody, CellLine, Company, ExperimentSession,
                                 IfResult, Member, Sample, Site, Target, WbResult)
    for M in (WbResult, IfResult, ExperimentSession, Sample, Antibody, CellLine,
              Company, Target):
        M.objects.using(DB).all().delete()

    site, _ = Site.objects.using(DB).get_or_create(
        short_code="LEI", defaults={"name": "Leicester"})
    user, _ = User.objects.using(DB).get_or_create(username="benchuser")
    member, _ = Member.objects.using(DB).get_or_create(
        user_id=user.pk, defaults={"site": site, "role": "experimenter",
                                   "is_active": True, "display_name": "Bench User"})
    target = Target.objects.using(DB).create(gene_name="STMN2", protein_name="Stathmin-2")
    company = Company.objects.using(DB).create(name="abcam")
    ab = Antibody.objects.using(DB).create(
        target=target, company=company, catalogue_number="A-STMN2-1",
        lot_number="RUN7-A1", concentration=1000)
    wt = CellLine.objects.using(DB).create(name="SH-SY5Y", genotype="WT", site=site)
    ko = CellLine.objects.using(DB).create(
        name="SH-SY5Y STMN2 KO", genotype="KO", target=target, site=site, parent_line=wt)
    session = ExperimentSession.objects.using(DB).create(
        procedure_type="WB", target=target, experimenter=member, date=_date(2026, 8, 1),
        site=site, cell_line_wt=wt, cell_line_ko=ko,
        status=ExperimentSession.SessionStatus.PLANNED)
    other = ExperimentSession.objects.using(DB).create(
        procedure_type="WB", target=target, experimenter=member, date=_date(2026, 8, 2),
        site=site, status=ExperimentSession.SessionStatus.PLANNED)
    return {"site": site, "member": member, "target": target, "ab": ab,
            "session": session, "other": other}


def _sheet(session):
    from pipeline.services.planning import generate_bench_sheet
    buf = io.BytesIO()
    generate_bench_sheet(session).save(buf)
    buf.seek(0)
    import openpyxl
    return openpyxl.load_workbook(buf)


def _upload(wb, name="filled.xlsx"):
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    buf.name = name
    return buf


def _header_row(ws):
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if any(str(c or "").strip().lower() == "ab#" for c in row):
            return i, [("" if c is None else str(c)) for c in row]
    raise AssertionError("no header row")


def test_the_concentration_column_names_the_unit_it_holds(bench):
    """`Antibody.concentration` is µg/mL and the column said mg/mL over 1000."""
    ws = _sheet(bench["session"])["WB"]
    _i, header = _header_row(ws)
    assert "Conc. (µg/mL)" in header
    assert "Conc. (mg/mL)" not in header


@pytest.mark.parametrize("proc", ["WB", "IP", "IF", "FC"])
def test_every_bench_sheet_says_which_session_it_is_for(bench, proc):
    """All four, not just the one the guard was written against.

    The IF plate map did not: its row 2 was the KO vial's C-number alone, so it
    was the one bench sheet with no session number in it — and a guard reading a
    stamp that is not there is a guard that always passes. Session 500's plate
    map was accepted by session 499 without a word, under a panel promising
    "both files carry this session's number".
    """
    session = bench["session"]
    session.procedure_type = proc
    session.save(using=DB)
    wb = _sheet(session)
    text = " ".join(str(c or "")
                    for ws in wb.worksheets
                    for r in ws.iter_rows(values_only=True) for c in r)
    assert f"Session #{session.pk}" in text


def test_the_if_plate_map_keeps_the_ko_vial_c_number_beside_the_stamp(bench):
    """Row 2 gained the stamp; it must not have lost what it already said."""
    from pipeline.models import CellLineVial
    session = bench["session"]
    CellLineVial.objects.using(DB).create(cell_line=session.cell_line_ko, c_number=1202)
    session.procedure_type = "IF"
    session.save(using=DB)

    row2 = " ".join(str(c or "") for c in next(
        _sheet(session).worksheets[0].iter_rows(min_row=2, max_row=2, values_only=True)))
    assert f"Session #{session.pk}" in row2
    assert "C-1202" in row2


def test_an_if_plate_map_fed_to_a_western_blot_session_is_refused(bench):
    """The owner's run-11 finding, end to end.

    Session 500's IF plate map, uploaded into session 499's western blot: it
    matched all three antibodies and offered *Record these*. `Used dilution` is
    the one column label WB and IF share, so the plate's `1in500_T` was read as
    a WB dilution and every row looked filled in.
    """
    from pipeline.models import ExperimentSession
    from pipeline.services import bench_results

    wb_session = bench["session"]
    if_session = ExperimentSession.objects.using(DB).create(
        procedure_type="IF", target=bench["target"], experimenter=bench["member"],
        date=_date(2026, 8, 4), site=bench["site"],
        cell_line_wt=wb_session.cell_line_wt, cell_line_ko=wb_session.cell_line_ko,
        status=ExperimentSession.SessionStatus.PLANNED)

    plate = _sheet(if_session)
    ws = plate.worksheets[0]
    i, header = _header_row(ws)
    ws.cell(row=i + 1, column=header.index("Specific signal") + 1, value="Yes")

    # Read the way the view reads it: with the *target* session's procedure.
    parsed = bench_results.parse(_upload(plate), wb_session.procedure_type)
    assert parsed["sheet_session_id"] == if_session.pk
    assert parsed["sheet_procedure"] == "IF"

    out = bench_results.plan(wb_session, parsed)
    assert out["ok"] is False
    assert f"#{if_session.pk}" in out["error"]
    assert bench_results.apply(wb_session, parsed)["ok"] is False


def test_a_sheet_of_another_procedure_is_refused_even_with_no_stamp(bench):
    """The stamp is the precise guard; the columns are the one that still works.

    Plate maps downloaded before the stamp existed carry no session number, and
    a sheet somebody typed never will — so the guard cannot rest on the stamp
    alone.
    """
    from pipeline.models import ExperimentSession
    from pipeline.services import bench_results

    if_session = ExperimentSession.objects.using(DB).create(
        procedure_type="IF", target=bench["target"], experimenter=bench["member"],
        date=_date(2026, 8, 4), site=bench["site"],
        status=ExperimentSession.SessionStatus.PLANNED)
    plate = _sheet(if_session)
    ws = plate.worksheets[0]
    i, header = _header_row(ws)
    ws.cell(row=i + 1, column=header.index("Specific signal") + 1, value="Yes")
    ws.cell(row=2, column=1, value="C-1202")  # a pre-stamp plate map

    parsed = bench_results.parse(_upload(plate), "WB")
    assert parsed["sheet_session_id"] is None
    out = bench_results.plan(bench["session"], parsed)
    assert out["ok"] is False
    assert "Immunofluorescence" in out["error"] or "IF" in out["error"]


def test_a_sheet_that_names_no_procedure_is_not_refused(bench):
    """Only files that are unmistakably something else. A hand-made sheet of
    `Ab#` and `Used dilution` belongs to no procedure in particular, and the
    parser has always accepted the variants on Riham's own sheets."""
    from pipeline.services import bench_results
    assert bench_results.sheet_procedure(["Ab#", "CatNumber", "Used dilution"]) is None
    assert bench_results.sheet_procedure(["Ab#", "Signal", "Rating"]) == "WB"
    assert bench_results.sheet_procedure(["Ab#", "Specific signal"]) == "IF"
    # Two procedures' worth is somebody's own combined sheet, not one of ours.
    assert bench_results.sheet_procedure(["Ab#", "Rating", "Specific signal"]) is None


def test_a_sheet_with_no_stamp_says_it_could_not_be_checked(bench):
    """Not a refusal — but the panel promises both files carry the session's
    number, and a promise that cannot be checked must not read as one that
    passed."""
    from pipeline.services import bench_results
    wb = _sheet(bench["session"])
    ws = wb["WB"]
    i, header = _header_row(ws)
    ws.cell(row=i + 1, column=header.index("Signal") + 1, value="a band")
    ws.cell(row=2, column=1, value="")  # an older download, before the stamp

    parsed = bench_results.parse(_upload(wb), "WB")
    assert parsed["sheet_session_id"] is None
    out = bench_results.plan(bench["session"], parsed)
    assert out["ok"] is True and out["summary"]["matched"] == 1
    assert f"#{bench['session'].pk}" in out["unstamped"]

    # …and a sheet that *is* stamped says nothing, or the note means nothing.
    stamped = bench_results.parse(_upload(_sheet(bench["session"])), "WB")
    assert bench_results.plan(bench["session"], stamped)["unstamped"] == ""


def test_a_sheet_uploaded_into_the_wrong_session_is_refused(bench):
    """Two sheets for one gene differ only in the readings written on them."""
    from pipeline.services import bench_results
    wb = _sheet(bench["session"])
    ws = wb["WB"]
    i, header = _header_row(ws)
    ws.cell(row=i + 1, column=header.index("Signal") + 1, value="single band")

    parsed = bench_results.parse(_upload(wb), "WB")
    assert parsed["sheet_session_id"] == bench["session"].pk

    out = bench_results.plan(bench["other"], parsed)
    assert out["ok"] is False
    assert f"#{bench['session'].pk}" in out["error"]
    assert f"#{bench['other'].pk}" in out["error"]

    # …and the commit re-checks it: a preview is not a permission slip.
    written = bench_results.apply(bench["other"], parsed)
    assert written["ok"] is False


def test_an_invented_column_is_named_counted_and_kept(bench):
    """The run-7 finding: three values, no message, nothing stored."""
    from pipeline.models import ExperimentSession
    from pipeline.services import bench_results

    wb = _sheet(bench["session"])
    ws = wb["WB"]
    i, header = _header_row(ws)
    extra = len(header) + 1
    ws.cell(row=i, column=extra, value="Owner notes")
    ws.cell(row=i + 1, column=header.index("Signal") + 1, value="single band")
    ws.cell(row=i + 1, column=extra, value="HV bench book p.21")

    parsed = bench_results.parse(_upload(wb), "WB")
    unknown = parsed["unknown"]
    assert len(unknown) == 1
    assert unknown[0]["header"] == "owner notes"
    # snake_cased, so it lands where every other condition lives and can be
    # round-tripped by key — the same rule the workbook importer holds.
    assert unknown[0]["key"] == "owner_notes"
    assert unknown[0]["value"] == "HV bench book p.21"

    out = bench_results.plan(bench["session"], parsed)
    assert out["ok"] and out["summary"]["unknown_columns"] == 1

    res = bench_results.apply(bench["session"], parsed)
    assert res["ok"] and res["conditions_stored"] == ["owner_notes"]
    session = ExperimentSession.objects.using(DB).get(pk=bench["session"].pk)
    assert session.session_conditions["owner_notes"] == "HV bench book p.21"


def test_rows_that_disagree_all_survive_into_the_one_condition(bench):
    """A condition belongs to the session, so several rows share one — but
    sharing is not the same as keeping the first and discarding the rest.

    The owner's run-11 finding, on a nine-well plate map: the first row is not a
    summary of the plate, it is the first well. Both importers used to keep it
    and count the other eight as dropped, which named the loss without stopping
    it. Nothing anybody typed goes missing now.
    """
    from pipeline.models import Antibody, ExperimentSession
    from pipeline.services import bench_results

    Antibody.objects.using(DB).create(
        target=bench["target"], company=bench["ab"].company,
        catalogue_number="A-STMN2-2", concentration=500)

    wb = _sheet(bench["session"])
    ws = wb["WB"]
    i, header = _header_row(ws)
    extra = len(header) + 1
    ws.cell(row=i, column=extra, value="Owner notes")
    for n, note in enumerate(("p.21", "p.22")):
        ws.cell(row=i + 1 + n, column=header.index("Signal") + 1, value="a band")
        ws.cell(row=i + 1 + n, column=extra, value=note)

    parsed = bench_results.parse(_upload(wb), "WB")
    assert parsed["unknown"][0]["value"] == "p.21 | p.22"
    assert parsed["unknown"][0]["values"] == 2
    assert parsed["unknown"][0]["rows"] == 2

    # …and the write stores what the preview said it would. A preview is only
    # true if it is compared against the write, so both go through one fold.
    assert bench_results.apply(bench["session"], parsed)["ok"]
    stored = ExperimentSession.objects.using(DB).get(
        pk=bench["session"].pk).session_conditions
    assert stored["owner_notes"] == "p.21 | p.22"


def test_a_column_repeating_one_value_still_reads_as_one_value(bench):
    """Joining must not turn an ordinary column into a list of duplicates."""
    from pipeline.models import Antibody
    from pipeline.services import bench_results

    Antibody.objects.using(DB).create(
        target=bench["target"], company=bench["ab"].company,
        catalogue_number="A-STMN2-2", concentration=500)

    wb = _sheet(bench["session"])
    ws = wb["WB"]
    i, header = _header_row(ws)
    extra = len(header) + 1
    ws.cell(row=i, column=extra, value="Owner notes")
    for n in range(2):
        ws.cell(row=i + 1 + n, column=header.index("Signal") + 1, value="a band")
        ws.cell(row=i + 1 + n, column=extra, value="p.21")

    parsed = bench_results.parse(_upload(wb), "WB")
    assert parsed["unknown"][0]["value"] == "p.21"
    assert parsed["unknown"][0]["values"] == 1


def test_a_condition_recorded_elsewhere_is_not_wiped(bench):
    """Touch only the keys the sheet carries."""
    from pipeline.models import ExperimentSession
    from pipeline.services import bench_results

    session = bench["session"]
    session.session_conditions = {"lysis_buffer": "RIPA"}
    session.save(using=DB)

    wb = _sheet(session)
    ws = wb["WB"]
    i, header = _header_row(ws)
    extra = len(header) + 1
    ws.cell(row=i, column=extra, value="Owner notes")
    ws.cell(row=i + 1, column=header.index("Signal") + 1, value="a band")
    ws.cell(row=i + 1, column=extra, value="p.21")

    bench_results.apply(session, bench_results.parse(_upload(wb), "WB"))
    stored = ExperimentSession.objects.using(DB).get(pk=session.pk).session_conditions
    assert stored["lysis_buffer"] == "RIPA"
    assert stored["owner_notes"] == "p.21"


def test_the_sheets_own_identity_columns_are_not_reported_as_invented(bench):
    """Gene, Clonality, Clone, Host and the rest ship with the sheet."""
    from pipeline.services import bench_results
    wb = _sheet(bench["session"])
    ws = wb["WB"]
    i, header = _header_row(ws)
    ws.cell(row=i + 1, column=header.index("Signal") + 1, value="a band")
    parsed = bench_results.parse(_upload(wb), "WB")
    assert parsed["unknown"] == []


def test_results_still_record_from_a_filled_sheet(bench):
    from pipeline.models import WbResult
    from pipeline.services import bench_results
    wb = _sheet(bench["session"])
    ws = wb["WB"]
    i, header = _header_row(ws)
    ws.cell(row=i + 1, column=header.index("Signal") + 1, value="single band at 20 kDa")
    ws.cell(row=i + 1, column=header.index("Used dilution") + 1, value="1:1000")

    parsed = bench_results.parse(_upload(wb), "WB")
    res = bench_results.apply(bench["session"], parsed)
    assert res["ok"] and len(res["created"]) == 1
    row = WbResult.objects.using(DB).get(session=bench["session"])
    assert row.signal == "single band at 20 kDa" and row.dilution == "1:1000"


def test_a_bench_sheet_still_parses_after_the_workbook_sniff_read_it(bench):
    """One file input, two parsers — and the sniff consumes the stream.

    `views/session_bulk._looks_like_workbook` reads the upload to decide which
    parser to use. `parse_template` rewinds itself; `bench_results.parse` does
    not, so an exhausted upload would read as an empty sheet and be reported as
    *"the sheet is empty"* about a sheet with rows on it.
    """
    from pipeline.services import bench_results
    from pipeline.views.session_bulk import _looks_like_workbook

    wb = _sheet(bench["session"])
    ws = wb["WB"]
    i, header = _header_row(ws)
    ws.cell(row=i + 1, column=header.index("Signal") + 1, value="single band")
    upload = _upload(wb)

    assert _looks_like_workbook(upload) is False
    # …and the sheet is still readable afterwards, which is the whole point.
    upload.seek(0)
    parsed = bench_results.parse(upload, "WB")
    assert len(parsed["rows"]) == 1


def test_a_session_workbook_is_recognised_as_one(bench):
    from pipeline.services import session_template
    from pipeline.views.session_bulk import _looks_like_workbook, _workbook_is_this_session
    from pipeline.services import session_import
    import io as _io

    data = session_template.build_session_template(bench["session"])
    buf = _io.BytesIO(data)
    buf.name = "wb.xlsx"
    assert _looks_like_workbook(buf) is True

    buf.seek(0)
    parsed = session_import.parse_template(buf)
    assert _workbook_is_this_session(parsed, bench["session"]) == ""
    # …and the same file is refused against a different session.
    refusal = _workbook_is_this_session(parsed, bench["other"])
    assert f"#{bench['session'].pk}" in refusal
