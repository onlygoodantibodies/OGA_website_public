"""Sessions round-trip: download what the board shows, edit it, upload it back.

Rebuilt on the board's read layer rather than kept as a parallel implementation.
The old ``session_roundtrip`` carried its own copy of "the result fields" and its
own filters; this derives both from ``services/session_board.py``, so the
spreadsheet and the screen cannot show different columns for the same session.

The contract, deliberately narrow:

  * **Edit only.** Every row is matched by ``session_id`` + ``result_id``. Rows
    with no id, or an id that does not exist, are reported and skipped. Nothing
    is created and nothing is deleted — creating sessions stays a separate,
    explicit act.
  * **Fill-only-blank by default.** A value that would overwrite something
    already populated is listed as a conflict and left alone unless the user
    ticks "update existing values". A blank cell never clears a field.
  * **Preview then commit**, on the same ``parse → plan → apply`` shape as every
    other bulk path in the app.
"""
from __future__ import annotations

import io
from datetime import date, datetime

from django.db import transaction

from pipeline.models import ExperimentSession, Member, Site
from pipeline.services import session_board as board

DB = "pipeline_db"

# The two key columns. Locked because the upload matches on them — edit anything
# else and the row still lands on the right record.
KEY_COLS = ["session_id", "result_id"]

# Read-only context, so a filled sheet is readable without cross-referencing.
REF_COLS = ["gene", "procedure", "antibody", "catalogue", "company"]

# Session header fields a sheet may change, and how to set them.
SESSION_COLS = ["date", "experimenter", "site", "status", "comments",
                "protein_loading_ug", "cell_line_wt", "cell_line_ko"]

# Session columns are prefixed in the sheet because a session field and a result
# field can share a name — every result model has its own ``comments``, distinct
# from the session's. Two columns called "comments" is not a naming nicety: the
# parser keeps the last one, so edits to the session comment vanished silently.
SESSION_PREFIX = "session_"


def session_header(field: str) -> str:
    return f"{SESSION_PREFIX}{field}"


COLUMN_TIPS = {
    "session_id": "Which session this row belongs to. Do not edit — the upload matches on it.",
    "result_id": "Which result row this is. Do not edit — the upload matches on it. "
                 "Blank means the session has no result row for this antibody yet, so "
                 "readings typed on this line cannot be saved: add the antibody to the "
                 "session on the board first, then download this sheet again.",
    "gene": "The session's target. Read-only here.",
    "procedure": "WB, IP, IF or FC. Read-only — the result columns depend on it.",
    "antibody": "The antibody this result is for. Read-only here.",
    "session_date": "The date the experiment was run (YYYY-MM-DD).",
    "session_experimenter": "Who ran it. Must match a member's name exactly.",
    "session_site": "Which YCharOS site. Must match a site name exactly.",
    "session_status": "planned, in_progress, complete, failed, repeat_needed or cancelled.",
    "session_comments": "Free text about the session — not the per-antibody comment.",
    "session_protein_loading_ug": "Micrograms of protein per lane, for the figure legend.",
    "session_cell_line_wt": "Wild-type line used. Must match a cell line name exactly.",
    "session_cell_line_ko": "Knockout line used. Must match a cell line name exactly.",
    "comments": "Comment on this antibody's result, not on the session.",
    # Board-only columns. Never appear in the sheet, so build_export never looks
    # them up — they live here so the board and the sheet keep one tip registry.
    "cell_lines": "The wild-type and knockout lines this session compared. "
                  "Type a cell line name exactly as it is recorded.",
    # Two numbers, because they answer two questions and the column used to
    # give only the first while being headed with the second's word. Planning a
    # session writes a blank row per antibody, so "22 results" appeared over a
    # session nobody had run. A row with nothing written on it is not evidence
    # that the work was not done — gaps in a results section are normal — so
    # this counts what is recorded and draws no conclusion from the rest.
    "results": "How many antibodies this session covers, and how many of those "
               "rows have something written on them. Click to open them and "
               "edit them.",
}


def _text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, (datetime, date)):
        return v.date().isoformat() if isinstance(v, datetime) else v.isoformat()
    return str(v).strip()


def columns_for(procedure) -> list[str]:
    return (KEY_COLS + REF_COLS
            + [session_header(f) for f in SESSION_COLS]
            + board.result_field_names(procedure))


def _session_values(row) -> dict:
    """Session header values under their prefixed sheet names."""
    raw = {
        "date": row["date"], "experimenter": row["experimenter"],
        "site": row["site"], "status": row["status"],
        "comments": row["comments"],
        "protein_loading_ug": ("" if row["protein_loading_ug"] is None
                               else row["protein_loading_ug"]),
        "cell_line_wt": row["cell_line_wt"], "cell_line_ko": row["cell_line_ko"],
    }
    return {session_header(k): v for k, v in raw.items()}


def build_export(**filters) -> bytes:
    """One sheet per procedure present in the filtered set.

    The four procedures have different result columns, so they cannot share a
    sheet — the same reason the board opens results per row.
    """
    import openpyxl
    from openpyxl.comments import Comment

    rows = board.board_rows(**filters)
    by_proc: dict[str, list] = {}
    for r in rows:
        by_proc.setdefault(r["procedure"], []).append(r)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    for proc in board.PROCEDURES:
        session_rows = by_proc.get(proc)
        if not session_rows:
            continue
        ws = wb.create_sheet(proc)
        cols = columns_for(proc)
        for i, name in enumerate(cols, start=1):
            c = ws.cell(1, i, name)
            tip = COLUMN_TIPS.get(name)
            if tip:
                c.comment = Comment(tip, "OGA")
            if name in KEY_COLS:
                c.value = f"{name} (do not edit)"

        line = 2
        for srow in session_rows:
            session = ExperimentSession.objects.using(DB).filter(pk=srow["id"]).first()
            if session is None:
                continue
            detail = board.results_for(session)
            svals = _session_values(srow)
            # A session with no results still gets a row, so its header fields
            # can be edited from the sheet.
            entries = detail["rows"] or [None]
            for entry in entries:
                values = {
                    "session_id": srow["id"],
                    "result_id": entry["id"] if entry else "",
                    "gene": srow["gene"], "procedure": proc,
                    "antibody": entry["antibody"] if entry else "",
                    "catalogue": entry["catalogue"] if entry else "",
                    "company": entry["company"] if entry else "",
                    **svals,
                }
                if entry:
                    values.update(entry["values"])
                for i, name in enumerate(cols, start=1):
                    ws.cell(line, i, values.get(name, ""))
                line += 1
        ws.freeze_panes = "A2"

    if not wb.sheetnames:
        wb.create_sheet("Sessions").cell(1, 1, "No sessions matched these filters.")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _norm(h) -> str:
    return str(h or "").strip().lower().replace(" (do not edit)", "")


def parse_upload(f) -> dict:
    """Read every sheet whose name is a procedure. Returns {ok, rows, errors}."""
    import openpyxl
    try:
        wb = openpyxl.load_workbook(f, data_only=True)
    except Exception as e:
        return {"ok": False, "error": f"could not read that file ({e})", "rows": []}

    rows, errors = [], []
    for name in wb.sheetnames:
        proc = name.strip().upper()
        if proc not in board.RESULT_MODELS:
            continue
        ws = wb[name]
        grid = list(ws.iter_rows(values_only=True))
        if not grid:
            continue
        header = [_norm(h) for h in grid[0]]
        known = set(columns_for(proc))
        for n, raw in enumerate(grid[1:], start=2):
            record = {header[i]: raw[i] for i in range(min(len(header), len(raw)))
                      if header[i]}
            if not any(_text(v) for v in record.values()):
                continue
            sid = _text(record.get("session_id"))
            if not sid.isdigit():
                errors.append(f"{name} row {n}: no session_id — skipped")
                continue
            record = {k: v for k, v in record.items() if k in known}
            rows.append({"sheet": name, "line": n, "procedure": proc,
                         "session_id": int(sid),
                         "result_id": (int(_text(record.get("result_id")))
                                       if _text(record.get("result_id")).isdigit()
                                       else None),
                         "values": record})
    return {"ok": True, "rows": rows, "errors": errors}


def _resolve_named(model, value, field="name"):
    v = _text(value)
    if not v:
        return None
    return model.objects.using(DB).filter(**{f"{field}__iexact": v}).first()


def _incoming(record) -> dict:
    """Only cells the sheet actually filled. A blank never clears a field."""
    return {k: _text(v) for k, v in record["values"].items()
            if k not in KEY_COLS and k not in REF_COLS and _text(v) != ""}


def _current_session_value(session, field) -> str:
    if field == "experimenter":
        return str(session.experimenter) if session.experimenter_id else ""
    if field == "site":
        return session.site.name if session.site_id else ""
    if field in ("cell_line_wt", "cell_line_ko"):
        line = getattr(session, field)
        return line.name if line else ""
    if field == "date":
        return session.date.isoformat() if session.date else ""
    return _text(getattr(session, field, ""))


def plan(parsed: dict) -> dict:
    """Read-only. Says exactly what would change, and what disagrees."""
    if not parsed.get("ok"):
        return parsed

    items, fills, conflicts, missing, dropped = [], 0, 0, 0, 0
    for record in parsed["rows"]:
        session = (ExperimentSession.objects.using(DB)
                   .select_related("experimenter", "site", "cell_line_wt", "cell_line_ko")
                   .filter(pk=record["session_id"]).first())
        if session is None:
            missing += 1
            items.append({**_label(record), "error": "no session with that id"})
            continue

        result = None
        if record["result_id"]:
            model = board.RESULT_MODELS[record["procedure"]]
            result = model.objects.using(DB).filter(
                pk=record["result_id"], session_id=session.pk).first()
            if result is None:
                missing += 1
                items.append({**_label(record),
                              "error": "no result row with that id on that session"})
                continue

        result_fields = set(board.result_field_names(record["procedure"]))
        changes = []
        # Readings typed onto a row that has no `result_id` — see `no_result_row`
        # below. Collected rather than skipped, because dropping them silently is
        # what made a whole row of work vanish behind "0 changes".
        homeless = []
        for header, incoming in _incoming(record).items():
            if header.startswith(SESSION_PREFIX):
                field = header[len(SESSION_PREFIX):]
                if field not in SESSION_COLS:
                    continue
                current = _current_session_value(session, field)
                scope = "session"
            elif header in result_fields:
                if result is None:
                    homeless.append(header)
                    continue
                field = header
                current = _text(getattr(result, field, ""))
                scope = "result"
            else:
                continue
            if current == incoming:
                continue
            kind = "fill" if not current else "conflict"
            if kind == "fill":
                fills += 1
            else:
                conflicts += 1
            changes.append({"scope": scope, "field": field, "from": current,
                            "to": incoming, "kind": kind})
        if homeless:
            # **A reading with nowhere to go is a refusal, not a no-op.** This
            # sheet only ever edits: a row is matched on `session_id` +
            # `result_id`, and a row typed in by hand — or exported from a
            # session that has no results yet, where `result_id` is blank — has
            # no result row to write to. Every one of those values used to be
            # dropped without a word, and because a row with nothing else on it
            # produced no item at all, the preview said `0 changes` about a line
            # holding a dilution, a signal and a rating.
            #
            # Named on the row, counted in the summary, and pointed at the one
            # surface that can create the row — the same shape as the
            # concentration and C-number refusals, which also keep what they can
            # and say what they cannot take.
            # A `warning`, not an `error`: `apply` skips a row whose error is set,
            # and the session-level cells on this row are perfectly writable. The
            # row goes in without the readings, deliberately — the same bargain
            # `bulk_antibodies` strikes with a concentration it cannot convert.
            dropped += len(homeless)
            items.append({**_label(record), "changes": changes,
                          "no_result_row": homeless,
                          "warning": (
                              f"{len(homeless)} reading(s) on this row cannot be "
                              f"recorded — the result_id cell is empty, and this "
                              f"sheet only edits result rows that already exist. "
                              f"Add the antibody to session #{record['session_id']} "
                              f"on the sessions board, then download the sheet "
                              f"again and it will carry a result_id.")})
        elif changes:
            items.append({**_label(record), "changes": changes})

    return {"ok": True, "items": items, "errors": parsed.get("errors", []),
            "summary": {"rows": len(parsed["rows"]), "changed_rows": len(items),
                        "fills": fills, "conflicts": conflicts, "unmatched": missing,
                        # Readings the sheet carried and this path cannot store.
                        "dropped_readings": dropped}}


def _label(record) -> dict:
    return {"sheet": record["sheet"], "line": record["line"],
            "session_id": record["session_id"], "result_id": record["result_id"]}


def apply(parsed: dict, *, apply_overwrites: bool = False) -> dict:
    """Write the plan. One transaction: a failure part way leaves nothing."""
    planned = plan(parsed)
    if not planned.get("ok"):
        return planned

    sessions_updated = results_updated = skipped = 0
    with transaction.atomic(using=DB):
        for item in planned["items"]:
            if item.get("error"):
                skipped += 1
                continue
            session = ExperimentSession.objects.using(DB).filter(
                pk=item["session_id"]).first()
            if session is None:
                skipped += 1
                continue
            model = None
            result = None
            if item["result_id"]:
                proc = session.procedure_type
                model = board.RESULT_MODELS.get(proc)
                result = model.objects.using(DB).filter(
                    pk=item["result_id"], session_id=session.pk).first() if model else None

            touched_session = touched_result = False
            for change in item["changes"]:
                if change["kind"] == "conflict" and not apply_overwrites:
                    continue
                field, value = change["field"], change["to"]
                if change["scope"] == "result":
                    if result is None:
                        continue
                    setattr(result, field, value)
                    touched_result = True
                else:
                    _set_session_field(session, field, value)
                    touched_session = True
            if touched_session:
                session.save(using=DB)
                sessions_updated += 1
            if touched_result:
                result.save(using=DB)
                results_updated += 1

    return {"ok": True, "sessions_updated": sessions_updated,
            "results_updated": results_updated, "skipped": skipped,
            # **At the save, not only at the check.** A reader who pressed
            # through the preview still has to be told what did not land, or the
            # success line is the last word on a row that lost its readings.
            "dropped_readings": planned["summary"]["dropped_readings"],
            "plan": planned}


def _set_session_field(session, field, value):
    """Same rules as the board's inline edit, so a sheet cannot do what the page
    refuses. Unresolvable names are left alone rather than blanking the field."""
    if field == "date":
        session.date = value
    elif field == "status":
        if value in {v for v, _ in ExperimentSession.SessionStatus.choices}:
            session.status = value
    elif field == "comments":
        session.comments = value
    elif field == "protein_loading_ug":
        try:
            session.protein_loading_ug = float(value)
        except (TypeError, ValueError):
            pass
    elif field == "site":
        site = _resolve_named(Site, value)
        if site:
            session.site_id = site.pk
    elif field == "experimenter":
        member = _resolve_named(Member, value, field="display_name")
        if member:
            session.experimenter_id = member.pk
    elif field in ("cell_line_wt", "cell_line_ko"):
        from pipeline.models import CellLine
        line = _resolve_named(CellLine, value)
        if line:
            setattr(session, f"{field}_id", line.pk)
