"""Sessions round-trip: download the board, edit the sheet, upload it back.

The contract is deliberately narrow and these pin every part of it: edit only
(never create, never delete), matched on the two id columns, fill-only-blank
unless overwrites are asked for, a blank never clears, and a re-upload after a
timeout changes nothing.
"""
from __future__ import annotations

import io
from datetime import date

import openpyxl
from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client, TestCase

from pipeline.models import (Antibody, Company, ExperimentSession, Member, Site,
                             Target, WbResult)
from pipeline.services import session_io as sio

DB = "pipeline_db"


class SessionRoundTripTests(TestCase):
    databases = {"default", "pipeline_db", "academy_db"}

    def setUp(self):
        self.site = Site.objects.create(name="Leicester", short_code="LEI")
        self.site2 = Site.objects.create(name="McGill", short_code="MCG")
        for alias in ("academy_db", DB):
            u = User(username="vera")
            u.set_password("pw")
            u.save(using=alias)
        pu = User.objects.using(DB).get(username="vera")
        self.member = Member.objects.create(
            user_id=pu.pk, site_id=self.site.pk, role="experimenter",
            is_active=True, display_name="Vera Ruiz Moleon")
        target = Target.objects.create(protein_name="SOD", gene_name="SOD1")
        company = Company.objects.create(name="Proteintech")
        self.antibody = Antibody.objects.create(
            target=target, company=company, catalogue_number="12A8")
        self.session = ExperimentSession.objects.create(
            procedure_type="WB", target=target, experimenter=self.member,
            site=self.site, date=date(2026, 2, 18), status="complete",
            comments="")
        self.result = WbResult.objects.create(
            session=self.session, antibody=self.antibody,
            signal="clean band", rating="")
        self.client = Client()
        self.assertTrue(self.client.login(username="vera", password="pw"))

    # ── download ───────────────────────────────────────────────────────

    def _download(self, query=""):
        response = self.client.get(f"/pipeline/sessions/board/export/?{query}")
        self.assertEqual(response.status_code, 200)
        return openpyxl.load_workbook(io.BytesIO(response.content))

    def test_download_has_one_sheet_per_procedure_present(self):
        wb = self._download()
        self.assertEqual(wb.sheetnames, ["WB"])
        header = [str(c.value) for c in wb["WB"][1]]
        self.assertTrue(header[0].startswith("session_id"))
        self.assertTrue(header[1].startswith("result_id"))
        self.assertIn("signal", header)
        self.assertIn("session_comments", header)

    def test_download_honours_the_boards_filters(self):
        ExperimentSession.objects.create(
            procedure_type="IF", target=self.session.target,
            experimenter=self.member, site=self.site, date=date(2026, 3, 1),
            status="planned")
        self.assertEqual(self._download().sheetnames, ["WB", "IF"])
        self.assertEqual(self._download("procedure=WB").sheetnames, ["WB"])

    # ── upload ─────────────────────────────────────────────────────────

    def _sheet_with(self, **overrides):
        """The downloaded sheet, with some cells changed — a real round-trip.

        Edits the row for ``self.result``, wherever the board's ordering put it,
        rather than assuming it is the first.
        """
        wb = self._download()
        ws = wb["WB"]
        header = [sio._norm(c.value) for c in ws[1]]
        rid = header.index("result_id") + 1
        line = next(r for r in range(2, ws.max_row + 1)
                    if str(ws.cell(r, rid).value or "") == str(self.result.pk))
        for field, value in overrides.items():
            ws.cell(line, header.index(field) + 1, value)
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def _preview(self, data):
        r = self.client.post("/pipeline/sessions/board/upload/preview/",
                             {"file": SimpleUploadedFile("s.xlsx", data)})
        self.assertEqual(r.status_code, 200, r.content[:300])
        return r.json()

    def _commit(self, data, overwrites=False):
        r = self.client.post("/pipeline/sessions/board/upload/commit/",
                             {"file": SimpleUploadedFile("s.xlsx", data),
                              "apply_overwrites": "true" if overwrites else "false"})
        self.assertEqual(r.status_code, 200, r.content[:300])
        return r.json()

    def test_filling_a_blank_field_applies_without_confirmation(self):
        data = self._sheet_with(rating="4", session_comments="repeat of the 12 Feb run")
        plan = self._preview(data)
        self.assertEqual(plan["summary"]["fills"], 2)
        self.assertEqual(plan["summary"]["conflicts"], 0)

        self._commit(data)
        self.result.refresh_from_db()
        self.session.refresh_from_db()
        self.assertEqual(self.result.rating, "4")
        self.assertEqual(self.session.comments, "repeat of the 12 Feb run")

    def test_overwriting_a_populated_field_needs_the_toggle(self):
        data = self._sheet_with(signal="smeary")
        plan = self._preview(data)
        self.assertEqual(plan["summary"]["conflicts"], 1)

        self._commit(data)  # without the toggle
        self.result.refresh_from_db()
        self.assertEqual(self.result.signal, "clean band")

        self._commit(data, overwrites=True)
        self.result.refresh_from_db()
        self.assertEqual(self.result.signal, "smeary")

    def test_a_blank_cell_never_clears_a_field(self):
        data = self._sheet_with(signal="")
        self._commit(data, overwrites=True)
        self.result.refresh_from_db()
        self.assertEqual(self.result.signal, "clean band")

    def test_names_resolve_on_upload(self):
        self._commit(self._sheet_with(session_site="McGill"), overwrites=True)
        self.session.refresh_from_db()
        self.assertEqual(self.session.site_id, self.site2.pk)

    def test_an_unresolvable_name_leaves_the_field_alone(self):
        """Better a stale value than a blanked one."""
        self._commit(self._sheet_with(session_site="Atlantis"), overwrites=True)
        self.session.refresh_from_db()
        self.assertEqual(self.session.site_id, self.site.pk)

    def test_upload_never_creates_or_deletes(self):
        before = (ExperimentSession.objects.using(DB).count(),
                  WbResult.objects.using(DB).count())
        self._commit(self._sheet_with(rating="4"))
        after = (ExperimentSession.objects.using(DB).count(),
                 WbResult.objects.using(DB).count())
        self.assertEqual(before, after)

    def test_a_row_whose_ids_do_not_exist_is_reported_not_written(self):
        data = self._sheet_with(session_id=999999)
        plan = self._preview(data)
        self.assertEqual(plan["summary"]["unmatched"], 1)
        result = self._commit(data)
        self.assertEqual(result["sessions_updated"], 0)
        self.assertEqual(result["results_updated"], 0)

    def test_a_result_id_from_another_session_is_refused(self):
        other = ExperimentSession.objects.create(
            procedure_type="WB", target=self.session.target,
            experimenter=self.member, site=self.site, date=date(2026, 5, 1),
            status="planned")
        data = self._sheet_with(session_id=other.pk)
        plan = self._preview(data)
        self.assertEqual(plan["summary"]["unmatched"], 1)

    def test_re_uploading_the_same_file_is_a_no_op(self):
        """The user who times out will retry."""
        data = self._sheet_with(rating="4")
        first = self._commit(data)
        self.assertEqual(first["results_updated"], 1)
        second = self._commit(data)
        self.assertEqual(second["results_updated"], 0)
        self.assertEqual(second["sessions_updated"], 0)

    def test_an_unreadable_file_is_a_message_not_a_crash(self):
        r = self.client.post("/pipeline/sessions/board/upload/preview/",
                             {"file": SimpleUploadedFile("s.xlsx", b"not a workbook")})
        self.assertEqual(r.status_code, 400)
        self.assertIn("could not read", r.json()["error"])

    def test_the_sheet_and_the_screen_show_the_same_result_columns(self):
        """One source of truth — the model — not a written-out list that drifts."""
        from pipeline.services import session_board as board
        wb = self._download()
        header = [sio._norm(c.value) for c in wb["WB"][1]]
        for field in board.result_field_names("WB"):
            self.assertIn(field, header)


class AReadingWithNoResultRowIsRefusedNotDroppedTests(TestCase):
    """The sheet matches on ``session_id`` + ``result_id`` and only ever edits.

    A missing ``session_id`` was reported. A missing ``result_id`` was **silent**:
    the result columns were skipped with a bare ``continue``, and a row carrying
    nothing else produced no item at all — so a line holding a dilution, a signal
    and a rating previewed as ``0 rows would change``, with no error and nothing
    in the errors list.

    Two ways to get there and neither is exotic: a session with no results yet
    exports with the cell blank, and adding a row by hand for an antibody that is
    not in the session is the obvious thing to do with a spreadsheet.

    So it is named per row, counted in the summary **and at the save**, and the
    row's writable half still lands — the same bargain a concentration that
    cannot be converted strikes.
    """

    databases = {"default", "pipeline_db", "academy_db"}

    def setUp(self):
        self.site = Site.objects.create(name="Leicester", short_code="LEI")
        for alias in ("academy_db", DB):
            u = User(username="vera")
            u.set_password("pw")
            u.save(using=alias)
        pu = User.objects.using(DB).get(username="vera")
        self.member = Member.objects.create(
            user_id=pu.pk, site_id=self.site.pk, role="experimenter",
            is_active=True, display_name="Vera Ruiz Moleon")
        self.target = Target.objects.create(protein_name="SOD", gene_name="SOD1")
        # Planned, no results — so the export writes a row with a blank result_id.
        self.session = ExperimentSession.objects.create(
            procedure_type="WB", target=self.target, experimenter=self.member,
            site=self.site, date=date(2026, 2, 18), status="planned", comments="")
        self.client = Client()
        self.assertTrue(self.client.login(username="vera", password="pw"))

    def _parsed(self, **values):
        return {"ok": True, "errors": [], "rows": [{
            "sheet": "WB", "line": 2, "procedure": "WB",
            "session_id": self.session.pk, "result_id": None,
            "values": {"session_id": self.session.pk, "result_id": "",
                       "gene": "SOD1", "procedure": "WB", **values},
        }]}

    def test_a_session_with_no_results_exports_a_blank_result_id(self):
        """The path a person reaches without doing anything unusual."""
        r = self.client.get("/pipeline/sessions/board/export/")
        wb = openpyxl.load_workbook(io.BytesIO(r.content))
        ws = wb["WB"]
        header = [sio._norm(c.value) for c in ws[1]]
        rid = header.index("result_id") + 1
        self.assertEqual(ws.max_row, 2, "the session should still get its row")
        self.assertIn(ws.cell(2, rid).value, ("", None))

    def test_the_readings_are_named_and_counted(self):
        plan = sio.plan(self._parsed(dilution="1:1000", signal="specific band",
                                     rating="pass"))
        self.assertEqual(plan["summary"]["dropped_readings"], 3)
        item = plan["items"][0]
        self.assertEqual(sorted(item["no_result_row"]),
                         ["dilution", "rating", "signal"])
        self.assertIn("result_id", item["warning"])
        self.assertIn("sessions board", item["warning"],
                      "a refusal must say where the control is")

    def test_the_writable_half_of_the_row_still_lands(self):
        """A `warning`, not an `error` — `apply` skips a row whose error is set,
        and the session cells on this row are perfectly writable."""
        parsed = self._parsed(dilution="1:1000",
                              session_comments="ran on the new gel")
        out = sio.apply(parsed)
        self.session.refresh_from_db()
        self.assertEqual(self.session.comments, "ran on the new gel")
        self.assertEqual(out["sessions_updated"], 1)
        self.assertEqual(out["skipped"], 0)

    def test_the_save_says_it_too(self):
        """At the save, not only at the check: whoever pressed through the
        preview is the person who most needs telling."""
        out = sio.apply(self._parsed(dilution="1:1000", rating="pass"))
        self.assertEqual(out["dropped_readings"], 2)

    def test_nothing_is_reported_when_there_is_nothing_to_drop(self):
        out = sio.apply(self._parsed(session_comments="just the header"))
        self.assertEqual(out["dropped_readings"], 0)
        self.assertEqual(sio.plan(self._parsed())["summary"]["dropped_readings"], 0)

    def test_the_column_tip_explains_a_blank(self):
        self.assertIn("Blank means", sio.COLUMN_TIPS["result_id"])
