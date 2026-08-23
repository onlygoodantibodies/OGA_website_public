"""Failure modes that have actually bitten: slow external APIs, and timeouts.

UniProt and SciCrunch are reached over the network from a request thread. They
have gone slow before, and dev often cannot reach them at all. Three rules keep
that from turning into a broken page or a half-written database, and each one is
pinned here:

  1. **The board never touches the network.** Reading, saving and exporting
     targets must not depend on an external service being up.
  2. **A slow or dead API degrades, never raises.** A paste still lands; the
     record is created bare and enriched later.
  3. **No external call inside an open transaction.** A 10-second UniProt call
     inside ``transaction.atomic`` holds a PostgreSQL transaction open for ten
     seconds. ``target_list_io.apply`` is deliberately built this way — targets
     are created bare and ``enrich_targets_from_uniprot`` fills them afterwards —
     and this pins it so a future edit cannot quietly reintroduce the coupling.

Plus the two properties that make a timeout survivable rather than destructive:
a commit is all-or-nothing, and re-uploading after a timeout is a no-op.
"""
from __future__ import annotations

import io
from contextlib import contextmanager
from unittest import mock

import requests
from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import transaction
from django.test import Client, TestCase

from pipeline.models import Member, Site, Target, TargetNomination

DB = "pipeline_db"

# Every place the pipeline can reach out. Patched wholesale rather than
# per-service so a newly added caller is caught too.
_HTTP_CALLS = ("requests.get", "requests.post", "requests.request")


@contextmanager
def no_network(reason="this code path must not touch the network"):
    """Make any outbound HTTP call a test failure."""
    def boom(*args, **kwargs):
        raise AssertionError(f"{reason} (called with {args[:1]})")

    with mock.patch.multiple("requests", get=boom, post=boom, request=boom):
        yield


@contextmanager
def network_raises(exc):
    """Every outbound call fails the way a dead or slow endpoint fails."""
    def boom(*args, **kwargs):
        raise exc

    with mock.patch.multiple("requests", get=boom, post=boom, request=boom):
        yield


def _atomic_depth():
    conn = transaction.get_connection(DB)
    return len(conn.savepoint_ids) + (1 if conn.in_atomic_block else 0)


@contextmanager
def assert_no_network_in_transaction():
    """Fail if an outbound call happens inside a transaction opened by the code
    under test.

    This is the one that costs real money in production: the connection is held
    for the duration of the call, so a slow UniProt becomes a slow database.

    Depth is measured *relative to entry*, not absolutely. Django's ``TestCase``
    wraps every test in a transaction of its own, so ``in_atomic_block`` is
    already true before the code under test runs — an absolute check fires on
    everything and therefore proves nothing.
    """
    baseline = _atomic_depth()

    def guard(*args, **kwargs):
        if _atomic_depth() > baseline:
            raise AssertionError(
                "outbound HTTP call made inside an open transaction on "
                f"{DB} — a slow endpoint would hold the connection open")
        raise requests.exceptions.ConnectionError("offline in tests")

    with mock.patch.multiple("requests", get=guard, post=guard, request=guard):
        yield


def _member_client(testcase, site):
    for alias in ("academy_db", DB):
        u = User(username="carl")
        u.set_password("pw")
        u.save(using=alias)
    pu = User.objects.using(DB).get(username="carl")
    Member.objects.create(user_id=pu.pk, site_id=site.pk, role="admin", is_active=True)
    client = Client()
    testcase.assertTrue(client.login(username="carl", password="pw"))
    return client


def _carl_workbook(rows, *, header_row=5):
    """A workbook shaped like the real target list."""
    import openpyxl
    from pipeline.services.tests.test_target_board import CARL_HEADERS
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Complete target list"
    ws.cell(1, 26, "legend")
    ws.cell(4, 7, "Target information")
    for i, h in enumerate(CARL_HEADERS, start=1):
        ws.cell(header_row, i, h)
    for r, row in enumerate(rows, start=header_row + 1):
        for i, v in enumerate(row, start=1):
            ws.cell(r, i, v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _row(gene, site="McGill"):
    return ["NIH", "AMP-AD", "Funding available", "2020", "", site, "",
            gene, "", "", "", "", "", "", "", "", "", "", "", ""]


class GuardsAreRealTests(TestCase):
    """The guards above are only worth having if they actually fire.

    A no-network assertion that silently passes because the patch missed is
    worse than no test at all — it reads as proof of something never checked.
    """
    databases = {"default", "pipeline_db"}

    def test_no_network_catches_an_outbound_call(self):
        with self.assertRaises(AssertionError):
            with no_network("deliberate"):
                requests.get("https://rest.uniprot.org/")

    def test_transaction_guard_fires_only_inside_a_transaction(self):
        with assert_no_network_in_transaction():
            # Outside a transaction the call is merely offline, not a failure.
            with self.assertRaises(requests.exceptions.ConnectionError):
                requests.get("https://rest.uniprot.org/")

        with assert_no_network_in_transaction():
            with self.assertRaises(AssertionError):
                with transaction.atomic(using=DB):
                    requests.get("https://rest.uniprot.org/")


class BoardIsOfflineSafeTests(TestCase):
    """Rule 1 — a UniProt outage must not be able to take the board down."""
    databases = {"default", "pipeline_db", "academy_db"}

    def setUp(self):
        self.site = Site.objects.create(name="McGill", short_code="MCG")
        self.client = _member_client(self, self.site)
        self.target = Target.objects.create(protein_name="Synuclein", gene_name="SNCA")

    def test_reading_the_board_makes_no_outbound_call(self):
        with no_network("loading board rows"):
            response = self.client.get("/pipeline/targets/board/rows/")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["count"], 1)

    def test_saving_a_cell_makes_no_outbound_call(self):
        with no_network("saving a board cell"):
            response = self.client.post(
                "/pipeline/targets/board/patch/",
                {"target_id": self.target.pk, "field": "essential_gene", "value": "NO"})
        self.assertEqual(response.status_code, 200, response.content[:200])
        self.target.refresh_from_db()
        self.assertEqual(self.target.essential_gene, "NO")

    def test_exporting_makes_no_outbound_call(self):
        with no_network("exporting the board"):
            response = self.client.get("/pipeline/targets/board/export/")
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response["Content-Disposition"].endswith('.xlsx"'))

    def test_upload_preview_and_commit_make_no_outbound_call(self):
        data = _carl_workbook([_row("MAPT")])
        with no_network("previewing an upload"):
            preview = self.client.post("/pipeline/targets/board/upload/preview/",
                                       {"file": SimpleUploadedFile("c.xlsx", data),
                                        "default_site": "McGill"})
        self.assertEqual(preview.status_code, 200, preview.content[:200])

        with no_network("committing an upload"):
            commit = self.client.post("/pipeline/targets/board/upload/commit/",
                                      {"file": SimpleUploadedFile("c.xlsx", data),
                                       "default_site": "McGill"})
        self.assertEqual(commit.status_code, 200, commit.content[:200])
        self.assertTrue(Target.objects.using(DB).filter(gene_name__iexact="MAPT").exists())


class SlowApiDegradesTests(TestCase):
    """Rule 2 — a slow or dead endpoint must not raise, and must not block a write."""
    databases = {"default", "pipeline_db", "academy_db"}

    def _resolve(self, gene="SNCA"):
        from pipeline.services.targets import resolve_or_create_target
        return resolve_or_create_target(gene)

    def test_uniprot_timeout_still_creates_the_target(self):
        with network_raises(requests.exceptions.Timeout("timed out after 10s")):
            target, created = self._resolve("SNCA")
        self.assertTrue(created)
        self.assertEqual(target.gene_name, "SNCA")
        self.assertIsNone(target.uniprot_id)

    def test_uniprot_connection_error_still_creates_the_target(self):
        with network_raises(requests.exceptions.ConnectionError("no route to host")):
            target, created = self._resolve("MAPT")
        self.assertTrue(created)
        self.assertEqual(target.gene_name, "MAPT")

    def test_a_hanging_endpoint_is_bounded_by_a_timeout_argument(self):
        """Every outbound call must pass timeout= — without it a hang is forever."""
        captured = {}

        def capture(*args, **kwargs):
            captured.update(kwargs)
            raise requests.exceptions.Timeout("bounded")

        with mock.patch("requests.get", side_effect=capture):
            self._resolve("SOD1")
        self.assertIn("timeout", captured,
                      "UniProt was called without a timeout — a hang would never return")
        self.assertLessEqual(captured["timeout"], 30)

    def test_a_second_resolve_after_an_outage_reuses_the_bare_target(self):
        """Retrying after a timeout must not create a duplicate."""
        with network_raises(requests.exceptions.Timeout("t")):
            first, created_first = self._resolve("TARDBP")
            second, created_second = self._resolve("TARDBP")
        self.assertTrue(created_first)
        self.assertFalse(created_second)
        self.assertEqual(first.pk, second.pk)
        self.assertEqual(Target.objects.using(DB).filter(gene_name__iexact="TARDBP").count(), 1)


class NoNetworkInsideTransactionTests(TestCase):
    """Rule 3 — the commit path must not hold a connection open across an API call."""
    databases = {"default", "pipeline_db", "academy_db"}

    def setUp(self):
        self.site = Site.objects.create(name="McGill", short_code="MCG")
        self.client = _member_client(self, self.site)

    def test_target_upload_commit_never_calls_out_mid_transaction(self):
        from pipeline.services import target_list_io as tio
        data = _carl_workbook([_row("SNCA"), _row("MAPT"), _row("SOD1")])
        parsed = tio.parse_workbook(io.BytesIO(data))
        with assert_no_network_in_transaction():
            result = tio.apply(parsed, default_site="McGill")
        self.assertTrue(result.get("ok", True), result)
        self.assertEqual(Target.objects.using(DB).count(), 3)


class TimeoutIsSurvivableTests(TestCase):
    """A commit that dies half way must leave nothing behind, and retrying must
    be safe — those two together are what make a timeout an inconvenience
    rather than a corruption."""
    databases = {"default", "pipeline_db", "academy_db"}

    def setUp(self):
        self.site = Site.objects.create(name="McGill", short_code="MCG")
        self.client = _member_client(self, self.site)
        self.data = _carl_workbook([_row("SNCA"), _row("MAPT"), _row("SOD1")])

    def _commit(self):
        return self.client.post("/pipeline/targets/board/upload/commit/",
                                {"file": SimpleUploadedFile("c.xlsx", self.data),
                                 "default_site": "McGill"})

    def test_a_failure_part_way_through_writes_nothing(self):
        from pipeline.services import target_list_io as tio
        parsed = tio.parse_workbook(io.BytesIO(self.data))
        real_save = TargetNomination.save
        calls = {"n": 0}

        def explode(self, *args, **kwargs):
            calls["n"] += 1
            if calls["n"] == 2:
                raise RuntimeError("connection reset mid-commit")
            return real_save(self, *args, **kwargs)

        with mock.patch.object(TargetNomination, "save", explode):
            with self.assertRaises(RuntimeError):
                tio.apply(parsed, default_site="McGill")

        self.assertEqual(Target.objects.using(DB).count(), 0,
                         "a failed commit left targets behind — it is not atomic")
        self.assertEqual(TargetNomination.objects.using(DB).count(), 0)

    def test_re_uploading_the_same_file_is_a_no_op(self):
        """The user who times out will retry. Retrying must not duplicate."""
        self.assertEqual(self._commit().status_code, 200)
        after_first = (Target.objects.using(DB).count(),
                       TargetNomination.objects.using(DB).count())

        self.assertEqual(self._commit().status_code, 200)
        after_second = (Target.objects.using(DB).count(),
                        TargetNomination.objects.using(DB).count())

        self.assertEqual(after_first, after_second,
                         f"re-upload changed the database {after_first} → {after_second}")
