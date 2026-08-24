"""All four boards download what you are looking at, and take it back.

Targets and sessions had download *and* upload. Antibodies and cell lines had
download only, and sent you to the whole-dataset page for the way back — a
different page, a different shape, and every record in the consortium rather
than the rows you had filtered to. The field test never got to test either half.

Two things are pinned. A download carries the grid's own filters, so the file you
edit is the rows you chose. And the upload goes through the same parse and the
same commit as a paste, so there is one write path per entity rather than one for
typing and another for files.
"""
from __future__ import annotations

import io

from django.test import TestCase

from pipeline.models import (Antibody, CellLine, Company, Member, Site, Target)
from pipeline.tests_timeouts import DB, _member_client


def _sheet(rows):
    """A tiny .xlsx: header row plus data, the shape the templates produce."""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    buf.name = "edited.xlsx"
    return buf


def _read(xlsx_bytes):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), read_only=True)
    return [list(r) for r in wb.active.iter_rows(values_only=True)]


class DownloadMatchesTheGridTests(TestCase):
    """The download reimplemented the retired search page's filters, so a board
    filtered to one gene handed you the whole dataset."""

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)
        self.stmn2 = Target.objects.using(DB).create(gene_name="STMN2")
        self.elp3 = Target.objects.using(DB).create(gene_name="ELP3")
        company = Company.objects.using(DB).create(name="Proteintech")
        for target, cat in ((self.stmn2, "10586-1-AP"), (self.elp3, "24523-1-AP")):
            Antibody.objects.using(DB).create(
                target_id=target.pk, company_id=company.pk,
                catalogue_number=cat, site_id=self.site.pk)
            CellLine.objects.using(DB).create(
                name=f"HAP1 {target.gene_name} KO", genotype="KO",
                target_id=target.pk, site_id=self.site.pk)

    def test_the_antibody_download_honours_the_gene_filter(self):
        resp = self.client.get("/pipeline/antibodies/export/", {"gene": "STMN2"})
        self.assertEqual(resp.status_code, 200)
        body = "\n".join(str(r) for r in _read(resp.content))
        self.assertIn("10586-1-AP", body)
        self.assertNotIn("24523-1-AP", body)

    def test_the_cell_line_download_honours_the_gene_filter(self):
        resp = self.client.get("/pipeline/cell-lines/export/", {"gene": "STMN2"})
        self.assertEqual(resp.status_code, 200)
        body = "\n".join(str(r) for r in _read(resp.content))
        self.assertIn("HAP1 STMN2 KO", body)
        self.assertNotIn("HAP1 ELP3 KO", body)

    def test_an_unfiltered_download_still_gives_everything(self):
        resp = self.client.get("/pipeline/antibodies/export/")
        body = "\n".join(str(r) for r in _read(resp.content))
        self.assertIn("10586-1-AP", body)
        self.assertIn("24523-1-AP", body)

    def test_both_boards_offer_the_way_back(self):
        for url in ("/pipeline/antibodies/board/", "/pipeline/cell-lines/board/"):
            with self.subTest(url=url):
                body = self.client.get(url).content.decode()
                self.assertIn("Upload a sheet", body)
                self.assertIn("Download these", body)


class UploadGoesThroughTheSamePreviewAndCommitTests(TestCase):
    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)
        self.target = Target.objects.using(DB).create(gene_name="STMN2")

    def test_an_uploaded_sheet_previews_per_row(self):
        f = _sheet([["gene", "catalogue", "company", "lot"],
                    ["STMN2", "10586-1-AP", "Proteintech", "20051"],
                    ["STMN2", "NBP1-49461", "Bio-Techne", ""]])
        resp = self.client.post("/pipeline/import/upload/antibodies/", {"file": f})
        self.assertEqual(resp.status_code, 200)
        data = resp.json()
        self.assertEqual(len(data["items"]), 2)
        for item in data["items"]:
            self.assertTrue(item["status"])

    def test_the_templates_own_example_row_is_dropped_on_upload(self):
        """The blank template, downloaded and posted back with one real row.

        Twelfth field test, listing what to keep: *"the e.g. example row in the
        template is skipped on upload"*. It was pinned in
        `services/example_row.py` and at the paste parsers, and **not on this
        door** — which is the split this repo keeps paying for, a rule enforced
        in one path and not the other. Abcam's `ab138501` being created as a
        real antibody is exactly the failure the gene page's own panel warns
        about in words.

        The real template, not a hand-built sheet: the marker lives in its first
        cell, and a fixture that spells it out tests the fixture.
        """
        import openpyxl
        cases = [
            ("antibodies", ["", "10586-1-AP", "Proteintech"],
             "/pipeline/antibodies/bulk/commit/",
             lambda: Antibody.objects.using(DB)
                     .filter(catalogue_number="ab138501").exists()),
            ("cell-lines", ["HAP1", "", "WT"],
             "/pipeline/cell-lines/bulk/commit/",
             lambda: CellLine.objects.using(DB).filter(catalogue_number="HZGHC001").exists()),
        ]
        for kind, row, commit_url, example_exists in cases:
            with self.subTest(kind=kind):
                blank = self.client.get(f"/pipeline/import/template/{kind}/")
                self.assertEqual(blank.status_code, 200)
                wb = openpyxl.load_workbook(io.BytesIO(blank.content))
                ws = wb.worksheets[0]
                self.assertIn("e.g.", str(ws.cell(row=2, column=1).value or "").lower(),
                              "row 2 of the template no longer says it is an example")
                ws.append(row)
                buf = io.BytesIO()
                wb.save(buf)
                buf.seek(0)
                buf.name = f"{kind}.xlsx"

                # `default_gene`, the way the gene page sends it — its own
                # instruction is to leave the gene column blank.
                data = self.client.post(
                    f"/pipeline/import/upload/{kind}/",
                    {"file": buf, "default_gene": "STMN2"}).json()
                self.assertEqual(
                    len(data["items"]), 1,
                    f"{kind}: the example row was read as a record — {data['items']}")

                # `text` still carries the example line, and that is fine — both
                # ends run the same parser, which is the whole point of handing
                # the text back. What must be true is at the far end.
                resp = self.client.post(
                    commit_url, {"text": data["text"], "dry_run": False,
                                 "default_gene": "STMN2"},
                    content_type="application/json")
                self.assertEqual(resp.status_code, 200)
                self.assertEqual(len(resp.json()["result"]["created"]), 1,
                                 f"{kind}: the commit created the example row too")
                self.assertFalse(example_exists(),
                                 f"{kind}: the template's example became a record")

    def test_the_preview_returns_the_text_so_the_save_commits_what_was_shown(self):
        """Without this the board would need its own commit endpoint for files —
        a second write path for the same job, free to drift from the paste box."""
        f = _sheet([["gene", "catalogue"], ["STMN2", "10586-1-AP"]])
        data = self.client.post("/pipeline/import/upload/antibodies/",
                                {"file": f}).json()
        self.assertIn("text", data)
        self.assertIn("10586-1-AP", data["text"])

        resp = self.client.post(
            "/pipeline/antibodies/bulk/commit/",
            {"text": data["text"], "dry_run": False},
            content_type="application/json")
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(len(resp.json()["result"]["created"]), 1)
        self.assertTrue(Antibody.objects.using(DB)
                        .filter(catalogue_number="10586-1-AP").exists())

    def test_a_cell_line_sheet_round_trips(self):
        f = _sheet([["name", "gene", "genotype", "parent"],
                    ["HAP1", "", "WT", ""],
                    ["HAP1 STMN2 KO", "STMN2", "KO", "HAP1"]])
        data = self.client.post("/pipeline/import/upload/cell-lines/",
                                {"file": f}).json()
        self.assertEqual(len(data["items"]), 2)
        resp = self.client.post(
            "/pipeline/cell-lines/bulk/commit/",
            {"text": data["text"], "dry_run": False},
            content_type="application/json")
        self.assertEqual(resp.status_code, 200)
        names = set(CellLine.objects.using(DB).values_list("name", flat=True))
        self.assertEqual(names, {"HAP1", "HAP1 STMN2 KO"})

    def test_a_wild_type_row_keeps_its_blank_gene(self):
        """The rule that costs people a second, wrong line if they get it wrong.
        A round trip must not quietly fill it in."""
        f = _sheet([["name", "gene", "genotype"], ["HAP1", "", "WT"]])
        data = self.client.post("/pipeline/import/upload/cell-lines/",
                                {"file": f}).json()
        self.client.post("/pipeline/cell-lines/bulk/commit/",
                         {"text": data["text"], "dry_run": False},
                         content_type="application/json")
        line = CellLine.objects.using(DB).get(name="HAP1")
        self.assertIsNone(line.target_id)

    def test_an_empty_file_is_refused_in_words(self):
        resp = self.client.post("/pipeline/import/upload/antibodies/",
                                {"file": _sheet([])})
        self.assertEqual(resp.status_code, 400)
        self.assertIn("empty", resp.json()["error"])

    def test_no_file_is_refused_rather_than_500ing(self):
        resp = self.client.post("/pipeline/import/upload/antibodies/", {})
        self.assertEqual(resp.status_code, 400)
        self.assertIn("no file", resp.json()["error"])


class TheRoundTripDoesNotRehomeAnotherSitesRowsTests(TestCase):
    """The half that was missing, and the reason it mattered.

    Both downloads left out ``site``. A row is one *vial*, and site is half of
    what makes it that vial, so an upload had nothing to read and stamped every
    row with whoever uploaded the file. SOD1 had 29 antibodies on file, 22 of them
    McGill's; a Leicester user downloading them and putting the sheet straight
    back was offered 22 **new** Leicester rows. The preview was honest — they
    would have been new — which is exactly why the file had to carry the column.
    """

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        self.leicester = Site.objects.using(DB).create(
            name="Leicester", short_code="LEI")
        self.mcgill = Site.objects.using(DB).create(name="McGill", short_code="MCG")
        self.client = _member_client(self, self.leicester)
        self.target = Target.objects.using(DB).create(gene_name="SOD1")
        self.company = Company.objects.using(DB).create(name="Proteintech")
        # McGill's vial. Nothing about it belongs to the Leicester user below.
        self.theirs = Antibody.objects.using(DB).create(
            target_id=self.target.pk, company_id=self.company.pk,
            catalogue_number="10269-1-AP", lot_number="69112",
            site_id=self.mcgill.pk)
        self.their_line = CellLine.objects.using(DB).create(
            name="HAP1 SOD1 KO", genotype="KO", target_id=self.target.pk,
            company_id=self.company.pk, site_id=self.mcgill.pk)

    # ── the column exists at all ────────────────────────────────────────────
    def test_the_antibody_download_names_the_site(self):
        rows = _read(self.client.get("/pipeline/antibodies/export/").content)
        self.assertIn("site", [str(c).lower() for c in rows[0]])
        self.assertIn("McGill", [str(c) for c in rows[1]])

    def test_the_cell_line_download_names_the_site(self):
        rows = _read(self.client.get("/pipeline/cell-lines/export/").content)
        self.assertIn("site", [str(c).lower() for c in rows[0]])
        self.assertIn("McGill", [str(c) for c in rows[1]])

    # ── and the upload honours it ───────────────────────────────────────────
    def test_putting_another_sites_antibody_back_updates_it_rather_than_copying_it(self):
        f = _sheet([["gene", "catalogue", "company", "lot", "site"],
                    ["SOD1", "10269-1-AP", "Proteintech", "69112", "McGill"]])
        data = self.client.post("/pipeline/import/upload/antibodies/",
                                {"file": f}).json()
        self.assertEqual([i["status"] for i in data["items"]], ["update"])

        self.client.post("/pipeline/antibodies/bulk/commit/",
                         {"text": data["text"], "dry_run": False},
                         content_type="application/json")
        rows = Antibody.objects.using(DB).filter(catalogue_number="10269-1-AP")
        self.assertEqual(rows.count(), 1, "the round trip duplicated the vial")
        self.assertEqual(rows.first().site_id, self.mcgill.pk)

    def test_putting_another_sites_cell_line_back_updates_it_rather_than_copying_it(self):
        f = _sheet([["name", "gene", "genotype", "supplier", "site"],
                    ["HAP1 SOD1 KO", "SOD1", "KO", "Proteintech", "McGill"]])
        data = self.client.post("/pipeline/import/upload/cell-lines/",
                                {"file": f}).json()
        self.assertEqual([i["status"] for i in data["items"]], ["update"])
        self.client.post("/pipeline/cell-lines/bulk/commit/",
                         {"text": data["text"], "dry_run": False},
                         content_type="application/json")
        rows = CellLine.objects.using(DB).filter(name="HAP1 SOD1 KO")
        self.assertEqual(rows.count(), 1, "the round trip duplicated the line")
        self.assertEqual(rows.first().site_id, self.mcgill.pk)

    def test_a_blank_site_still_means_whoever_is_pasting(self):
        """The behaviour before there was a column, and the right default for
        someone entering their own bench's reagents."""
        f = _sheet([["gene", "catalogue", "company", "site"],
                    ["SOD1", "AB-NEW-1", "Proteintech", ""]])
        data = self.client.post("/pipeline/import/upload/antibodies/",
                                {"file": f}).json()
        self.client.post("/pipeline/antibodies/bulk/commit/",
                         {"text": data["text"], "dry_run": False},
                         content_type="application/json")
        made = Antibody.objects.using(DB).get(catalogue_number="AB-NEW-1")
        self.assertEqual(made.site_id, self.leicester.pk)

    def test_a_site_that_is_not_on_file_is_refused_by_name(self):
        """Not silently treated as the uploader's own — that is the re-homing this
        exists to stop — and the refusal lists the sites that do exist."""
        f = _sheet([["gene", "catalogue", "company", "site"],
                    ["SOD1", "AB-NEW-2", "Proteintech", "McGil"]])
        data = self.client.post("/pipeline/import/upload/antibodies/",
                                {"file": f}).json()
        item = data["items"][0]
        self.assertEqual(item["status"], "blocked")
        self.assertIn("McGil", item["note"])
        self.assertIn("McGill", item["note"])
        self.assertIn("Leicester", item["note"])

        self.client.post("/pipeline/antibodies/bulk/commit/",
                         {"text": data["text"], "dry_run": False},
                         content_type="application/json")
        self.assertFalse(Antibody.objects.using(DB)
                         .filter(catalogue_number="AB-NEW-2").exists())

    def test_the_preview_names_the_site_each_row_would_land_at(self):
        """The check a person actually makes. It has to be per row, because a
        downloaded sheet legitimately spans benches."""
        f = _sheet([["gene", "catalogue", "company", "site"],
                    ["SOD1", "10269-1-AP", "Proteintech", "McGill"],
                    ["SOD1", "AB-NEW-3", "Proteintech", ""]])
        data = self.client.post("/pipeline/import/upload/antibodies/",
                                {"file": f}).json()
        self.assertEqual([i["site"] for i in data["items"]],
                         ["McGill", "Leicester"])
        self.assertEqual(data["summary"]["sites"], ["Leicester", "McGill"])

    def test_an_unrecognised_site_never_creates_one(self):
        before = Site.objects.using(DB).count()
        f = _sheet([["gene", "catalogue", "company", "site"],
                    ["SOD1", "AB-NEW-4", "Proteintech", "Univeristy of Nowhere"]])
        data = self.client.post("/pipeline/import/upload/antibodies/",
                                {"file": f}).json()
        self.client.post("/pipeline/antibodies/bulk/commit/",
                         {"text": data["text"], "dry_run": False},
                         content_type="application/json")
        self.assertEqual(Site.objects.using(DB).count(), before)

    def test_the_site_column_is_in_the_template_the_add_grid_uses(self):
        """One list behind the Excel template, the Add grid and the parser — a
        column the parser honours but the template omits is a column nobody
        knows exists."""
        from pipeline.views.imports import columns_and_example
        for kind in ("antibodies", "cell-lines"):
            with self.subTest(kind=kind):
                self.assertIn("site", columns_and_example(kind)[0])
