"""Regressions from the 30 Jul 2026 boards field test.

Each test here pins one thing the field test found, so the same defect cannot
come back quietly. The findings themselves are recorded in PLATFORM_ROADMAP.md;
what matters here is that these are all *observable from the page or the
endpoint*, not internal details — a leaked template comment and a 500 dressed as
a success are both things only a rendered response shows.
"""
import io
import json
import os
import re
import shutil
import subprocess
import tempfile
from datetime import date
from decimal import Decimal
from pathlib import Path
from unittest import mock, skipUnless

from django.conf import settings
from django.test import SimpleTestCase, TestCase

from pipeline.models import (
    CellLine, CellLineVial, Site, Target, TargetNomination)
from pipeline.tests_feasibility_add import _uniprot_says
from pipeline.tests_timeouts import DB, _member_client


class NoTemplateWrapsAShortFormCommentTests(SimpleTestCase):
    """`{# … #}` is single-line, and a wrapped one is not a comment at all.

    Django's short form ends at the first `#}` **on the same line**; there is no
    multi-line version. So a comment somebody wrapped for readability is not
    parsed as a comment — it is rendered, delimiters and all, straight onto the
    page. `{% comment %}` is the one that spans lines.

    This is the fourth time it has been found and the third place it has been
    fixed. The boards leaked one, then the public gene page did, and each was
    answered with a test that renders *those* pages and greps the response — so
    the review queue leaked two more onto a page the owner was standing in front
    of, above the Release button, in a red panel about publishing to the public
    website. Three tests, each pinning the surface that had already failed, and
    the next surface unguarded every time.

    Reading the source instead of rendering pages is what makes it complete: it
    covers every template in the repo including the ones that need a login, a
    fixture or a query string to reach, and the ones nobody has written yet. It
    is also about four hundred times cheaper than the three it subsumes.
    """

    #: Anything that is not ours to edit. `/site-packages/` rather than a list
    #: of venv names: the sweep walks BASE_DIR, so a contributor who ran the
    #: conventional `python -m venv venv` had it find grappelli's and ckeditor's
    #: vendored TinyMCE and fail on a third-party file they cannot edit. Only
    #: `.venv` was skipped, though .gitignore has always listed four venv names.
    SKIP = ("/site-packages/", "/.venv/", "/staticfiles/", "/node_modules/",
            "/benchmarks/", "/pipeline_skeleton/")

    def _templates(self):
        root = Path(settings.BASE_DIR)
        for path in sorted(root.rglob("*.html")):
            if any(part in f"/{path.relative_to(root)}" for part in self.SKIP):
                continue
            yield path

    def test_no_template_wraps_a_short_form_comment(self):
        offenders = []
        for path in self._templates():
            source = path.read_text(errors="ignore")
            for match in re.finditer(re.escape("{" + "#"), source):
                line = source.count("\n", 0, match.start()) + 1
                rest = source[match.start():]
                end = rest.find("#" + "}")
                if end == -1:
                    offenders.append(f"{path.name}:{line} — never closed")
                elif "\n" in rest[:end]:
                    offenders.append(f"{path.name}:{line} — closes on a later line")
        self.assertEqual(
            offenders, [],
            "These render as visible page text rather than being stripped. "
            "Use {% comment %} for anything that does not fit on one line:\n  "
            + "\n  ".join(offenders))


class BoardsRenderCleanlyTests(TestCase):
    """Nothing meant for a developer's eyes reaches a scientist's screen."""

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)

    def test_no_template_comment_leaks_onto_a_board(self):
        """Django's short comment form does not span lines: a wrapped one renders
        as text. It also ends at the *first* closing delimiter, so one that names
        the delimiters cuts itself short and prints its own tail.

        This test looked only for the opening delimiter and so passed over the
        second leak, which showed a bare tail with no opening brace in it. Both
        halves are checked now, on every board.
        """
        for url in ("/pipeline/antibodies/board/", "/pipeline/cell-lines/board/",
                    "/pipeline/sessions/board/", "/pipeline/targets/board/"):
            for delim in ("{" + "#", "#" + "}"):
                with self.subTest(url=url, delim=delim):
                    resp = self.client.get(url)
                    # 200, or an unauthenticated redirect would pass this vacuously.
                    self.assertEqual(resp.status_code, 200)
                    body = resp.content.decode()
                    self.assertNotIn(
                        delim, body, f"a template comment leaked onto {url}")

    def test_an_empty_cell_shows_a_dash_and_not_the_markup_for_one(self):
        """`<span class="text-gray-300">—</span>` printed as visible words on
        three boards and on the session result cards, because the empty-cell
        placeholder was markup and the two places that draw it — esc() and
        textContent — both treat what they are given as text.

        This one lives entirely in the browser, so what is pinned is the source:
        an editable cell's placeholder must be the text form. There is no server
        response that shows the defect, and a guard that cannot run is worse than
        an honest source check.
        """
        js = (Path(settings.BASE_DIR) / "pipeline/static/pipeline/board.js").read_text()
        self.assertIn("const EMPTY = '—';", js,
                      "board.js needs a text-only empty placeholder")
        self.assertNotIn(
            "opts.empty) || DASH", js,
            "cell() is handing markup to an escaping context again")

    def test_sessions_board_keeps_results_on_the_board(self):
        """The create confirmation used to link to /pipeline/session/<id>/ — the
        page the board replaces — which is how the field test ended up recording
        every result on the legacy surface."""
        resp = self.client.get("/pipeline/sessions/board/")
        self.assertEqual(resp.status_code, 200)
        body = resp.content.decode()
        self.assertIn("data-open-results", body)
        self.assertNotIn("Open the session to record results", body)


class BulkCommitRefusesInJsonTests(TestCase):
    """A failed paste is a refusal the page can read, never a raw 500."""

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)

    def _post(self, url, payload):
        return self.client.post(url, payload, content_type="application/json")

    def test_cell_line_commit_failure_is_json_not_a_500(self):
        with mock.patch("pipeline.services.bulk_cell_lines.apply",
                        side_effect=RuntimeError("boom")):
            resp = self._post("/pipeline/cell-lines/bulk/commit/",
                              '{"text": "name\\nHAP1", "dry_run": false}')
        self.assertEqual(resp.status_code, 400)
        data = resp.json()
        self.assertIs(data["ok"], False)
        # The page needs a sentence to show; "boom" is for the log only.
        self.assertIn("Check the board", data["error"])
        self.assertNotIn("boom", data["error"])

    def test_antibody_commit_failure_is_json_not_a_500(self):
        with mock.patch("pipeline.services.bulk_antibodies.apply",
                        side_effect=RuntimeError("boom")):
            resp = self._post("/pipeline/antibodies/bulk/commit/",
                              '{"text": "gene\\tcatalogue\\nSOD1\\t123", '
                              '"dry_run": false}')
        self.assertEqual(resp.status_code, 400)
        self.assertIs(resp.json()["ok"], False)


class PastePreviewIsPerRowTests(TestCase):
    """A count cannot be checked; the row about to be overwritten must be named.

    The preview table is drawn client-side from ``items``, so what is pinned here
    is the contract it draws from: one item per data row, each carrying a status.
    """

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)

    def test_antibody_preview_reports_a_status_for_every_row(self):
        Target.objects.using(DB).create(gene_name="STMN2")
        resp = self.client.post(
            "/pipeline/antibodies/bulk/parse/",
            json.dumps({"text": "gene\tcatalogue\n"
                                "STMN2\t10586-1-AP\n"
                                "STMN2\tNBP1-49461\n"}),
            content_type="application/json")
        self.assertEqual(resp.status_code, 200)
        items = resp.json()["items"]
        self.assertEqual(len(items), 2)
        self.assertEqual([i["catalogue"] for i in items],
                         ["10586-1-AP", "NBP1-49461"])
        for i in items:
            self.assertTrue(i["status"], "every row needs a status to render")

    def test_cell_line_preview_reports_a_status_for_every_row(self):
        resp = self.client.post(
            "/pipeline/cell-lines/bulk/parse/",
            json.dumps({"text": "name\tgenotype\nHAP1\tWT\nSH-SY5Y\tWT\n"}),
            content_type="application/json")
        self.assertEqual(resp.status_code, 200)
        items = resp.json()["items"]
        self.assertEqual(len(items), 2)
        self.assertEqual([i["name"] for i in items], ["HAP1", "SH-SY5Y"])
        for i in items:
            self.assertTrue(i["status"])


class ResultColumnGroupingTests(TestCase):
    """Folding the method away must never lose a field.

    The results panel shows the answer first and folds the method (fixative,
    microscope, buffers) away, because 21 columns in one row needs a sideways
    scrollbar that then hides the session table's own columns. The risk in
    grouping is a hand-kept list going stale, so what is pinned is that every
    field the model has still lands in exactly one group.
    """
    databases = {"academy_db", "pipeline_db"}

    def test_every_result_field_is_in_exactly_one_group(self):
        from pipeline.services.session_board import (RESULT_MODELS,
                                                     result_column_group,
                                                     result_field_names)
        for proc in RESULT_MODELS:
            fields = result_field_names(proc)
            self.assertTrue(fields, f"{proc} has no result fields")
            groups = {f: result_column_group(proc, f) for f in fields}
            with self.subTest(procedure=proc):
                self.assertEqual(set(groups), set(fields))
                self.assertTrue(set(groups.values()) <= {"reading", "method"})
                # Every procedure must show *something* without expanding.
                self.assertIn("reading", groups.values())

    def test_an_unknown_field_defaults_to_method_not_dropped(self):
        """A field added to a result model appears without touching this file."""
        from pipeline.services.session_board import result_column_group
        self.assertEqual(result_column_group("IF", "brand_new_field"), "method")

    def test_comments_stay_in_the_reading_group(self):
        """Every result model has its own comments, distinct from the session's,
        and it is where a scientist writes what actually happened."""
        from pipeline.services.session_board import (RESULT_MODELS,
                                                     result_column_group)
        for proc in RESULT_MODELS:
            self.assertEqual(result_column_group(proc, "comments"), "reading")


class VialMatchingTests(TestCase):
    """An antibody is the product; a row is one vial of it.

    Leicester and Montreal can hold the same catalogue number from the same lot
    and still have physically different vials, tested and written up separately.
    The paste matcher used to resolve on the product alone, so the second site's
    paste edited the first site's record instead of recording their own.
    """

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        from pipeline.models import Company
        self.leicester = Site.objects.using(DB).create(name="Leicester",
                                                       short_code="LEI")
        self.montreal = Site.objects.using(DB).create(name="Montreal",
                                                      short_code="MTL")
        self.client = _member_client(self, self.leicester)
        self.target = Target.objects.using(DB).create(gene_name="STMN2")
        self.company = Company.objects.using(DB).create(name="Proteintech")

    def _member_at(self, site):
        class _M:
            site_id = site.pk
        return _M()

    def _row(self, lot=""):
        return {"gene": "STMN2", "catalogue": "10586-1-AP",
                "company": "Proteintech", "lot": lot}

    def _apply(self, row, member):
        from pipeline.services import bulk_antibodies
        return bulk_antibodies.apply([row], member=member)

    def _antibodies(self):
        from pipeline.models import Antibody
        return list(Antibody.objects.using(DB)
                    .filter(target_id=self.target.pk).order_by("pk"))

    def test_a_second_site_records_its_own_vial(self):
        self._apply(self._row(lot="LOT-A"), self._member_at(self.leicester))
        out = self._apply(self._row(lot="LOT-A"), self._member_at(self.montreal))
        self.assertEqual(len(out["created"]), 1, "Montreal must get its own row")
        abs_ = self._antibodies()
        self.assertEqual(len(abs_), 2)
        self.assertEqual({a.site_id for a in abs_},
                         {self.leicester.pk, self.montreal.pk})
        # Same product, same lot — different vials.
        self.assertEqual({a.lot_number for a in abs_}, {"LOT-A"})

    def test_a_second_lot_at_one_site_is_a_second_vial(self):
        me = self._member_at(self.leicester)
        self._apply(self._row(lot="LOT-A"), me)
        out = self._apply(self._row(lot="LOT-B"), me)
        self.assertEqual(len(out["created"]), 1)
        self.assertEqual({a.lot_number for a in self._antibodies()},
                         {"LOT-A", "LOT-B"})

    def test_re_pasting_the_same_vial_updates_it(self):
        me = self._member_at(self.leicester)
        self._apply(self._row(lot="LOT-A"), me)
        out = self._apply(self._row(lot="LOT-A"), me)
        self.assertEqual(out["created"], [])
        self.assertEqual(len(out["updated"]), 1)
        self.assertEqual(len(self._antibodies()), 1)

    def test_a_lot_fills_a_blank_one_rather_than_forking(self):
        """A blank lot is "nobody wrote it down", not "a different vial" — else
        the row with the lot and the row with the results drift apart."""
        me = self._member_at(self.leicester)
        self._apply(self._row(), me)
        out = self._apply(self._row(lot="LOT-A"), me)
        self.assertEqual(out["created"], [])
        abs_ = self._antibodies()
        self.assertEqual(len(abs_), 1)
        self.assertEqual(abs_[0].lot_number, "LOT-A")

    def test_a_paste_with_no_lot_updates_rather_than_making_a_lotless_twin(self):
        me = self._member_at(self.leicester)
        self._apply(self._row(lot="LOT-A"), me)
        out = self._apply(self._row(), me)
        self.assertEqual(out["created"], [])
        self.assertEqual(len(self._antibodies()), 1)

    def test_the_preview_agrees_with_what_the_write_does(self):
        """The one thing a preview must never do is disagree — the count of
        creates it promises has to be the count that happens."""
        from pipeline.services import bulk_antibodies
        self._apply(self._row(lot="LOT-A"), self._member_at(self.leicester))
        mtl = self._member_at(self.montreal)
        row = self._row(lot="LOT-A")
        planned = bulk_antibodies.plan([row], member=mtl)
        self.assertEqual([i["status"] for i in planned], ["create"])
        out = self._apply(row, mtl)
        self.assertEqual(len(out["created"]), 1)

    def test_a_row_with_no_site_is_adopted_not_duplicated(self):
        """Most of the live table predates anyone recording a site. Without this
        the first paste after the rule changed would duplicate all of it."""
        from pipeline.models import Antibody
        legacy = Antibody.objects.using(DB).create(
            target_id=self.target.pk, company_id=self.company.pk,
            catalogue_number="10586-1-AP", site_id=None)
        out = self._apply(self._row(lot="LOT-A"),
                          self._member_at(self.leicester))
        self.assertEqual(out["created"], [])
        abs_ = self._antibodies()
        self.assertEqual(len(abs_), 1)
        self.assertEqual(abs_[0].pk, legacy.pk)
        self.assertEqual(abs_[0].lot_number, "LOT-A")

    def test_only_one_site_can_adopt_an_unsited_row(self):
        from pipeline.models import Antibody
        Antibody.objects.using(DB).create(
            target_id=self.target.pk, company_id=self.company.pk,
            catalogue_number="10586-1-AP", site_id=None)
        self._apply(self._row(lot="LOT-A"), self._member_at(self.leicester))
        out = self._apply(self._row(lot="LOT-A"), self._member_at(self.montreal))
        self.assertEqual(len(out["created"]), 1)
        self.assertEqual(len(self._antibodies()), 2)

    def test_a_session_attaches_to_your_own_sites_vial(self):
        """Once a product has a vial per site, "the antibody" is ambiguous, and
        picking the oldest row would hand Leicester Montreal's tube."""
        from pipeline.services import bulk_sessions
        self._apply(self._row(lot="LOT-A"), self._member_at(self.montreal))
        self._apply(self._row(lot="LOT-B"), self._member_at(self.leicester))
        ab, created = bulk_sessions.resolve_or_create(
            self.target, {"antibody": "10586-1-AP", "company": "Proteintech"},
            self._member_at(self.leicester))
        self.assertFalse(created)
        self.assertEqual(ab.site_id, self.leicester.pk)


class SiteAllocationTests(TestCase):
    """A target's site: typed names are checked, and one can be changed."""

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.other = Site.objects.using(DB).create(name="McGill", short_code="MCG")
        self.client = _member_client(self, self.site)
        self.target = Target.objects.using(DB).create(gene_name="STMN2")

    def _patch_site(self, value):
        return self.client.post("/pipeline/targets/board/patch/", {
            "target_id": self.target.pk, "field": "site", "value": value})

    def _noms(self):
        return list(TargetNomination.objects.using(DB)
                    .filter(target_id=self.target.pk))

    def test_a_site_can_be_set_on_a_target_with_no_nomination(self):
        """Every target added from the feasibility page starts with none."""
        resp = self._patch_site("Leicester")
        self.assertEqual(resp.status_code, 200)
        self.assertIs(resp.json()["ok"], True)
        self.assertEqual([n.site_id for n in self._noms()], [self.site.pk])

    def test_a_site_can_be_changed(self):
        """Reallocation — the task the field test could not complete, because
        the cell went read-only as soon as a site was set."""
        self._patch_site("Leicester")
        resp = self._patch_site("McGill")
        self.assertEqual(resp.status_code, 200)
        self.assertEqual([n.site_id for n in self._noms()], [self.other.pk])

    def test_an_unknown_site_name_is_refused_not_invented(self):
        """A typo used to create a real Site, which then showed up in the site
        filter of all four boards for good."""
        before = Site.objects.using(DB).count()
        resp = self._patch_site("Leicster")
        self.assertEqual(resp.status_code, 400)
        body = resp.json()
        self.assertIs(body["ok"], False)
        self.assertIn("no site called 'Leicster'", body["error"])
        # It names the sites that do exist, so the fix is obvious.
        self.assertIn("Leicester", body["error"])
        self.assertEqual(Site.objects.using(DB).count(), before)
        self.assertEqual(self._noms(), [])

    def test_adding_a_target_from_feasibility_nominates_it_at_your_site(self):
        # The press confirms the gene against UniProt now — a gene it cannot
        # confirm is not created, the same rule the bulk door has always held.
        # A test about *nominations* has to say what UniProt answered, or it is
        # really a test about dev's blocked network.
        with _uniprot_says(gene_name="ELP3", protein_name="Elongator 3"):
            resp = self.client.post(
                "/pipeline/feasibility/add/",
                json.dumps({"gene_name": "ELP3", "protein_name": "Elongator 3"}),
                content_type="application/json")
        self.assertEqual(resp.status_code, 200)
        self.assertIs(resp.json()["created"], True)
        target = Target.objects.using(DB).get(gene_name="ELP3")
        noms = list(TargetNomination.objects.using(DB).filter(target_id=target.pk))
        self.assertEqual([n.site_id for n in noms], [self.site.pk])
        # Unfunded is the honest starting state — nobody has said otherwise.
        self.assertFalse(noms[0].funded)


class BoardRowsAreJsonTests(TestCase):
    """Every board row must be JSON, on every board.

    The second field test opened the cell lines board and saw "No cell lines
    match these filters" over 559 intact rows. The rows endpoint was returning
    500, because ``row_for`` handed ``JsonResponse`` a ``CellLine`` object:
    ``arrived_with_ko`` is a ForeignKey to the KO line a wild-type was shipped
    alongside, not a yes/no flag. Only wild-type rows carry it, so filtering to
    knockouts worked and opening the board did not — and one such row took the
    whole response down.

    So the pin is the general rule rather than that one field: no board row may
    contain anything ``json`` cannot serialise. Each board is exercised through
    its real endpoint with a populated row, because the defect only appears once
    the value is not null.
    """

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)

    def _wt_ko_pair(self):
        """The row shape that broke it: a wild-type with no gene, paired to a KO."""
        target = Target.objects.using(DB).create(gene_name="ELP3")
        wt = CellLine.objects.using(DB).create(
            name="HAP1", genotype="WT", site_id=self.site.pk)
        ko = CellLine.objects.using(DB).create(
            name="HAP1 ELP3 KO", genotype="KO", target_id=target.pk,
            parent_line_id=wt.pk, site_id=self.site.pk)
        wt.arrived_with_ko_id = ko.pk
        wt.save(using=DB, update_fields=["arrived_with_ko"])
        return wt, ko

    def test_the_cell_lines_board_loads_with_a_wild_type_line_in_it(self):
        wt, ko = self._wt_ko_pair()
        resp = self.client.get("/pipeline/cell-lines/board/rows/")
        self.assertEqual(resp.status_code, 200,
                         "a wild-type line must not take the board down")
        rows = resp.json()["rows"]
        self.assertEqual({r["name"] for r in rows}, {wt.name, ko.name})
        # And the wild-type row's empty Gene column is correct, not missing data.
        self.assertEqual([r["gene"] for r in rows if r["name"] == wt.name], [""])

    def test_the_unfiltered_cell_lines_board_is_the_default_a_curator_gets(self):
        """No filter, no query string — exactly how the board opens."""
        self._wt_ko_pair()
        resp = self.client.get("/pipeline/cell-lines/board/rows/")
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.json()["count"], 2)

    def test_every_board_row_is_json_serialisable(self):
        self._wt_ko_pair()
        for url in ("/pipeline/cell-lines/board/rows/",
                    "/pipeline/antibodies/board/rows/",
                    "/pipeline/targets/board/rows/",
                    "/pipeline/sessions/board/rows/"):
            with self.subTest(url=url):
                resp = self.client.get(url)
                self.assertEqual(resp.status_code, 200)
                # JsonResponse already serialised it; re-serialising the parsed
                # rows proves nothing snuck through a custom encoder either.
                json.dumps(resp.json()["rows"])

    def test_arrived_with_ko_is_not_offered_as_a_flag_to_toggle(self):
        """It is a ForeignKey. Listed among the boolean fields, the patch endpoint
        would try to store True in it."""
        from pipeline.services import cell_line_board
        self.assertNotIn("arrived_with_ko", cell_line_board.BOOLEAN_FIELDS)
        self.assertNotIn("arrived_with_ko", cell_line_board.EDITABLE_FIELDS)


class CNumberIsANumberTests(TestCase):
    """C-number is an IntegerField behind a text cell.

    Emptying the cell sent "" and int("") raises, so the board answered "could
    not save that — the value has been left unchanged" and a wrong C-number could
    not be cleared at all.
    """

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)
        self.line = CellLine.objects.using(DB).create(
            name="HAP1", genotype="WT", c_number=631, site_id=self.site.pk)

    def _patch(self, value):
        return self.client.post("/pipeline/cell-lines/board/patch/", {
            "target_id": self.line.pk, "field": "c_number", "value": value})

    def _fresh(self):
        return CellLine.objects.using(DB).get(pk=self.line.pk)

    def test_a_c_number_can_be_cleared(self):
        resp = self._patch("")
        self.assertEqual(resp.status_code, 200)
        self.assertIs(resp.json()["ok"], True)
        self.assertIsNone(self._fresh().c_number)

    def test_a_c_number_can_be_changed(self):
        resp = self._patch(" 742 ")
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(self._fresh().c_number, 742)

    def test_a_c_number_that_is_not_a_number_says_so(self):
        """``C742`` used to be refused here and quietly accepted by the paste
        box, which read it as 742. Run 11 showed the disagreement was worse than
        either answer: ``services/c_number.py`` is one reader for both doors, and
        it takes the label as it is written on the tube.

        What must still be refused is a string that merely *contains* digits —
        that is the one that put two knockouts under the same made-up number.
        """
        self.assertEqual(self._patch("C742").status_code, 200)
        self.assertEqual(self._fresh().c_number, 742)

        resp = self._patch("C-RUN11-01")
        self.assertEqual(resp.status_code, 400)
        body = resp.json()
        self.assertIs(body["ok"], False)
        # Not the generic "could not save that", and not "not a number" either:
        # `C-RUN11-01` *is* a number, several of them, which is the whole bug.
        self.assertIn("plain number", body["error"])
        self.assertEqual(self._fresh().c_number, 742)


@skipUnless(shutil.which("node"), "needs node to parse the boards' JavaScript")
class BoardScriptsParseTests(TestCase):
    """Each board's behaviour is written in a <script> block inside its Django
    template, where a stray brace or an unclosed template literal is invisible
    to ``manage.py check`` and to every other test here: the page still renders
    200, and the grid simply never loads. So the scripts are parsed.
    """
    databases = {"academy_db", "pipeline_db"}

    # target_detail is here because it grew OGABoard pop-outs of its own: it is
    # a board surface now, whatever its filename says. The bench-workbook upload
    # is a partial included by two of these, so its script never appears in
    # either file's own source — it has to be named here or it is unchecked.
    # cropper.html does not extend base.html and is not an OGABoard surface, so
    # it sat outside this sweep — and it is the file in the app with the most
    # inline JavaScript by a wide margin. A stray brace there fails exactly the
    # same way: 200, no grid, nothing in any other test to say so.
    # data_io.html is here for the same reason target_detail is: it is several
    # hundred lines of inline JavaScript that builds the whole field picker, and
    # it calls OGABoard for its download receipts, so it is a board surface
    # whatever its filename says.
    # recommendations.html builds every card, every thumbnail and every inline
    # handler in inline JS and does not use OGABoard at all, which is why it was
    # not read as one of these — a stray brace there fails identically: 200, no
    # cards, and nothing anywhere to say so.
    TEMPLATES = ("target_board.html", "antibody_board.html",
                 "cell_line_board.html", "session_board.html",
                 "target_detail.html", "_bench_workbook_upload.html",
                 "user_board.html", "cropper.html", "data_io.html",
                 "recommendations.html", "review_queue.html",
                 # outcomes.html is recommendations.html's shape exactly: every
                 # card, every axis button and every redraw built in inline JS
                 # with no OGABoard, so a stray brace gives 200, no cards, and
                 # nothing anywhere to say so.
                 "outcomes.html")

    def _inline_scripts(self, name):
        path = (Path(settings.BASE_DIR) / "pipeline/templates/pipeline" / name)
        source = path.read_text()
        for m in re.finditer(r"<script(?![^>]*\bsrc=)[^>]*>(.*?)</script>",
                             source, re.S):
            js = re.sub(r"\{%.*?%\}", "", m.group(1), flags=re.S)
            yield re.sub(r"\{\{.*?\}\}", "x", js, flags=re.S)

    def _check(self, js, label):
        with tempfile.NamedTemporaryFile("w", suffix=".js", delete=False) as fh:
            fh.write(js)
            tmp = fh.name
        try:
            done = subprocess.run(["node", "--check", tmp],
                                  capture_output=True, text=True)
        finally:
            os.unlink(tmp)
        self.assertEqual(done.returncode, 0,
                         f"{label} does not parse:\n{done.stderr}")

    def test_every_board_template_script_parses(self):
        for name in self.TEMPLATES:
            for i, js in enumerate(self._inline_scripts(name)):
                with self.subTest(template=name, script=i):
                    self._check(js, f"{name} script {i}")

    def test_board_js_parses(self):
        path = Path(settings.BASE_DIR) / "pipeline/static/pipeline/board.js"
        self._check(path.read_text(), "board.js")

    def test_no_page_declares_one_function_name_twice(self):
        """Two `function foo(){}` in one script is valid JS and a live defect.

        The later declaration wins for the whole scope, silently, so the earlier
        one is never called no matter where it sits in the file. `data_io.html`
        had two functions called `chip` — one tagging a field in the download
        picker, one drawing a count in the upload preview — and every field in
        the picker was therefore labelled by the wrong one: `id` read
        `key: undefined`, and the `+ db` tag explained in the legend underneath
        was drawn on nothing at all.

        `manage.py check` passes, `node --check` passes and the page renders
        200 — the same shape as `setCommitLabel` in board.js, which shipped
        declared in one function and called from another. This is the cheapest
        thing that can fail on it: a browser is not needed to read two `function`
        keywords with one name between them.
        """
        for name in self.TEMPLATES:
            for i, js in enumerate(self._inline_scripts(name)):
                names = re.findall(r"^\s*function\s+([A-Za-z_$][\w$]*)\s*\(",
                                   js, re.M)
                dupes = sorted({n for n in names if names.count(n) > 1})
                with self.subTest(template=name, script=i):
                    self.assertEqual(
                        dupes, [],
                        f"{name} script {i} declares {dupes} more than once — "
                        "the last declaration wins and the earlier one is dead.")

    def test_a_blank_first_column_survives_the_paste_box(self):
        """``pastedBlock`` keeps a leading tab; ``trim()`` ate it on line 1 only.

        A gene's page tells you to leave the gene column blank, so every line of
        every paste made there begins with an empty first column — and
        ``textarea.value.trim()`` strips leading whitespace from the *string*,
        which is the first line's first tab and nothing else. Row 1 slid one
        column left, rows 2 and 3 were perfect, and the field test that found it
        only noticed because the shift pushed a concentration into ``site``.

        Run rather than grepped: this loads board.js in node and calls the
        function. It is still the weaker half — it says the helper is right, not
        that ``tsv()`` calls it — which is what
        ``tests_browser_board.py::test_a_paste_whose_first_column_is_blank_
        keeps_its_columns`` is for, and why that one is worth a browser.
        """
        path = Path(settings.BASE_DIR) / "pipeline/static/pipeline/board.js"
        script = (f"global.window = {{}};\n"
                  f"require({str(path)!r});\n"
                  "const pb = window.OGABoard.pastedBlock;\n"
                  "console.log(JSON.stringify({\n"
                  "  blankFirst: pb('\\tab1\\tAbcam\\n\\tab2\\tAbcam'),\n"
                  "  padded: pb('\\n\\n\\tX\\tY\\n\\n'),\n"
                  "  allTabRows: pb('\\t\\t\\n\\tX\\tY\\n\\t\\t'),\n"
                  "  empty: pb(''),\n"
                  "}));\n")
        with tempfile.NamedTemporaryFile("w", suffix=".js", delete=False) as fh:
            fh.write(script)
            tmp = fh.name
        try:
            done = subprocess.run(["node", tmp], capture_output=True, text=True)
        finally:
            os.unlink(tmp)
        self.assertEqual(done.returncode, 0, done.stderr)
        got = json.loads(done.stdout)

        # The bug, exactly: both rows keep the empty first column.
        self.assertEqual(got["blankFirst"], "\tab1\tAbcam\n\tab2\tAbcam")
        # And what the trim was there for still happens — blank lines off the
        # ends, including a line that is only tabs, which is a row with nothing
        # in any column rather than a row whose first column is blank.
        self.assertEqual(got["padded"], "\tX\tY")
        self.assertEqual(got["allTabRows"], "\tX\tY")
        self.assertEqual(got["empty"], "")


class ScreensAgreeWithEachOtherTests(TestCase):
    """The same fact must not be shown two different ways one screen apart."""

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.other = Site.objects.using(DB).create(name="McGill", short_code="MCG")
        self.client = _member_client(self, self.site)
        self.target = Target.objects.using(DB).create(gene_name="STMN2")
        TargetNomination.objects.using(DB).create(
            target_id=self.target.pk, site_id=self.site.pk)

    def test_the_target_page_and_the_board_name_the_same_site(self):
        """The board read the nomination; the page read Target.site, a
        denormalised leftover of the Access import that nothing writes. So the
        board said Leicester and the page said "Site —" for one target."""
        rows = self.client.get("/pipeline/targets/board/rows/").json()["rows"]
        on_board = next(r for r in rows if r["gene"] == "STMN2")["sites"]
        self.assertEqual(on_board, ["Leicester"])

        page = self.client.get(f"/pipeline/target/{self.target.pk}/")
        self.assertEqual(page.status_code, 200)
        body = page.content.decode()
        self.assertEqual(page.context["nominated_sites"], on_board)
        self.assertIn("Leicester", body)

    def test_a_target_nobody_has_nominated_says_so_rather_than_dash(self):
        """…and now offers the empty cells that would record one.

        The page used to state "not nominated by any site yet" and stop, which
        was honest and was also the case where "there is nowhere to record it"
        was most true. The Funding panel draws a blank nomination block for it,
        and `has_nomination` is what tells the renderer to say nothing has been
        recorded rather than to draw a row of dashes.
        """
        bare = Target.objects.using(DB).create(gene_name="ELP3")
        page = self.client.get(f"/pipeline/target/{bare.pk}/")
        self.assertEqual(page.status_code, 200)
        self.assertEqual(page.context["nominated_sites"], [])
        row = page.context["board_row"]
        self.assertEqual(row["nominations"], [])
        self.assertIs(row["has_nomination"], False)

    def test_the_back_link_goes_where_it_says_it_goes(self):
        """"Back to dashboard" pointed at pipeline:dashboard, which redirects to
        the target board — the app's own link did not describe where it went."""
        page = self.client.get(f"/pipeline/target/{self.target.pk}/")
        body = page.content.decode()
        self.assertNotIn("Back to dashboard", body)
        self.assertIn("Back to the target board", body)


class AnAliasHasAHomeOnScreenTests(TestCase):
    """A gene's other names live in two columns and only one was ever drawn.

    The Downloads & uploads round trip writes `Target.aliases` — that is the
    route the page invites, "fill in missing values across many genes". A field
    test did exactly that: the preview named the record, the field, the old
    value and the new one, the save reported it, and then the gene page went on
    reading "Also known as: ANKTM1" from `alternative_name` alone. The value was
    real (the search box found it and said which field answered) and there was
    nowhere on any screen to see it, which from outside is indistinguishable
    from the upload having quietly done nothing.
    """

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)
        self.target = Target.objects.using(DB).create(
            gene_name="TRPA1", protein_name="Transient receptor potential A1",
            alternative_name="ANKTM1", aliases="COWORK RUN15 alias probe")

    def test_the_gene_page_shows_an_alias_added_through_the_spreadsheet(self):
        page = self.client.get(f"/pipeline/target/{self.target.pk}/")
        self.assertEqual(page.status_code, 200)
        body = page.content.decode()
        self.assertIn("COWORK RUN15 alias probe", body)
        # …beside the UniProt synonym, not instead of it.
        self.assertIn("ANKTM1", body)

    def test_the_targets_board_shows_it_too(self):
        """Fifty genes is what the spreadsheet route is for, and the board is the
        only screen that shows fifty at once. `row_for` has carried
        `alternative` since it was written and drew it nowhere."""
        rows = self.client.get("/pipeline/targets/board/rows/").json()["rows"]
        row = next(r for r in rows if r["gene"] == "TRPA1")
        self.assertEqual(row["other_names"], ["ANKTM1", "COWORK RUN15 alias probe"])

    def test_the_gene_page_has_a_slot_for_it_when_it_is_empty(self):
        """An empty field with no label reads as "this page does not hold that".
        The subtitle disappears when there is nothing; the labelled row does not.
        """
        bare = Target.objects.using(DB).create(gene_name="ELP3")
        page = self.client.get(f"/pipeline/target/{bare.pk}/")
        self.assertEqual(page.context["other_names"], [])
        self.assertIn("Also Known As", page.content.decode())


class UploadDefaultsToYourOwnSiteTests(TestCase):
    """A Leicester curator uploading a sheet with no site column nominated every
    row to McGill, because the box was hard-coded to it and it is the default."""

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        self.leicester = Site.objects.using(DB).create(name="Leicester",
                                                       short_code="LEI")
        self.mcgill = Site.objects.using(DB).create(name="McGill", short_code="MCG")
        self.client = _member_client(self, self.leicester)

    def test_the_default_site_is_the_signed_in_users_own(self):
        resp = self.client.get("/pipeline/targets/board/")
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.context["my_site"], "Leicester")
        body = resp.content.decode()
        # McGill is still *offered* — it is a real site. It is no longer the
        # value the box arrives holding.
        self.assertIn('<option value="Leicester" selected>', body)
        self.assertNotIn('<option value="McGill" selected>', body)

    def test_the_site_is_chosen_from_the_list_not_typed(self):
        """Free text accepted "Atlantis Institute" without comment, though the
        filter directly above it is a five-item dropdown of the real sites."""
        body = self.client.get("/pipeline/targets/board/").content.decode()
        self.assertIn('<select id="default-site"', body)
        self.assertNotIn('<input type="text" id="default-site"', body)
        for site in ("Leicester", "McGill"):
            self.assertIn(f'<option value="{site}"', body)


class PreviewNamesTheVialTests(TestCase):
    """Two rows differing only in lot both read "10586-1-AP · STMN2" — one "new",
    one "will be updated" — and the lot, the only reason the outcomes differ, was
    not shown anywhere. The preview was truthful and unreadable at once."""

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)
        self.target = Target.objects.using(DB).create(gene_name="STMN2")

    def _parse(self, text):
        resp = self.client.post("/pipeline/antibodies/bulk/parse/",
                                json.dumps({"text": text}),
                                content_type="application/json")
        self.assertEqual(resp.status_code, 200)
        return resp.json()["items"]

    def test_every_previewed_row_carries_its_whole_identity(self):
        items = self._parse(
            "gene\tcatalogue\tcompany\tlot\n"
            "STMN2\t10586-1-AP\tProteintech\t20051\n"
            "STMN2\t10586-1-AP\tProteintech\t99999-NEWLOT\n")
        self.assertEqual(len(items), 2)
        self.assertEqual([i["lot"] for i in items], ["20051", "99999-NEWLOT"])
        for i in items:
            self.assertEqual(i["company"], "Proteintech")
            self.assertEqual(i["site"], "Leicester")

    def test_the_row_that_will_be_updated_says_which_vial_it_matched(self):
        from pipeline.models import Antibody
        from pipeline.services.cropper import db as cdb
        company = cdb.resolve_company("Proteintech", "10586-1-AP", create=True)
        Antibody.objects.using(DB).create(
            target_id=self.target.pk, company_id=company.pk,
            catalogue_number="10586-1-AP", lot_number="20051",
            site_id=self.site.pk)
        items = self._parse(
            "gene\tcatalogue\tcompany\tlot\n"
            "STMN2\t10586-1-AP\tProteintech\t20051\n"
            "STMN2\t10586-1-AP\tProteintech\t99999-NEWLOT\n")
        by_lot = {i["lot"]: i for i in items}
        self.assertEqual(by_lot["20051"]["status"], "update")
        self.assertIn("20051", by_lot["20051"]["note"])
        self.assertEqual(by_lot["99999-NEWLOT"]["status"], "create")
        self.assertEqual(by_lot["99999-NEWLOT"]["note"], "")


class ControlsSayWhatTheyChangeTests(TestCase):
    """The only button in a target row was called "Click to change" and nothing
    else. The field test clicked it expecting the site cell and silently flipped a
    gene's funding on the shared master list."""

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)

    def _rendered(self, url):
        resp = self.client.get(url)
        self.assertEqual(resp.status_code, 200)
        return resp.content.decode()

    def test_no_control_on_a_board_is_named_only_click_to_change(self):
        for url in ("/pipeline/targets/board/", "/pipeline/antibodies/board/",
                    "/pipeline/cell-lines/board/"):
            with self.subTest(url=url):
                body = self._rendered(url)
                self.assertNotIn('title="Click to change"', body)
                # Every toggle states the field and the value it holds.
                self.assertIn("aria-pressed=", body)

    def test_the_funded_toggle_names_the_field_and_the_next_value(self):
        body = self._rendered("/pipeline/targets/board/")
        self.assertIn("Click to mark it ", body)
        self.assertIn("aria-label=\"Funded:", body)

    def test_the_ko_validated_reason_cell_says_it_is_for_a_reason(self):
        """Column header "KO VALIDATED", a greyed word "validated" for the tick,
        and an unlabelled box under it. Nothing said that box takes the reason."""
        body = self._rendered("/pipeline/cell-lines/board/")
        self.assertIn("why not? add a reason", body)
        self.assertIn("how was it confirmed?", body)

    def test_a_wild_type_line_says_its_empty_gene_is_deliberate(self):
        body = self._rendered("/pipeline/cell-lines/board/")
        self.assertIn("no gene — wild type", body)
        self.assertIn("recorded once, with no gene attached", body)

    def test_the_add_panels_state_the_rule_a_guide_reader_would_not_need(self):
        antibodies = self._rendered("/pipeline/antibodies/board/")
        self.assertIn("different lot", antibodies)
        self.assertIn("second row", antibodies)
        cell_lines = self._rendered("/pipeline/cell-lines/board/")
        self.assertIn("once, with no gene", cell_lines)


class AddingATargetSaysWhatElseItWroteTests(TestCase):
    """One click on "Add to Pipeline" writes the target, a nomination at your own
    site, and "not funded". It reported the first of the three, so the other two
    were things you found out from the board row afterwards — and the report
    listed them among the rules a person who never opens a guide gets wrong."""

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)

    def test_the_confirmation_names_the_site_and_the_funding_state(self):
        with _uniprot_says(gene_name="STMN2", protein_name="Stathmin 2"):
            resp = self.client.post(
                "/pipeline/feasibility/add/",
                json.dumps({"gene_name": "STMN2", "protein_name": "Stathmin 2"}),
                content_type="application/json")
        self.assertEqual(resp.status_code, 200)
        body = resp.json()
        self.assertIs(body["created"], True)
        self.assertEqual(body["site"], "Leicester")
        self.assertIn("Leicester", body["nominated_note"])
        self.assertIn("not funded", body["nominated_note"])


class MultiSiteNominationsAreVisibleWhereTheBoardSendsYouTests(TestCase):
    """The board's site cell goes read-only for a gene nominated more than once
    and used to say "edit on the target page" — a page that had no editor and did
    not even list the nominations. The claim is now what is actually there."""

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        self.leicester = Site.objects.using(DB).create(name="Leicester",
                                                       short_code="LEI")
        self.mcgill = Site.objects.using(DB).create(name="McGill", short_code="MCG")
        self.client = _member_client(self, self.leicester)
        self.target = Target.objects.using(DB).create(gene_name="STMN2")
        for site in (self.leicester, self.mcgill):
            TargetNomination.objects.using(DB).create(
                target_id=self.target.pk, site_id=site.pk)

    def test_the_board_no_longer_promises_an_editor_that_is_not_there(self):
        body = self.client.get("/pipeline/targets/board/").content.decode()
        self.assertNotIn("edit on the target page", body)
        self.assertIn("see them on the target page", body)

    def test_the_target_page_lists_every_nomination(self):
        resp = self.client.get(f"/pipeline/target/{self.target.pk}/")
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(len(resp.context["nominations"]), 2)
        # …and hands the page each of them as its own editable record. The panel
        # is drawn from `board_row` rather than rendered server-side, so the
        # site names alone prove nothing — every active site is in the page as a
        # dropdown option. What matters is that both nominations arrive with an
        # id, which is what a cell edit posts back to write the right one.
        noms = resp.context["board_row"]["nominations"]
        self.assertEqual(sorted(n["site"] for n in noms), ["Leicester", "McGill"])
        self.assertTrue(all(n["id"] for n in noms))

    def test_both_sites_are_named_in_the_summary(self):
        resp = self.client.get(f"/pipeline/target/{self.target.pk}/")
        self.assertEqual(resp.context["nominated_sites"], ["Leicester", "McGill"])


class OverviewCountsNominationsTests(TestCase):
    """Overview is linked from the hub and the nav, and it read Target.site too.

    Two consequences: the SITE column printed "—" for every target added since
    nominations became the record, and the cross-site summary counted a site's
    targets only once that site had entered an antibody — so a site that had
    started work but not yet bought reagents showed as doing nothing.
    """

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        self.site = Site.objects.using(DB).create(
            name="Leicester", short_code="LEI", is_active=True)
        self.client = _member_client(self, self.site)
        self.target = Target.objects.using(DB).create(
            gene_name="STMN2", status="in_progress")
        TargetNomination.objects.using(DB).create(
            target_id=self.target.pk, site_id=self.site.pk)

    def test_the_site_column_names_the_nominating_site(self):
        resp = self.client.get("/pipeline/overview/")
        self.assertEqual(resp.status_code, 200)
        active = {t.gene_name: t.site_codes for t in resp.context["active_targets"]}
        self.assertEqual(active["STMN2"], "LEI")

    def test_a_site_with_a_nomination_and_no_antibodies_still_counts(self):
        resp = self.client.get("/pipeline/overview/")
        summaries = {s["site"].name: s for s in resp.context["site_summaries"]}
        self.assertIn("Leicester", summaries)
        self.assertEqual(summaries["Leicester"]["active_targets"], 1)
        self.assertEqual(summaries["Leicester"]["antibodies"], 0)


# ══════════════════════════════════════════════════════════════════════════════
# Third field test, 31 Jul 2026. Two genes onboarded from nothing on the live
# site; four real defects, one of which stopped the job.
# ══════════════════════════════════════════════════════════════════════════════

class TwoPanelsOnOnePageDoNotShareTheirControlsTests(TestCase):
    """The blocker. A gene's page carries the Add antibodies pop-out *and* the Add
    cell lines pop-out, and ``OGABoard.newEntry`` gave every element it built a
    fixed id — ``ne-table``, ``ne-add-row``, ``ne-paste``. Two panels meant two
    elements per id, ``getElementById`` returned the first, and so every button on
    the cell-lines panel drove the antibodies panel. From the page: + Add row did
    nothing, the Paste tab would not open, Check these produced no output. On every
    gene page, for cell lines, with no way to tell why.

    Pinned at the source, because the defect only exists once two panels are on
    one page and no server response shows it.
    """
    databases = {"academy_db", "pipeline_db"}

    def test_new_entry_namespaces_the_elements_it_builds(self):
        js = (Path(settings.BASE_DIR) / "pipeline/static/pipeline/board.js").read_text()
        body = js[js.index("function newEntry("):js.index("function uploadPanel(")]
        self.assertIn("panelSeq", body,
                      "newEntry must number its panels so two can coexist")
        self.assertNotIn("'ne-", body, "a hardcoded element id is back in newEntry")
        self.assertNotIn('"ne-', body, "a hardcoded element id is back in newEntry")

    def test_every_pop_out_builder_namespaces_itself(self):
        """uploadPanel and identityDialog have the same exposure — a board can
        legitimately hold more than one of each."""
        js = (Path(settings.BASE_DIR) / "pipeline/static/pipeline/board.js").read_text()
        for fn in ("function newEntry(", "function uploadPanel(",
                   "function identityDialog("):
            with self.subTest(fn=fn):
                start = js.index(fn)
                self.assertIn("++panelSeq", js[start:start + 900],
                              f"{fn} does not take a namespace of its own")

    def test_the_gene_page_really_does_mount_two_panels(self):
        """The condition that made it fail. If this ever stops being true the test
        above is guarding nothing."""
        html = (Path(settings.BASE_DIR)
                / "pipeline/templates/pipeline/target_detail.html").read_text()
        self.assertIn("ab-add-panel", html)
        self.assertIn("cl-add-panel", html)


class AnExactSupplierNameBeatsTheGuessTests(TestCase):
    """A paste of "Bio-Techne (Novus Biologicals)" previewed as typed and saved as
    "Bio-Techne (R&D Systems)".

    Both are real, separate suppliers that display as Bio-Techne, and the resolver
    picks between them by catalogue prefix — deliberately, because a pasted vendor
    of just "Bio-Techne" is ambiguous and the catalogue is the only thing that
    disambiguates it. The rule was applied to unambiguous input too, so somebody
    who named the brand in full had it overruled. A preview showing one supplier
    and the save writing another is the one thing a preview must never do.
    """

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        from pipeline.models import Company
        self.novus = Company.objects.using(DB).create(
            name="Bio-Techne (Novus Biologicals)")
        self.rnd = Company.objects.using(DB).create(
            name="Bio-Techne (R&D Systems)")

    def _resolve(self, vendor, catalogue):
        from pipeline.services.cropper import db as cdb
        return cdb.resolve_company(vendor, catalogue, create=False, db=DB)

    def test_the_full_brand_name_is_taken_at_its_word(self):
        got = self._resolve("Bio-Techne (Novus Biologicals)", "AF3130")
        self.assertEqual(got.pk, self.novus.pk)

    def test_the_other_full_brand_name_too(self):
        got = self._resolve("Bio-Techne (R&D Systems)", "NB100-1234")
        self.assertEqual(got.pk, self.rnd.pk)

    def test_a_bare_bio_techne_is_still_split_by_catalogue(self):
        """The rule that exists for a reason keeps working."""
        self.assertEqual(self._resolve("Bio-Techne", "NB100-1234").pk, self.novus.pk)
        self.assertEqual(self._resolve("Bio-Techne", "AF3130").pk, self.rnd.pk)

    def test_the_preview_shows_the_supplier_that_will_be_saved(self):
        """Not the one that was typed. When the resolver substitutes, the paste
        preview has to say so before anything is written.

        This lived in the antibodies board's template until run 4 found the gene
        page still echoing the typed name — so it is `OGABoard.supplierLabel` now
        and this checks the shared one. See
        ``OnePlaceDecidesWhatASupplierWillSayTests``.
        """
        js = (Path(settings.BASE_DIR)
              / "pipeline/static/pipeline/board.js").read_text()
        self.assertIn("company_resolved", js)
        self.assertIn("you typed", js)


class TheBenchSheetIsTheSessionsOwnAntibodiesTests(TestCase):
    """A bench sheet listed every vial recorded for the gene, not the ones in the
    session: a 3-antibody WB session printed 4 rows, a 2-antibody IP session
    printed 4, a 1-antibody FC session printed 3. A row on a bench sheet is an
    invitation to write a reading on it, so an extra row is a result for an
    experiment nobody ran.
    """

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        from pipeline.models import (Antibody, Company, ExperimentSession,
                                     Member, WbResult)
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)
        self.target = Target.objects.using(DB).create(gene_name="STMN2")
        company = Company.objects.using(DB).create(name="Proteintech")
        self.abs = [
            Antibody.objects.using(DB).create(
                target_id=self.target.pk, company_id=company.pk,
                catalogue_number=f"CAT-{i}", site_id=self.site.pk)
            for i in range(4)
        ]
        member = Member.objects.using(DB).get(site_id=self.site.pk)
        self.session = ExperimentSession.objects.using(DB).create(
            target_id=self.target.pk, procedure_type="WB", site_id=self.site.pk,
            experimenter_id=member.pk, date=date(2026, 7, 31))
        self.member = member
        # Two of the four are actually in the session.
        for ab in self.abs[:2]:
            WbResult.objects.using(DB).create(
                session_id=self.session.pk, antibody_id=ab.pk)

    def _rows(self):
        from pipeline.services import planning
        return [ab.catalogue_number for ab in planning._session_antibodies(self.session)]

    def test_the_sheet_lists_only_the_session_s_antibodies(self):
        self.assertEqual(sorted(self._rows()), ["CAT-0", "CAT-1"])

    def test_a_session_with_no_results_yet_falls_back_and_says_so(self):
        """A sheet has to have rows to be worth printing, so an empty session
        still gets the gene's vials — labelled as a picking list, not as the
        session's own antibodies."""
        from pipeline.models import ExperimentSession
        from pipeline.services import planning
        empty = ExperimentSession.objects.using(DB).create(
            target_id=self.target.pk, procedure_type="IP", site_id=self.site.pk,
            experimenter_id=self.member.pk, date=date(2026, 7, 31))
        self.assertEqual(len(planning._session_antibodies(empty)), 4)
        self.assertIn("no antibodies recorded",
                      planning._session_info_line(empty))

    def test_a_populated_session_is_not_labelled_as_a_picking_list(self):
        from pipeline.services import planning
        self.assertNotIn("no antibodies recorded",
                         planning._session_info_line(self.session))

    def test_the_downloaded_workbook_has_one_row_per_session_antibody(self):
        """End to end, because the row count is what the tester counted."""
        from pipeline.services import planning
        ws = planning.generate_bench_sheet(self.session).active
        # Title, info line, headers, then the data rows.
        cats = [ws.cell(row=r, column=4).value for r in range(4, ws.max_row + 1)]
        self.assertEqual(sorted(c for c in cats if c), ["CAT-0", "CAT-1"])


class ARefusedSiteNamesTheOnesOnFileTests(TestCase):
    """"no site called 'Leicster'" tells you it is wrong and not what is right, and
    the only place to find the spelling of your own institution was the filter
    dropdown on a different part of the page. The target board already listed
    them; the other three did not."""

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        from pipeline.models import (Antibody, Company, ExperimentSession,
                                     Member)
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.mcgill = Site.objects.using(DB).create(name="McGill", short_code="MCG")
        self.client = _member_client(self, self.site)
        self.target = Target.objects.using(DB).create(gene_name="STMN2")
        company = Company.objects.using(DB).create(name="Proteintech")
        self.antibody = Antibody.objects.using(DB).create(
            target_id=self.target.pk, company_id=company.pk,
            catalogue_number="CAT-1", site_id=self.site.pk)
        self.line = CellLine.objects.using(DB).create(
            name="HAP1 STMN2 KO", genotype="KO", target_id=self.target.pk,
            site_id=self.site.pk)
        member = Member.objects.using(DB).get(site_id=self.site.pk)
        self.session = ExperimentSession.objects.using(DB).create(
            target_id=self.target.pk, procedure_type="WB", site_id=self.site.pk,
            experimenter_id=member.pk, date=date(2026, 7, 31))

    def _refusal(self, url, pk, key="target_id"):
        resp = self.client.post(url, {key: pk, "field": "site",
                                      "value": "Leicster"})
        self.assertEqual(resp.status_code, 400, url)
        return resp.json()["error"]

    def test_every_board_lists_the_sites_on_file(self):
        cases = [
            ("/pipeline/antibodies/board/patch/", self.antibody.pk, "target_id"),
            ("/pipeline/cell-lines/board/patch/", self.line.pk, "target_id"),
            ("/pipeline/sessions/board/patch/", self.session.pk, "session_id"),
        ]
        for url, pk, key in cases:
            with self.subTest(url=url):
                msg = self._refusal(url, pk, key)
                self.assertIn("Leicster", msg)
                self.assertIn("Leicester", msg)
                self.assertIn("McGill", msg)

    def test_a_refused_site_is_never_created(self):
        before = Site.objects.using(DB).count()
        self._refusal("/pipeline/antibodies/board/patch/", self.antibody.pk)
        self.assertEqual(Site.objects.using(DB).count(), before)

    def test_a_short_code_is_accepted_as_well_as_the_name(self):
        resp = self.client.post("/pipeline/antibodies/board/patch/",
                                {"target_id": self.antibody.pk, "field": "site",
                                 "value": "MCG"})
        self.assertEqual(resp.status_code, 200)
        self.antibody.refresh_from_db()
        self.assertEqual(self.antibody.site_id, self.mcgill.pk)


class CountsAndVerbsAgreeTests(TestCase):
    """"1 recorded result already point at this row" and "1 knockout were made
    from it". Both sit next to a number, which is where a grammar slip costs the
    most: it makes a careful reader distrust the count."""
    databases = {"academy_db", "pipeline_db"}

    def _template(self, name):
        return (Path(settings.BASE_DIR) / "pipeline/templates/pipeline" / name).read_text()

    def test_the_antibody_dialog_agrees_with_its_count(self):
        html = self._template("antibody_board.html")
        self.assertNotIn("already point at this row", html)
        self.assertIn("'point' : 'points'", html)

    def test_the_cell_line_dialog_agrees_with_its_count(self):
        html = self._template("cell_line_board.html")
        self.assertNotIn("were made from it`", html)
        self.assertIn("' was' : 's were'", html)


class TheBoardsPointAtPagesThatExistTests(TestCase):
    """A footnote on the cell lines board read "Supplier goes through
    ``Company.resolve``, so it is edited on the cell line page" — an internal
    function name a scientist cannot act on, and a page retired on 31 Jul 2026.
    Between them they made an editable field look uneditable.
    """

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)

    BOARDS = ("/pipeline/antibodies/board/", "/pipeline/cell-lines/board/",
              "/pipeline/sessions/board/", "/pipeline/targets/board/")

    def _prose(self, url):
        """The page with its scripts removed — what a reader actually reads.

        The scripts are excluded deliberately: a board's behaviour lives in an
        inline ``<script>``, and the notes in it explain *why* the code is the way
        it is, naming the functions and the retired pages that made it so. Those
        are for whoever edits the file next. What must not name them is the prose.
        """
        body = self.client.get(url).content.decode()
        return re.sub(r"<script.*?</script>", "", body, flags=re.S)

    def test_no_board_names_an_internal_function(self):
        for url in self.BOARDS:
            with self.subTest(url=url):
                prose = self._prose(url)
                for leak in ("Company.resolve", "resolve_or_create_target",
                             "find_vial", "row_for", "unique_antibody_per_site_lot"):
                    self.assertNotIn(leak, prose, f"{leak} is on screen at {url}")

    def test_no_board_sends_you_to_a_retired_page(self):
        for url in self.BOARDS:
            with self.subTest(url=url):
                prose = self._prose(url)
                for dead in ("the cell line page", "the antibody page",
                             "the session page"):
                    self.assertNotIn(dead, prose,
                                     f"{url} points at a retired page")

    def test_a_cell_line_s_supplier_can_actually_be_changed(self):
        """The footnote's claim has to be true. Supplier is identity — it goes
        through ``Company.resolve`` — so it belongs in the dialog, and until now
        it was in neither the grid nor the dialog, which is to say nowhere."""
        from pipeline.models import Company
        from pipeline.services import identity
        line = CellLine.objects.using(DB).create(
            name="HAP1", genotype="WT", site_id=self.site.pk)
        Company.objects.using(DB).create(name="Horizon Discovery")
        identity.change_cell_line_identity(line, {
            "name": "HAP1", "genotype": "WT", "gene": "", "parent": "",
            "company": "horizon discovery"})
        line.refresh_from_db()
        self.assertEqual(line.company.name, "Horizon Discovery")
        self.assertIn("company", identity.cell_line_identity(line))


class TheFeasibilityVerdictIsInWordsTests(TestCase):
    """The headline was a coloured circle whose amber label was "??" — which is
    what a page prints when it has failed. Amber is also the commonest verdict, so
    the summary of an otherwise excellent report looked broken for every gene
    tried, including the page's own suggested examples."""
    databases = {"academy_db", "pipeline_db"}

    def test_the_amber_verdict_is_not_two_question_marks(self):
        html = (Path(settings.BASE_DIR)
                / "pipeline/templates/pipeline/feasibility.html").read_text()
        self.assertNotIn("label: '??'", html)

    def test_every_signal_carries_a_sentence(self):
        html = (Path(settings.BASE_DIR)
                / "pipeline/templates/pipeline/feasibility.html").read_text()
        self.assertIn('id="summary-verdict"', html)
        self.assertEqual(html.count("verdict: '"), 4,
                         "each of green/amber/red/unknown needs its own wording")


# ══════════════════════════════════════════════════════════════════════════════
# Fourth field test, 31 Jul 2026. Eight of nine claims held; the bench-sheet
# round trip worked end to end for the first time. These are what did not.
# ══════════════════════════════════════════════════════════════════════════════

class AWildTypeNeverTakesThePagesGeneTests(TestCase):
    """The run's one HIGH, and the app disagreeing with itself in three places.

    A gene's page pastes with ``default_gene`` set to that gene, because every
    row on that page is about it — except the wild-type parent, which is the one
    row that must have no gene at all. One HAP1 WT serves every knockout made
    from it, so a gene on it forks a second parental line the next time somebody
    works on a different gene, and the two are not obviously duplicates.

    The panel says so above the grid. The identity endpoint refuses the same state
    with a 400. The bulk write path did it anyway: the tester left the gene column
    blank, exactly as instructed, and got ``SH-SY5Y WT`` with ``gene = STMN2``.
    """

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)
        Target.objects.using(DB).create(gene_name="STMN2")

    HEADER = "name\tgene\tgenotype\tparent"

    def _parse(self, *rows, default_gene="STMN2"):
        from pipeline.services import bulk_cell_lines as bulk
        return bulk.parse("\n".join((self.HEADER,) + rows), default_gene=default_gene)

    def test_the_default_gene_skips_a_wild_type_row(self):
        rows = self._parse("SH-SY5Y WT\t\tWT\t")
        self.assertEqual(rows[0]["gene"], "")

    def test_the_default_gene_still_reaches_a_knockout_row(self):
        rows = self._parse("SH-SY5Y STMN2 KO\t\tKO\tSH-SY5Y WT")
        self.assertEqual(rows[0]["gene"], "STMN2")

    def test_a_row_inferred_as_wild_type_is_skipped_too(self):
        """Genotype is often left blank and inferred. A row with no genotype, no
        parent and no KO in its name is a wild type, and must be treated as one
        *before* the gene is defaulted onto it."""
        rows = self._parse("SH-SY5Y\t\t\t")
        self.assertEqual(rows[0]["gene"], "")

    def test_a_row_inferred_as_a_knockout_still_gets_the_gene(self):
        rows = self._parse("SH-SY5Y STMN2 KO\t\t\t")
        self.assertEqual(rows[0]["gene"], "STMN2")

    def test_the_gene_pages_paste_creates_a_wild_type_with_no_gene(self):
        """End to end through the endpoint the gene page posts to, because that is
        where it went wrong — the service was never the thing under test."""
        from pipeline.models import CellLine
        payload = {"text": f"{self.HEADER}\nSH-SY5Y WT\t\tWT\t",
                   "default_gene": "STMN2", "dry_run": False}
        resp = self.client.post("/pipeline/cell-lines/bulk/commit/",
                                json.dumps(payload), content_type="application/json")
        self.assertEqual(resp.status_code, 200)
        line = CellLine.objects.using(DB).get(name="SH-SY5Y WT")
        self.assertIsNone(line.target_id, "a wild type was stamped with the page's gene")
        self.assertEqual(line.genotype, "WT")

    def test_the_preview_shows_the_wild_type_with_no_gene(self):
        """The preview is what a person checks. It said `· WT · STMN2`."""
        payload = {"text": f"{self.HEADER}\nSH-SY5Y WT\t\tWT\t",
                   "default_gene": "STMN2"}
        resp = self.client.post("/pipeline/cell-lines/bulk/parse/",
                                json.dumps(payload), content_type="application/json")
        self.assertEqual(resp.json()["items"][0]["gene"], "")

    def test_typing_a_gene_onto_a_new_wild_type_is_refused_in_the_same_words(self):
        """A rule the identity dialog enforces and the paste box does not is a
        rule nobody can rely on."""
        from pipeline.services import bulk_cell_lines as bulk
        rows = bulk.parse(f"{self.HEADER}\nSH-SY5Y WT\tSTMN2\tWT\t")
        item = bulk.plan(rows)[0]
        self.assertEqual(item["status"], "blocked")
        self.assertIn("recorded once, with no gene", item["note"])

    def test_a_wild_type_already_on_file_with_a_gene_can_still_be_updated(self):
        """Legacy rows exist. Re-uploading a downloaded sheet of them must not be
        refused — the refusal is about *creating* a second wrong line."""
        from pipeline.models import CellLine
        from pipeline.services import bulk_cell_lines as bulk
        target = Target.objects.using(DB).get(gene_name="STMN2")
        CellLine.objects.using(DB).create(name="OLD WT", genotype="WT",
                                          target_id=target.pk, site_id=self.site.pk)
        rows = bulk.parse(f"{self.HEADER}\nOLD WT\tSTMN2\tWT\t")
        self.assertEqual(bulk.plan(rows)[0]["status"], "update")


class ACellLinePreviewSaysWhatItMatchedTests(TestCase):
    """`HAP1 WT [COWORK RUN4]` — a name nothing on file had — previewed as
    "already on file — will be updated", with nothing to say what it had matched
    or whether saving would rename somebody else's row. The matcher is right (it
    matched on supplier and catalogue, which is how a second parental line gets
    prevented); the preview was silent about it. The antibody preview has named
    the vial it matched since run 2."""

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        from pipeline.models import Company
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)
        self.horizon = Company.objects.using(DB).create(name="Horizon Discovery")
        self.line = CellLine.objects.using(DB).create(
            name="HAP1", genotype="WT", company_id=self.horizon.pk,
            catalogue_number="C631", cellosaurus_id="CVCL_0030",
            site_id=self.site.pk)

    def _plan(self, row):
        from pipeline.services import bulk_cell_lines as bulk
        header = "name\tgenotype\tsupplier\tcatalogue\tcellosaurus"
        return bulk.plan(bulk.parse(f"{header}\n{row}"), member=None)[0]

    def test_a_match_on_supplier_and_catalogue_says_so(self):
        item = self._plan("HAP1 WT [COWORK RUN4]\tWT\tHorizon Discovery\tC631\t")
        self.assertEqual(item["status"], "update")
        self.assertIn("HAP1", item["note"])
        self.assertIn("C631", item["note"])

    def test_it_says_the_name_on_file_is_kept(self):
        """The tester's open question, and the answer is that an update never
        renames — only a create sets the name."""
        item = self._plan("HAP1 WT [COWORK RUN4]\tWT\tHorizon Discovery\tC631\t")
        self.assertIn("name on file stays", item["note"])

    def test_a_match_on_cellosaurus_names_the_accession(self):
        item = self._plan("Something else\tWT\t\t\tCVCL_0030")
        self.assertIn("CVCL_0030", item["note"])

    def test_a_genuinely_new_row_says_nothing(self):
        item = self._plan("ZZQQ-UNIQUE-R4\tWT\t\t\t")
        self.assertEqual(item["status"], "create")
        self.assertEqual(item["note"], "")

    def test_an_update_really_does_not_rename_the_row_on_file(self):
        from pipeline.services import bulk_cell_lines as bulk
        header = "name\tgenotype\tsupplier\tcatalogue"
        bulk.apply(bulk.parse(f"{header}\nHAP1 WT [COWORK RUN4]\tWT\tHorizon Discovery\tC631"))
        self.line.refresh_from_db()
        self.assertEqual(self.line.name, "HAP1")
        self.assertEqual(CellLine.objects.using(DB).count(), 1)


class ASiteNameInTheQueryStringIsNotA500Tests(TestCase):
    """`?site=Leicester` — the board's own URL, with `board` swapped for `export`
    — returned HTTP 500 from both exports. All four boards filtered with
    `site_id=<value>`, which is right for the id the form submits and a crash for
    anything else, so the rows endpoints went down the same way: a URL somebody
    copied and edited took the grid out and blamed the filters."""

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        from pipeline.models import Antibody, Company
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.mcgill = Site.objects.using(DB).create(name="McGill", short_code="MCG")
        self.client = _member_client(self, self.site)
        target = Target.objects.using(DB).create(gene_name="STMN2")
        company = Company.objects.using(DB).create(name="Proteintech")
        Antibody.objects.using(DB).create(
            target_id=target.pk, company_id=company.pk,
            catalogue_number="MINE", site_id=self.site.pk)
        Antibody.objects.using(DB).create(
            target_id=target.pk, company_id=company.pk,
            catalogue_number="THEIRS", site_id=self.mcgill.pk)
        CellLine.objects.using(DB).create(name="HAP1", genotype="WT",
                                          site_id=self.site.pk)

    ROWS = ("/pipeline/antibodies/board/rows/", "/pipeline/cell-lines/board/rows/",
            "/pipeline/sessions/board/rows/", "/pipeline/targets/board/rows/")
    EXPORTS = ("/pipeline/antibodies/export/", "/pipeline/cell-lines/export/")

    def test_every_rows_endpoint_accepts_a_site_name(self):
        for url in self.ROWS:
            with self.subTest(url=url):
                resp = self.client.get(url, {"site": "Leicester"})
                self.assertEqual(resp.status_code, 200)

    def test_every_rows_endpoint_accepts_a_short_code(self):
        for url in self.ROWS:
            with self.subTest(url=url):
                self.assertEqual(
                    self.client.get(url, {"site": "LEI"}).status_code, 200)

    def test_both_exports_accept_a_site_name(self):
        for url in self.EXPORTS:
            with self.subTest(url=url):
                self.assertEqual(
                    self.client.get(url, {"site": "Leicester"}).status_code, 200)

    def test_the_id_the_filter_form_submits_still_works(self):
        resp = self.client.get("/pipeline/antibodies/board/rows/",
                               {"site": self.site.pk})
        cats = {r["catalogue"] for r in resp.json()["rows"]}
        self.assertEqual(cats, {"MINE"})

    def test_a_name_filters_to_the_same_rows_as_its_id(self):
        by_name = self.client.get("/pipeline/antibodies/board/rows/",
                                  {"site": "Leicester"}).json()["rows"]
        self.assertEqual({r["catalogue"] for r in by_name}, {"MINE"})

    def test_a_site_that_does_not_exist_matches_nothing_rather_than_everything(self):
        """Honest: the filter named a site, so showing every row would be a lie
        in the other direction."""
        resp = self.client.get("/pipeline/antibodies/board/rows/",
                               {"site": "Atlantis Institute"})
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.json()["rows"], [])


class ThePickingListSaysSoOnEverySheetTests(TestCase):
    """The caveat went on the WB/IP/FC sheet and not the IF plate map, whose row 2
    is the KO vial's C-number — so the one zero-result session on the site handed
    back a picking list with `C-472` where the explanation should have been."""

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        from pipeline.models import (Antibody, Company, ExperimentSession,
                                     IfResult, Member)
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)
        self.member = Member.objects.using(DB).get(site_id=self.site.pk)
        self.target = Target.objects.using(DB).create(gene_name="STMN2")
        company = Company.objects.using(DB).create(name="Proteintech")
        self.abs = [Antibody.objects.using(DB).create(
            target_id=self.target.pk, company_id=company.pk,
            catalogue_number=f"CAT-{i}", site_id=self.site.pk) for i in range(3)]
        self.empty = ExperimentSession.objects.using(DB).create(
            target_id=self.target.pk, procedure_type="IF", site_id=self.site.pk,
            experimenter_id=self.member.pk, date=date(2026, 7, 31))
        self.filled = ExperimentSession.objects.using(DB).create(
            target_id=self.target.pk, procedure_type="IF", site_id=self.site.pk,
            experimenter_id=self.member.pk, date=date(2026, 7, 31))
        IfResult.objects.using(DB).create(session_id=self.filled.pk,
                                          antibody_id=self.abs[0].pk)

    def _row2(self, session):
        from pipeline.services import planning
        ws = planning.generate_bench_sheet(session).active
        return str(ws.cell(row=2, column=1).value or "")

    def test_the_plate_map_says_when_the_list_is_the_genes_not_the_sessions(self):
        self.assertIn("no antibodies recorded", self._row2(self.empty))

    def test_a_session_with_its_own_antibodies_is_not_labelled_that_way(self):
        self.assertNotIn("no antibodies recorded", self._row2(self.filled))

    def test_the_plate_map_allocates_wells_to_the_sessions_antibodies_only(self):
        """The reason it matters here: the plate map books two wells per antibody
        per permeabilisation, so a gene-wide list reserves wells for antibodies
        nobody is testing."""
        from pipeline.services import planning
        self.assertEqual(
            [a.catalogue_number for a in planning._session_antibodies(self.filled)],
            ["CAT-0"])


class TheGenePagesBadgeIsDerivedTests(TestCase):
    """The badge top-right read **Not Started** in the same viewport as a strip
    showing two procedures complete and six antibodies on file. It was the stored
    ``Target.status`` — written once when the row is created and advanced by
    nothing, the same trap as ``Target.site``."""

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        from pipeline.models import (Antibody, Company, ExperimentSession,
                                     Member, TargetNomination, WbResult)
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)
        self.member = Member.objects.using(DB).get(site_id=self.site.pk)
        self.target = Target.objects.using(DB).create(
            gene_name="STMN2", status="not_started")
        TargetNomination.objects.using(DB).create(
            target_id=self.target.pk, site_id=self.site.pk)
        company = Company.objects.using(DB).create(name="Proteintech")
        ab = Antibody.objects.using(DB).create(
            target_id=self.target.pk, company_id=company.pk,
            catalogue_number="CAT-1", site_id=self.site.pk)
        session = ExperimentSession.objects.using(DB).create(
            target_id=self.target.pk, procedure_type="WB", site_id=self.site.pk,
            experimenter_id=self.member.pk, date=date(2026, 7, 31))
        # With a reading on it. This row used to be blank, and the completion
        # assertion below still passed — because `procedure_summary` counted
        # result *rows*, which is the defect batch six of the field tests found
        # from the other end: a planned session's twenty-two blank rows read as
        # twenty-two results and marked the western blot complete. The fixture
        # is what this class is about — a gene somebody has actually worked on
        # — so it now carries a reading rather than an empty row.
        WbResult.objects.using(DB).create(session_id=session.pk, antibody_id=ab.pk,
                                          signal="single band ~20 kDa")

    def test_the_page_does_not_call_a_worked_gene_not_started(self):
        body = self.client.get(f"/pipeline/target/{self.target.pk}/").content.decode()
        self.assertEqual(self.target.status, "not_started")
        self.assertNotIn("Not Started", body)
        self.assertIn("Testing under way", body)

    def test_the_procedure_summary_ticks_from_sessions_not_from_the_status(self):
        resp = self.client.get(f"/pipeline/target/{self.target.pk}/")
        summary = resp.context["procedure_summary"]
        self.assertTrue(summary["WB"]["complete"])
        self.assertFalse(summary["FC"]["complete"])

    def test_headline_walks_the_whole_way_up(self):
        from pipeline.services import gene_progress
        def label(**done):
            steps = [{"key": k, "done": v} for k, v in done.items()]
            return gene_progress.headline(steps)["label"]
        self.assertEqual(label(nominated=False), "Not started")
        self.assertEqual(label(nominated=True), "Nominated, not started")
        self.assertEqual(label(nominated=True, antibodies=True),
                         "Reagents being gathered")
        self.assertEqual(label(nominated=True, antibodies=True, app_WB=True),
                         "Testing under way")
        self.assertEqual(label(nominated=True, app_WB=True, reported=True),
                         "Reported")


class OnePlaceDecidesWhatASupplierWillSayTests(TestCase):
    """Run 3 fixed the antibodies board's preview to show the supplier that will
    be *saved*. Run 4 pasted on a gene's page and got the typed name back, because
    the fix lived in one page's template. `OGABoard.supplierLabel` is the one
    place now, and nothing may hand-roll it again."""
    databases = {"academy_db", "pipeline_db"}

    STATIC = Path(settings.BASE_DIR) / "pipeline/static/pipeline/board.js"
    TEMPLATES = ("antibody_board.html", "cell_line_board.html",
                 "target_detail.html", "session_board.html")

    def test_board_js_owns_the_resolved_supplier_label(self):
        js = self.STATIC.read_text()
        self.assertIn("function supplierLabel(", js)
        self.assertIn("supplierLabel,", js, "it has to be exported to be used")

    def test_no_template_rebuilds_it(self):
        for name in self.TEMPLATES:
            path = Path(settings.BASE_DIR) / "pipeline/templates/pipeline" / name
            with self.subTest(template=name):
                self.assertNotIn("you typed", path.read_text(),
                                 f"{name} is hand-rolling supplierLabel again")

    def test_every_preview_that_names_a_supplier_uses_it(self):
        """The two that name one: the antibodies board and a gene's page."""
        for name in ("antibody_board.html", "target_detail.html"):
            path = Path(settings.BASE_DIR) / "pipeline/templates/pipeline" / name
            with self.subTest(template=name):
                self.assertIn("OGABoard.supplierLabel", path.read_text())


class EscapeAndTheBannerBehaveTheSameEverywhereTests(TestCase):
    """Two run-4 findings that live entirely in the browser, pinned at the source.

    Escape cancels an inline cell edit, which is the habit the boards teach; the
    identity dialog ignored it and only Cancel would close it. And a banner above
    a tall grid can be scrolled off-screen, which turns a refused edit into one
    that looks accepted — the worst possible failure for a save.

    Run 11 reported the identity dialog ignoring Escape again. **Refuted in a
    real browser**: opened from the antibodies board, one `Escape` puts `hidden`
    back on `idty3-scrim`, and the same keystroke closes the Add panel. Driver,
    not app — the fourth shape in a row where a keystroke was reported as having
    no effect and had one.
    """
    databases = {"academy_db", "pipeline_db"}

    STATIC = Path(settings.BASE_DIR) / "pipeline/static/pipeline/board.js"

    def test_escape_closes_every_pop_out(self):
        """Five of them now — `renumberPanel` is the newest, and a modal you
        cannot dismiss with the keystroke every other modal on the page takes is
        exactly the shape three field tests have already misread as broken.

        Counting rather than naming, so a sixth pop-out has to answer for
        itself: the failure this pins is a panel that was *added* without it.
        And `renumberPanel` earned its place here twice over — its first version
        passed `escapeCloses` the mount *element* where a predicate belongs,
        which parses cleanly and throws only when somebody presses the key.
        """
        js = self.STATIC.read_text()
        self.assertIn("function escapeCloses(", js)
        for fn in ("newEntry", "uploadPanel", "identityDialog", "deleteDialog",
                   "renumberPanel"):
            with self.subTest(pop_out=fn):
                block = js[js.index(f"function {fn}("):]
                end = block.find("\n  function ", 1)
                self.assertIn("escapeCloses(", block[:end if end > 0 else None])
        self.assertEqual(js.count("escapeCloses("), 6,
                         "one definition plus the five pop-outs")

    def test_the_banner_announces_itself_and_scrolls_into_view(self):
        js = self.STATIC.read_text()
        body = js[js.index("function banner("):js.index("async function requestJson(")]
        self.assertIn("role", body)
        self.assertIn("bringIntoView(", body)

    def test_every_message_that_can_sit_below_the_fold_scrolls_itself_in(self):
        """The rule the banner earned, applied to the surfaces that share its
        shape.

        `uploadPanel` writes its receipt into the bottom element of a
        `fixed inset-0 overflow-y-auto` modal — the same tall scroller the grid
        is, and the twelfth field test pressed Save there and reported *"no
        Saved, no row count, no show them on this page"* about two uploads that
        had both written perfectly. A save that wrote and announced nothing is
        the same failure as a refusal that shows nothing, one direction over.

        Source-level because it is cheap and the browser test beside it
        (`test_an_upload_says_what_it_saved`) is the one that proves the wiring:
        this catches a *fifth* message area being added without it.
        """
        js = self.STATIC.read_text()
        self.assertIn("function bringIntoView(", js)
        block = js[js.index("function uploadPanel("):]
        end = block.find("\n  /* ─")
        block = block[:end if end > 0 else None]
        # Both directions: the receipt, and the refusal.
        self.assertEqual(block.count("bringIntoView("), 2,
                         "uploadPanel must scroll in both its receipt and its "
                         "refusal")
        self.assertIn('role="status"', block)


class TheExportReferenceIsFilledInForEveryRowTests(TestCase):
    """`ab #` was the Access-era sequential number and nothing issued one, so it
    was blank for exactly the rows a person had just created — useless as a
    reference for the newest work.

    Both halves are pinned here. A record created now is given its bench's next
    A-number and the sheet prints `A-1`; a record from before that has no number
    and the cell is **empty**, because the record-id fallback that used to fill
    it put a number in front of a reader that was not the vial's number at
    all."""

    databases = {"pipeline_db", "academy_db"}

    def test_a_new_antibody_gets_a_reference(self):
        from pipeline.models import Antibody, Company
        import io
        import openpyxl
        site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        client = _member_client(self, site)
        target = Target.objects.using(DB).create(gene_name="STMN2")
        company = Company.objects.using(DB).create(name="Proteintech")
        ab = Antibody.objects.using(DB).create(
            target_id=target.pk, company_id=company.pk,
            catalogue_number="BRAND-NEW", site_id=site.pk)
        self.assertEqual(ab.ab_number, 1, "the site's first antibody is A-1")
        resp = client.get("/pipeline/antibodies/export/")
        wb = openpyxl.load_workbook(io.BytesIO(resp.content), read_only=True)
        rows = [list(r) for r in wb.active.iter_rows(values_only=True)]
        self.assertEqual(str(rows[1][0]), "A-1")

    def test_a_row_with_no_number_exports_an_empty_cell(self):
        """The half the first field test found, and the reason this class no
        longer promises every row a reference.

        The column used to fall back to the record id so no cell was ever
        blank. That put `4550` under `ab #` for a vial the board called *not
        numbered* — a value that is not the vial's identifier, that a person
        would write on a tube, and that the upload ignores on the way back. A
        blank says the true thing."""
        from pipeline.models import Antibody, Company
        import io
        import openpyxl
        from pipeline.services import lab_numbers
        site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        client = _member_client(self, site)
        target = Target.objects.using(DB).create(gene_name="STMN2")
        company = Company.objects.using(DB).create(name="Proteintech")
        with lab_numbers.suspended():
            ab = Antibody.objects.using(DB).create(
                target_id=target.pk, company_id=company.pk,
                catalogue_number="OLD-ROW", site_id=site.pk)
        self.assertIsNone(ab.ab_number)
        resp = client.get("/pipeline/antibodies/export/")
        wb = openpyxl.load_workbook(io.BytesIO(resp.content), read_only=True)
        rows = [list(r) for r in wb.active.iter_rows(values_only=True)]
        self.assertIn(rows[1][0], (None, ""),
                      "a record id under `ab #` reads as this vial's number")
        self.assertNotIn(str(ab.pk), str(rows[1][0] or ""))


class ThePreviewNamesTheSupplierThatWillBeStoredTests(TestCase):
    """Run 5, and the third time this preview has been wrong in a new way.

    ``OGABoard.supplierLabel`` says ``X (you typed "Y")`` when the resolved
    supplier differs from the typed one, and it is the one place that decides —
    runs 3 and 4 saw to that. What neither pinned is the *value* it compares
    against, and run 5 found it announcing a substitution that did not happen and
    staying silent on one that did, in a single paste:

      * typed ``abcam`` → previewed ``Abcam (you typed "abcam")`` → stored
        ``abcam``. ``plan`` reported ``Company.display_name``, the public-facing
        spelling, while every reader of the record — ``antibody_board.row_for``
        among them — shows ``Company.name``.
      * typed ``Bio-Techne`` with an ``NB`` catalogue → previewed ``Bio-Techne``
        → stored ``Bio-Techne (Novus Biologicals)``. ``plan`` asked
        ``resolve_company(create=False)``, which answers "which supplier row
        exists *now*" and returns nothing when the write is going to create one.
        Silence on exactly the catalogue-prefix path that can file a vial under
        the wrong vendor.

    So the preview must be compared against what ``apply`` writes, not against a
    spelling or a lookup that asks a different question.
    """

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.member = _member_client(self, self.site)
        Target.objects.using(DB).create(gene_name="STMN2")

    def _member_row(self):
        from pipeline.models import Member
        return Member.objects.using(DB).first()

    def _resolved(self, company, catalogue):
        from pipeline.services import bulk_antibodies as bulk
        rows = [{"gene": "STMN2", "catalogue": catalogue, "company": company}]
        return bulk.plan(rows, member=self._member_row())[0]["company_resolved"]

    def _stored(self, company, catalogue):
        from pipeline.models import Antibody
        from pipeline.services import bulk_antibodies as bulk
        rows = [{"gene": "STMN2", "catalogue": catalogue, "company": company}]
        bulk.apply(rows, member=self._member_row())
        ab = Antibody.objects.using(DB).get(catalogue_number=catalogue)
        return ab.company.name if ab.company_id else ""

    def test_a_display_spelling_is_not_announced_as_a_substitution(self):
        """The supplier on file is `abcam`; `Abcam` is only how it is displayed.
        Nothing about the row changes, so the preview must say nothing."""
        from pipeline.models import Company
        Company.objects.using(DB).create(name="abcam", display_name="Abcam")
        self.assertEqual(self._resolved("abcam", "CAT-1"), "abcam")

    def test_the_preview_matches_what_the_write_stores_for_a_known_supplier(self):
        from pipeline.models import Company
        Company.objects.using(DB).create(name="abcam", display_name="Abcam")
        self.assertEqual(self._resolved("abcam", "CAT-1"),
                         self._stored("abcam", "CAT-1"))

    def test_a_brand_the_write_will_create_is_named_in_the_preview(self):
        """No Bio-Techne row on file at all: the write creates
        `Bio-Techne (Novus Biologicals)` from the `NB` prefix, so the preview
        has to say so *before* it happens."""
        self.assertEqual(self._resolved("Bio-Techne", "NB100-1234"),
                         "Bio-Techne (Novus Biologicals)")

    def test_the_other_brand_too(self):
        self.assertEqual(self._resolved("Bio-Techne", "AF3130"),
                         "Bio-Techne (R&D Systems)")

    def test_the_preview_matches_what_the_write_stores_for_a_new_brand(self):
        self.assertEqual(self._resolved("Bio-Techne", "NB100-1234"),
                         self._stored("Bio-Techne", "NB100-1234"))

    def test_an_unremarkable_new_supplier_is_reported_as_typed(self):
        """Nothing to resolve and nothing to announce — the annotation must not
        start firing on every new vendor either."""
        self.assertEqual(self._resolved("Proteintech", "CAT-2"), "Proteintech")


class TheGeneTemplateRoundTripKeepsItsConditionsTests(TestCase):
    """A per-gene template tab wrote its protocol conditions as bare column
    names, and a protocol condition is very often named after the result field it
    describes — IP's ``bead_type``, ``lysis_buffer``, ``gel``, ``membrane``,
    ``ecl``, ``detection_system``; IF's ``blocking``, ``permeabilisation``. Every
    one of those appeared **twice** in the header row.

    That is the bug ``session_io.SESSION_PREFIX`` already exists to prevent one
    surface over: ``parse_template`` builds ``{header: cell}`` per row, so a
    duplicated name keeps whichever copy is rightmost — the blank result column —
    and ``_condition_cols`` then discards the name as a result field. The
    pre-filled protocol value is dropped twice over, silently, on a sheet whose
    own instructions promise a blank never clears a value.

    The sessions board already writes conditions as ``cond:<key>``. The template
    uses the same convention now, so a condition can never collide with a result.
    """

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        from pipeline.models import Antibody, Company, ProtocolTemplate
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)
        self.target = Target.objects.using(DB).create(gene_name="STMN2")
        company = Company.objects.using(DB).create(name="Proteintech")
        Antibody.objects.using(DB).create(
            target_id=self.target.pk, company_id=company.pk,
            catalogue_number="CAT-1", site_id=self.site.pk)
        # A protocol whose conditions are named after IP result fields, which is
        # what the real Montreal/Leicester templates do.
        ProtocolTemplate.objects.using(DB).create(
            name="IP standard", procedure_type="IP", site_id=self.site.pk,
            is_default=True,
            conditions={"bead_type": "Protein A/G", "lysis_buffer": "RIPA"})

    def _columns(self, proc="IP"):
        from pipeline.services import session_template as tmpl
        columns, _rows, _name = tmpl._tab_plan(self.target, proc, self.site.pk)
        return columns

    def test_no_column_name_appears_twice(self):
        cols = self._columns()
        dupes = sorted({c for c in cols if cols.count(c) > 1})
        self.assertEqual(dupes, [], f"duplicate headers in the IP tab: {dupes}")

    def test_a_condition_is_written_under_its_own_prefix(self):
        cols = self._columns()
        self.assertIn("cond:bead_type", cols)
        self.assertIn("cond:lysis_buffer", cols)

    def test_the_result_field_of_the_same_name_survives_alongside_it(self):
        cols = self._columns()
        self.assertIn("bead_type", cols)
        self.assertIn("lysis_buffer", cols)

    def test_the_importer_reads_a_prefixed_condition_back_as_a_condition(self):
        from pipeline.services import session_import
        header = self._columns()
        self.assertIn("cond:bead_type", session_import._condition_cols("IP", header))
        self.assertNotIn("bead_type", session_import._condition_cols("IP", header))

    def test_a_condition_and_its_result_field_keep_separate_values(self):
        """The whole point: the protocol said Protein A/G, the scientist wrote
        what they actually used, and both have to arrive."""
        from pipeline.services import session_import
        header = self._columns()
        row = {h: "" for h in header}
        row.update({"session_ref": "IP · STMN2 (new)", "gene": "STMN2",
                    "antibody": "CAT-1", "company": "Proteintech",
                    "cond:bead_type": "Protein A/G", "bead_type": "Dynabeads M-280"})
        parsed = {"IP": {"header": header, "rows": [row]}}
        res = session_import.apply_import(parsed, uploader=self._member_row())
        self.assertTrue(res.get("ok"), res)
        from pipeline.models import ExperimentSession, IpResult
        session = ExperimentSession.objects.using(DB).latest("id")
        self.assertEqual(session.session_conditions.get("bead_type"), "Protein A/G")
        self.assertEqual(IpResult.objects.using(DB).latest("id").bead_type,
                         "Dynabeads M-280")

    def test_an_older_sheet_without_the_prefix_still_imports_its_conditions(self):
        """Sheets already downloaded are out there. An unprefixed column that is
        not a context or result field is still a condition."""
        from pipeline.services import session_import
        self.assertIn("owner notes",
                      session_import._condition_cols("IP", ["gene", "owner notes"]))

    def _member_row(self):
        from pipeline.models import Member
        return Member.objects.using(DB).first()


class TheGeneTemplateNamesTheWildTypeParentTests(TestCase):
    """``cell_line_wt`` was blank on every row of every tab, and it is blank for
    the reason run 5's C1 verified: **a wild type has no gene**, so it is never in
    ``target.cell_lines`` and ``_first_genotype(cell_lines, "WT")`` can only ever
    return ``None``. The one correct rule made the lookup that ignored it useless.

    The WT is reachable — it is the knockout's ``parent_line``.
    """

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        from pipeline.models import Antibody, Company
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)
        self.target = Target.objects.using(DB).create(gene_name="STMN2")
        company = Company.objects.using(DB).create(name="Proteintech")
        Antibody.objects.using(DB).create(
            target_id=self.target.pk, company_id=company.pk,
            catalogue_number="CAT-1", site_id=self.site.pk)
        self.wt = CellLine.objects.using(DB).create(
            name="SH-SY5Y WT", genotype="WT", site_id=self.site.pk)
        CellLine.objects.using(DB).create(
            name="SH-SY5Y STMN2 KO", genotype="KO", target_id=self.target.pk,
            parent_line_id=self.wt.pk, site_id=self.site.pk)

    def test_the_wild_type_column_is_filled_from_the_knockouts_parent(self):
        """…and named the way the app names it, site included.

        The bare name used to be written. A bare `SH-SY5Y` is shared by five rows
        and a bare `HAP1` by hundreds — the convention puts the gene in its own
        column — so a name alone is not an identity, and resolving one landed a
        Leicester session on McGill's knockout. `services/cell_lines.py` reads
        either spelling back, so sheets written before this still import; what the
        label buys is that a sheet written today cannot be ambiguous tomorrow.
        """
        from pipeline.services import session_template as tmpl
        columns, rows, _ = tmpl._tab_plan(self.target, "WB", self.site.pk)
        row = dict(zip(columns, rows[0]))
        self.assertEqual(row["cell_line_wt"], "SH-SY5Y WT — Leicester")
        self.assertEqual(row["cell_line_ko"], "SH-SY5Y STMN2 KO — Leicester")

        # And it round-trips: what the sheet says resolves to the row it came from.
        from pipeline.services import cell_lines as clines
        found, err = clines.resolve(row["cell_line_wt"], genotype="WT")
        self.assertIsNone(err)
        self.assertEqual(found.pk, self.wt.pk)


class TheSessionTemplateUploadIsReachableTests(TestCase):
    """The per-gene bench workbook has been a one-way door for five runs.

    ``session_import`` parses it, ``plan_import``/``apply_import`` are written and
    tested, and ``session_template_upload_preview``/``_commit`` are routed. No
    page has ever posted to them: the gene page and the sessions board offer the
    download and nothing offers the upload, so the only route back in was the
    sessions round-trip sheet, which matches on ``session_id`` and rejects every
    row of a template with *no session_id — skipped*.

    A page that hands out a file it cannot take back is worse than one that never
    offered it, because the sheet's own instructions promise the upload.
    """

    databases = {"pipeline_db", "academy_db"}

    DOWNLOAD = "/pipeline/session/template/"
    PREVIEW = "/pipeline/session/template/upload/preview/"
    COMMIT = "/pipeline/session/template/upload/commit/"

    def setUp(self):
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)
        self.target = Target.objects.using(DB).create(gene_name="STMN2")

    def _pages(self):
        return {
            "the gene's own page": f"/pipeline/target/{self.target.pk}/",
            "the sessions board": "/pipeline/sessions/board/?gene=STMN2",
        }

    def test_every_page_offering_the_download_offers_the_upload(self):
        """Rendered, not grepped from the template source: the wiring is in an
        include, and what matters is what reaches the screen."""
        for where, url in self._pages().items():
            with self.subTest(page=where):
                body = self.client.get(url).content.decode()
                self.assertIn(self.DOWNLOAD, body,
                              f"{where} no longer offers the workbook at all")
                self.assertIn(self.PREVIEW, body,
                              f"{where} hands out a workbook it cannot take back")
                self.assertIn(self.COMMIT, body)

    def test_the_button_is_there_to_open_it(self):
        for where, url in self._pages().items():
            with self.subTest(page=where):
                self.assertIn('id="bench-upload-btn"',
                              self.client.get(url).content.decode())

    def test_the_sessions_board_offers_neither_without_a_gene(self):
        """The workbook is per-gene and the export refuses without one, so the
        upload must not appear where the download does not."""
        body = self.client.get("/pipeline/sessions/board/").content.decode()
        self.assertNotIn(self.PREVIEW, body)
        self.assertNotIn('id="bench-upload-btn"', body)

    def test_the_endpoints_accept_what_the_page_posts(self):
        """A wired-up button that 404s is the same dead end with extra steps."""
        for url in (self.PREVIEW, self.COMMIT):
            with self.subTest(url=url):
                resp = self.client.post(url, {})
                self.assertEqual(resp.status_code, 400)
                self.assertIn("no file", resp.json()["error"])


class AKnockoutIsConfirmedByItsCellLinesTests(TestCase):
    """``Target.ko_validated`` is another dead Access column, in the company of
    ``Target.status``, ``site``, ``project`` and ``granting_agency``: no write
    path sets it. The gene page's progress strip derives KO confirmation from the
    cell lines and said *KO confirmed — 1 of 1 confirmed*, while TARGET
    INFORMATION a few centimetres above read **No**, from the stored field.

    Same page, same question, two answers — which is precisely what run 4's badge
    finding was, one field over.
    """

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)
        self.target = Target.objects.using(DB).create(gene_name="STMN2")
        wt = CellLine.objects.using(DB).create(
            name="SH-SY5Y WT", genotype="WT", site_id=self.site.pk)
        self.ko = CellLine.objects.using(DB).create(
            name="SH-SY5Y STMN2 KO", genotype="KO", target_id=self.target.pk,
            parent_line_id=wt.pk, site_id=self.site.pk, ko_validated=True)

    def _verdict(self):
        """The Yes/No the KO Validated row actually prints, whitespace collapsed."""
        body = self.client.get(f"/pipeline/target/{self.target.pk}/").content.decode()
        after = body[body.index("KO Validated") + len("KO Validated"):]
        text = re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", after)).strip()
        return text[:80]

    def test_a_confirmed_knockout_is_not_reported_as_unconfirmed(self):
        """The stored field is False and the cell line says otherwise."""
        self.assertFalse(Target.objects.using(DB).get(pk=self.target.pk).ko_validated)
        self.assertTrue(self._verdict().startswith("Yes"), self._verdict())

    def test_a_gene_with_no_confirmed_knockout_still_says_so(self):
        self.ko.ko_validated = False
        self.ko.save(using=DB)
        self.assertTrue(self._verdict().startswith("No"), self._verdict())

    def test_the_count_beside_it_agrees_with_its_verb(self):
        """Next to a number is where a grammar slip costs most."""
        self.assertIn("1 of 1 line confirmed", self._verdict())


class ARefusedLotNamesTheVialItClashesWithTests(TestCase):
    """Two ways to push a row onto another row's identity, two very different
    messages. Retyping the *catalogue* in the identity dialog names the row it
    collided with, gives its id and says to merge them. Editing the *lot* in the
    grid — the same five-part key, the same refusal — gave
    ``Could not save that — the value has been left unchanged.``

    The board's patch handler wraps the save in a bare ``except Exception``, so
    the ``IntegrityError`` from ``unique_antibody_per_site_lot`` arrives as the
    generic message. A refusal that names no alternative is half a message, and
    this one names nothing at all.
    """

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        from pipeline.models import Antibody, Company
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)
        target = Target.objects.using(DB).create(gene_name="STMN2")
        company = Company.objects.using(DB).create(name="Proteintech")
        common = dict(target_id=target.pk, company_id=company.pk,
                      catalogue_number="CAT-1", site_id=self.site.pk)
        self.first = Antibody.objects.using(DB).create(lot_number="LOT-A", **common)
        self.second = Antibody.objects.using(DB).create(lot_number="LOT-B", **common)

    def test_editing_a_lot_onto_an_existing_vial_names_it(self):
        resp = self.client.post(
            "/pipeline/antibodies/board/patch/",
            {"antibody_id": self.second.pk, "field": "lot_number", "value": "LOT-A"})
        self.assertEqual(resp.status_code, 400)
        error = resp.json()["error"]
        self.assertIn("CAT-1", error)
        self.assertIn(str(self.first.pk), error)
        self.assertIn("merge", error.lower())

    def test_the_lot_is_left_as_it_was(self):
        from pipeline.models import Antibody
        self.client.post(
            "/pipeline/antibodies/board/patch/",
            {"antibody_id": self.second.pk, "field": "lot_number", "value": "LOT-A"})
        self.assertEqual(
            Antibody.objects.using(DB).get(pk=self.second.pk).lot_number, "LOT-B")

    def test_an_ordinary_lot_edit_still_saves(self):
        from pipeline.models import Antibody
        resp = self.client.post(
            "/pipeline/antibodies/board/patch/",
            {"antibody_id": self.second.pk, "field": "lot_number", "value": "LOT-C"})
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(
            Antibody.objects.using(DB).get(pk=self.second.pk).lot_number, "LOT-C")


class TheBenchWorkbookGoesBackInTests(TestCase):
    """§2(a) of the run-5 brief, end to end: download the per-gene workbook, fill
    it in as a scientist would — invented columns and all — and upload it.

    Every part of this existed except the last one. `session_import` parses the
    workbook, `apply_import` writes it, both endpoints are routed, and no page
    posted to either, so the brief's three questions were all blocked at step 1.
    This drives the endpoints the page now posts to, on a file the app itself
    produced.
    """
    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        from pipeline.models import Antibody, Company, ProtocolTemplate
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)
        self.target = Target.objects.using(DB).create(gene_name="STMN2")
        company = Company.objects.using(DB).create(name="Proteintech")
        self.ab = Antibody.objects.using(DB).create(
            target_id=self.target.pk, company_id=company.pk,
            catalogue_number="CAT-1", site_id=self.site.pk)
        wt = CellLine.objects.using(DB).create(
            name="SH-SY5Y WT", genotype="WT", site_id=self.site.pk)
        CellLine.objects.using(DB).create(
            name="SH-SY5Y STMN2 KO", genotype="KO", target_id=self.target.pk,
            parent_line_id=wt.pk, site_id=self.site.pk)
        ProtocolTemplate.objects.using(DB).create(
            name="WB standard", procedure_type="WB", site_id=self.site.pk,
            is_default=True,
            # `gel` is also a WB *result* field — the collision this whole
            # round trip used to lose. `lysis_buffer` is in the procedure's
            # condition registry, so it is the one a cell edit may touch.
            conditions={"gel": "4-20% WedgeWell", "lysis_buffer": "RIPA"})

    def _filled_workbook(self, **extra):
        """The real download, filled in on the WB tab, as bytes."""
        import io
        import openpyxl
        from pipeline.services import session_template as tmpl
        wb = openpyxl.load_workbook(io.BytesIO(tmpl.build_gene_template("STMN2")))
        ws = wb["WB"]
        header = [c.value for c in ws[1]]
        for name, value in {"signal": "21 kDa band", "rating": "5",
                            "dilution": "1:1000", **extra}.items():
            if name not in header:
                header.append(name)
                ws.cell(1, len(header), name)
            ws.cell(2, header.index(name) + 1, value)
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def _upload(self, url, data):
        from django.core.files.uploadedfile import SimpleUploadedFile
        return self.client.post(url, {"file": SimpleUploadedFile(
            "session_template_STMN2.xlsx", data,
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})

    def test_the_preview_counts_the_session_it_would_create(self):
        d = self._upload("/pipeline/session/template/upload/preview/",
                         self._filled_workbook()).json()
        self.assertTrue(d["ok"], d)
        self.assertEqual(d["counts"]["sessions"], 1)
        self.assertEqual(d["counts"]["results"], 1)

    def test_the_preview_writes_nothing(self):
        from pipeline.models import ExperimentSession
        self._upload("/pipeline/session/template/upload/preview/",
                     self._filled_workbook())
        self.assertEqual(ExperimentSession.objects.using(DB).count(), 0)

    def test_the_commit_creates_the_session_and_its_results(self):
        from pipeline.models import ExperimentSession, WbResult
        d = self._upload("/pipeline/session/template/upload/commit/",
                         self._filled_workbook()).json()
        self.assertTrue(d["ok"], d)
        self.assertEqual(d["sessions_created"], 1)
        self.assertEqual(d["results_created"], 1)
        session = ExperimentSession.objects.using(DB).get()
        self.assertEqual(session.procedure_type, "WB")
        result = WbResult.objects.using(DB).get()
        self.assertEqual(result.signal, "21 kDa band")
        self.assertEqual(result.antibody_id, self.ab.pk)

    def test_the_protocol_condition_arrives_as_a_condition(self):
        """`gel` is both a WB result field and this protocol's condition. It used
        to appear twice in the header and lose its value on the way back."""
        from pipeline.models import ExperimentSession
        self._upload("/pipeline/session/template/upload/commit/",
                     self._filled_workbook())
        session = ExperimentSession.objects.using(DB).get()
        self.assertEqual(session.session_conditions.get("gel"), "4-20% WedgeWell")

    def test_an_invented_column_is_kept(self):
        """The brief's step 3: a column the app has never heard of, with a value
        in it, survives the upload."""
        from pipeline.models import ExperimentSession
        self._upload("/pipeline/session/template/upload/commit/",
                     self._filled_workbook(**{"Owner notes": "ran with Ana"}))
        session = ExperimentSession.objects.using(DB).get()
        # `owner_notes`, not `owner notes`. Every key the condition registry
        # knows is snake_case; a promoted column arrived lower-cased with its
        # spaces intact, so it was the one key in the dict nothing round-tripping
        # by key would survive. Run 6, F3b.
        self.assertEqual(session.session_conditions.get("owner_notes"),
                         "ran with Ana")

    def test_editing_one_condition_does_not_wipe_the_invented_column(self):
        """The other half of step 3, and the reason `session_conditions` is
        touched key by key rather than rebuilt."""
        from pipeline.models import ExperimentSession
        self._upload("/pipeline/session/template/upload/commit/",
                     self._filled_workbook(**{"Owner notes": "ran with Ana"}))
        session = ExperimentSession.objects.using(DB).get()
        resp = self.client.post("/pipeline/sessions/board/patch/",
                                {"session_id": session.pk,
                                 "field": "cond:lysis_buffer",
                                 "value": "Pierce IP Lysis Buffer"})
        self.assertEqual(resp.status_code, 200, resp.content)
        session.refresh_from_db()
        self.assertEqual(session.session_conditions.get("lysis_buffer"),
                         "Pierce IP Lysis Buffer")
        self.assertEqual(session.session_conditions.get("owner_notes"),
                         "ran with Ana")
        # And the condition the registry does not know is still there, unedited.
        self.assertEqual(session.session_conditions.get("gel"), "4-20% WedgeWell")

    def test_the_tabs_you_did_not_fill_in_do_not_become_experiments(self):
        """A workbook has a tab per application because you cannot know in advance
        which the week will bring, so three of the four come back untouched — with
        their context rows still pre-filled.

        Uploading one filled tab created **four** sessions, each marked *complete*,
        three of them with a blank result row against a real antibody. That is a
        record of an experiment nobody ran, which is the same defect the bench
        sheet had in run 3, one surface over and marked complete this time.
        """
        from pipeline.models import ExperimentSession
        d = self._upload("/pipeline/session/template/upload/commit/",
                         self._filled_workbook()).json()
        self.assertEqual(d["sessions_created"], 1)
        self.assertEqual(
            [s.procedure_type for s in ExperimentSession.objects.using(DB).all()],
            ["WB"])

    def test_the_preview_says_so_before_the_commit_does(self):
        """A preview that promised four sessions and a commit that made one would
        be the same disagreement in the other direction."""
        d = self._upload("/pipeline/session/template/upload/preview/",
                         self._filled_workbook()).json()
        self.assertEqual([s["procedure"] for s in d["sessions"]], ["WB"])

    def _read_me(self):
        import io
        import openpyxl
        from pipeline.services import session_template as tmpl
        wb = openpyxl.load_workbook(io.BytesIO(tmpl.build_gene_template("STMN2")))
        return "\n".join(" ".join(str(c) for c in row if c)
                         for row in wb["How to use"].iter_rows(values_only=True))

    def test_the_instruction_tab_does_not_promise_an_update(self):
        """That tab is the only guidance somebody has at the bench, and it said
        *"Upload adds/updates only"* about an upload that did not then exist —
        and which, now it does, only ever creates. Promising an update would send
        somebody to upload a corrected sheet expecting it to replace a reading."""
        text = self._read_me()
        self.assertNotIn("adds/updates", text)
        self.assertIn("never edits or deletes", text)
        self.assertIn("recorded twice", text)

    def test_the_instruction_tab_explains_the_two_things_that_would_surprise_you(self):
        text = self._read_me()
        self.assertIn("cond:", text)
        self.assertIn("is not recorded", text)

    def test_a_row_naming_an_unknown_antibody_is_reported_not_guessed(self):
        import io
        import openpyxl
        from pipeline.services import session_template as tmpl
        wb = openpyxl.load_workbook(io.BytesIO(tmpl.build_gene_template("STMN2")))
        ws = wb["WB"]
        header = [c.value for c in ws[1]]
        ws.cell(2, header.index("antibody") + 1, "NOT-A-CATALOGUE")
        ws.cell(2, header.index("signal") + 1, "something")
        buf = io.BytesIO()
        wb.save(buf)
        d = self._upload("/pipeline/session/template/upload/preview/",
                         buf.getvalue()).json()
        self.assertEqual(d["counts"]["sessions"], 0)
        self.assertTrue(any("NOT-A-CATALOGUE" in b["label"] for b in d["blocked"]),
                        d["blocked"])


# ═══════════════════════════════════════════════════════════════════════════
# The sixth field test, 31 Jul 2026.
#
# Its verdict was that the app very often does the right thing without telling
# you, and most of what follows pins a screen rather than a write: a panel that
# counted a record and then said there was none, a download that produced an
# excellent document and showed nothing, a preview that reported ten conditions
# and never mentioned the column whose values it was about to throw away.
#
# Two findings were REFUTED by driving a real browser against a
# LiveServerTestCase — Enter does commit an inline edit, and "Include cancelled"
# does apply — so there is deliberately nothing here for either. The harness
# needs playwright, which is not in requirements.txt; see the cowork-field-test
# skill, §4.
# ═══════════════════════════════════════════════════════════════════════════


class AWildTypeIsNeverInItsGenesQuerysetTests(TestCase):
    """One root cause, two of the run's findings, and the rule behind both.

    A wild type is recorded once, **with no gene**, so
    ``CellLine.objects.filter(target=…)`` can never return one. Three readers
    assumed it could, and the two that a person sees both failed by *omitting*
    the wild type rather than by raising — which is why neither was noticed until
    somebody read the screen against the database.
    """

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        from pipeline.models import Antibody, Company, ExperimentSession, Member
        from django.contrib.auth.models import User
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)
        pu = User.objects.using(DB).get(username="carl")
        self.member = Member.objects.using(DB).get(user_id=pu.pk)
        self.target = Target.objects.using(DB).create(
            gene_name="ELP3", protein_name="Elongator complex protein 3")
        self.company = Company.objects.using(DB).create(name="Horizon Discovery")
        # The parent has no gene. That is the rule, not an omission.
        self.wt = CellLine.objects.using(DB).create(
            name="HAP1", genotype="WT", site_id=self.site.pk)
        self.ko = CellLine.objects.using(DB).create(
            name="HAP1 ELP3 KO", genotype="KO", target_id=self.target.pk,
            parent_line_id=self.wt.pk, site_id=self.site.pk,
            company_id=self.company.pk, catalogue_number="HZGHC-ELP3")
        Antibody.objects.using(DB).create(
            target_id=self.target.pk, company_id=self.company.pk,
            catalogue_number="A-ELP3-001", site_id=self.site.pk)
        ExperimentSession.objects.using(DB).create(
            target_id=self.target.pk, site_id=self.site.pk, procedure_type="WB",
            date=date(2026, 7, 31), experimenter_id=self.member.pk,
            cell_line_wt_id=self.wt.pk, cell_line_ko_id=self.ko.pk)

    def test_the_cell_lines_panel_lists_the_line_it_counted(self):
        """The panel read "CELL LINES (1)" above "No cell lines recorded for this
        target." — and it did so *because* the knockout was linked to its parent
        correctly. Pairing walked the WT lines in the target's own queryset,
        where a parent structurally never is, so a properly linked KO fell out of
        both loops and the count and the list disagreed. Doing the right thing
        made the row vanish.
        """
        resp = self.client.get(f"/pipeline/target/{self.target.pk}/")
        self.assertEqual(resp.context["cell_line_count"], 1)
        self.assertEqual(len(resp.context["cell_line_pairs"]), 1)
        self.assertNotIn("No cell lines recorded for this target.",
                         resp.content.decode())

    def test_every_counted_line_is_shown(self):
        """The invariant, over every genotype. `other` was computed and then
        dropped on the floor, which was the same defect one branch along."""
        CellLine.objects.using(DB).create(
            name="HAP1 ELP3 het", genotype="other", target_id=self.target.pk,
            site_id=self.site.pk)
        resp = self.client.get(f"/pipeline/target/{self.target.pk}/")
        shown = set()
        for pair in resp.context["cell_line_pairs"]:
            for line in (pair.get("ko"), pair.get("other")):
                if line is not None:
                    shown.add(line.pk)
            # A WT parent is shown alongside its KO but is not on this gene's
            # list, so it does not count towards the header's number.
        self.assertEqual(len(shown), resp.context["cell_line_count"])

    def test_the_wild_type_parent_is_shown_even_though_it_has_no_gene(self):
        resp = self.client.get(f"/pipeline/target/{self.target.pk}/")
        pair = resp.context["cell_line_pairs"][0]
        self.assertIsNotNone(pair["wt"], "the KO's parent was not fetched")
        self.assertEqual(pair["wt"].pk, self.wt.pk)

    def test_the_report_names_the_wild_type_cell_line(self):
        """Table 1 listed only the knockout and every figure legend fell back to
        the literal placeholder "[WT cell line]" — two paragraphs above a Method
        section that named HAP1 correctly, in a document that leaves the
        building."""
        import docx
        from pipeline.services.report_generator import generate_report
        doc = docx.Document(generate_report(self.target.pk))
        text = "\n".join(p.text for p in doc.paragraphs)
        self.assertNotIn("[WT cell line]", text)
        table_text = "\n".join(
            " | ".join(c.text for c in row.cells)
            for t in doc.tables for row in t.rows)
        self.assertIn("HAP1", table_text)
        self.assertRegex(table_text, r"HAP1\s*\|\s*WT")

    def test_the_gene_page_lists_the_sites_wild_type_parentals(self):
        """A WT can never appear under a gene, so the only way to discover that
        your site already holds one was to attempt the mistake and read the
        preview's refusal. The panel says which ones are on file."""
        CellLine.objects.using(DB).create(
            name="SH-SY5Y", genotype="WT", site_id=self.site.pk)
        resp = self.client.get(f"/pipeline/target/{self.target.pk}/")
        body = resp.content.decode()
        self.assertEqual(sorted(resp.context["wt_parentals"]), ["HAP1", "SH-SY5Y"])
        self.assertIn("Wild-type parentals already on file", body)
        self.assertIn("SH-SY5Y", body)

    def test_the_parent_cell_offers_them_instead_of_asking_you_to_retype_one(self):
        """A list you read and a box you type into are two halves of one job.

        The values are what `bulk_cell_lines.resolve_parent` accepts — a bare
        name or a C-number — and never `cell_lines.label()`'s "HAP1 — Leicester",
        which that parser would refuse. One call behind the sentence and the
        dropdown, so they cannot disagree about what this site has.
        """
        resp = self.client.get(f"/pipeline/target/{self.target.pk}/")
        options = resp.context["wt_options"]
        self.assertEqual([o["value"] for o in options],
                         sorted(resp.context["wt_parentals"]))
        for option in options:
            self.assertNotIn("—", option["value"])
        body = resp.content.decode()
        self.assertIn('id="wt-options"', body)
        self.assertIn("suggestions: {parent: json('wt-options')}", body)

    def test_a_wild_type_can_be_started_from_the_gene_page(self):
        """The panel says a knockout needs its parent, and a wild type can never
        be listed under a gene — so with none on file nothing said where the
        first one comes from. The button opens the same grid with the two cells
        that make a row a wild type already set: genotype WT, and no gene."""
        body = self.client.get(f"/pipeline/target/{self.target.pk}/").content.decode()
        self.assertIn('id="cl-add-wt-btn"', body)
        self.assertIn("cellLines.open({genotype: 'WT', gene: 'NA'})", body)


class ADownloadSaysItHappenedTests(TestCase):
    """Generate Report produced a 40 kB Data Note and the page said nothing.

    The tester filed the app's best output as a broken feature, corrected it
    only by listing the Downloads folder, and got that wrong too. A plain
    ``<a href>`` to a file is invisible: the browser writes it and the page does
    not move. Pinned at the source as well as in the response, because what is
    wrong is a thing that does not happen.
    """

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)
        self.target = Target.objects.using(DB).create(gene_name="ELP3")

    def test_both_downloads_on_the_gene_page_ask_for_a_receipt(self):
        body = self.client.get(f"/pipeline/target/{self.target.pk}/").content.decode()
        self.assertIn('data-receipt="report-receipt"', body)
        self.assertIn('data-receipt="workbook-receipt"', body)
        self.assertIn('id="report-receipt"', body)
        self.assertIn('id="workbook-receipt"', body)

    def test_the_shared_helper_exists_and_is_wired(self):
        js = (Path(settings.BASE_DIR) / "pipeline/static/pipeline/board.js").read_text()
        self.assertIn("function downloadWithReceipt", js)
        self.assertIn("downloadReceipts", js)
        body = self.client.get(f"/pipeline/target/{self.target.pk}/").content.decode()
        self.assertIn("OGABoard.downloadReceipts()", body)

    def test_the_reports_panel_says_what_it_is_for(self):
        """"No reports generated for this target." under a button called Generate
        Report is what made a working feature read as a broken one. The panel
        tracks the *published* record; a draft never lands in it."""
        body = self.client.get(f"/pipeline/target/{self.target.pk}/").content.decode()
        self.assertNotIn("No reports generated for this target.", body)
        self.assertIn("Nothing published for this target yet", body)
        self.assertIn("draft", body)

    def test_generating_a_report_does_not_demote_a_published_one(self):
        """`status` is a lifecycle and two curated paths set it to published — a
        DOI typed on the target board, and Carl's workbook. It was stamped
        'generated' unconditionally, so pressing a button that reads as read-only
        and gives no sign it wrote anything quietly walked the record backwards.
        """
        from pipeline.models import Report
        report = Report.objects.using(DB).create(
            target_id=self.target.pk, status=Report.ReportStatus.PUBLISHED,
            zenodo_doi="https://doi.org/10.5281/zenodo.1")
        resp = self.client.get(f"/pipeline/target/{self.target.pk}/generate-report/")
        self.assertEqual(resp.status_code, 200)
        report.refresh_from_db()
        self.assertEqual(report.status, Report.ReportStatus.PUBLISHED)

    def test_a_draft_report_row_still_advances(self):
        from pipeline.models import Report
        report = Report.objects.using(DB).create(
            target_id=self.target.pk, status=Report.ReportStatus.DRAFT)
        self.client.get(f"/pipeline/target/{self.target.pk}/generate-report/")
        report.refresh_from_db()
        self.assertEqual(report.status, Report.ReportStatus.GENERATED)

    def test_the_report_is_not_written_to_a_path_named_after_the_gene_alone(self):
        """Two people asking for one gene's report at the same moment wrote to
        the same /tmp path and each could be handed the other's half-written
        file."""
        src = (Path(settings.BASE_DIR) / "pipeline/views/dashboard.py").read_text()
        self.assertIn("TemporaryDirectory", src)


class ThePanelSaysWhatTheDraftWouldContainTests(TestCase):
    """Generate Report on an empty target hands back 39 KB of placeholder.

    The twelfth field test pressed it on a sandbox TRPA1 with no antibodies, no
    cell lines and no sessions, and got a Data Note whose abstract said *"we have
    characterized 0 antibodies"*, two header-only tables and a title reading
    *"…for use in [no application has a recorded result yet]"*. The document is
    doing the right thing — the bracketed gaps are the convention the whole
    generator uses, and a draft that invented a buffer would be far worse — but
    the gene page knew all of it before the press and said nothing.

    It is not gated the way the feasibility page's Add button is: that one
    writes to the master list and cannot be taken back, this one costs a file in
    Downloads. It owes a manifest, not a wall.
    """

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        from pipeline.models import Member
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)
        self.member = Member.objects.using(DB).get(user__username="carl")
        self.target = Target.objects.using(DB).create(gene_name="TRPA1")

    def _page(self):
        return self.client.get(f"/pipeline/target/{self.target.pk}/").content.decode()

    def test_an_empty_target_says_so_before_the_press(self):
        body = self._page()
        self.assertIn("Nothing has been recorded against this gene yet", body)
        self.assertIn("skeleton", body)

    def test_and_the_press_is_made_deliberate_rather_than_removed(self):
        """A download writes nothing, so it is confirmed, never disabled —
        somebody wanting the empty skeleton is entitled to it."""
        body = self._page()
        self.assertIn("data-confirm=", body)
        self.assertIn("Download it anyway?", body)
        js = (Path(settings.BASE_DIR) / "pipeline/static/pipeline/board.js").read_text()
        self.assertIn("a.dataset.confirm", js)
        # Still the same anchor, so the receipt rule is not lost to the confirm.
        self.assertIn('data-receipt="report-receipt"', body)

    def test_a_gene_with_readings_gets_a_manifest_and_no_confirmation(self):
        from pipeline.models import (Antibody, Company, ExperimentSession,
                                     WbResult)
        company = Company.objects.using(DB).create(name="Abcam")
        ab = Antibody.objects.using(DB).create(
            target_id=self.target.pk, company_id=company.pk,
            catalogue_number="ab1", site_id=self.site.pk)
        session = ExperimentSession.objects.using(DB).create(
            target_id=self.target.pk, procedure_type="WB", site_id=self.site.pk,
            experimenter_id=self.member.pk, date=date(2026, 8, 1))
        WbResult.objects.using(DB).create(
            session_id=session.pk, antibody_id=ab.pk, dilution="1:1000")
        body = self._page()
        self.assertIn("The draft would be built from", body)
        self.assertIn("1 antibody", body)
        self.assertIn("1 recorded reading", body)
        self.assertNotIn("data-confirm=", body)

    def test_the_count_is_the_documents_own_and_not_the_row_count(self):
        """A session carrying one blank result row is not an experiment.

        `_get_sessions` has dropped those since a report claimed three
        antibodies characterized over three empty sessions. The panel above
        counts result *rows* for its procedure ticks, so reading that number
        here would print a count the file then contradicts — two answers to one
        question on one screen, which is the shape that reads as data loss.
        """
        from pipeline.models import (Antibody, Company, ExperimentSession,
                                     WbResult)
        from pipeline.services import report_generator
        company = Company.objects.using(DB).create(name="Abcam")
        ab = Antibody.objects.using(DB).create(
            target_id=self.target.pk, company_id=company.pk,
            catalogue_number="ab1", site_id=self.site.pk)
        session = ExperimentSession.objects.using(DB).create(
            target_id=self.target.pk, procedure_type="WB", site_id=self.site.pk,
            experimenter_id=self.member.pk, date=date(2026, 8, 1))
        WbResult.objects.using(DB).create(
            session_id=session.pk, antibody_id=ab.pk)   # nobody wrote on it
        draft = report_generator.draft_contents(self.target)
        self.assertEqual(draft["readings"], 0)
        self.assertEqual(draft["sessions"], 0)
        self.assertTrue(draft["empty"])
        self.assertIn("Nothing has been recorded against this gene yet",
                      self._page())


class ConcentrationKeepsItsUnitTests(TestCase):
    """"1.0 mg/mL" was stored as the number 1.

    The parser scraped the digits and dropped the unit, into a field that means
    µg/mL — a thousandfold error with nothing on any screen to catch it by, and
    the report's own Table 2 heading said µg/µl, which is a second thousandfold
    the other way. One reader for every write path now.
    """

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        from pipeline.models import Antibody, Company
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)
        self.target = Target.objects.using(DB).create(gene_name="ELP3")
        self.company = Company.objects.using(DB).create(name="abcam")
        self.ab = Antibody.objects.using(DB).create(
            target_id=self.target.pk, company_id=self.company.pk,
            catalogue_number="A-ELP3-001", site_id=self.site.pk)

    def test_a_unit_is_converted_not_discarded(self):
        from pipeline.services import concentration
        self.assertEqual(concentration.parse("1.0 mg/mL")[0], Decimal("1000.0"))
        self.assertEqual(concentration.parse("1 ug/ul")[0], Decimal("1000"))
        self.assertEqual(concentration.parse("200 ug/mL")[0], Decimal("200"))
        # A bare number is the stored unit — that is what every sheet this app
        # exports contains, so it must not suddenly mean something else.
        self.assertEqual(concentration.parse("1.0")[0], Decimal("1.0"))
        self.assertIsNone(concentration.parse("")[0])

    def test_a_unit_it_cannot_convert_is_refused_by_name(self):
        value, err = __import__(
            "pipeline.services.concentration", fromlist=["parse"]).parse("1.0 mM")
        self.assertIsNone(value)
        self.assertIn("1.0 mM", err)
        self.assertIn("µg/mL", err)

    def test_the_board_cell_converts_rather_than_calling_it_not_a_number(self):
        resp = self.client.post(
            "/pipeline/antibodies/board/patch/",
            {"target_id": self.ab.pk, "field": "concentration", "value": "1.0 mg/mL"})
        self.assertEqual(resp.status_code, 200, resp.content)
        self.ab.refresh_from_db()
        self.assertEqual(self.ab.concentration, Decimal("1000.000"))

    def test_a_paste_converts_and_the_preview_says_when_it_cannot(self):
        from pipeline.services import bulk_antibodies
        rows = bulk_antibodies.parse(
            "gene\tcatalogue\tcompany\tconcentration\n"
            "ELP3\tA-ELP3-777\tabcam\t1.0 mg/mL\n"
            "ELP3\tA-ELP3-778\tabcam\t3 mM\n")
        items = bulk_antibodies.plan(rows)
        notes = " ".join(i["note"] for i in items)
        self.assertNotIn("1.0 mg/mL", notes, "a convertible unit needs no warning")
        self.assertIn("3 mM", notes, "an unconvertible unit went in silently")

    def test_the_preview_and_the_write_agree_on_the_number(self):
        """The Data I/O upload kept its own copy of the digit-scrape, so once the
        write learned to convert, the diff on screen and the value stored
        disagreed by a thousand — the one thing a preview may never do."""
        from pipeline.services.dataset import _incoming_antibody_values
        values = dict((f, comparable) for f, _l, _new, comparable
                      in _incoming_antibody_values({"concentration": "1.0 mg/mL"}))
        self.assertEqual(values["concentration"], Decimal("1000.0"))

    def test_the_report_table_states_the_unit_the_field_is_stored_in(self):
        """Table 2's heading said µg/µl about a field stored in µg/mL — a second
        thousandfold, in the one table that leaves the building."""
        import docx
        from pipeline.services.report_generator import generate_report
        doc = docx.Document(generate_report(self.target.pk))
        headings = "\n".join(c.text for t in doc.tables for c in t.rows[0].cells)
        self.assertIn("µg/mL", headings)
        self.assertNotIn("µg/µl", headings)

    def test_the_pasted_column_names_its_unit(self):
        from pipeline.views.imports import columns_and_example
        columns, _example = columns_and_example("antibodies")
        self.assertTrue(any("ug/mL" in c for c in columns), columns)


class ARefusalNamesTheSupplierTheRecordIsFiledUnderTests(TestCase):
    """`display_name` is the public website's spelling; every screen that reads a
    pipeline record shows `.name`.

    So the clash banner announced "already a row … from **Abcam**" about a row
    stored as `abcam` — the exact substitution the paste preview had just
    corrected you away from. One helper answers it now, and it lives beside
    `resolved_company_name`, which answers the other half of the same question.
    """

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        from pipeline.models import Antibody, Company
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)
        self.target = Target.objects.using(DB).create(gene_name="ELP3")
        self.company = Company.objects.using(DB).create(
            name="abcam", display_name="Abcam")
        common = dict(target_id=self.target.pk, company_id=self.company.pk,
                      catalogue_number="A-ELP3-C6A", site_id=self.site.pk)
        self.first = Antibody.objects.using(DB).create(lot_number="RUN6-A4", **common)
        self.second = Antibody.objects.using(DB).create(lot_number="RUN6-A6", **common)

    def test_the_clash_banner_names_the_stored_spelling(self):
        resp = self.client.post(
            "/pipeline/antibodies/board/patch/",
            {"target_id": self.second.pk, "field": "lot_number", "value": "RUN6-A4"})
        self.assertEqual(resp.status_code, 400)
        error = resp.json()["error"]
        self.assertIn("from abcam", error)
        self.assertNotIn("from Abcam", error)
        self.assertIn(str(self.first.pk), error, "the banner must name the row")

    def test_the_croppers_preview_answers_the_write_not_the_lookup(self):
        """Both halves of the run-5 supplier bug were still here, one page over:
        `display_name`, and a value set only when a row already exists — so a
        bare "Bio-Techne" said nothing on exactly the catalogue-prefix path that
        can file a vial under the wrong vendor."""
        src = (Path(settings.BASE_DIR) / "pipeline/views/cropper.py").read_text()
        self.assertIn("resolved_company_name", src)
        self.assertNotIn("existing.display_name or existing.name", src)


class ThePreviewSaysWhetherTheParentResolvedTests(TestCase):
    """A knockout's parent is the whole point of the record, and the preview said
    only "new" whether it had matched or not.

    The tester's diagnosis — "the PARENT column wants a C-number" — is wrong, and
    that is the finding: the column is keyed by *name*, and the only statement of
    any convention was a placeholder on row 1 that disappears as you type.
    """

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.other = Site.objects.using(DB).create(name="McGill", short_code="MCG")
        self.client = _member_client(self, self.site)
        self.target = Target.objects.using(DB).create(gene_name="ELP3")

    def _plan(self, parent):
        from pipeline.models import Member
        from pipeline.services import bulk_cell_lines
        member = Member.objects.using(DB).get(site_id=self.site.pk)
        rows = bulk_cell_lines.parse(
            "name\tgene\tgenotype\tparent\n"
            f"HAP1 ELP3 KO\tELP3\tKO\t{parent}\n")
        return bulk_cell_lines.plan(rows, member=member)

    def test_a_parent_that_matches_nothing_says_so_before_the_save(self):
        note = self._plan("HAP1")[0]["note"]
        self.assertIn("no cell line called 'HAP1'", note)
        self.assertIn("no parent linked", note)

    def test_the_refusal_names_the_parentals_on_file(self):
        CellLine.objects.using(DB).create(
            name="SH-SY5Y", genotype="WT", site_id=self.site.pk)
        note = self._plan("HAP1")[0]["note"]
        self.assertIn("SH-SY5Y", note)

    def test_your_own_sites_line_wins_over_another_sites(self):
        """Two labs really do both call their parental HAP1 — which is why the
        sessions board showed "HAP1 / HAP1". The lookup was unscoped, three lines
        below one that is scoped on purpose, so it could pick either."""
        from pipeline.services import bulk_cell_lines
        theirs = CellLine.objects.using(DB).create(
            name="HAP1", genotype="WT", site_id=self.other.pk)
        mine = CellLine.objects.using(DB).create(
            name="HAP1", genotype="WT", site_id=self.site.pk)
        self.assertLess(theirs.pk, mine.pk, "the unscoped lookup would take this one")
        parent, note = bulk_cell_lines.resolve_parent("HAP1", site_id=self.site.pk)
        self.assertEqual(parent.pk, mine.pk)
        self.assertIn("your site", note)

    def test_another_sites_line_is_matched_but_named(self):
        from pipeline.services import bulk_cell_lines
        theirs = CellLine.objects.using(DB).create(
            name="HAP1", genotype="WT", site_id=self.other.pk)
        parent, note = bulk_cell_lines.resolve_parent("HAP1", site_id=self.site.pk)
        self.assertEqual(parent.pk, theirs.pk)
        self.assertIn("McGill", note)

    def test_the_preview_asks_the_same_question_as_the_write(self):
        """`apply` knew all along and put its answer in a note no preview reads.
        One helper, called by both, or the two drift."""
        src = (Path(settings.BASE_DIR)
               / "pipeline/services/bulk_cell_lines.py").read_text()
        self.assertEqual(src.count("def resolve_parent("), 1)
        self.assertGreaterEqual(src.count("resolve_parent("), 3,
                                "plan and apply must both call it")


class AParentMayBeARowOfTheSamePasteTests(TestCase):
    """One helper called by both was still not the same question, because the two
    are asked at different moments.

    ``apply`` sorts wild types ahead of knockouts on purpose, so a WT and the
    knockouts made from it go in together — which is what both workbooks teach
    and what a gene's page asks for. The preview asked the database as it stands
    *before* the paste, where that WT does not exist, and answered from whatever
    else shares the name.

    Adding a Leicester ``U2OS`` wild type and a Leicester ``TRPA1`` knockout in
    one paste therefore previewed as *parent "U2OS" is McGill's line — check that
    is the one you mean*, and then saved onto the Leicester line, correctly. The
    save's own note said ``is your site's line``. Two notes about one write,
    disagreeing, and the one a person reads before committing was the wrong one.

    Same family as the supplier preview: a preview is only true if it asks what
    the write will ask.
    """

    databases = {"pipeline_db", "academy_db"}

    PASTE = ("name\tgene\tgenotype\tparent\n"
             "U2OS\tTRPA1\tKO\tU2OS\n"
             "U2OS\tNA\tWT\t\n")

    def setUp(self):
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.other = Site.objects.using(DB).create(name="McGill", short_code="MCG")
        self.client = _member_client(self, self.site)
        Target.objects.using(DB).create(gene_name="TRPA1")
        # The U2OS already on file, at the other site — what the stale warning
        # named. 17 real U2OS rows exist across the consortium.
        self.theirs = CellLine.objects.using(DB).create(
            name="U2OS", genotype="WT", site_id=self.other.pk)

    def _member(self):
        from pipeline.models import Member
        return Member.objects.using(DB).get(site_id=self.site.pk)

    def _plan(self, text=None):
        from pipeline.services import bulk_cell_lines
        rows = bulk_cell_lines.parse(text or self.PASTE)
        return bulk_cell_lines.plan(rows, member=self._member())

    def test_the_preview_names_the_row_that_supplies_the_parent(self):
        note = self._plan()[0]["note"]
        self.assertIn("row 2 of this paste", note)
        self.assertNotIn("McGill", note,
                         "the warning named a line the save will not use")

    def test_the_save_does_what_the_preview_said(self):
        from pipeline.services import bulk_cell_lines
        rows = bulk_cell_lines.parse(self.PASTE)
        bulk_cell_lines.apply(rows, member=self._member())
        ko = CellLine.objects.using(DB).get(genotype="KO", name="U2OS")
        self.assertIsNotNone(ko.parent_line_id, "the knockout lost its parent")
        self.assertEqual(ko.parent_line.site_id, self.site.pk)
        self.assertNotEqual(ko.parent_line.pk, self.theirs.pk)

    def test_the_warning_still_fires_when_the_wild_type_is_not_in_the_paste(self):
        """The case it was written for. Softening that is the opposite bug — a
        knockout quietly parented onto another institution's line."""
        note = self._plan("name\tgene\tgenotype\tparent\n"
                          "U2OS\tTRPA1\tKO\tU2OS\n")[0]["note"]
        self.assertIn("McGill", note)
        self.assertIn("check that is the one you mean", note)

    def test_a_blocked_wild_type_row_is_not_offered_as_a_parent(self):
        """`pending` is what will be *written*, not what was typed. A WT row the
        preview refuses never reaches the database, so a knockout naming it must
        still be told the parent is missing."""
        note = self._plan("name\tgene\tgenotype\tparent\n"
                          "HeLa\tTRPA1\tKO\tHeLa\n"
                          "HeLa\tTRPA1\tWT\t\n")[0]["note"]
        self.assertNotIn("of this paste", note)
        self.assertIn("no cell line called 'HeLa'", note)

    def test_a_pending_wild_type_is_named_when_nothing_is_on_file(self):
        """Otherwise the refusal reads "no cell line called 'HeLa'" directly over
        the row that is about to create exactly that."""
        note = self._plan("name\tgene\tgenotype\tparent\n"
                          "HeLa\tTRPA1\tKO\tHeLa\n"
                          "HeLa\tNA\tWT\t\n")[0]["note"]
        self.assertIn("row 2 of this paste", note)

    def test_the_parent_note_still_precedes_the_per_cell_refusals(self):
        """Moving the lookup into a second pass must not reorder the message."""
        note = self._plan("name\tgene\tgenotype\tparent\tc number\n"
                          "U2OS\tTRPA1\tKO\tU2OS\tC-RUN11-01\n"
                          "U2OS\tNA\tWT\t\t1111\n")[0]["note"]
        self.assertLess(note.index("parent"), note.index("C-RUN11-01"))

    def test_the_preview_rows_stay_json(self):
        """The second pass carries the C-number notes on the item between passes.
        A working key left on the row would 500 the whole response — the
        `arrived_with_ko` failure, one dict over."""
        import json
        for it in self._plan():
            self.assertNotIn("_cnum_notes", it)
            json.dumps({k: v for k, v in it.items() if k != "row"})


class TheGridStatesItsConventionsWhereTheyStayTests(TestCase):
    """A placeholder is hidden the moment its cell has a value.

    So the only statement of what `c number` or `parent` wants vanished as
    soon as you typed, and rows 2 and beyond never had one — which is how a
    blank `parent` placeholder between two numeric ones got read as "parent
    wants a C-number".

    The answer was a legend printed under the grid, and it was half right: the
    words stayed put, but they were the *same example a second time*, and the
    placeholders were still on row 0 reading as a filled-in first row. Both are
    one pinned `e.g.` row now — see `tests_conventions.py`.

    The parent convention itself was also wrong, which no amount of restating
    would have fixed: the lab writes C-numbers, 145 times out of 169.
    """

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)

    def test_the_example_stays_visible_as_a_row_of_its_own(self):
        js = (Path(settings.BASE_DIR) / "pipeline/static/pipeline/board.js").read_text()
        self.assertIn("function exampleRow", js)
        self.assertIn("exampleRow(cols, cfg.example)", js)
        self.assertNotIn("exampleLegend", js)

    def test_the_parent_convention_names_both_forms_it_is_written_in(self):
        for url in ("/pipeline/cell-lines/board/",):
            body = self.client.get(url).content.decode()
            self.assertIn("C-number", body,
                          f"{url} does not say how a parent is named")
            self.assertIn("<b>name</b>", body,
                          f"{url} does not say how a parent is named")
            self.assertNotIn("not a\n              C-number", body,
                             f"{url} still refuses the convention on file")

    def test_the_save_button_is_disabled_with_a_reason_rather_than_absent(self):
        """A first-time user faced a filled grid offering only Check these and
        Clear, with nothing to suggest a save button existed at all."""
        js = (Path(settings.BASE_DIR) / "pipeline/static/pipeline/board.js").read_text()
        self.assertIn("Check the rows first", js)
        self.assertNotIn('class="hidden px-4 py-2 rounded-lg bg-ycharos-600', js)

    def test_both_create_grids_take_a_note_as_the_rows_are_made(self):
        """Tagging a batch meant creating it and then editing every row, one cell
        at a time — twenty-four of them on this run."""
        from pipeline.views.imports import columns_and_example
        for kind in ("antibodies", "cell-lines"):
            with self.subTest(kind=kind):
                self.assertIn("comments", columns_and_example(kind)[0])

    def test_a_pasted_comment_reaches_the_record(self):
        from pipeline.models import Company, Member
        from pipeline.services import bulk_antibodies, bulk_cell_lines
        from pipeline.models import Antibody
        Company.objects.using(DB).create(name="abcam")
        target = Target.objects.using(DB).create(gene_name="ELP3")
        member = Member.objects.using(DB).get(site_id=self.site.pk)
        rows = bulk_antibodies.parse(
            "gene\tcatalogue\tcompany\tcomments\n"
            "ELP3\tA-ELP3-901\tabcam\t[COWORK RUN6] arc step 3\n")
        bulk_antibodies.apply(rows, member=member)
        ab = Antibody.objects.using(DB).get(catalogue_number="A-ELP3-901")
        self.assertIn("[COWORK RUN6]", ab.comments)

        rows = bulk_cell_lines.parse(
            "name\tgene\tgenotype\tcomments\n"
            "HAP1 ELP3 KO\tELP3\tKO\t[COWORK RUN6] from the Feb order\n")
        bulk_cell_lines.apply(rows, member=member)
        line = CellLine.objects.using(DB).get(name="HAP1 ELP3 KO")
        self.assertIn("[COWORK RUN6]", line.origin_comments)


class TwoLabsHAP1AreTellableApartTests(TestCase):
    """The WT dropdown offered "HAP1 (Wild Type)" twice.

    One was McGill's, one was ours, and the only cue was that Leicester does not
    use C-numbers — which you would have to already know. The KO list beside it
    was never ambiguous, which is what made the WT list read as a bug. The
    sessions board then rendered the pair as "HAP1 / HAP1".
    """

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        from pipeline.models import ExperimentSession, Member
        from django.contrib.auth.models import User
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.other = Site.objects.using(DB).create(name="McGill", short_code="MCG")
        self.client = _member_client(self, self.site)
        pu = User.objects.using(DB).get(username="carl")
        self.member = Member.objects.using(DB).get(user_id=pu.pk)
        self.target = Target.objects.using(DB).create(gene_name="ELP3")
        self.mine = CellLine.objects.using(DB).create(
            name="HAP1", genotype="WT", site_id=self.site.pk)
        CellLine.objects.using(DB).create(
            name="HAP1", genotype="WT", site_id=self.other.pk)
        self.ko = CellLine.objects.using(DB).create(
            name="HAP1 ELP3 KO", genotype="KO", target_id=self.target.pk,
            site_id=self.site.pk, parent_line_id=self.mine.pk)
        self.session = ExperimentSession.objects.using(DB).create(
            target_id=self.target.pk, site_id=self.site.pk, procedure_type="WB",
            date=date(2026, 7, 31), experimenter_id=self.member.pk,
            cell_line_wt_id=self.mine.pk, cell_line_ko_id=self.ko.pk)

    def test_the_dropdown_payload_carries_the_site(self):
        """The payload is two lists now, one per box — the page used to fetch
        one and split it by genotype in the browser, which is how a line that
        is not a knockout of anything reached the KO box
        (`cell_lines.session_options`). The site is still in every label.

        Each label now carries the line's own C-number as well as its batches
        (run 17). Both HAP1s read `[C-1]`, and that is right rather than a
        collision: C-numbers are issued per site, so Leicester's C-1 and
        McGill's C-1 are different lines — which is the whole reason the site is
        the half that tells them apart.
        """
        resp = self.client.get(
            f"/pipeline/session/new/?ajax=cell_lines&target={self.target.pk}")
        self.assertEqual(resp.status_code, 200, resp.content[:200])
        data = resp.json()
        mine = CellLine.objects.using(DB).get(pk=self.mine.pk)
        theirs = CellLine.objects.using(DB).get(name="HAP1", site_id=self.other.pk)
        wt = [c["label"] for c in data["wt"]]
        self.assertEqual(sorted(wt), sorted([
            f"HAP1 [C-{mine.c_number}] — Leicester",
            f"HAP1 [C-{theirs.c_number}] — McGill"]))
        ko = CellLine.objects.using(DB).get(pk=self.ko.pk)
        self.assertEqual([c["label"] for c in data["ko"]],
                         [f"HAP1 ELP3 KO [C-{ko.c_number}] — Leicester"])

    def test_the_sessions_board_names_which_lab_each_line_is(self):
        from pipeline.services import session_board
        row = session_board.row_for(
            session_board.board_queryset().get(pk=self.session.pk))
        self.assertEqual(row["cell_line_wt"], "HAP1 — Leicester")
        self.assertEqual(row["cell_line_ko"], "HAP1 ELP3 KO — Leicester")

    def test_a_knockout_stored_under_the_parents_bare_name_is_still_told_apart(self):
        """The case the live data actually had, and the one the site alone does
        not fix: run 6's knockout was stored as plain `HAP1`, so the WT and the
        KO made from it printed the same word twice. Adding only the site turned
        "HAP1 / HAP1" into "HAP1 — Leicester" twice — longer, no clearer. The
        session form's dropdown has appended the gene to KO rows since it was
        written; the board does the same now."""
        from pipeline.services import session_board
        bare = CellLine.objects.using(DB).create(
            name="HAP1", genotype="KO", target_id=self.target.pk,
            site_id=self.site.pk, parent_line_id=self.mine.pk)
        self.session.cell_line_ko_id = bare.pk
        self.session.save(using=DB)
        row = session_board.row_for(
            session_board.board_queryset().get(pk=self.session.pk))
        self.assertEqual(row["cell_line_wt"], "HAP1 — Leicester")
        self.assertEqual(row["cell_line_ko"], "HAP1 ELP3 KO — Leicester")
        self.assertNotEqual(row["cell_line_wt"], row["cell_line_ko"])

    def test_a_gene_already_in_the_name_is_not_said_twice(self):
        from pipeline.services import session_board
        row = session_board.row_for(
            session_board.board_queryset().get(pk=self.session.pk))
        self.assertEqual(row["cell_line_ko"], "HAP1 ELP3 KO — Leicester")

    def test_naming_the_line_does_not_cost_a_query_per_row(self):
        from django.test.utils import CaptureQueriesContext
        from django.db import connections
        from pipeline.services import session_board
        for _ in range(6):
            from pipeline.models import ExperimentSession
            ExperimentSession.objects.using(DB).create(
                target_id=self.target.pk, site_id=self.site.pk, procedure_type="IP",
                date=date(2026, 7, 31), experimenter_id=self.member.pk,
                cell_line_wt_id=self.mine.pk, cell_line_ko_id=self.ko.pk)
        with CaptureQueriesContext(connections[DB]) as ctx:
            session_board.board_rows()
        self.assertLess(len(ctx.captured_queries), 12,
                        "the WT/KO labels are costing a query per row")


class TheWorkbookKeepsEveryValueOfAnInventedColumnTests(TestCase):
    """The bench workbook kept an invented column — and quietly kept one value.

    Three antibody rows carried three different page references under a column
    called "Owner notes"; the first was stored as a session-level condition and
    the other two were dropped, with the preview reporting "10 conditions" and
    the save message reporting nothing.

    Run 7 fixed the silence: the preview named the column and counted what it
    was about to drop. Run 11 said the obvious next thing about the same
    behaviour on a nine-well plate map — the first row is not a summary of the
    plate, it is the first well — so a condition now holds every distinct value
    the column carried. Keeping the column was right; keeping one of it was not.
    """
    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        from pipeline.models import Antibody, Company, ExperimentSession, Member
        from django.contrib.auth.models import User
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)
        pu = User.objects.using(DB).get(username="carl")
        self.member = Member.objects.using(DB).get(user_id=pu.pk)
        self.target = Target.objects.using(DB).create(gene_name="STMN2")
        company = Company.objects.using(DB).create(name="Proteintech")
        for cat in ("CAT-1", "CAT-2"):
            Antibody.objects.using(DB).create(
                target_id=self.target.pk, company_id=company.pk,
                catalogue_number=cat, site_id=self.site.pk)
        wt = CellLine.objects.using(DB).create(
            name="SH-SY5Y", genotype="WT", site_id=self.site.pk)
        CellLine.objects.using(DB).create(
            name="SH-SY5Y STMN2 KO", genotype="KO", target_id=self.target.pk,
            parent_line_id=wt.pk, site_id=self.site.pk)

    def _workbook(self, per_row):
        """The real download, with an invented column carrying a different value
        on each antibody row."""
        import io
        import openpyxl
        from pipeline.services import session_template as tmpl
        wb = openpyxl.load_workbook(io.BytesIO(tmpl.build_gene_template("STMN2")))
        ws = wb["WB"]
        header = [c.value for c in ws[1]]
        header.append("Owner notes")
        ws.cell(1, len(header), "Owner notes")
        for i, value in enumerate(per_row, start=2):
            ws.cell(i, header.index("signal") + 1, "21 kDa band")
            ws.cell(i, len(header), value)
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def _post(self, url, data):
        from django.core.files.uploadedfile import SimpleUploadedFile
        return self.client.post(url, {"file": SimpleUploadedFile(
            "session_template_STMN2.xlsx", data,
            content_type="application/vnd.openxmlformats-officedocument"
                         ".spreadsheetml.sheet")})

    def test_the_preview_names_the_column_and_what_the_condition_will_hold(self):
        d = self._post("/pipeline/session/template/upload/preview/",
                       self._workbook(["p.14", "p.15"])).json()
        self.assertTrue(d["ok"], d)
        unknown = d["sessions"][0]["unknown_columns"]
        self.assertEqual(len(unknown), 1)
        self.assertEqual(unknown[0]["header"], "Owner notes")
        self.assertEqual(unknown[0]["value"], "p.14 | p.15")
        self.assertEqual(unknown[0]["values"], 2)
        self.assertEqual(d["counts"]["unknown_columns"], 1)

    def test_a_column_whose_rows_agree_reads_as_one_value(self):
        """Joining must not turn an ordinary column into a list of duplicates."""
        d = self._post("/pipeline/session/template/upload/preview/",
                       self._workbook(["p.14", "p.14"])).json()
        unknown = d["sessions"][0]["unknown_columns"]
        self.assertEqual(unknown[0]["value"], "p.14")
        self.assertEqual(unknown[0]["values"], 1)

    def test_the_promoted_key_is_snake_case_like_every_other(self):
        from pipeline.models import ExperimentSession
        self._post("/pipeline/session/template/upload/commit/",
                   self._workbook(["p.14", "p.15"]))
        session = ExperimentSession.objects.using(DB).get()
        # Both rows, not the first: naming a loss is not preventing one.
        self.assertEqual(session.session_conditions.get("owner_notes"), "p.14 | p.15")
        self.assertNotIn("owner notes", session.session_conditions)

    def test_the_preview_says_the_gene_already_has_a_session(self):
        """Planning a session through the wizard and then downloading the
        workbook from the same page looks like one task in two halves. It is
        two: the sheet cannot fill a planned session, and made a second."""
        from pipeline.models import ExperimentSession
        planned = ExperimentSession.objects.using(DB).create(
            target_id=self.target.pk, site_id=self.site.pk, procedure_type="WB",
            date=date(2026, 7, 31), experimenter_id=self.member.pk,
            status=ExperimentSession.SessionStatus.PLANNED)
        d = self._post("/pipeline/session/template/upload/preview/",
                       self._workbook(["p.14"])).json()
        existing = d["sessions"][0]["existing"]
        self.assertEqual(existing["total"], 1)
        self.assertEqual([p["id"] for p in existing["planned"]], [planned.pk])

    def test_the_upload_panel_shows_all_three_and_does_not_ask_for_a_reload(self):
        body = self.client.get(
            f"/pipeline/target/{self.target.pk}/").content.decode()
        self.assertIn("not recognised", body)
        self.assertIn("always creates new sessions", body)
        # The rendered message, not the note beside it: a comment inside an
        # inline script is for whoever edits the file next, and this one quotes
        # the wording it replaced.
        self.assertNotIn(">Reload the page to see them.</p>", body)
        self.assertIn(">Show them</button>", body)

    def test_an_empty_condition_does_not_print_an_example_as_a_value(self):
        """SESSION CONDITIONS printed "NUMBER OF GELS / e.g. 3" for a condition
        nobody had filled in, so a reader could not tell three gels from nobody
        having said."""
        body = self.client.get("/pipeline/sessions/board/").content.decode()
        self.assertIn("function emptyPrompt", body)
        self.assertNotIn("empty: c.placeholder || 'add'", body)


class TheWrongSessionsBenchSheetIsRefusedAtTheDoorTests(TestCase):
    """Run 11, at the endpoint the owner actually posted to.

    Two sessions of NR3C2 — a western blot and an immunofluorescence — both made
    from one workbook. The IF plate map was fed to the WB session and the check
    answered *"3 antibodies read from the sheet"* and offered **Record these**.

    Worth pinning here as well as in the service tests, because this door does
    two things first that could each swallow the guard: it sniffs the upload to
    choose between two parsers (which reads the stream to the end), and it parses
    with the *session's* procedure rather than the sheet's.
    """

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        import datetime
        from django.contrib.auth.models import User
        from pipeline.models import Antibody, Company, ExperimentSession, Member
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)
        pu = User.objects.using(DB).get(username="carl")
        self.member = Member.objects.using(DB).get(user_id=pu.pk)
        self.target = Target.objects.using(DB).create(gene_name="NR3C2")
        company = Company.objects.using(DB).create(name="Proteintech")
        for cat in ("CAT-1", "CAT-2", "CAT-3"):
            Antibody.objects.using(DB).create(
                target_id=self.target.pk, company_id=company.pk,
                catalogue_number=cat, site_id=self.site.pk)
        wt = CellLine.objects.using(DB).create(
            name="HeLa", genotype="WT", site_id=self.site.pk)
        ko = CellLine.objects.using(DB).create(
            name="HeLa NR3C2 KO B4", genotype="KO", target_id=self.target.pk,
            parent_line_id=wt.pk, site_id=self.site.pk)
        common = dict(target_id=self.target.pk, site_id=self.site.pk,
                      experimenter_id=self.member.pk, date=datetime.date(2026, 8, 4),
                      cell_line_wt_id=wt.pk, cell_line_ko_id=ko.pk)
        self.wb_session = ExperimentSession.objects.using(DB).create(
            procedure_type="WB", **common)
        self.if_session = ExperimentSession.objects.using(DB).create(
            procedure_type="IF", **common)

    def _sheet_bytes(self, session):
        import io
        from pipeline.services.planning import generate_bench_sheet
        buf = io.BytesIO()
        generate_bench_sheet(session).save(buf)
        return buf.getvalue()

    def _upload_to(self, session, data):
        from django.core.files.uploadedfile import SimpleUploadedFile
        return self.client.post(
            f"/pipeline/session/{session.pk}/results/upload/",
            {"file": SimpleUploadedFile(
                "bench_sheet.xlsx", data,
                content_type="application/vnd.openxmlformats-officedocument"
                             ".spreadsheetml.sheet")}).json()

    def test_the_if_plate_map_is_refused_by_the_western_blot_session(self):
        d = self._upload_to(self.wb_session, self._sheet_bytes(self.if_session))
        self.assertIs(d["ok"], False)
        self.assertIn(f"#{self.if_session.pk}", d["error"])
        self.assertNotIn("items", d)

    def test_the_if_plate_map_is_still_accepted_by_its_own_session(self):
        """The guard has to refuse the wrong sheet without refusing the right
        one — a check that says no to everything is not a check."""
        d = self._upload_to(self.if_session, self._sheet_bytes(self.if_session))
        self.assertIs(d["ok"], True)
        self.assertEqual(d["kind"], "bench")
        self.assertEqual(d["sheet_session_id"], self.if_session.pk)
        self.assertEqual(d["sheet_procedure"], "IF")
        self.assertEqual(d["unstamped"], "")

    def test_the_western_blot_sheet_is_refused_by_the_if_session(self):
        d = self._upload_to(self.if_session, self._sheet_bytes(self.wb_session))
        self.assertIs(d["ok"], False)
        self.assertIn(f"#{self.wb_session.pk}", d["error"])


class TheSearchBoxReachesWhatItPromisesTests(TestCase):
    """The box said "catalogue number" and a cell line's catalogue number matched
    nothing.

    Four hand-written field lists had drifted apart: antibodies matched on clone
    id but not lot or comments, sessions matched on comments, cell lines on name
    alone. From outside that reads as broken rather than narrow.
    """

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        from pipeline.models import Antibody, Company, Member
        from django.contrib.auth.models import User
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)
        pu = User.objects.using(DB).get(username="carl")
        self.member = Member.objects.using(DB).get(user_id=pu.pk)
        self.target = Target.objects.using(DB).create(gene_name="ELP3")
        company = Company.objects.using(DB).create(name="Horizon Discovery")
        Antibody.objects.using(DB).create(
            target_id=self.target.pk, company_id=company.pk,
            catalogue_number="A-ELP3-001", lot_number="RUN6-A1",
            site_id=self.site.pk)
        CellLine.objects.using(DB).create(
            name="HAP1 ELP3 KO", genotype="KO", target_id=self.target.pk,
            site_id=self.site.pk, company_id=company.pk,
            catalogue_number="HZGHC-ELP3-RUN6", lot_number="RUN6-L1")

    def _keys(self, term):
        from pipeline.services import find
        return {g["key"]: g for g in find.find(term)["groups"]}

    def test_an_antibody_is_found_by_its_lot(self):
        self.assertIn("antibodies", self._keys("RUN6-A1"))

    def test_a_cell_line_is_found_by_its_catalogue_number(self):
        groups = self._keys("HZGHC-ELP3-RUN6")
        self.assertIn("cell_lines", groups)
        self.assertTrue(groups["cell_lines"]["rows"][0]["why"],
                        "a cell-line hit must say which field matched")

    def test_a_cell_line_is_found_by_its_lot(self):
        self.assertIn("cell_lines", self._keys("RUN6-L1"))

    def test_a_run_tag_finds_both_kinds_of_record(self):
        groups = self._keys("RUN6")
        self.assertIn("antibodies", groups)
        self.assertIn("cell_lines", groups)

    def test_the_footer_describes_what_is_actually_searched(self):
        body = self.client.get("/pipeline/find/?q=RUN6").content.decode()
        self.assertNotIn("antibodies (catalogue, RRID, clone), cell lines", body)
        self.assertIn("lot", body)

    def test_the_search_still_makes_no_outbound_call(self):
        from pipeline.tests_timeouts import no_network
        with no_network("searching"):
            self.client.get("/pipeline/find/?q=RUN6")


class ASessionIsFoundByItsNumberTests(TestCase):
    """A session's number is the one identifier this app prints on paper, and
    the search box could not find it.

    It is stamped on all four bench sheets, on every row of the per-gene
    workbook (`WB · NR3C2 #504`), and it is what the app quotes when a sheet is
    uploaded into the wrong session — so `#504` is on the desk of anybody
    recording a result. Searching `501` returned two genes matched on a UniProt
    ID, thirty-nine antibodies matched on catalogue substrings, and no Sessions
    section at all.
    """

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        from pipeline.models import ExperimentSession, Member
        from django.contrib.auth.models import User

        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)
        pu = User.objects.using(DB).get(username="carl")
        self.member = Member.objects.using(DB).get(user_id=pu.pk)
        self.target = Target.objects.using(DB).create(gene_name="NR3C2")
        self.session = ExperimentSession.objects.using(DB).create(
            target_id=self.target.pk, site_id=self.site.pk,
            procedure_type="WB", date=date(2026, 8, 5),
            experimenter_id=self.member.pk)

    def _keys(self, term):
        from pipeline.services import find
        return {g["key"]: g for g in find.find(term)["groups"]}

    def test_the_bare_number_finds_it(self):
        groups = self._keys(str(self.session.pk))
        self.assertIn("sessions", groups)
        self.assertEqual(groups["sessions"]["rows"][0]["url"],
                         f"/pipeline/sessions/board/?open={self.session.pk}")

    def test_the_number_as_the_sheets_write_it_finds_it(self):
        """Every artefact prints it with the hash, so that is what gets typed."""
        self.assertIn("sessions", self._keys(f"#{self.session.pk}"))

    def test_the_hit_says_the_number_is_why(self):
        """The title is the gene and the procedure, so a number match is not
        visible in it — and a hit that explains nothing reads as a hit on
        nothing. Widening the query means widening `_why` in the same commit."""
        row = self._keys(str(self.session.pk))["sessions"]["rows"][0]
        self.assertIn("session number", row["why"])

    def test_the_boards_own_search_box_agrees(self):
        """Two search boxes in one app answering differently reads as one of
        them being broken, with no way to tell which from outside — the whole
        reason `services/find.py` owns the `Q` builders."""
        from pipeline.services import session_board

        rows = session_board.board_rows(q=str(self.session.pk))
        self.assertEqual([r["id"] for r in rows], [self.session.pk])

    def test_a_number_too_big_for_the_column_matches_nothing(self):
        """`pk` is a 32-bit integer and PostgreSQL raises on a value outside it
        rather than matching nothing, so a long run of digits typed into the box
        in the chrome would 500 every page it sits on."""
        self.assertNotIn("sessions", self._keys("99999999999999999999"))
        self.assertEqual(
            self.client.get("/pipeline/find/?q=99999999999999999999").status_code,
            200)


class ActiveMeansWorkHasStartedTests(TestCase):
    """Overview's "Active" read `Target.status`, which nothing advances.

    So a site that added two targets, ran seven sessions and recorded
    twenty-two results watched its figure stay on 6 all afternoon, and neither
    gene appeared in the ACTIVE TARGETS list. The funnel three lines above it on
    the same page had the right answer already.
    """

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        from pipeline.models import Antibody, Company
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)
        self.company = Company.objects.using(DB).create(name="abcam")
        # A gene added the way feasibility adds one: status untouched.
        self.worked_on = Target.objects.using(DB).create(gene_name="ELP3")
        TargetNomination.objects.using(DB).create(
            target_id=self.worked_on.pk, site_id=self.site.pk)
        Antibody.objects.using(DB).create(
            target_id=self.worked_on.pk, company_id=self.company.pk,
            catalogue_number="A-ELP3-001", site_id=self.site.pk)

    def test_a_gene_with_antibodies_on_file_is_active(self):
        resp = self.client.get("/pipeline/overview/")
        genes = {t.gene_name for t in resp.context["active_targets"]}
        self.assertIn("ELP3", genes)
        self.assertGreaterEqual(resp.context["overall_stats"]["active"], 1)

    def test_the_sites_own_figure_moves_too(self):
        resp = self.client.get("/pipeline/overview/")
        summaries = {s["site"].name: s for s in resp.context["site_summaries"]}
        self.assertEqual(summaries["Leicester"]["active_targets"], 1)

    def test_a_reported_gene_is_no_longer_active(self):
        from pipeline.models import Report
        Report.objects.using(DB).create(target_id=self.worked_on.pk)
        resp = self.client.get("/pipeline/overview/")
        genes = {t.gene_name for t in resp.context["active_targets"]}
        self.assertNotIn("ELP3", genes)

    def test_the_page_no_longer_calls_it_funded(self):
        """It never counted funding. Saying so sent the reader looking for a
        funding flag to set, which is not what would have moved the number."""
        body = self.client.get("/pipeline/overview/").content.decode()
        self.assertNotIn("Funded and in progress", body)


class TheTargetBoardSaysHowASecondSiteGetsAGeneTests(TestCase):
    """The board promises a duplicate flag for "the same gene on more than one
    site's list" and never said where such a row comes from.

    The first answer was that a site nominates for itself, from Check
    feasibility. It was true and it was a description of a limitation: neither
    door let you name the site, so nobody could allocate a gene to a bench —
    the owner's ask, 4 Aug. Add targets carries a site now, so the sentence
    points at the control on this page rather than at another one, and the SITES
    cell keeps its own job: typing over it *moves* a gene rather than adding it,
    on a shared master list, with a green save flash and no wording.
    """

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)

    def test_the_board_says_where_a_second_nomination_comes_from(self):
        body = self.client.get("/pipeline/targets/board/").content.decode()
        self.assertIn("second site's list", body)
        self.assertIn("Add targets", body)

    def test_the_sites_cell_says_that_typing_over_it_moves_the_gene(self):
        body = self.client.get("/pipeline/targets/board/").content.decode()
        self.assertIn("moves this gene off", body)


class TheIdentityDialogTakesEnterTests(TestCase):
    """Enter saves in every cell on every board, and did nothing at all in the
    one pop-out where the change is most deliberate.

    Same shape as the Escape gap already fixed on this dialog: the grid teaches
    the habit and the dialog ignored it. No response shows this, so it is pinned
    at the source, which is the weaker option — the alternative is a guard that
    cannot fail.
    """

    databases = {"pipeline_db", "academy_db"}

    def test_enter_in_an_identity_field_saves(self):
        js = (Path(settings.BASE_DIR) / "pipeline/static/pipeline/board.js").read_text()
        self.assertIn("async function saveChange()", js)
        self.assertIn("el('save').addEventListener('click', saveChange)", js)
        self.assertIn("input[data-f]", js)


class TheIdentityDialogGoesThroughTheSupplierWriterTests(TestCase):
    """One identity edit could retire the Bio-Techne split for everybody.

    The dialog called `Company.resolve` directly, which matches on the canonical
    key and **creates** when nothing matches — and `canonical_key("Bio-Techne")`
    matches neither bracketed row. So correcting a supplier to a bare
    "Bio-Techne" minted a third Company; and because an exact name wins over
    every heuristic in `resolve_company`, that row then answered every later
    paste and the catalogue-prefix split became unreachable, site-wide, for good.

    Found by the triage rather than by the run — the field test hit the ambiguity
    (F14) but never opened the dialog on it.
    """

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        from pipeline.models import Antibody, Company
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)
        self.target = Target.objects.using(DB).create(gene_name="ELP3")
        self.novus = Company.objects.using(DB).create(
            name="Bio-Techne (Novus Biologicals)")
        self.rnd = Company.objects.using(DB).create(name="Bio-Techne (R&D Systems)")
        self.other = Company.objects.using(DB).create(name="abcam")
        self.ab = Antibody.objects.using(DB).create(
            target_id=self.target.pk, company_id=self.other.pk,
            catalogue_number="NB100-1234", site_id=self.site.pk)

    def _save(self, **extra):
        return self.client.post("/pipeline/antibodies/board/identity/save/", dict(
            {"antibody_id": self.ab.pk, "catalogue_number": self.ab.catalogue_number,
             "company": "Bio-Techne", "gene": "ELP3"}, **extra))

    def test_a_bare_bio_techne_is_split_by_the_catalogue_not_created(self):
        from pipeline.models import Company
        resp = self._save()
        self.assertEqual(resp.status_code, 200, resp.content)
        self.ab.refresh_from_db()
        # NB… is Novus. The point is which of the two, not that one was chosen.
        self.assertEqual(self.ab.company_id, self.novus.pk)
        self.assertFalse(
            Company.objects.using(DB).filter(name__iexact="Bio-Techne").exists(),
            "a third bare Bio-Techne row was created, which disables the split")

    def test_a_non_nb_catalogue_goes_to_the_other_brand(self):
        resp = self._save(catalogue_number="AF1234")
        self.assertEqual(resp.status_code, 200, resp.content)
        self.ab.refresh_from_db()
        self.assertEqual(self.ab.company_id, self.rnd.pk)

    def test_an_ordinary_vendor_still_resolves_case_insensitively(self):
        """`resolve_company` ends in `Company.resolve` on the ordinary path, so
        nothing else about supplier dedup may change."""
        from pipeline.models import Company
        before = Company.objects.using(DB).count()
        self._save(company="ABCAM")
        self.ab.refresh_from_db()
        self.assertEqual(self.ab.company_id, self.other.pk)
        self.assertEqual(Company.objects.using(DB).count(), before)

    def test_a_cell_lines_supplier_takes_the_same_route(self):
        src = (Path(settings.BASE_DIR) / "pipeline/services/identity.py").read_text()
        self.assertNotIn("Company.resolve(company_name", src,
                         "an identity path still bypasses the brand split")


class ARedrawNeverPaintsAStaleSnapshotTests(TestCase):
    """Two cell saves on one row produce two whole-row snapshots, and the screen
    showed whichever *reply* landed last rather than whichever *write* did.

    Holding a redraw while another cell is open made that deterministic rather
    than merely racy — the parked, older payload is replayed at the end of the
    second cell's save, which is after the newer one has been painted. So editing
    two cells in a row quickly put the first edit's old value back on screen,
    saved and invisible. The inverse of the rule the file exists for: nothing on
    screen may imply a save that did not happen, and nothing may hide one that
    did. Browser-only, so pinned at the source.
    """

    databases = {"pipeline_db", "academy_db"}

    def test_a_redraw_is_stamped_and_a_stale_one_is_dropped(self):
        js = (Path(settings.BASE_DIR) / "pipeline/static/pipeline/board.js").read_text()
        self.assertIn("const painted = new Map()", js)
        self.assertIn("function replaceRow(id, data, n)", js)
        self.assertIn("if (n <= (painted.get(String(id)) || 0)) return;", js)

    def test_the_stamp_is_taken_before_the_request_is_awaited(self):
        """The order the saves were *started* in is what decides which snapshot
        is newer — taking it after the await would stamp them in reply order,
        which is the bug."""
        js = (Path(settings.BASE_DIR) / "pipeline/static/pipeline/board.js").read_text()
        save = js[js.index("async function save(body, id) {"):]
        save = save[:save.index("\n    }")]
        self.assertLess(save.index("++seq"), save.index("await patch"))

    def test_a_held_redraw_carries_its_stamp(self):
        js = (Path(settings.BASE_DIR) / "pipeline/static/pipeline/board.js").read_text()
        self.assertIn("heldRedraw.set(String(id), {data, n})", js)
        self.assertIn("replaceRow(id, held.data, held.n)", js)


class TheTwoBenchSheetsAreToldApartTests(TestCase):
    """Two artefacts, near-identical names, opposite jobs.

    The per-gene **bench workbook** captures a week nobody has recorded and
    always creates new sessions; the per-session **bench sheet** on the sessions
    board belongs to one session that exists and records into it. The gene page
    offered the first a few centimetres below a list of that gene's planned
    sessions and distinguished them nowhere, so the run planned a session and
    then made a second one filling in what looked like its sheet.
    """

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)
        self.target = Target.objects.using(DB).create(gene_name="ELP3")

    def test_the_gene_page_says_which_sheet_does_which(self):
        body = self.client.get(f"/pipeline/target/{self.target.pk}/").content.decode()
        self.assertIn("becomes a <em>new</em> session", body)
        self.assertIn("bench sheet", body)


class AHandTypedConditionColumnReachesItsKeyTests(TestCase):
    """The real reason the promoted key had to be snake_case.

    A site with no default protocol template gets a workbook shipping **no**
    `cond:` columns at all, so a scientist adds their own — "Lysis Buffer",
    "Gel Chemistry". Lower-casing alone gave `"lysis buffer"`, which is not the
    registry key `lysis_buffer`, so the value was stored where nothing reads it
    and `report_generator` substituted its default into a generated Method
    paragraph. Snake-casing the promoted key is what connects them.
    """

    databases = {"pipeline_db", "academy_db"}

    def test_a_spaced_title_case_header_lands_on_the_registry_key(self):
        from pipeline.services import session_import
        from pipeline.services.session_board import condition_field_names
        for header, key in (("Lysis Buffer", "lysis_buffer"),
                            ("Gel Chemistry", "gel_chemistry"),
                            ("cond:lysis_buffer", "lysis_buffer")):
            with self.subTest(header=header):
                self.assertEqual(session_import._condition_key(header), key)
        # And that key is one the app actually reads, not just a tidier string.
        self.assertIn("lysis_buffer", condition_field_names("WB"))


class EveryPasteSurfaceOffersItsTemplateTests(TestCase):
    """A blank spreadsheet, wherever you can paste one in.

    ``views/imports.py::import_template`` builds it from ``columns_and_example``
    — the same list behind each grid's headings and behind the paste parser —
    and was **routed and reachable from nowhere**: not one template or script
    mentioned the URL. So the boards could be typed into or pasted into, and the
    file most people would rather work in existed and could not be got at.

    The mirror of the rule about an export with no importer. A routed endpoint is
    not a reachable one, and only a test that reads the pages can tell the
    difference.
    """

    databases = {"pipeline_db", "academy_db"}

    # Every surface with a paste box, and the import kind its columns come from.
    #
    # `feasibility.html` is deliberately **not** here any more. Its Bulk Add
    # Targets box is a paste surface, but the owner's review said of its
    # template: "An excel sheet to upload a one column list is not needed." A
    # spreadsheet whose only column is a gene symbol is a worse way to type gene
    # symbols than the box directly above it, and it is the one template that
    # carries no convention worth stating — which is what the others are for.
    # The targets template itself stays, offered by the target board.
    SURFACES = {
        "antibody_board.html": "antibodies",
        "cell_line_board.html": "cell-lines",
        "session_board.html": "sessions",
        "target_board.html": "targets",
        "target_detail.html": "antibodies",     # plus cell-lines, checked below
    }

    # A pop-out that legitimately has no blank spreadsheet, and why. Keyed by
    # (template, kind of mount) — each of these files mounts one of each.
    #
    # An exemption is a *decision*, so it is written down next to the thing it
    # exempts. An empty reason is not allowed, which is what stops this dict
    # becoming the place failures go to be silenced.
    EXEMPT_MOUNTS = {
        ("user_board.html", "newEntry"):
            "There is no `people` import kind and there should not be: a working "
            "account is three rows in two databases (services/members.py::grant), "
            "not a spreadsheet row.",
        ("_bench_workbook_upload.html", "uploadPanel"):
            "Its counterpart is the per-gene bench workbook — pre-filled, stamped "
            "with the gene and the session number. A blank version of it would "
            "carry neither, and `session_import` refuses a sheet that names no "
            "session. The download is offered on both host pages.",
        ("target_detail.html", "newEntry"):
            "A factory: it forwards a per-kind cfg the caller supplies, and both "
            "callers pass a templateUrl. Pinned by the count test below.",
        ("target_detail.html", "uploadPanel"):
            "Same factory shape as the Add panel above it.",
    }

    # `OGABoard.newEntry(` / `uploadPanel(` through to its closing paren, so a
    # config spread over two arguments is read whole.
    _MOUNT = re.compile(r"OGABoard\.(newEntry|uploadPanel)\(")

    def _source(self, name):
        return (Path(settings.BASE_DIR) / "pipeline/templates/pipeline" / name).read_text()

    def _mounts(self, source):
        """`(kind_of_mount, its config source)` for every pop-out a page mounts."""
        for m in self._MOUNT.finditer(source):
            start, depth, i = m.end() - 1, 0, m.end() - 1
            while i < len(source):
                if source[i] == "(":
                    depth += 1
                elif source[i] == ")":
                    depth -= 1
                    if depth == 0:
                        break
                i += 1
            yield m.group(1), source[start:i + 1]

    def test_every_pop_out_that_mounts_carries_its_template(self):
        """The file-granular test above passes on one link anywhere in the page.

        That is exactly how a gap survived: `session_board.html` satisfied it
        with its Add panel's link while its **upload** panel had none, and
        nothing could tell the two apart from outside. This asks each mount.
        """
        root = Path(settings.BASE_DIR) / "pipeline/templates/pipeline"
        seen = 0
        for path in sorted(root.glob("*.html")):
            for kind, config in self._mounts(path.read_text()):
                seen += 1
                with self.subTest(template=path.name, mount=kind):
                    if "templateUrl" in config:
                        continue
                    reason = self.EXEMPT_MOUNTS.get((path.name, kind), "")
                    self.assertTrue(
                        reason,
                        f"{path.name}'s {kind} takes rows and offers no blank "
                        f"spreadsheet. Add `templateUrl`, or add it to "
                        f"EXEMPT_MOUNTS with the reason.")
        self.assertGreaterEqual(seen, 10,
                                "the mount scanner stopped finding pop-outs — "
                                "check the regex before believing this passed")

    def test_the_gene_pages_two_factories_are_called_with_a_template_each(self):
        """`target_detail.html` mounts its panels through two factories, so the
        per-mount test above cannot see the links — they are at the call sites.
        Four panels, four templates: antibodies and cell lines, add and upload."""
        source = self._source("target_detail.html")
        self.assertEqual(source.count("templateUrl:"), 4, source.count("templateUrl:"))

    def test_the_gene_page_offers_its_blank_sheets_without_opening_a_pop_out(self):
        """A link only `newEntry` renders is a link you must already have
        decided to use. Somebody standing at "Antibodies (0)", working out how
        to record twenty vials, cannot see that a spreadsheet is an option — and
        this page's own panels tell you to "bring it back with the upload
        button", which is the half that was already on the page."""
        site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        target = Target.objects.using(DB).create(gene_name="TRPA1")
        client = _member_client(self, site)
        body = client.get(f"/pipeline/target/{target.pk}/").content.decode()
        # The links, in the page itself rather than in a script that mounts later.
        static = re.sub(r"(?s)<script.*?</script>", "", body)
        for kind in ("antibodies", "cell-lines"):
            self.assertIn(f"/pipeline/import/template/{kind}/", static,
                          f"the {kind} panel's blank sheet is behind a pop-out")
        # And each says it arrived, like every other download here.
        for receipt in ("ab-template-receipt", "cl-template-receipt"):
            self.assertIn(receipt, static)

    def test_the_sessions_upload_panel_offers_the_sheet_it_asks_for(self):
        """It said "Download the sheet, edit it, upload it back" and gave the
        reader nothing to click — the Download button was outside the panel it
        had just told them to open.

        And the sheet it needs is **not** the blank template the Add panel on
        the same board carries: this endpoint matches on `session_id` +
        `result_id` and only ever edits, so a blank sheet reads as "0 rows" with
        no error. Two surfaces on one board, two artefacts.
        """
        source = self._source("session_board.html")
        panel = source[source.index('id="upload-panel"'):source.index('id="filters"')]
        self.assertIn("upload-dl-btn", panel,
                      "the upload panel names a sheet it does not offer")
        self.assertNotIn("import_template", panel,
                         "a blank template here parses as 0 rows — this "
                         "endpoint only edits rows the app has stamped")
        # One function behind both doors, so the panel cannot lose the filters.
        self.assertEqual(source.count("addEventListener('click', downloadBoard)"), 2)
        self.assertEqual(source.count("function downloadBoard()"), 1)

    # A board page that must offer the blank sheet in the page itself, and the
    # kind it is. The gene page has its own test above; these are the boards.
    STATIC_LINK_PAGES = {
        "cell_line_board.html": "cell-lines",
        "antibody_board.html": "antibodies",
        "target_board.html": "targets",
    }

    # Written down rather than left out, same rule as EXEMPT_MOUNTS above.
    EXEMPT_PAGES = {
        "session_board.html":
            "Two surfaces, two artefacts. Its Add panel takes the blank template "
            "and its Upload panel only ever edits rows the app has stamped, so a "
            "blank sheet there reads as 0 rows with no error. A link in the "
            "header would sit between them naming neither.",
    }

    def test_each_board_offers_its_blank_sheet_without_opening_a_pop_out(self):
        """A link only a pop-out renders is a link you must already have decided
        to use.

        The gene page was fixed for this reason and the boards were not, so
        somebody looking at the cell lines board and wondering whether they
        could work in Excel had no way to find out that they could — while
        uOttawa, who had the file, were emailing to ask what a column meant.
        Both halves of one gap: the sheet is reachable, and it explains itself
        (`imports.TEMPLATE_NOTES`).
        """
        root = Path(settings.BASE_DIR) / "pipeline/templates/pipeline"
        for name, kind in self.STATIC_LINK_PAGES.items():
            with self.subTest(template=name):
                source = (root / name).read_text()
                # Outside every <script>, so a mounted panel's copy cannot
                # satisfy it — that is exactly what hid this.
                static = re.sub(r"(?s)<script.*?</script>", "", source)
                self.assertIn(f"kind='{kind}'", static,
                              f"{name}'s blank sheet is behind a pop-out")
                # And it says it arrived, like every other download here.
                self.assertIn("data-receipt=", static)
        for name, reason in self.EXEMPT_PAGES.items():
            self.assertTrue(reason, name)

    def test_the_header_template_links_are_wired_to_a_receipt(self):
        """`data-receipt` is half of it: `downloadReceipts()` is what turns the
        attribute into a line of text. The page marked one and never called it
        on data_io.html, and the attribute did nothing at all."""
        root = Path(settings.BASE_DIR) / "pipeline/templates/pipeline"
        for name in self.STATIC_LINK_PAGES:
            with self.subTest(template=name):
                source = (root / name).read_text()
                if 'data-receipt="header-template-receipt"' not in source:
                    continue
                self.assertIn('id="header-template-receipt"', source,
                              f"{name} marks a receipt anchor with nowhere to "
                              f"write the receipt")
                self.assertIn("OGABoard.downloadReceipts()", source,
                              f"{name} marks a receipt anchor and never wires it")

    def test_every_paste_surface_names_the_template_endpoint(self):
        for name, kind in self.SURFACES.items():
            with self.subTest(template=name):
                source = self._source(name)
                self.assertIn("pipeline:import_template", source,
                              f"{name} has a paste box and no downloadable template")
                self.assertIn(f"kind='{kind}'", source,
                              f"{name} links a template of the wrong kind")

    def test_the_gene_page_offers_both_of_its_kinds(self):
        """It mounts two Add panels — antibodies and cell lines — so it needs two."""
        source = self._source("target_detail.html")
        for kind in ("antibodies", "cell-lines"):
            self.assertIn(f"kind='{kind}'", source)

    def test_every_kind_linked_actually_builds(self):
        from pipeline.views.imports import columns_and_example
        for kind in sorted(set(self.SURFACES.values()) | {"cell-lines"}):
            with self.subTest(kind=kind):
                columns, example = columns_and_example(kind)
                self.assertTrue(columns, f"no columns for '{kind}'")
                self.assertEqual(len(columns), len(example),
                                 f"'{kind}' example row does not match its columns")

    def test_the_target_board_grid_and_its_template_share_one_column_list(self):
        """The one board whose columns were a literal in its view was also the
        only one with no template. Same source now, so they cannot disagree."""
        from pipeline.views.imports import columns_and_example
        site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        client = _member_client(self, site)
        resp = client.get("/pipeline/targets/board/")
        self.assertEqual(resp.status_code, 200)
        columns, example = columns_and_example("targets")
        self.assertEqual(list(resp.context["new_columns"]), columns)
        self.assertEqual(list(resp.context["new_example"]), example)

    def test_a_download_link_says_it_arrived(self):
        """A bare <a href> writes the file and leaves the page silent — the
        finding that had a field test file the app's best output as broken."""
        board_js = (Path(settings.BASE_DIR) / "pipeline/static/pipeline/board.js").read_text()
        self.assertIn("data-receipt", board_js)
        # Both pop-outs render the link, and both wire it after mounting: the
        # markup postdates any page-load pass. Three occurrences: the definition
        # and the two call sites.
        self.assertEqual(board_js.count("templateLink(cfg,"), 3)
        self.assertGreaterEqual(board_js.count("downloadReceipts(cfg.mount)"), 2)


class BulkAddTargetsPreviewsFirstTests(TestCase):
    """Bulk Add Targets was the one paste box that wrote straight through.

    A hundred symbols in, targets out, and the first thing you learned about a
    typo was a permanent junk row with a nomination hanging off it. It had no
    preview because a preview needs UniProt, and the lookups were being done
    **while writing**, so a slow API could time the request out mid-batch.

    Two endpoints now: the check owns the network and is bounded by a deadline,
    the commit owns the writes and makes no call at all. Both shapes are pinned
    because two surfaces post here — the feasibility panel sends previewed rows,
    and the target board's Add panel sends bare genes, since its own preview is
    deliberately offline.
    """

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)

    def _fake_uniprot(self, table):
        from pipeline.services import bulk_targets
        return mock.patch.object(
            bulk_targets.uniprot, "lookup_gene",
            side_effect=lambda g: table.get(g.upper(), {"found": False, "error": "no such gene"}))

    @staticmethod
    def _found(gene):
        return {"found": True, "gene_name": gene, "protein_name": f"{gene} protein",
                "uniprot_id": "P00001", "mass_kda": 16.0, "gene_synonyms": []}

    def test_the_check_writes_nothing(self):
        before = Target.objects.using(DB).count()
        with self._fake_uniprot({"SOD1": self._found("SOD1")}):
            resp = self.client.post(
                "/pipeline/feasibility/bulk-check/",
                data=json.dumps({"genes": "SOD1"}), content_type="application/json")
        self.assertEqual(resp.status_code, 200)
        body = resp.json()
        self.assertTrue(body["ok"])
        self.assertEqual(body["rows"][0]["status"], "new")
        self.assertEqual(Target.objects.using(DB).count(), before)

    def test_the_commit_creates_what_the_check_showed(self):
        with self._fake_uniprot({"SOD1": self._found("SOD1")}):
            rows = self.client.post(
                "/pipeline/feasibility/bulk-check/",
                data=json.dumps({"genes": "SOD1"}),
                content_type="application/json").json()["rows"]

        # No UniProt patch here on purpose: if the commit calls out, the real
        # network guard in this environment makes it fail loudly rather than
        # quietly passing on a stub.
        resp = self.client.post(
            "/pipeline/feasibility/bulk-add/",
            data=json.dumps({"rows": rows}), content_type="application/json")
        self.assertEqual(resp.status_code, 200)
        body = resp.json()
        self.assertEqual(len(body["created"]), 1)
        self.assertTrue(Target.objects.using(DB).filter(gene_name="SOD1").exists())

    def test_the_target_boards_bare_gene_shape_still_works(self):
        """Its Add panel previews with the offline check endpoint, so it has no
        UniProt data to hand back and posts gene names. `tests_timeouts` pins
        that a board makes no outbound call, so that is not going to change."""
        with self._fake_uniprot({"TARDBP": self._found("TARDBP")}):
            resp = self.client.post(
                "/pipeline/feasibility/bulk-add/",
                data=json.dumps({"genes": "TARDBP"}), content_type="application/json")
        self.assertEqual(resp.status_code, 200)
        body = resp.json()
        # The old response keys the board renders are still there.
        self.assertEqual(len(body["added"]), 1)
        self.assertEqual(body["not_found"], [])
        self.assertIn("skipped", body)
        self.assertTrue(Target.objects.using(DB).filter(gene_name="TARDBP").exists())

    def test_a_gene_uniprot_cannot_confirm_is_reported_and_not_created(self):
        with self._fake_uniprot({}):
            resp = self.client.post(
                "/pipeline/feasibility/bulk-add/",
                data=json.dumps({"genes": "NOTAGENE"}), content_type="application/json")
        body = resp.json()
        self.assertEqual(body["created"], [])
        self.assertEqual(len(body["not_found"]), 1)
        self.assertFalse(Target.objects.using(DB).filter(gene_name="NOTAGENE").exists())

    def test_the_page_offers_the_check_before_the_save(self):
        """A save button is disabled with a reason, not hidden — and the copy has
        to say why a long list takes a moment, since that is the whole reason
        this box had no preview for seven field tests."""
        resp = self.client.get("/pipeline/feasibility/")
        self.assertEqual(resp.status_code, 200)
        html = resp.content.decode()
        self.assertIn("bulk-check-btn", html)
        self.assertIn("Check these", html)
        self.assertIn('id="bulk-add-btn" disabled', html)
        self.assertIn("UniProt lookup", html)


class ANominationIsNotASkippedRowTests(TestCase):
    """Run 8, F1 — the highest-cost finding of the run, and a counting bug.

    Bulk Add Targets reported `✓ 2 added · ↷ 4 skipped` about a press that wrote
    **four** records: two new targets, and Leicester nominations on two of
    McGill's. The preview had warned on both rows (`already on the list at
    McGill · will be added to your site's list`); only the result summary
    disagreed with it, and "skipped" is exactly the word that stops somebody
    writing a record down. Those two nominations were the only records the run
    left that its own `[COWORK RUN8]` search could not reach.

    `apply` has returned `nominated` all along. Nothing rendered it, and the
    view lumped every on-file row under `skipped` whether or not one had just
    been written for it.
    """

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.other = Site.objects.using(DB).create(name="McGill", short_code="MCG")
        self.client = _member_client(self, self.site)
        # A gene already on the consortium's list, nominated by somebody else.
        self.theirs = Target.objects.using(DB).create(
            gene_name="LRRK2", protein_name="LRRK2 protein")
        TargetNomination.objects.using(DB).create(
            target_id=self.theirs.pk, site_id=self.other.pk, funded=False)
        # A gene that is already ours — nothing to do for this one.
        self.mine = Target.objects.using(DB).create(
            gene_name="MAPT", protein_name="MAPT protein")
        TargetNomination.objects.using(DB).create(
            target_id=self.mine.pk, site_id=self.site.pk, funded=False)

    def _add(self, genes):
        return self.client.post(
            "/pipeline/feasibility/bulk-add/",
            data=json.dumps({"rows": [{"gene": g, "status": "on_file",
                                       "target_id": t.pk, "gene_name": g}
                                      for g, t in genes]}),
            content_type="application/json").json()

    def test_a_row_that_gained_a_nomination_is_not_reported_as_skipped(self):
        body = self._add([("LRRK2", self.theirs)])
        self.assertEqual(body["nominated_existing"], ["LRRK2"],
                         "the write has to be named in the response")
        self.assertNotIn("LRRK2", body["skipped"],
                         "a row a nomination was written for was not skipped")

    def test_a_row_that_was_already_ours_is_still_reported_as_skipped(self):
        body = self._add([("MAPT", self.mine)])
        self.assertEqual(body["nominated_existing"], [])
        self.assertEqual(body["skipped"], ["MAPT"])

    def test_the_two_kinds_are_counted_apart_in_one_press(self):
        body = self._add([("LRRK2", self.theirs), ("MAPT", self.mine)])
        self.assertEqual(body["nominated_existing"], ["LRRK2"])
        self.assertEqual(body["skipped"], ["MAPT"])
        self.assertEqual(body["site"], "Leicester",
                         "the summary says whose list the row went on")
        self.assertTrue(
            TargetNomination.objects.using(DB)
            .filter(target_id=self.theirs.pk, site_id=self.site.pk).exists())

    def test_a_newly_created_target_is_not_double_counted(self):
        """Creating a target nominates it too, so `nominated` includes it. The
        summary must not then report the same gene under both headings."""
        rows = [{"gene": "SOD1", "status": "new", "gene_name": "SOD1",
                 "target_id": None, "protein_name": "SOD1 protein",
                 "uniprot_id": "P00441"}]
        body = self.client.post(
            "/pipeline/feasibility/bulk-add/",
            data=json.dumps({"rows": rows}),
            content_type="application/json").json()
        self.assertEqual([c["gene"] for c in body["created"]], ["SOD1"])
        self.assertEqual(body["nominated_existing"], [],
                         "a gene reported as added must not be reported again")

    def test_both_doors_render_the_nomination_count(self):
        """The feasibility panel and the target board's Add panel post to one
        endpoint and each renders the answer itself, so a fix in one is a fix for
        one page."""
        feasibility = self.client.get("/pipeline/feasibility/").content.decode()
        board = self.client.get("/pipeline/targets/board/").content.decode()
        for name, html in (("feasibility", feasibility), ("target board", board)):
            self.assertIn("nominated_existing", html,
                          f"{name} does not read the nominations it just wrote")


class BothAddDoorsChooseTheSiteTests(TestCase):
    """Owner's ask, 4 Aug: *"I need to be able to allocate site (Leicester,
    Ontario (new), McGill, Cornell etc.) here and on the targets board."*

    Both bulk doors wrote every nomination at the **adder's own** site, so a
    consortium coordinator could not put a gene on another bench's list from
    anywhere in the app — and somebody whose account was on the wrong site could
    not stop it filing their work there. The single-gene Add box on the same
    feasibility page had had the dropdown since Carl Laflamme asked for it; the
    bulk box a few centimetres below it had not, which is the two-doors failure
    this suite keeps finding.
    """

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.other = Site.objects.using(DB).create(name="Ontario", short_code="ONT")
        self.client = _member_client(self, self.site)
        self.target = Target.objects.using(DB).create(
            gene_name="TRPA1", protein_name="TRPA1 protein")

    def _add(self, **extra):
        rows = [{"gene": "TRPA1", "status": "on_file", "gene_name": "TRPA1",
                 "target_id": self.target.pk}]
        return self.client.post(
            "/pipeline/feasibility/bulk-add/",
            data=json.dumps({"rows": rows, **extra}),
            content_type="application/json")

    def test_the_check_answers_about_the_site_it_was_asked_about(self):
        body = self.client.post(
            "/pipeline/feasibility/bulk-check/",
            data=json.dumps({"genes": "TRPA1", "site": self.other.pk}),
            content_type="application/json").json()
        self.assertEqual(body["site"], "Ontario")
        self.assertTrue(body["rows"][0]["will_nominate"])

    def test_the_save_writes_the_nomination_at_the_chosen_site(self):
        body = self._add(site=self.other.pk).json()
        self.assertEqual(body["site"], "Ontario",
                         "the result says whose list it went on")
        self.assertTrue(TargetNomination.objects.using(DB)
                        .filter(target_id=self.target.pk, site_id=self.other.pk).exists())
        self.assertFalse(TargetNomination.objects.using(DB)
                         .filter(target_id=self.target.pk, site_id=self.site.pk).exists(),
                         "it must not also land on the adder's own list")

    def test_no_site_chosen_still_means_your_own(self):
        self._add()
        self.assertTrue(TargetNomination.objects.using(DB)
                        .filter(target_id=self.target.pk, site_id=self.site.pk).exists())

    def test_a_site_that_is_not_on_file_is_refused_and_names_the_ones_that_are(self):
        """Not silently swapped for the member's. Nothing is written yet when
        this is asked, so refusing costs a press — and a batch quietly re-homed
        is what the whole control exists to prevent."""
        resp = self._add(site=999999)
        self.assertEqual(resp.status_code, 400)
        self.assertIn("Ontario", resp.json()["error"])
        self.assertFalse(TargetNomination.objects.using(DB)
                         .filter(target_id=self.target.pk).exists())

    def test_both_doors_render_a_site_control_in_their_bulk_panel(self):
        """A control on one door and not the other is the shape this pair keeps
        failing in — the funder landed on both, the supplier preview did not."""
        pages = {"feasibility": ("/pipeline/feasibility/", "bulk-site"),
                 "target board": ("/pipeline/targets/board/", "ne-site")}
        for name, (url, control) in pages.items():
            html = self.client.get(url).content.decode()
            self.assertIn(f'id="{control}"', html,
                          f"{name}'s bulk add panel has no site box")
            self.assertIn("Ontario", html,
                          f"{name}'s site box does not list the sites on file")


class ARefusedConcentrationIsCountedAtTheSaveTests(TestCase):
    """Run 8, F3 — the preview refuses `3 mM` by name and the row saves anyway.

    That the row is still written is deliberate: an odd concentration is no
    reason to discard an otherwise good vial. What was missing is the build's own
    standard — *say what you are about to drop, and count it* — which the
    unrecognised-column preview two panels away has met since run 6. Run 8 read
    the per-row refusal, pressed **Create them** with nothing else said, and got
    antibody A-ELP3-R8D with no concentration at all.
    """

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)
        Target.objects.using(DB).create(gene_name="ELP3", protein_name="ELP3 protein")

    ROWS = ("catalogue\tcompany\tgene\tconcentration\n"
            "A-ELP3-R8C\tabcam\tELP3\t200\n"
            "A-ELP3-R8D\tabcam\tELP3\t3 mM\n")

    def test_the_preview_counts_the_rows_whose_value_will_not_be_stored(self):
        from pipeline.services import bulk_antibodies
        rows = bulk_antibodies.parse(self.ROWS)
        summary = bulk_antibodies.summarize(bulk_antibodies.plan(rows, False))
        self.assertEqual(summary["no_concentration"], 1,
                         "a refusal nothing adds up is a refusal a button outlives")

    def test_the_save_names_what_it_dropped(self):
        from pipeline.services import bulk_antibodies
        rows = bulk_antibodies.parse(self.ROWS)
        result = bulk_antibodies.apply(rows)
        dropped = result["no_concentration"]
        self.assertEqual([d["catalogue"] for d in dropped], ["A-ELP3-R8D"])
        self.assertEqual(dropped[0]["typed"], "3 mM",
                         "say the value back, or the reader cannot find the cell")

    def test_a_convertible_unit_is_not_counted_as_dropped(self):
        from pipeline.services import bulk_antibodies
        rows = bulk_antibodies.parse(
            "catalogue\tcompany\tgene\tconcentration\nA-ELP3-R8B\tabcam\tELP3\t1.0 mg/mL\n")
        items = bulk_antibodies.plan(rows, False)
        self.assertEqual(bulk_antibodies.summarize(items)["no_concentration"], 0)

    def test_a_blank_concentration_is_not_a_refusal(self):
        from pipeline.services import bulk_antibodies
        rows = bulk_antibodies.parse(
            "catalogue\tcompany\tgene\tconcentration\nA-ELP3-R8E\tabcam\tELP3\t\n")
        items = bulk_antibodies.plan(rows, False)
        self.assertEqual(bulk_antibodies.summarize(items)["no_concentration"], 0)

    def test_every_antibody_paste_surface_reads_the_count(self):
        target = Target.objects.using(DB).get(gene_name="ELP3")
        pages = {
            "antibodies board": self.client.get("/pipeline/antibodies/board/"),
            "gene page": self.client.get(f"/pipeline/target/{target.pk}/"),
        }
        for name, resp in pages.items():
            self.assertEqual(resp.status_code, 200, name)
            self.assertIn("concentration", resp.content.decode().lower(), name)
            # `OGABoard.saveNotes` is the one writer every save box calls now —
            # it composes the concentration refusal, the C-number refusal and
            # the lab numbers a press has just issued. Three messages of this
            # kind have each had to be wired into all seven render sites by
            # hand; composing them means the fourth cannot reach six of seven.
            self.assertIn("OGABoard.saveNotes", resp.content.decode(),
                          f"{name} does not repeat the refusal at the save")


class TheQuickSessionPanelSaysWhichCellLineItPickedTests(TestCase):
    """Run 8, F2 / claim 8 — refuted claim, surviving defect.

    Site preference means a bare `HAP1` at Leicester resolves to Leicester's
    HAP1, which is right and is the fix for run 7's cross-institution control.
    What run 8 found is that nothing says so: two lines are called HAP1, one was
    chosen, and the only way to learn which was to create the session and reopen
    the row. On a board whose subject is the WT/KO comparison that has to be a
    line of text.
    """

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.other = Site.objects.using(DB).create(name="McGill", short_code="MCG")
        self.client = _member_client(self, self.site)
        self.target = Target.objects.using(DB).create(
            gene_name="ELP3", protein_name="ELP3 protein")
        self.mine = CellLine.objects.using(DB).create(
            name="HAP1", genotype="WT", site_id=self.site.pk)
        self.theirs = CellLine.objects.using(DB).create(
            name="HAP1", genotype="WT", site_id=self.other.pk)

    def _check(self, wt="", ko=""):
        return self.client.post(
            "/pipeline/session/plan/parse/",
            data=json.dumps({"gene": "ELP3", "text": "", "rows": [],
                             "cell_line_wt": wt, "cell_line_ko": ko}),
            content_type="application/json").json()

    def test_the_check_answers_about_the_cell_lines_at_all(self):
        body = self._check(wt="HAP1")
        self.assertIn("cell_lines", body,
                      "the check reported only the antibody rows")

    def test_a_shared_name_says_which_row_it_resolved_to(self):
        wt = self._check(wt="HAP1")["cell_lines"]["wt"]
        self.assertEqual(wt["status"], "resolved")
        self.assertEqual(wt["label"], "HAP1 — Leicester")
        self.assertEqual(wt["shared"], 2)
        self.assertIn("2 cell lines are called", wt["note"])

    def test_an_unambiguous_name_does_not_nag(self):
        CellLine.objects.using(DB).filter(pk=self.theirs.pk).delete()
        wt = self._check(wt="HAP1")["cell_lines"]["wt"]
        self.assertEqual(wt["label"], "HAP1 — Leicester")
        self.assertEqual(wt["note"], "", "one match is not a choice to explain")

    def test_a_name_nothing_answers_to_says_it_will_be_created(self):
        wt = self._check(wt="NOTALINE-R8")["cell_lines"]["wt"]
        self.assertEqual(wt["status"], "new")
        self.assertIn("created", wt["note"])

    def test_a_genuine_tie_is_refused_before_the_save_not_after(self):
        """Two lines of the same name at sites that are both strangers. `resolve`
        refuses; the check has to show that refusal rather than let a save fail."""
        third = Site.objects.using(DB).create(name="Montreal", short_code="MTL")
        CellLine.objects.using(DB).create(name="U2OS", genotype="WT",
                                          site_id=self.other.pk)
        CellLine.objects.using(DB).create(name="U2OS", genotype="WT",
                                          site_id=third.pk)
        wt = self._check(wt="U2OS")["cell_lines"]["wt"]
        self.assertEqual(wt["status"], "error")
        self.assertIn("U2OS", wt["error"])
        self.assertIn("McGill", wt["error"], "a refusal names the alternatives")

    def test_nothing_is_written_by_the_check(self):
        before = CellLine.objects.using(DB).count()
        self._check(wt="NOTALINE-R8", ko="ALSO-NOT-A-LINE")
        self.assertEqual(CellLine.objects.using(DB).count(), before)

    def test_the_panel_prints_it(self):
        html = self.client.get("/pipeline/sessions/board/").content.decode()
        self.assertIn("cell_lines", html)
        self.assertIn("cellLineLine", html)


class TheSearchBoxReachesACellLinesOwnNotesTests(TestCase):
    """Run 8, F7 (and run 7's, unchanged) — `[COWORK RUN8]` found 7 antibodies
    and 7 sessions and neither of the two cell lines carrying the same string.

    A cell line has no `comments` column: its free text is `origin_comments` and
    `ko_validation_notes`, so the widening that reached antibody and session
    comments missed this group entirely. Two runs have now tagged their rows and
    then failed to find them.
    """

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)
        target = Target.objects.using(DB).create(
            gene_name="STMN2", protein_name="Stathmin-2")
        CellLine.objects.using(DB).create(
            name="SH-SY5Y STMN2 KO", genotype="KO", target_id=target.pk,
            site_id=self.site.pk, origin_comments="[COWORK RUN8]")
        CellLine.objects.using(DB).create(
            name="HAP1 ELP3 KO", genotype="KO", site_id=self.site.pk,
            ko_validation_notes="[COWORK RUN8] confirmed by WB")

    def _groups(self, term):
        from pipeline.services import find
        return {g["key"]: g for g in find.find(term)["groups"]}

    def test_a_tag_in_the_origin_notes_is_found(self):
        groups = self._groups("COWORK RUN8")
        self.assertIn("cell_lines", groups,
                      "a tag a run wrote on a cell line has to be findable")
        self.assertEqual(groups["cell_lines"]["count"], 2)

    def test_the_result_says_which_field_answered(self):
        rows = {r["title"]: r for r in self._groups("COWORK RUN8")["cell_lines"]["rows"]}
        self.assertIn("origin notes", rows["SH-SY5Y STMN2 KO"]["why"])
        self.assertIn("validation notes", rows["HAP1 ELP3 KO"]["why"])

    def test_the_page_reaches_them_too(self):
        resp = self.client.get("/pipeline/find/?q=COWORK+RUN8")
        self.assertEqual(resp.status_code, 200)
        self.assertIn("SH-SY5Y STMN2 KO", resp.content.decode())


class AReportDoesNotAssertAProcedureNobodyRanTests(TestCase):
    """Run 8, F6 — ELP3's draft announced three antibodies characterized for
    western blot over three WB sessions with no reading between them.

    One had no result rows; two carried a single blank row each, created while
    testing the quick New session panel. The abstract named western blot, a full
    Method paragraph was written and a figure legend printed. `session_import`
    already refuses to *create* such a session from a workbook — *"a tab or a row
    nobody wrote on is not an experiment"* — and the workbook's own How-to-use tab
    states the rule. This is that same sentence at the far end of the pipeline,
    where the draft goes out to a reader.
    """

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        from pipeline.models import (Antibody, Company, ExperimentSession,
                                     IpResult, Member, WbResult)
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)
        self.target = Target.objects.using(DB).create(
            gene_name="ELP3", protein_name="Elongator complex protein 3")
        company = Company.objects.using(DB).create(name="abcam")
        self.ab = Antibody.objects.using(DB).create(
            target_id=self.target.pk, company_id=company.pk,
            catalogue_number="A-ELP3-R8A", site_id=self.site.pk)
        member = Member.objects.using(DB).get(site_id=self.site.pk)

        def session(proc):
            return ExperimentSession.objects.using(DB).create(
                target_id=self.target.pk, procedure_type=proc,
                site_id=self.site.pk, experimenter_id=member.pk,
                date=date(2026, 8, 1))

        # WB: planned only, plus a session whose one result row is entirely blank.
        self.wb_planned = session("WB")
        self.wb_blank = session("WB")
        WbResult.objects.using(DB).create(
            session_id=self.wb_blank.pk, antibody_id=self.ab.pk)
        # IP: an actual reading.
        self.ip_real = session("IP")
        IpResult.objects.using(DB).create(
            session_id=self.ip_real.pk, antibody_id=self.ab.pk,
            enrichment="yes", ip_assessment="pass")

    def _grouped(self):
        from pipeline.services import report_generator
        return report_generator._get_sessions(self.target)

    def test_a_procedure_whose_every_row_is_blank_is_not_claimed(self):
        self.assertNotIn("WB", self._grouped(),
                         "a blank result row is not a western blot")

    def test_a_procedure_with_a_reading_survives(self):
        grouped = self._grouped()
        self.assertIn("IP", grouped)
        self.assertEqual([s.pk for s in grouped["IP"]], [self.ip_real.pk])

    def test_a_comment_alone_does_not_make_it_a_run_session(self):
        """This pinned the opposite until run 19: a typed comment counted as
        "written on", consistent with `session_import._has_result`. Then the
        sessions board's Add panel put its per-row NOTES into `comments` at
        *planning* time, and a session nobody had run read "2 results — every
        row has something recorded" the moment it was created. A row cannot say
        whether its comment was written before or after the blot, and the
        error that misleads is the planned one reading as done — so a comment
        is not a reading anywhere (`session_board.NOT_A_READING`), and this
        report keeps not naming a procedure that has no reading behind it. A
        reading beside the comment still counts."""
        from pipeline.models import WbResult
        WbResult.objects.using(DB).filter(session_id=self.wb_blank.pk).update(
            comments="faint band, repeat next week")
        self.assertNotIn("WB", self._grouped())
        WbResult.objects.using(DB).filter(session_id=self.wb_blank.pk).update(
            signal="faint band")
        self.assertIn("WB", self._grouped())

    def test_the_written_draft_does_not_name_the_unrun_procedure(self):
        import tempfile
        from docx import Document
        from pipeline.services import report_generator
        with tempfile.TemporaryDirectory() as tmp:
            path = report_generator.generate_report(
                self.target.pk, output_path=os.path.join(tmp, "elp3.docx"))
            text = "\n".join(p.text for p in Document(path).paragraphs)
        lower = text.lower()
        self.assertIn("immunoprecipitation", lower)
        # Not a bare "western blot" search: an IP is *read out* by western blot,
        # so its own Method paragraph says so correctly. What must not appear is
        # WB claimed as an application characterized — the abstract's list, and
        # the Method section that only a WB session gets.
        self.assertNotIn("antibody screening by western blot", lower)
        abstract = next(line for line in lower.splitlines()
                        if "here we have characterized" in line)
        self.assertIn("for immunoprecipitation using", abstract)
        self.assertNotIn("western blot", abstract)

    def test_a_target_with_no_readings_at_all_leaves_a_named_gap(self):
        """Now reachable: every session blank means no procedure at all, and
        "characterized 1 antibodies for  using" is worse than a named gap."""
        import tempfile
        from docx import Document
        from pipeline.models import IpResult
        IpResult.objects.using(DB).filter(session_id=self.ip_real.pk).delete()
        with tempfile.TemporaryDirectory() as tmp:
            path = report_generator_generate(self.target.pk, tmp)
            text = "\n".join(p.text for p in Document(path).paragraphs)
        self.assertIn("[no application has a recorded result yet]", text)


def report_generator_generate(target_pk, tmpdir):
    from pipeline.services import report_generator
    return report_generator.generate_report(
        target_pk, output_path=os.path.join(tmpdir, "report.docx"))


class ReportTableThreeHasOneRowPerProcedureTests(TestCase):
    """Run 8, F5 — Table 3 printed "Western blot" three times for one session.

    One secondary came off the session's conditions and two off the per-result
    `secondary_ab` values. All three are true; a reader who has not seen the
    sessions reads three western blots. A procedure is run once here, and the
    secondaries it used belong inside its row.
    """

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        from pipeline.models import (Antibody, Company, ExperimentSession,
                                     Member, WbResult)
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)
        self.target = Target.objects.using(DB).create(gene_name="STMN2")
        company = Company.objects.using(DB).create(name="Proteintech")
        member = Member.objects.using(DB).get(site_id=self.site.pk)
        self.session = ExperimentSession.objects.using(DB).create(
            target_id=self.target.pk, procedure_type="WB", site_id=self.site.pk,
            experimenter_id=member.pk, date=date(2026, 8, 1),
            session_conditions={"secondary_ab": "HRP-conjugated (RGAR001/RGAM001)",
                                "secondary_dilution": "1:10000"})
        for i, sec in enumerate(("HRP anti-rabbit (RGAR001)",
                                 "HRP anti-mouse (RGAM001)")):
            ab = Antibody.objects.using(DB).create(
                target_id=self.target.pk, company_id=company.pk,
                catalogue_number=f"A-STMN2-R8{i}", site_id=self.site.pk)
            WbResult.objects.using(DB).create(
                session_id=self.session.pk, antibody_id=ab.pk,
                secondary_ab=sec, secondary_ab_dilution="1:10000", signal="yes")

    def test_one_western_blot_is_one_row(self):
        from pipeline.services import report_generator
        grouped = report_generator._get_sessions(self.target)
        secondaries = report_generator._get_secondary_antibodies(grouped)
        self.assertEqual(len(secondaries), 3, "all three are still collected")
        by_proc = report_generator._group_by_procedure(secondaries)
        self.assertEqual(list(by_proc), ["WB"])
        self.assertEqual(len(by_proc["WB"]), 3)

    def test_the_row_lists_every_secondary_it_used(self):
        import tempfile
        from docx import Document
        with tempfile.TemporaryDirectory() as tmp:
            doc = Document(report_generator_generate(self.target.pk, tmp))
        row = self._table3_row(doc)
        self.assertIn("HRP anti-rabbit (RGAR001)", row[1])
        self.assertIn("HRP anti-mouse (RGAM001)", row[1])
        self.assertEqual(row[1].count(";"), 2)

    def test_the_columns_line_up_positionally(self):
        """Run 9: three antibodies against two dilutions.

        Each column was de-duplicated on its own, so two identical per-row
        dilutions collapsed into one entry and the lists stopped corresponding.
        Both were true and a reader could not tell which dilution belonged to
        which antibody. Position N must mean the same antibody in every column.
        """
        import tempfile
        from docx import Document
        with tempfile.TemporaryDirectory() as tmp:
            doc = Document(report_generator_generate(self.target.pk, tmp))
        row = self._table3_row(doc)
        antibodies = [p.strip() for p in row[1].split(";")]
        dilutions = [p.strip() for p in row[2].split(";")]
        self.assertEqual(len(antibodies), 3)
        self.assertEqual(len(dilutions), len(antibodies),
                         "a reader cannot map a short list onto a long one")
        # The two per-result rows carried the same dilution; both entries survive
        # so the third antibody's is not silently attributed to the first.
        self.assertEqual(dilutions[1], dilutions[2])

    def _table3_row(self, doc):
        table3 = [t for t in doc.tables
                  if t.rows and t.rows[0].cells[0].text.strip() == "Procedure"]
        self.assertTrue(table3, "Table 3 was not written")
        self.assertEqual(len(table3[0].rows), 2, "one procedure, one row")
        return [c.text.strip() for c in table3[0].rows[1].cells]

    def test_the_written_table_names_the_procedure_once(self):
        import tempfile
        from docx import Document
        from pipeline.services import report_generator
        with tempfile.TemporaryDirectory() as tmp:
            path = report_generator_generate(self.target.pk, tmp)
            doc = Document(path)
        table3 = [t for t in doc.tables
                  if t.rows and t.rows[0].cells[0].text.strip() == "Procedure"]
        self.assertTrue(table3, "Table 3 was not written")
        procedures = [r.cells[0].text.strip() for r in table3[0].rows[1:]]
        self.assertEqual(procedures, ["Western blot"],
                         "one session, one row — not one row per secondary")


class AnExportSpeaksAtClickTimeTests(TestCase):
    """Run 8, F4 — **refuted**, and pinned so it is not re-filed.

    The report says pressing an export leaves the screen doing nothing for up to
    40 seconds. It does not: `downloadWithReceipt` writes `Preparing the file…`
    as its **first statement**, before the fetch is issued. Driven in a real
    browser (Playwright, `StaticLiveServerTestCase`), clicking the Generate
    Report anchor and reading the receipt element **in the same JavaScript turn**
    returned `'Preparing the file…'` on both the report and workbook links, with
    the green `Downloaded …` line replacing it once the bytes arrived. The 40
    seconds run 8 measured is time-to-*completion*, not time-to-first-feedback.

    This is the weaker check — a browser is what settled it, and a browser
    harness cannot live in the repo because playwright is not in
    requirements.txt. What it can pin is that the reassurance is still said
    before anything is awaited, and that every export anchor is still wired to
    the mechanism that says it.

    The *latency* half of F4 stands and is not a code defect this pins: see the
    action register.
    """

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)
        self.target = Target.objects.using(DB).create(
            gene_name="STMN2", protein_name="Stathmin-2")

    def _board_js(self):
        return (Path(settings.BASE_DIR) / "pipeline" / "static" / "pipeline"
                / "board.js").read_text()

    def test_the_reassurance_comes_before_the_await(self):
        js = self._board_js()
        body = js[js.index("async function downloadWithReceipt"):]
        body = body[:body.index("\n  function downloadReceipts")]
        said = body.index("Preparing the file")
        awaited = body.index("await fetch")
        self.assertLess(said, awaited,
                        "the page must speak before it waits, not after")

    def test_every_export_anchor_is_wired_to_it(self):
        """A `data-receipt` anchor nobody wires is a plain href — the silent
        download the whole mechanism exists to stop."""
        import re
        pages = {
            "gene page": f"/pipeline/target/{self.target.pk}/",
            "sessions board": "/pipeline/sessions/board/",
            "feasibility": "/pipeline/feasibility/",
        }
        for name, url in pages.items():
            html = self.client.get(url).content.decode()
            if 'data-receipt=' not in html:
                continue
            self.assertIn("downloadReceipts", html,
                          f"{name} has receipt anchors nothing wires")
            # Every receipt names an element that exists on the same page. One
            # that does not makes `say()` a no-op and the download silent again.
            for target_id in set(re.findall(r'data-receipt="([a-z0-9\-]+)"', html)):
                self.assertIn(f'id="{target_id}"', html,
                              f"{name}: nothing on the page is called {target_id}")


class ACellLineThatIsNotYoursSaysSoTests(TestCase):
    """Run 9 — the one refuted sub-case, and the only finding with a data edge.

    Site preference only *speaks* when two rows compete. 17 rows are called
    `U2OS`, 16 of them knockouts, so a wild-type slot narrows them to exactly one
    — which looked unambiguous and resolved silently to McGill's line. A
    Leicester user typing a name that is not theirs got another institution's
    reagent attached with a one-line, warning-free receipt.

    Whose it is has to be said whether or not there was a choice to make.
    """

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        from pipeline.models import Company
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.other = Site.objects.using(DB).create(name="McGill", short_code="MCG")
        self.client = _member_client(self, self.site)
        self.target = Target.objects.using(DB).create(
            gene_name="ELP3", protein_name="ELP3 protein")
        # One U2OS wild type, at McGill, among many McGill knockouts of that name.
        CellLine.objects.using(DB).create(
            name="U2OS", genotype="WT", site_id=self.other.pk)
        for i in range(3):
            t = Target.objects.using(DB).create(gene_name=f"GENE{i}")
            CellLine.objects.using(DB).create(
                name="U2OS", genotype="KO", target_id=t.pk, site_id=self.other.pk)
        # Two HAP1 wild types, one of them ours — the case that already worked.
        CellLine.objects.using(DB).create(
            name="HAP1", genotype="WT", site_id=self.site.pk)
        CellLine.objects.using(DB).create(
            name="HAP1", genotype="WT", site_id=self.other.pk)

    def _check(self, wt):
        return self.client.post(
            "/pipeline/session/plan/parse/",
            data=json.dumps({"gene": "ELP3", "text": "", "rows": [],
                             "date": "2026-08-02", "procedure_type": "WB",
                             "cell_line_wt": wt}),
            content_type="application/json").json()["cell_lines"]["wt"]

    def test_another_sites_line_is_named_as_theirs(self):
        wt = self._check("U2OS")
        self.assertEqual(wt["status"], "resolved")
        self.assertEqual(wt["label"], "U2OS — McGill")
        self.assertFalse(wt["mine"])
        self.assertIn("McGill's, not yours", wt["note"])

    def test_it_counts_every_row_of_that_name_not_just_the_matching_kind(self):
        """The narrowing is what hid this: 4 rows are called U2OS, one is a WT.
        Saying "1 cell line is called U2OS" would be true of the slot and
        misleading about the database."""
        wt = self._check("U2OS")
        self.assertEqual(wt["shared"], 1, "one wild type of that name")
        self.assertEqual(wt["shared_any"], 4, "four rows carry the name")
        self.assertIn("4 cell lines are called", wt["note"])

    def test_your_own_line_still_reads_as_yours(self):
        wt = self._check("HAP1")
        self.assertEqual(wt["label"], "HAP1 — Leicester")
        self.assertTrue(wt["mine"])
        self.assertIn("this is your site's", wt["note"])
        self.assertNotIn("not yours", wt["note"])

    def test_an_unshared_line_of_your_own_says_nothing(self):
        CellLine.objects.using(DB).create(
            name="SK-N-SH", genotype="WT", site_id=self.site.pk)
        wt = self._check("SK-N-SH")
        self.assertTrue(wt["mine"])
        self.assertEqual(wt["note"], "", "one row, and it is yours — no news")


class TheCheckNamesAMissingDateTests(TestCase):
    """Run 9 — the check was silent about an empty Date and **Create them** then
    refused with `date is required and must be YYYY-MM-DD; got ''`.

    A check exists to say what will happen before anything is written; a required
    field it stays quiet about is a refusal deferred to the save.
    """

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)
        Target.objects.using(DB).create(gene_name="ELP3", protein_name="ELP3 protein")

    def _check(self, **kw):
        body = {"gene": "ELP3", "text": "", "rows": [], "procedure_type": "WB"}
        body.update(kw)
        return self.client.post("/pipeline/session/plan/parse/",
                                data=json.dumps(body),
                                content_type="application/json").json()

    def test_an_empty_date_is_named_on_the_check(self):
        blocking = self._check(date="")["blocking"]
        self.assertTrue(any("Date is empty" in b for b in blocking), blocking)

    def test_an_unparseable_date_is_named_on_the_check(self):
        blocking = self._check(date="2 August 2026")["blocking"]
        self.assertTrue(any("YYYY-MM-DD" in b for b in blocking), blocking)

    def test_a_good_date_blocks_nothing(self):
        self.assertEqual(self._check(date="2026-08-02")["blocking"], [])

    def test_the_panel_renders_it(self):
        html = self.client.get("/pipeline/sessions/board/").content.decode()
        self.assertIn("blocking", html)


class OneSearchDefinitionForTheBoxAndTheBoardTests(TestCase):
    """Run 9 — `RUN9` in the cell-lines board's own Search box returned nothing
    while `/pipeline/find/?q=RUN9` returned both rows.

    The same drift this module was written to fix, one level down: the board and
    the nav box each carried their own hand-written `Q()` list. Two search boxes
    in one app disagreeing about what "search" means reads as one of them being
    broken, and there is no way to tell which from outside.
    """

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        from pipeline.models import Antibody, Company, ExperimentSession, Member
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)
        self.target = Target.objects.using(DB).create(
            gene_name="STMN2", protein_name="Stathmin-2")
        company = Company.objects.using(DB).create(name="Synthego")
        CellLine.objects.using(DB).create(
            name="SH-SY5Y STMN2 KO", genotype="KO", target_id=self.target.pk,
            site_id=self.site.pk, company_id=company.pk,
            catalogue_number="SYN-STMN2-RUN9", origin_comments="[COWORK RUN9]")
        Antibody.objects.using(DB).create(
            target_id=self.target.pk, company_id=company.pk,
            catalogue_number="A-STMN2-R9A", site_id=self.site.pk,
            lot_number="LOT-RUN9", comments="[COWORK RUN9]")
        member = Member.objects.using(DB).get(site_id=self.site.pk)
        self.session = ExperimentSession.objects.using(DB).create(
            target_id=self.target.pk, procedure_type="WB", site_id=self.site.pk,
            experimenter_id=member.pk, date=date(2026, 8, 2))

    def _board(self, path, q):
        return self.client.get(f"{path}?q={q}").json()

    def _find_count(self, term, key):
        from pipeline.services import find
        for g in find.find(term)["groups"]:
            if g["key"] == key:
                return g["count"]
        return 0

    def test_the_cell_lines_board_reaches_a_catalogue_number(self):
        rows = self._board("/pipeline/cell-lines/board/rows/", "RUN9")
        self.assertEqual(rows["count"], 1,
                         "the board's own box found nothing the nav box finds")
        self.assertEqual(self._find_count("RUN9", "cell_lines"), 1)

    def test_the_cell_lines_board_reaches_the_origin_notes(self):
        rows = self._board("/pipeline/cell-lines/board/rows/", "COWORK+RUN9")
        self.assertEqual(rows["count"], 1)

    def test_the_antibodies_board_reaches_a_lot_number(self):
        rows = self._board("/pipeline/antibodies/board/rows/", "LOT-RUN9")
        self.assertEqual(rows["count"], 1, "the board did not reach lot number")
        self.assertEqual(self._find_count("LOT-RUN9", "antibodies"), 1)

    def test_each_board_agrees_with_the_nav_box(self):
        """The real invariant. Both must answer the same question."""
        for term in ("RUN9", "STMN2", "Synthego", "COWORK"):
            for path, key in (
                ("/pipeline/cell-lines/board/rows/", "cell_lines"),
                ("/pipeline/antibodies/board/rows/", "antibodies"),
                ("/pipeline/targets/board/rows/", "targets"),
                ("/pipeline/sessions/board/rows/", "sessions"),
            ):
                self.assertEqual(
                    self._board(path, term)["count"],
                    self._find_count(term, key),
                    f"{path} and the nav box disagree about {term!r}")


class ATagOnAResultRowIsFindableTests(TestCase):
    """Run 8 and run 9 both tagged every result row and both cleanup lists had to
    warn that the search would not reach them. A reading's comment is part of its
    session's record."""

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        from pipeline.models import (Antibody, Company, ExperimentSession,
                                     Member, WbResult)
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)
        target = Target.objects.using(DB).create(gene_name="STMN2")
        company = Company.objects.using(DB).create(name="Proteintech")
        member = Member.objects.using(DB).get(site_id=self.site.pk)
        self.session = ExperimentSession.objects.using(DB).create(
            target_id=target.pk, procedure_type="WB", site_id=self.site.pk,
            experimenter_id=member.pk, date=date(2026, 8, 2))
        for i in range(3):
            ab = Antibody.objects.using(DB).create(
                target_id=target.pk, company_id=company.pk,
                catalogue_number=f"CAT-{i}", site_id=self.site.pk)
            WbResult.objects.using(DB).create(
                session_id=self.session.pk, antibody_id=ab.pk,
                comments="[COWORK RUN9] band as expected")

    def _group(self, term):
        from pipeline.services import find
        return {g["key"]: g for g in find.find(term)["groups"]}.get("sessions")

    def test_a_result_comment_is_reached_through_its_session(self):
        group = self._group("COWORK RUN9")
        self.assertIsNotNone(group, "a tag on a reading found nothing")
        self.assertEqual(group["count"], 1)

    def test_three_matching_readings_are_one_session_not_three(self):
        """Four joined tables without `.distinct()` counts the session once per
        matching reading, and the headline number is the first thing checked
        against the list."""
        self.assertEqual(self._group("COWORK RUN9")["count"], 1)

    def test_the_row_says_a_reading_was_what_matched(self):
        row = self._group("COWORK RUN9")["rows"][0]
        self.assertIn("a result row's comments", row["why"])
        self.assertIn(f"open={self.session.pk}", row["url"])

    def test_the_sessions_board_agrees(self):
        rows = self.client.get("/pipeline/sessions/board/rows/?q=COWORK+RUN9").json()
        self.assertEqual(rows["count"], 1)


class ASaveReceiptSurvivesLongEnoughToReadTests(TestCase):
    """Run 9 — the gene page reloaded 900 ms after a save.

    That fixed the counts in the panels below and **destroyed the receipt**. Run
    9 needed an 80 ms poller to catch *"1 row saved without a concentration:
    A-STMN2-R9D (3 mM)"* — the one line saying a value had been dropped, added in
    run 8 precisely because the save used to be silent. A message a human cannot
    read is not a message.

    Same family as *never tell the reader to reload*: here the page did it for
    you, which is worse, because there is nothing left on screen to explain what
    went past.
    """

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)
        self.target = Target.objects.using(DB).create(
            gene_name="STMN2", protein_name="Stathmin-2")

    def test_the_gene_page_does_not_reload_itself_after_a_save(self):
        html = self.client.get(f"/pipeline/target/{self.target.pk}/").content.decode()
        script = "\n".join(re.findall(r"<script[^>]*>(.*?)</script>", html, re.S))
        self.assertNotRegex(
            script, r"setTimeout\(\s*\(\)\s*=>\s*window\.location\.reload",
            "the page reloads out from under the result line")

    def test_it_offers_the_reader_the_refresh_instead(self):
        html = self.client.get(f"/pipeline/target/{self.target.pk}/").content.decode()
        self.assertIn("afterResult", html)
        self.assertIn("Show them on this page", html)

    def test_board_js_gives_a_panel_somewhere_to_put_it(self):
        js = (Path(settings.BASE_DIR) / "pipeline" / "static" / "pipeline"
              / "board.js").read_text()
        self.assertIn("cfg.afterResult", js)


class BothTargetDoorsCountWhatTheyWillWriteTests(TestCase):
    """Run 9 — the feasibility door reads `Add 5 to pipeline`; the targets board's
    Add panel, posting to the same endpoint, still said `Create them`.

    Run 8's counting fix landed on one of two doors, which is the shape
    `board.js` exists to stop. `newEntry` takes a `commitLabel` now, so a board
    that has a meaningful count says it and the three that do not are unchanged.
    """

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)

    def test_the_shared_panel_supports_a_counted_label(self):
        js = self._board_js()
        self.assertIn("cfg.commitLabel", js)
        self.assertIn("COMMIT_DEFAULT", js,
                      "a board with no count must keep the default wording")

    def test_the_helper_is_called_only_where_it_is_declared(self):
        """**The test that should have existed first.**

        Run 9 shipped this fix with the call site inside `uploadPanel` while the
        helper was declared inside `newEntry` — so `newEntry`'s own preview never
        relabelled its button (the targets board went on saying `Create them`,
        which run 10 refuted) and `uploadPanel` threw a `ReferenceError` on every
        preview. Both `manage.py check` and `node --check` pass on that: it is
        valid syntax, and the error only happens when the line runs.

        The test that let it through asserted the *file contained* the string —
        which it did, in the wrong function. Containment is not wiring. This
        checks the call sites fall inside the declaring function's source range,
        which is the specific mistake, cheaply. It is weaker than driving a
        browser (which is how the fix was actually confirmed: `Create them` →
        `Add 4 to pipeline`, with no page errors); a browser harness cannot live
        here because playwright is not in `requirements.txt`.
        """
        js = self._board_js()
        start = js.index("  function newEntry(cfg) {")
        end = js.index("  function uploadPanel(cfg) {")
        self.assertLess(start, end)
        declared = js[start:end]
        self.assertIn("function setCommitLabel(d) {", declared,
                      "the helper moved — update this test's anchors")
        for name in ("setCommitLabel", "COMMIT_DEFAULT"):
            outside = js[:start].count(name) + js[end:].count(name)
            self.assertEqual(
                outside, 0,
                f"{name} is used outside the function that declares it — "
                "that is a ReferenceError at runtime, not a syntax error")

    def test_the_preview_actually_relabels_the_button(self):
        """`arm(true)` resets the label to the default, so the count has to be
        applied *after* it. Ordering, not presence — the two lines either side of
        this were both present when the button still said `Create them`."""
        js = self._board_js()
        block = js[js.index("        arm(true);"):]
        block = block[:block.index("      } finally {")]
        self.assertIn("setCommitLabel(d)", block,
                      "newEntry's preview never relabels its own button")

    def _board_js(self):
        return (Path(settings.BASE_DIR) / "pipeline" / "static" / "pipeline"
                / "board.js").read_text()

    def test_both_doors_read_one_helper_rather_than_agreeing_by_hand(self):
        """The arithmetic was written twice, with a comment on each copy saying
        it had to match the other. One copy now, in the shared file — which is
        what the two-doors rule has meant everywhere else in here."""
        for url in ("/pipeline/targets/board/", "/pipeline/feasibility/"):
            with self.subTest(url=url):
                html = self.client.get(url).content.decode()
                self.assertIn("OGABoard.targetAddSummary", html)
        js = self._board_js()
        self.assertIn("function targetAddSummary(", js)
        self.assertIn("targetAddSummary,", js, "it is not exported")

    def test_both_doors_offer_one_funder_and_one_project_for_the_batch(self):
        """Owner's ask: *"need to be able to add funder and project info for
        bulk targets (just 1 each)"*.

        On **both** doors, because they post to one endpoint and a control on
        one of them is the shape this pair has failed in three times — the
        count, the refusal, and now this. Picked from the rendered lists rather
        than typed: free text would mint a second "CIHR", the way a typed
        supplier minted a third Bio-Techne.
        """
        for url, ids in (("/pipeline/targets/board/",
                          ("ne-agency", "ne-project", "ne-funded")),
                         ("/pipeline/feasibility/",
                          ("bulk-agency", "bulk-project", "bulk-funded"))):
            with self.subTest(url=url):
                html = self.client.get(url).content.decode()
                for element_id in ids:
                    self.assertIn(element_id, html, f"{url} has no {element_id}")
                self.assertIn("granting_agency", html,
                              f"{url} does not send the funder")
                self.assertRegex(html, r'"?project"?\s*:\s*pick\(',
                                 f"{url} does not send the project")
                # Blanks only, and the page says so — a second paste of an
                # overlapping list must not read as re-filing somebody's genes.
                self.assertIn("only fill blanks", html)

    def test_the_board_says_it_takes_a_spreadsheet(self):
        """Owner's ask: *"signal the target board can accept xls uploads if you
        have lots of targets you need to add"*.

        Both buttons existed and neither named the other, so a person with sixty
        genes pasted them a screenful at a time. The sentence is in the header,
        where somebody with a long list is looking — not only inside a panel
        they have to open first.
        """
        html = self.client.get("/pipeline/targets/board/").content.decode()
        header = html.split('id="filters"', 1)[0]
        self.assertIn("spreadsheet", header)
        self.assertIn("Upload a sheet", header)
        self.assertIn("import/template/targets", header.replace("-", "/"))
        # And that link says the file arrived, like every other export here.
        self.assertIn('data-receipt="header-template-receipt"', header)
        self.assertIn('id="header-template-receipt"', header)

    def test_a_check_that_can_write_nothing_leaves_the_save_greyed(self):
        """**A check that succeeded is not a check that found something to
        write**, and the two doors disagreed about that for as long as both
        existed.

        The feasibility door greyed its button over a list that would create
        nothing. The targets board's Add panel armed `Create them` over the same
        list, because `newEntry` armed on any 200. The case that makes it matter
        is a UniProt outage — a documented, transient condition here — where
        every row comes back `unchecked`, the preview says *press Check these
        again*, and a live save button sat directly under that sentence offering
        to write "0 target(s) added".
        """
        js = self._board_js()
        self.assertIn("cfg.whyNotCommit", js)
        # It has to short-circuit the arm, not merely be consulted.
        block = js[js.index("        const nothing = cfg.whyNotCommit"):]
        block = block[:block.index("      } finally {")]
        self.assertIn("arm(false, nothing); return;", block)
        self.assertLess(block.index("arm(false, nothing)"), block.index("arm(true)"))
        html = self.client.get("/pipeline/targets/board/").content.decode()
        self.assertIn("whyNotCommit", html, "the targets door supplies no reason")

    def test_the_refusal_names_the_reason_it_actually_is(self):
        """Three ways a checked list writes nothing, and they are not the same
        sentence. The feasibility door greyed correctly and then said "already
        on Leicester's list" about genes UniProt had simply not answered for —
        the same mistake as arming, one layer down: right about the state,
        wrong about the cause."""
        js = self._board_js()
        block = js[js.index("function targetAddSummary("):]
        block = block[:block.index("\n  /* A concentration")]
        self.assertIn("UniProt did", block)
        self.assertIn("Check the spelling", block)
        self.assertIn("already on", block)
        # And the outage branch is asked first — an unchecked gene is also not
        # on your list, so the order is what makes the sentence true.
        self.assertLess(block.index("unchecked"), block.index("already on"))


class TheSaveMessageDoesNotPromiseAReloadTests(TestCase):
    """Run 10, F1 — the success box still ended with the word `Reloading…`.

    Run 9 removed the 900 ms auto-reload because it was destroying the line that
    names a dropped concentration; the word announcing that reload was left
    behind. It then sat on screen for the 156 seconds run 10 watched it, telling
    the reader to expect something that never came — directly above a button
    whose entire purpose is that there is no reload. Half a fix reads worse than
    none, because the page now contradicts itself.
    """

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)
        self.target = Target.objects.using(DB).create(
            gene_name="ELP3", protein_name="ELP3 protein")

    def test_no_panel_says_it_is_reloading(self):
        html = self.client.get(f"/pipeline/target/{self.target.pk}/").content.decode()
        self.assertNotIn("Reloading", html,
                         "the page announces a reload it no longer does")

    def test_the_button_that_replaced_it_is_still_there(self):
        html = self.client.get(f"/pipeline/target/{self.target.pk}/").content.decode()
        self.assertIn("Show them on this page", html)


class TheNominationCountIsTheSameNumberEverywhereTests(TestCase):
    """Run 10, F3 — one check, three different answers.

    Four genes: two new, two already on the consortium's list. The feasibility
    chip read `2 will go on Leicester's list`, its own button read `Add 4 to
    pipeline`, all four rows said "will be added to your site's list", and the
    targets board said `4 will go on Leicester's list`. The chip was the outlier
    — a reader believing it would have expected two nominations and got four.

    The forecast counts **genes that end up on your list**; the split between new
    and existing belongs in the *result*, where they are two different things
    that happened.
    """

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.other = Site.objects.using(DB).create(name="McGill", short_code="MCG")
        self.client = _member_client(self, self.site)
        from pipeline.models import TargetNomination
        for gene in ("SOD1", "TARDBP"):
            t = Target.objects.using(DB).create(gene_name=gene, protein_name=gene)
            TargetNomination.objects.using(DB).create(
                target_id=t.pk, site_id=self.other.pk, funded=False)

    def test_the_forecast_counts_every_gene_that_ends_up_yours(self):
        """`plan` is what both doors' counters read; four genes, four
        `will_nominate`, whether the target is new or somebody else's."""
        from pipeline.services import bulk_targets
        from pipeline.models import Member
        member = Member.objects.using(DB).get(site_id=self.site.pk)
        rows = [{"gene": "SOD1", "status": bulk_targets.ON_FILE,
                 "target_id": Target.objects.using(DB).get(gene_name="SOD1").pk},
                {"gene": "TARDBP", "status": bulk_targets.ON_FILE,
                 "target_id": Target.objects.using(DB).get(gene_name="TARDBP").pk},
                {"gene": "SEC61B", "status": bulk_targets.CREATES, "target_id": None},
                {"gene": "GOLGA2", "status": bulk_targets.CREATES, "target_id": None}]
        already = set()
        flags = [bulk_targets._needs_nomination(r, self.site.pk, already) for r in rows]
        self.assertEqual(flags, [True, True, True, True],
                         "all four genes end up on Leicester's list")

    def test_both_doors_use_the_same_arithmetic_for_the_button(self):
        """New targets counted once, not twice: a new target's own nomination is
        already inside `creating`.

        It is one function now (`OGABoard.targetAddSummary`) rather than two
        copies asserted to match — so this checks the rule where it lives and
        that both doors reach it, which is the only way two doors cannot drift.
        """
        js = (Path(settings.BASE_DIR) / "pipeline" / "static" / "pipeline"
              / "board.js").read_text()
        block = js[js.index("function targetAddSummary("):]
        block = block[:block.index("\n  /* A concentration")]
        self.assertIn("status !== 'new'", block,
                      "the shared counter double-counts a new target")
        for url in ("/pipeline/targets/board/", "/pipeline/feasibility/"):
            with self.subTest(url=url):
                html = self.client.get(url).content.decode()
                self.assertIn("OGABoard.targetAddSummary", html)
                self.assertNotIn(
                    "status !== 'new'", html,
                    "a second copy of the arithmetic has come back")

    def test_the_chip_does_not_exclude_new_targets(self):
        """The chip forecasts **genes that end up on your list**, new targets
        included — it was the one thing on the page counting something else."""
        html = self.client.get("/pipeline/feasibility/").content.decode()
        chip = html[html.index("const nominating ="):]
        chip = chip[:chip.index("\n\n")]
        self.assertNotIn("status !== 'new'", chip,
                         "the chip forecasts fewer genes than the button writes")
        self.assertIn("r.will_nominate", chip)


class AHalfTypedDateIsNotCalledEmptyTests(TestCase):
    """Run 10, claim 5's untestable third — worth fixing rather than filing.

    A native `<input type="date">` given `2 August 2026` holds
    `validity.badInput` with `value === ""`, so the string never leaves the
    browser and the server can only ever see an empty date. The check therefore
    said "Date is empty" to somebody looking at a box they had visibly typed
    into. True of what was sent, wrong about what they did — and the only place
    that can tell the difference is the client.
    """

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)

    def test_the_panel_asks_the_input_whether_it_is_half_typed(self):
        html = self.client.get("/pipeline/sessions/board/").content.decode()
        self.assertIn("validity", html)
        self.assertIn("badInput", html)

    def test_it_says_unfinished_rather_than_empty(self):
        html = self.client.get("/pipeline/sessions/board/").content.decode()
        self.assertIn("not finished", html)

    def test_the_server_still_owns_the_genuinely_empty_case(self):
        """The client only speaks for what the server cannot see. An empty box is
        still the server's answer, asked of the rule the writer uses."""
        from pipeline.views.session_bulk import _blocking
        self.assertTrue(any("Date is empty" in b
                            for b in _blocking({"procedure_type": "WB"})))
        self.assertEqual(_blocking({"date": "2026-08-02", "procedure_type": "WB"}), [])


# ---------------------------------------------------------------------------
# Run 11 — 2 Aug 2026
# ---------------------------------------------------------------------------

class ACNumberIsNeverInventedTests(TestCase):
    """``C-RUN11-01`` is not C-11, and the paste box used to think it was.

    ``bulk_cell_lines`` read the column with ``re.search(r"\\d+")``, which does
    not read a C-number so much as go looking for digits in whatever it is
    given. Run 11 typed ``C-RUN11-01`` and ``C-RUN11-02`` into two knockouts,
    got a preview calling both rows ``new`` with no comment on the value, and
    saved both as **11** — one number, two lines, indistinguishable from
    McGill's real C-11 and noticed only because a later screen printed
    ``[C-11]``.

    The grid had refused the same cell by name since it was written. One reader
    now (``services/c_number.py``), so the two doors cannot disagree; and the
    refusal is counted at the save, not only noted at the check.
    """

    databases = {"pipeline_db", "academy_db"}

    HEADER = "name\tgene\tgenotype\tparent\tc number\tsite"

    def setUp(self):
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)
        self.member = __import__(
            "pipeline.models", fromlist=["Member"]).Member.objects.using(DB).first()

    def _plan(self, cell):
        from pipeline.services import bulk_cell_lines as bulkcl
        rows = bulkcl.parse(
            f"{self.HEADER}\nHAP1 STMN2 KO\tSTMN2\tKO\tHAP1\t{cell}\tLeicester")
        return bulkcl.plan(rows, member=self.member)

    def test_a_label_it_cannot_read_is_refused_by_name_and_not_coerced(self):
        items = self._plan("C-RUN11-01")
        self.assertEqual(len(items), 1)
        it = items[0]
        self.assertIsNone(it["c_number"], "a digit run inside a label is not a C-number")
        self.assertIn("C-RUN11-01", it["note"])
        # "not a number" would be wrong about the case that matters: it *is* a
        # number, several of them. Say what a C-number is instead.
        self.assertIn("plain number", it["note"])

    def test_the_ordinary_ways_of_writing_one_still_read(self):
        for cell, expected in (("631", 631), ("C-631", 631), ("C631", 631),
                               (" c 631 ", 631), ("", None)):
            with self.subTest(cell=cell):
                it = self._plan(cell)[0]
                self.assertEqual(it["c_number"], expected)
                if expected is not None:
                    self.assertNotIn("C-number", it["note"])

    def test_the_check_counts_what_the_save_will_drop(self):
        from pipeline.services import bulk_cell_lines as bulkcl
        items = self._plan("C-RUN11-01")
        self.assertEqual(bulkcl.summarize(items)["no_c_number"], 1)

    def test_the_save_says_which_rows_went_in_without_one(self):
        """A row with an unreadable C-number is still written — deliberately, an
        odd batch label is no reason to discard a good cell line — so the save
        has to say so, or the reader presses the button and hears nothing."""
        from pipeline.services import bulk_cell_lines as bulkcl
        CellLine.objects.using(DB).create(name="HAP1", genotype="WT",
                                          site_id=self.site.pk)
        # The gene exists first: this paste no longer creates one (owner,
        # 5 Sep 2026), and what is under test here is the C-number.
        Target.objects.using(DB).create(gene_name="STMN2")
        rows = bulkcl.parse(
            f"{self.HEADER}\nHAP1 STMN2 KO\tSTMN2\tKO\tHAP1\tC-RUN11-01\tLeicester")
        out = bulkcl.apply(rows, member=self.member)
        self.assertEqual(len(out["created"]), 1)
        self.assertEqual([d["typed"] for d in out["no_c_number"]], ["C-RUN11-01"])
        line = CellLine.objects.using(DB).get(name="HAP1 STMN2 KO")
        self.assertIsNone(line.c_number, "a C-number nobody could read was invented")

    def test_the_grid_and_the_paste_box_answer_the_same_way(self):
        line = CellLine.objects.using(DB).create(
            name="HAP1", genotype="WT", c_number=631, site_id=self.site.pk)

        def patch(value):
            return self.client.post("/pipeline/cell-lines/board/patch/", {
                "target_id": line.pk, "field": "c_number", "value": value})

        # Accepted on both: C-742 is how it is written on the tube.
        self.assertEqual(patch("C-742").status_code, 200)
        self.assertEqual(CellLine.objects.using(DB).get(pk=line.pk).c_number, 742)
        # Refused on both, in the same words.
        resp = patch("C-RUN11-01")
        self.assertEqual(resp.status_code, 400)
        self.assertIn("plain number", resp.json()["error"])
        self.assertEqual(CellLine.objects.using(DB).get(pk=line.pk).c_number, 742)


def _excel_saved(path, sheets, active=1):
    """A workbook as **Excel** writes one: the tab that was selected on save is
    the active one, which may not be the first.

    openpyxl always writes ``activeTab="0"``, so a fixture written with the same
    library the app reads with cannot reach this at all — which is exactly why
    four field tests and the whole suite missed it. Patch the bytes.
    """
    import zipfile
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, rows in sheets:
        ws = wb.create_sheet(title)
        for row in rows:
            ws.append(row)
    wb.save(path)
    src = zipfile.ZipFile(path)
    items = {n: src.read(n) for n in src.namelist()}
    src.close()
    xml = items["xl/workbook.xml"].decode()
    xml = (re.sub(r'activeTab="\d+"', f'activeTab="{active}"', xml)
           if "activeTab" in xml
           else xml.replace("<workbookView", f'<workbookView activeTab="{active}"', 1))
    items["xl/workbook.xml"] = xml.encode()
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as out:
        for n, data in items.items():
            out.writestr(n, data)
    return path


class TheSheetWithTheDataIsTheOneThatIsReadTests(TestCase):
    """A ``Notes`` tab left selected in Excel used to replace the whole upload.

    Two readers keyed off ``wb.active`` — the Upload-a-sheet door on three
    boards, and a session's printable bench sheet — and ``wb.active`` is
    whichever tab was in front when the file was saved. Add a notes tab, save,
    upload: the app read the notes, found no rows, and said **nothing at all**.

    Both halves are pinned: the right sheet is read, and a file nothing could be
    read from is refused in words rather than answered with a blank panel.
    """

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)
        self.dir = tempfile.mkdtemp()

    def tearDown(self):
        shutil.rmtree(self.dir, ignore_errors=True)

    def _file(self, name, sheets, active=1):
        path = _excel_saved(os.path.join(self.dir, name), sheets, active=active)
        return open(path, "rb")

    def test_the_fixture_really_is_the_shape_the_bug_needs(self):
        """Guard the guard: if openpyxl ever wrote the active tab through, this
        whole class would pass while testing nothing."""
        import openpyxl
        with self._file("guard.xlsx", [("Cell lines", [["name"], ["HAP1"]]),
                                       ("Notes", [["my notes"]])]) as fh:
            wb = openpyxl.load_workbook(fh)
        self.assertEqual(wb.active.title, "Notes")
        self.assertEqual(wb.worksheets[0].title, "Cell lines")

    def test_an_upload_reads_the_data_sheet_not_the_selected_one(self):
        from pipeline.views.imports import _file_to_text
        with self._file("cl.xlsx", [
                ("Cell lines", [["name", "genotype", "site"],
                                ["HAP1", "WT", "Leicester"]]),
                ("Notes", [["my own notes about this batch"]])]) as fh:
            text, note = _file_to_text(fh, "cell-lines")
        self.assertIn("HAP1", text)
        self.assertNotIn("my own notes", text)
        # And it says which sheet it read, because a reader that picks silently
        # is one whose mistake cannot be found.
        self.assertIn("Cell lines", note)
        self.assertIn("Notes", note)

    def test_a_bench_sheet_reads_its_own_tab_too(self):
        from pipeline.services import bench_results
        with self._file("bench.xlsx", [
                ("Notes", [["reminder: order more ECL"]]),
                ("Bench sheet", [["Session #12"],
                                 ["Ab#", "CatNumber", "Used dilution"],
                                 ["1", "ab138501", "1:1000"]])], active=0) as fh:
            rows = bench_results._read_rows(fh)
        flat = [str(c) for r in rows for c in r if c is not None]
        self.assertIn("ab138501", flat)
        self.assertNotIn("reminder: order more ECL", flat)

    def test_a_file_nothing_could_be_read_from_is_refused_in_words(self):
        """The half that made this unfindable. The upload returned 200 and the
        panel rendered nothing, so a whole failed upload looked like a click
        that had not registered."""
        with self._file("empty.xlsx", [("Cell lines", []), ("Notes", [[""]])]) as fh:
            resp = self.client.post("/pipeline/import/upload/cell-lines/", {"file": fh})
        self.assertEqual(resp.status_code, 400)
        self.assertIn("No rows were read", resp.json()["error"])

    def test_the_preview_panel_never_renders_an_empty_string(self):
        """`out().innerHTML = cfg.renderPreview(d) || ''` was the shape of it —
        a preview with nothing to say wiped the panel."""
        source = (Path(settings.BASE_DIR)
                  / "pipeline/static/pipeline/board.js").read_text()
        self.assertNotIn("out().innerHTML = cfg.renderPreview(d) || ''", source,
                         "a preview with nothing to say still wipes the panel")
        self.assertIn("Nothing was read from that file", source)


class ASiteNamedInTheUrlActuallyFiltersTests(TestCase):
    """``?site=Leicester`` looked filtered and was not.

    ``sites.filter_by`` has taken a pk, a name or a short code since it was
    written, so the *server* narrowed correctly — but ``board.js`` rebuilds its
    query from the filter form, whose ``<select>`` carries pks, so a name
    matched no option, the box read *All sites*, and the very first rows fetch
    dropped the filter. The page looked filtered while the data was not: the
    same failure the ``?gene=`` rule already names, one control along.
    """

    databases = {"pipeline_db", "academy_db"}

    BOARDS = ("/pipeline/targets/board/", "/pipeline/antibodies/board/",
              "/pipeline/cell-lines/board/", "/pipeline/sessions/board/")

    def setUp(self):
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        Site.objects.using(DB).create(name="McGill", short_code="MCG")
        self.client = _member_client(self, self.site)

    def _selected(self, body):
        """The value of the site select's selected option."""
        block = re.search(r'<select name="site".*?</select>', body, re.S)
        self.assertIsNotNone(block, "no site select on the page")
        m = re.search(r'<option value="([^"]*)"[^>]*\bselected\b', block.group(0))
        return m.group(1) if m else ""

    def test_a_site_name_selects_the_same_option_its_pk_does(self):
        for url in self.BOARDS:
            for value in ("Leicester", "LEI", str(self.site.pk)):
                with self.subTest(url=url, site=value):
                    resp = self.client.get(url, {"site": value})
                    self.assertEqual(resp.status_code, 200)
                    self.assertEqual(self._selected(resp.content.decode()),
                                     str(self.site.pk))

    def test_a_site_that_is_not_on_file_is_shown_rather_than_dropped(self):
        """Silently widening to the whole consortium is the failure. Match
        nothing, and say what was asked for."""
        for url in self.BOARDS:
            with self.subTest(url=url):
                body = self.client.get(url, {"site": "Leicster"}).content.decode()
                self.assertEqual(self._selected(body), "Leicster")
                self.assertIn("not a site on file", body)


class OneListOfWhoCouldHaveRunItTests(TestCase):
    """Two doors to a session offered two different lists of experimenters.

    The step-by-step form added ``select_related('user')`` — an inner join on a
    non-nullable one-to-one — and ordered by ``user__first_name``, while the
    sessions board's quick panel listed every active member ordered by
    ``display_name``. A member whose row in pipeline_db's own ``auth_user``
    table is missing simply vanished from the form, with nothing on screen to
    say a name was absent.
    """

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)

    def test_both_doors_offer_the_same_people(self):
        from pipeline.services import session_board
        form = self.client.get("/pipeline/session/new/")
        board = self.client.get("/pipeline/sessions/board/")
        self.assertEqual(form.status_code, 200)
        self.assertEqual(board.status_code, 200)
        expected = [m.pk for m in session_board.filter_options()["experimenters"]]
        self.assertEqual([m.pk for m in form.context["experimenters"]], expected)
        self.assertEqual([m.pk for m in board.context["experimenters"]], expected)

    def test_the_shared_list_does_not_join_to_the_auth_table(self):
        """The join bought nothing — ``Member.__str__`` prefers ``display_name``
        and the site comes from its own ``select_related`` — and cost a name."""
        from pipeline.services import members
        self.assertNotIn("user", members.experimenters().query.select_related or {})


class TheConcentrationColumnIsOnTheBoardThatNamesItTests(TestCase):
    """"Add it on the board in µg/mL" pointed at a board with no such column.

    The value was editable, in the row payload, in the export and in the
    report's Table 2 — and drawn nowhere. So the one instruction the app gives
    for fixing a refused concentration could not be followed.
    """

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)

    def test_the_board_draws_the_column_the_message_sends_you_to(self):
        body = self.client.get("/pipeline/antibodies/board/").content.decode()
        self.assertIn("Conc.", body)
        source = (Path(settings.BASE_DIR)
                  / "pipeline/templates/pipeline/antibody_board.html").read_text()
        self.assertIn("cell(r, 'concentration'", source,
                      "the column has to be an editable cell, not a read-only one")

    def test_the_column_count_matches_the_header_count(self):
        """A colspan left behind makes the empty and failed rows span the wrong
        width — cosmetic, and the sort of thing that survives for releases.

        Asked of the **rendered page**, not the template source. Both numbers
        come from `services/board_columns.py` now, so counting literals in the
        file stopped being possible — and reading the page is the stronger
        question anyway: it is the width a person actually sees.
        """
        for url in ("/pipeline/antibodies/board/", "/pipeline/cell-lines/board/"):
            with self.subTest(url=url):
                body = self.client.get(url).content.decode()
                head = re.search(r"<thead.*?</thead>", body, re.S).group(0)
                headers = len(re.findall(r"<th\b", head))
                spans = {int(n) for n in re.findall(r'colspan="(\d+)"', body)}
                self.assertTrue(spans, f"{url} has no colspan to check")
                self.assertEqual(spans, {headers},
                                 f"{url}: {headers} headers, colspans {spans}")


class ABenchSheetSaysWhichSessionItIsForTests(TestCase):
    """"Both files carry this session's number" — and the bench sheet's name
    did not, so two for one gene and procedure collide in Downloads."""

    databases = {"pipeline_db", "academy_db"}

    def test_the_filename_carries_the_session_number(self):
        from pipeline.models import ExperimentSession
        from pipeline.services import planning
        site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        _member_client(self, site)
        member = __import__("pipeline.models",
                            fromlist=["Member"]).Member.objects.using(DB).first()
        target = Target.objects.using(DB).create(gene_name="STMN2")
        session = ExperimentSession.objects.using(DB).create(
            target_id=target.pk, site_id=site.pk, procedure_type="WB",
            experimenter_id=member.pk, date=date(2026, 8, 2))
        resp = planning.bench_sheet_to_response(session)
        self.assertIn(f"session_{session.pk}", resp["Content-Disposition"])


class TheReportDoesNotCallALineSuitableAgainstItsOwnNumberTests(TestCase):
    """The draft quoted the 2.5 log₂(TPM+1) cut-off and then declared the line
    suitable regardless — so run 11's STMN2 note called HAP1 suitable while the
    app's own feasibility page said HAP1 expresses STMN2 at 0.03.

    Both numbers come from one column. Only the document drew a conclusion from
    it, and it drew the wrong one.
    """

    databases = {"pipeline_db", "academy_db"}

    def _intro(self, expression):
        """The cell-line selection sentence for a target with this expression."""
        from docx import Document
        from pipeline.services import report_generator
        target = Target.objects.using(DB).create(
            gene_name="STMN2", protein_name="Stathmin-2",
            depmap_expression=expression)
        out = os.path.join(tempfile.mkdtemp(), "note.docx")
        report_generator.generate_report(target.pk, output_path=out)
        return "\n".join(p.text for p in Document(out).paragraphs)

    def test_a_line_below_the_cut_off_is_not_called_suitable(self):
        text = self._intro(Decimal("0.03"))
        self.assertIn("below", text)
        self.assertIn("[explain why this line was used]", text)
        self.assertNotIn("was identified as a suitable cell line", text)

    def test_a_line_above_the_cut_off_reads_as_it_always_did(self):
        self.assertIn("was identified as a suitable cell line",
                      self._intro(Decimal("11.13")))

    def test_an_unknown_expression_leaves_a_named_gap_not_a_verdict(self):
        text = self._intro(None)
        self.assertIn("[X.X]", text)
        self.assertIn("[confirm this line meets the cut-off]", text)
        self.assertNotIn("was identified as a suitable cell line", text)


class ACountAndTheListItCountsComeFromOneQueryTests(TestCase):
    """Run 11 read ``587 targets`` on the chip over 594 rows in the table.

    Not reproducible from the code — every board's rows endpoint returns
    ``count = len(rows)`` of the very list it draws, and each row renders one
    ``<tr>`` — so the count and the list cannot disagree by construction. This
    pins that construction, on all four boards, so the day one of them grows a
    second query for its count the test says so rather than a reader does.
    """

    databases = {"pipeline_db", "academy_db"}

    ROWS = ("/pipeline/targets/board/rows/", "/pipeline/antibodies/board/rows/",
            "/pipeline/cell-lines/board/rows/", "/pipeline/sessions/board/rows/")

    def setUp(self):
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)
        for gene in ("STMN2", "ELP3", "SOD1"):
            t = Target.objects.using(DB).create(gene_name=gene)
            TargetNomination.objects.using(DB).create(target_id=t.pk,
                                                      site_id=self.site.pk)
            CellLine.objects.using(DB).create(name=f"HAP1 {gene} KO", genotype="KO",
                                              target_id=t.pk, site_id=self.site.pk)

    def test_the_count_is_the_length_of_the_list(self):
        for url in self.ROWS:
            with self.subTest(url=url):
                data = self.client.get(url).json()
                self.assertEqual(data["count"], len(data["rows"]))

    def test_one_row_of_data_draws_one_table_row(self):
        """The other half of the arithmetic, and the only part a browser could
        get wrong: a ``rowHtml`` that opened two ``<tr>`` would put the table
        ahead of its own chip."""
        for name in ("target_board.html", "antibody_board.html",
                     "cell_line_board.html", "session_board.html"):
            with self.subTest(template=name):
                source = (Path(settings.BASE_DIR)
                          / "pipeline/templates/pipeline" / name).read_text()
                body = re.search(r"function rowHtml\(.*?\n  \}", source, re.S)
                self.assertIsNotNone(body, f"{name} has no rowHtml")
                self.assertEqual(len(re.findall(r"<tr\b", body.group(0))), 1,
                                 f"{name} draws more than one row per record")


class SigningOutIsNotSomethingALinkPreviewCanDoTests(TestCase):
    """A GET must not end a session.

    The twelfth field test fetched ``/pipeline/logout/`` in the background purely
    to read it and was signed out — no click, no confirmation. Anything that
    follows a link without a person deciding to does the same: a browser's link
    prefetcher, a chat client building an unfurl preview, a stray middle-click.
    A real user loses whatever they had half-typed and nothing says why.

    Both apps had it. Django's own ``LogoutView`` has been POST-only since 4.1
    for this reason and ``AcademyLogoutView`` had explicitly opted back out.
    """

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)

    def _signed_in(self):
        return "_auth_user_id" in self.client.session

    def test_a_get_to_the_pipeline_logout_does_not_sign_you_out(self):
        self.assertTrue(self._signed_in())
        response = self.client.get("/pipeline/logout/")
        self.assertEqual(response.status_code, 200)
        self.assertTrue(self._signed_in(),
                        "a GET to /pipeline/logout/ ended the session")

    def test_a_post_to_the_pipeline_logout_signs_you_out(self):
        response = self.client.post("/pipeline/logout/")
        self.assertEqual(response.status_code, 302)
        self.assertFalse(self._signed_in())

    def test_a_get_to_the_academy_logout_does_not_sign_you_out(self):
        """`academy:logout`, not `/accounts/logout/`.

        allauth's own logout has been safe all along — asking it this proves
        nothing and passes whatever the app does. The defect is in
        `AcademyLogoutView`, which overrode `http_method_names` to re-admit GET
        and is what the public site's nav actually links to.
        """
        from django.urls import reverse

        self.assertTrue(self._signed_in())
        response = self.client.get(reverse("academy:logout"))
        self.assertEqual(response.status_code, 200)
        self.assertTrue(self._signed_in(),
                        "a GET to the academy logout ended the session")

    def test_a_post_to_the_academy_logout_signs_you_out(self):
        from django.urls import reverse

        self.client.post(reverse("academy:logout"))
        self.assertFalse(self._signed_in())

    def test_the_nav_posts_rather_than_linking(self):
        """The ordinary path stays one click.

        Refusing the GET is only half of it: had the nav kept its ``<a href>``,
        pressing Logout would land on the confirmation every time, and the fix
        for a background fetch would have cost every real sign-out a click.
        """
        page = self.client.get("/pipeline/start/").content.decode()
        self.assertNotIn('<a href="/pipeline/logout/"', page)
        self.assertIn('action="/pipeline/logout/" method="post"', page)


class TheAccountPagesAreNotRawAllauthTests(TestCase):
    """Change password, change email and connections carry the site's chrome.

    They were django-allauth's own unstyled layout — Times New Roman on white, a
    bare "Menu:" list, no branding and no way back — and they are linked from the
    pipeline nav on *every* page, which made it the most visibly unfinished thing
    on the site.

    The fix is one layout override rather than three page copies, so this asks
    for all three: a page-specific fix would pass on one and leave the others,
    which is how the same defect has been shipped several times in this repo.
    """

    databases = {"pipeline_db", "academy_db"}

    # Reversed rather than typed, so the test follows allauth's own routing —
    # `socialaccount_connections` moved to /accounts/3rdparty/ and the old path
    # is a 301 that drops its query string.
    ACCOUNT_URL_NAMES = ("account_change_password", "account_email",
                         "socialaccount_connections")

    def setUp(self):
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)

    def test_every_account_page_carries_the_site_chrome(self):
        from django.urls import reverse

        for name in self.ACCOUNT_URL_NAMES:
            with self.subTest(page=name):
                page = self.client.get(reverse(name)).content.decode()
                self.assertNotIn("Menu:", page,
                                 f"{name} is still allauth's own layout")
                self.assertIn("bootstrap", page.lower(),
                              f"{name} has no stylesheet")
                self.assertIn("Only Good Antibodies", page,
                              f"{name} carries no branding")

    def test_the_url_the_field_test_used_still_lands_somewhere_styled(self):
        """The pipeline reached connections at the pre-rename path. It 301s, and
        a 301 to allauth's raw layout would look exactly like no fix at all."""
        page = self.client.get("/accounts/social/connections/",
                               follow=True).content.decode()
        self.assertNotIn("Menu:", page)
        self.assertIn("Only Good Antibodies", page)

    def test_every_account_page_offers_a_way_back(self):
        """Arriving from the pipeline, there was no route back but the Back
        button. The nav sends ?next=, which is also where the save returns."""
        from django.urls import reverse

        for name in self.ACCOUNT_URL_NAMES:
            with self.subTest(page=name):
                page = self.client.get(
                    reverse(name),
                    {"next": "/pipeline/sessions/board/"}).content.decode()
                self.assertIn("/pipeline/sessions/board/", page,
                              f"{name} offers no way back to where you came from")

    def test_the_pipeline_nav_says_where_it_came_from(self):
        page = self.client.get("/pipeline/start/").content.decode()
        self.assertIn("/accounts/password/change/?next=", page,
                      "the nav sends no ?next=, so the account page cannot "
                      "offer a way back into the pipeline")


class TheSignInPageDoesNotAdvertiseADeadRouteTests(TestCase):
    """Sign Up is not offered to somebody being sent to the pipeline.

    A self-registered Academy account reaches no pipeline page: access is three
    rows in two databases and only a superuser grants it, so the invitation reads
    as the way in and ends at "Access Denied". The body copy already knew this;
    the header did not — the same rule enforced in one of two places, which is
    the shape of most of the defects in this file.
    """

    databases = {"pipeline_db", "academy_db"}

    def test_the_pipeline_sign_in_does_not_offer_sign_up(self):
        from django.urls import reverse

        page = self.client.get(
            "/accounts/login/", {"next": "/pipeline/start/"}).content.decode()
        self.assertIn("Sign in to the YCharOS Pipeline", page)
        self.assertNotIn(reverse("academy:signup"), page,
                         "the pipeline sign-in page advertises Academy signup")

    def test_the_academy_sign_in_still_offers_sign_up(self):
        """The route is real for the Academy — it is only the pipeline it
        cannot reach, so hiding it everywhere would be the opposite mistake.

        This is the case the old wording got wrong in the other direction: with
        no ?next= at all the value is None, `'/pipeline/' not in None` raises, and
        a template swallows that as false — so Sign Up was hidden from the one
        page it belongs on.
        """
        from django.urls import reverse

        page = self.client.get("/accounts/login/").content.decode()
        self.assertIn(reverse("academy:signup"), page)

    def test_the_sign_in_field_says_what_to_type(self):
        """allauth labels the field "Login" when it accepts a username *or* an
        email — a label that names the page rather than the field. It knows the
        right words and puts them in the placeholder, which vanishes on typing.
        """
        page = self.client.get("/accounts/login/").content.decode()
        self.assertIn("Username or email", page)
        self.assertNotIn(">Login</label>", page)

    def test_signing_in_with_no_next_lands_on_the_academy(self):
        """A sign-in that works must not end on a 404.

        The hidden `next` field was rendered unconditionally over
        `redirect_field_value`, which is None with no ?next= — and a template
        prints None as the four characters `None`. So the form posted
        `next=None`, allauth read it as a relative URL, and a correct password
        landed on /academy/login/None. The pipeline nav's Login link carries no
        ?next=, so that was every sign-in taken from it.
        """
        from django.contrib.auth import get_user_model

        get_user_model().objects.create_user(username="reader", password="pw12345")
        page = self.client.get("/accounts/login/").content.decode()
        self.assertNotIn('value="None"', page,
                         "the sign-in form posts the string None as its redirect")

        signed_in = self.client.post("/academy/login/",
                                     {"login": "reader", "password": "pw12345"},
                                     follow=True)
        self.assertEqual(signed_in.status_code, 200)
        self.assertEqual(signed_in.redirect_chain[-1][0], "/academy/")


class ThePortfolioSaysWhatEachTableCountsTests(TestCase):
    """Three tables, three denominators, and two of them said so nowhere.

    The header counts targets; "By site" counts (site, target) pairs and "By
    funder" counts (funder, target) pairs. So the site column read 373 under a
    headline of 583 with no caveat — on the page a grant gets written from, where
    two disagreeing numbers on one screen read as records having been lost. The
    class table above them already carried its caveat.
    """

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        self.leicester = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.mcgill = Site.objects.using(DB).create(name="McGill", short_code="MCG")
        self.client = _member_client(self, self.leicester)

        # One gene both sites are pursuing, one only Leicester, one nobody has
        # nominated anywhere — the three cases the caveat has to describe.
        self.shared = Target.objects.using(DB).create(gene_name="STMN2")
        self.mine = Target.objects.using(DB).create(gene_name="ELP3")
        self.orphan = Target.objects.using(DB).create(gene_name="TRPA1")
        for site in (self.leicester, self.mcgill):
            TargetNomination.objects.using(DB).create(
                target_id=self.shared.pk, site_id=site.pk, funded=True)
        TargetNomination.objects.using(DB).create(
            target_id=self.mine.pk, site_id=self.leicester.pk, funded=False)

    def test_the_coverage_numbers_reconcile_the_site_table_to_the_header(self):
        from pipeline.services.target_board import portfolio

        p = portfolio()
        c, totals = p["coverage"], p["totals"]

        self.assertEqual(c["site_rows"], sum(x["total"] for x in p["by_site"].values()),
                         "the caveat's row count is not what the column adds up to")
        self.assertEqual(c["site_targets"] + c["site_missing"], totals["targets"],
                         "the targets the table can see plus the ones it cannot "
                         "do not account for the headline figure")
        self.assertEqual(c["site_rows"] - c["site_shared"], c["site_targets"],
                         "the double-counted rows are not accounted for")
        self.assertEqual(c["site_missing"], 1)   # TRPA1, nominated nowhere
        self.assertEqual(c["site_shared"], 1)    # STMN2, counted at two sites

    def test_the_coverage_numbers_reconcile_the_funder_table_to_the_header(self):
        from pipeline.services.target_board import portfolio

        p = portfolio()
        c, totals = p["coverage"], p["totals"]
        self.assertEqual(c["agency_rows"], sum(x["total"] for x in p["by_agency"].values()))
        self.assertEqual(c["agency_targets"] + c["agency_missing"], totals["targets"])

    def test_the_page_prints_the_caveat_under_both_tables(self):
        """Computing it and not drawing it is the same bug — this repo has shipped
        that one several times, most recently with the workbook's dropped rows."""
        page = self.client.get("/pipeline/targets/portfolio/").content.decode()
        body = page.split("By site", 1)[1]
        # "no nomination naming a site" until the thirteenth field test: 192 of
        # the targets that sentence was about *do* have a site, recorded by the
        # Access import and shown on Overview all along, so the column was wrong
        # rather than incomplete and the caveat was describing the wrong gap.
        self.assertIn("no site recorded anywhere", body)
        self.assertIn("no nomination naming a funder", body)
        self.assertIn("more than one site", body)


class TheStageChartIsDrawnInProportionTests(TestCase):
    """Bar heights are a share of the tallest, not a raw pixel count.

    Each bar's height was `{{ stage.count }}px` under `max-height: 200px`, so
    every stage over 200 drew at exactly the same height — Feasibility at 262
    rendering barely taller than Published at 134, when it should have been
    nearly double, and any two stages above the cap indistinguishable. A chart is
    read at a glance; that is the whole reason to draw one beside the numbers.
    """

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)
        # Stages of unequal height, or there is no proportion to check: an empty
        # pipeline draws every bar at zero and the test passes on nothing.
        for i in range(5):
            Target.objects.using(DB).create(gene_name=f"FEAS{i}")
        for i in range(2):
            Target.objects.using(DB).create(gene_name=f"PUB{i}", status="published")

    def test_no_bar_is_sized_in_raw_pixels(self):
        page = self.client.get("/pipeline/overview/").content.decode()
        self.assertNotIn("max-height: 200px", page,
                         "bar heights are still capped raw counts")

    def test_the_tallest_bar_is_the_biggest_count(self):
        """`stage_max` has to be the maximum, not the first or the last — a
        template cannot compute one, so getting it wrong here is invisible on the
        page until two stages happen to straddle the value."""
        from pipeline.views.dashboard import master_dashboard  # noqa: F401

        response = self.client.get("/pipeline/overview/")
        stages, stage_max = response.context["stages"], response.context["stage_max"]
        self.assertEqual(stage_max, max(max(s["count"] for s in stages), 1))
        self.assertGreaterEqual(stage_max, 1, "a zero would divide by zero")

    def test_a_bar_twice_the_count_is_twice_the_height(self):
        """The property the raw-pixel version broke, asserted on the rendered
        style rather than on the numbers behind it."""
        import re

        page = self.client.get("/pipeline/overview/").content.decode()
        heights = [int(m) for m in re.findall(r"height:\s*(\d+)%", page)]
        response = self.client.get("/pipeline/overview/")
        counts = [s["count"] for s in response.context["stages"]]
        self.assertEqual(len(heights), len(counts),
                         "one percentage height per stage")
        biggest = max(range(len(counts)), key=lambda i: counts[i])
        self.assertEqual(heights[biggest], 100,
                         "the tallest bar does not fill the track")


class OverviewDrawsAPageNotTheDatasetTests(TestCase):
    """358 active targets used to render as one 16,000px table.

    The same defect `services/board_page.py` was written for after the antibodies
    board handed a browser 3,225 rows and got a "page unresponsive" dialog. This
    page never got the fix, and it is the one a coordinator opens first.
    """

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)
        # More targets than fit on a page, each active by the page's own
        # definition: an antibody on file and no report.
        from pipeline.models import Antibody, Company

        company = Company.objects.using(DB).create(name="Abcam")
        for i in range(30):
            target = Target.objects.using(DB).create(gene_name=f"GENE{i:03d}")
            Antibody.objects.using(DB).create(
                target_id=target.pk, company_id=company.pk,
                catalogue_number=f"ab{i:05d}", site_id=self.site.pk)

    def test_the_table_draws_one_page(self):
        response = self.client.get("/pipeline/overview/")
        drawn = len(response.context["active_targets"].object_list)
        self.assertLess(drawn, 30, "every active target is still drawn at once")

    def test_the_count_above_the_table_is_the_whole_set(self):
        """Paginating the number as well turns 358 into 25 and reads as work
        having disappeared — the failure the boards' pager was built around."""
        response = self.client.get("/pipeline/overview/")
        self.assertEqual(response.context["active_total"], 30)
        self.assertIn("30 targets", response.content.decode())

    def test_a_second_page_holds_the_rest(self):
        first = self.client.get("/pipeline/overview/")
        second = self.client.get("/pipeline/overview/?page=2")
        drawn = [t.pk for t in first.context["active_targets"]]
        rest = [t.pk for t in second.context["active_targets"]]
        self.assertTrue(rest, "page 2 is empty")
        self.assertFalse(set(drawn) & set(rest), "page 2 repeats page 1")

    def test_a_nonsense_page_number_does_not_500(self):
        for value in ("0", "-1", "abc", "99999"):
            with self.subTest(page=value):
                response = self.client.get(f"/pipeline/overview/?page={value}")
                self.assertEqual(response.status_code, 200)

    def test_the_gene_name_is_a_link(self):
        """Styled link-blue and not a link: the row's onclick navigated, so it
        worked for a left-click and for nothing else — no middle-click, no new
        tab, no keyboard, no destination in the status bar."""
        target = Target.objects.using(DB).get(gene_name="GENE000")
        page = self.client.get("/pipeline/overview/").content.decode()
        self.assertIn(f'href="/pipeline/target/{target.pk}/"', page)


class TheGuideIsNotAddressedToOnePersonTests(TestCase):
    """Four board guides, and one of them was headed "a guide for Carl".

    It is served to every pipeline member at /pipeline/targets/guide/, and the
    other three already read "— a guide".

    The heading was the visible half. The body of the target guide was a reply
    to one reader throughout — *your workbook*, *at your request*, *you said all
    you need is funded and completed* — and the field test that reported the
    heading reported the tone, which is the part that keeps going stale: half of
    those sentences describe a moment (a spreadsheet not yet imported, a scope
    decision being negotiated) that has since passed, so a reader who takes them
    at face value is reading last month's app.

    What is pinned here is only what can be checked cheaply: no named reader,
    and no guide pointing at a page that was retired. The prose is the owner's
    to judge.
    """

    databases = {"pipeline_db", "academy_db"}

    GUIDES = ("TARGET_BOARD_GUIDE.md", "ANTIBODY_BOARD_GUIDE.md",
              "CELL_LINE_BOARD_GUIDE.md", "SESSION_BOARD_GUIDE.md")

    def _text(self, name):
        return (Path(settings.BASE_DIR) / name).read_text(encoding="utf-8")

    def test_no_guide_is_addressed_to_a_named_reader(self):
        for name in self.GUIDES:
            with self.subTest(guide=name):
                text = self._text(name)
                self.assertNotIn("for Carl", text)
                heading = text.splitlines()[0]
                self.assertTrue(heading.endswith("a guide"),
                                f"{name} is headed {heading!r}")

    def test_no_guide_is_written_to_one_person_s_spreadsheet(self):
        """"Your workbook" and "at your request" are addressed to whoever asked
        for the change, and every other reader is left working out whether a
        sentence is about them. The generic *you* — "leave a cell blank and the
        row is recorded at your own site" — is fine and is not what this asks
        about."""
        for name in self.GUIDES:
            with self.subTest(guide=name):
                text = self._text(name).lower()
                for phrase in ("your workbook", "your spreadsheet",
                               "your original file", "at your request",
                               "you said", "until your list is imported"):
                    self.assertNotIn(phrase, text,
                                     f"{name} is written to one reader: {phrase!r}")

    def test_no_guide_sends_you_to_a_page_that_was_retired(self):
        """The same rule the boards' own prose is held to, and it bites harder
        here: a guide is what somebody reads *because* they are already stuck,
        and thirteen pages went on 31 Jul 2026. The antibodies guide still said
        identity was fixed "on the antibody's own page" and the sessions guide
        sent you to "the session's own page" for a result row."""
        for name in self.GUIDES:
            with self.subTest(guide=name):
                text = self._text(name).lower()
                for dead in ("antibody's own page", "cell line page",
                             "session's own page", "still its own page",
                             "adding lines is still", "the antibody page"):
                    self.assertNotIn(dead, text,
                                     f"{name} points at a retired page: {dead!r}")


class AWrongUrlLandsSomewhereWithAWayHomeTests(TestCase):
    """There was no 404 template and no handler404 at all.

    So a mistyped address, an old bookmark, or a link into one of the thirteen
    pages retired on 31 Jul 2026 got Django's bare "Not Found" — different
    typeface, no branding, no way back. Django picks templates/404.html up by
    filename, but only when DEBUG is off, which is why this test says so.

    A 404 template that raises is silently replaced by the default, so the test
    that matters is that it renders at all.
    """

    databases = {"pipeline_db", "academy_db"}

    def test_a_missing_page_is_branded_and_offers_a_way_home(self):
        with self.settings(DEBUG=False, ALLOWED_HOSTS=["testserver"]):
            response = self.client.get("/no-such-page-anywhere/")
        self.assertEqual(response.status_code, 404)
        body = response.content.decode()
        self.assertIn("Only Good Antibodies", body)
        self.assertIn('href="/"', body)

    def test_a_retired_pipeline_url_is_pointed_at_the_hub(self):
        """The most likely way to reach a 404 here is a bookmark for a page that
        was retired in favour of a board, so that reader gets the hub rather than
        the public home page."""
        with self.settings(DEBUG=False, ALLOWED_HOSTS=["testserver"]):
            response = self.client.get("/pipeline/no-such-board/")
        self.assertEqual(response.status_code, 404)
        self.assertIn("/pipeline/start/", response.content.decode())

    def test_the_error_page_renders_without_a_request_context(self):
        """A 500 is rendered with an empty context and no context processors, by
        design — so anything dynamic in it fails at exactly the moment it is
        needed. Rendered here the way Django renders it."""
        from django.template.loader import get_template

        body = get_template("500.html").render({})
        self.assertIn("Only Good Antibodies", body)
        self.assertNotIn("{{", body)


class ProteinClassesAreNotBuriedUnderGeneFamiliesTests(TestCase):
    """The grant-writing table shows 22 curated classes, not 318 rows.

    `protein_class.classify` emits a "<PREFIX> family" label for every gene symbol
    that has one — RAB11A, RAB5C and RAB7A all become "RAB family" — and those
    went into the same table as the curated rubric. 296 families against 22 real
    classes, sorted by size, 241 of the families holding a single gene. The rows
    somebody would quote in a grant were unfindable.

    The families are kept, not dropped: `family_expertise` reads them for the
    "McGill has done most Rabs" warning on Add Targets.
    """

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        from pipeline.models import TargetClassification

        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)

        rab = Target.objects.using(DB).create(gene_name="RAB11A")
        gnaq = Target.objects.using(DB).create(gene_name="GNAQ")
        for target, label, source in (
            (rab, "RAB family", TargetClassification.Source.FAMILY),
            (rab, "Membrane", TargetClassification.Source.UNIPROT),
            (gnaq, "GPCR", TargetClassification.Source.UNIPROT),
            # A hand-added label that duplicates a derived one: both rows are
            # legitimate — the unique constraint is per (target, label, source).
            (gnaq, "GPCR", TargetClassification.Source.MANUAL),
        ):
            TargetClassification.objects.using(DB).create(
                target_id=target.pk, label=label, source=source)

    def test_a_family_is_not_listed_as_a_protein_class(self):
        from pipeline.services.target_board import portfolio

        p = portfolio()
        self.assertIn("Membrane", p["by_class"])
        self.assertIn("GPCR", p["by_class"])
        self.assertNotIn("RAB family", p["by_class"],
                         "gene families are still in the protein class table")
        self.assertIn("RAB family", p["by_family"],
                      "gene families were dropped rather than moved")

    def test_a_class_counts_targets_not_rows(self):
        """Two rows, one target. Counting rows inflated the exact figure that
        gets typed into a grant."""
        from pipeline.services.target_board import portfolio

        self.assertEqual(portfolio()["by_class"]["GPCR"]["total"], 1)

    def test_the_page_draws_the_families_separately(self):
        # Split on the <details> element, not on the words "Gene families" — the
        # intro paragraph above the class table says the families are listed
        # below, so partitioning on the phrase cuts the page in the wrong place
        # and the assertion passes or fails for the wrong reason.
        page = self.client.get("/pipeline/targets/portfolio/").content.decode()
        classes, marker, families = page.partition("<details")
        self.assertTrue(marker, "the families are not in a collapsed section")
        self.assertIn("GPCR", classes)
        self.assertNotIn("RAB family", classes,
                         "a family is drawn in the protein class table")
        self.assertIn("RAB family", families)


# ---------------------------------------------------------------------------
# Batch six of the Cowork field tests, 5 Aug 2026. Four boards, six findings,
# and every one of them is a screen asserting something the record does not say.
# ---------------------------------------------------------------------------


class AUnitIsNeverCaseTransformedTests(TestCase):
    """"The concentration column says MG/ML but the numbers are µg/mL."

    The board's header row is styled `uppercase`, and CSS uppercases `µ`
    (U+00B5 MICRO SIGN) to `Μ` (U+039C GREEK CAPITAL MU) — the same glyph as a
    Latin M in every font this site uses. So the one column whose whole history
    is a thousandfold unit error (`services/concentration.py`) was drawn with a
    heading a thousandfold out, over values of 1000 and 529.
    """

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)

    def test_the_heading_carries_no_micron_into_an_uppercased_span(self):
        from pipeline.services import board_columns

        for column in board_columns.registry("antibodies"):
            if column.board == board_columns.OFF or not column.label:
                continue
            self.assertNotIn(
                "µ", column.label,
                f"{column.key}'s board heading carries a µ, which the header "
                f"row's `uppercase` renders as a capital M")

    def test_the_unit_is_still_shown(self):
        """Dropping the µ would be the wrong fix — the column would then say
        nothing about its unit at all, on the field that has no unit column."""
        from pipeline.services import board_columns

        conc = next(c for c in board_columns.registry("antibodies")
                    if c.key == "concentration")
        self.assertEqual(conc.unit, "µg/mL")
        page = self.client.get("/pipeline/antibodies/board/").content.decode()
        self.assertIn('normal-case">(µg/mL)</span>', page,
                      "the unit is drawn inside the uppercased header row")


class AKnockoutThatFailedIsNotConfirmedTests(TestCase):
    """"A knockout line whose validation FAILED is shown with a green
    'validated' badge", and "on the unvalidated rows, the word 'validated' still
    appears".

    Both are the same defect: the board drew `ko_validated` — one bit — and
    called it *validated* in green when set and *validated* in grey when not.
    On the live data the bit and the reason beside it disagree on 54 rows
    (counted from the Access export in `services/ko_validation.py`), six of
    them a green confirmation directly above their own record saying the check
    failed.
    """

    databases = {"pipeline_db", "academy_db"}

    @classmethod
    def setUpTestData(cls):
        cls.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        cls.target = Target.objects.using(DB).create(gene_name="ACE")
        cls.failed = CellLine.objects.using(DB).create(
            name="SKNFI", genotype="KO", target_id=cls.target.pk,
            site_id=cls.site.pk, ko_validated=True,
            ko_validation_notes="Failed-Decreased protein level")
        cls.good = CellLine.objects.using(DB).create(
            name="HAP1", genotype="KO", target_id=cls.target.pk,
            site_id=cls.site.pk, ko_validated=True,
            ko_validation_notes="Confirmed")
        cls.untold = CellLine.objects.using(DB).create(
            name="U2OS", genotype="KO", target_id=cls.target.pk,
            site_id=cls.site.pk, ko_validated=False)
        cls.unticked = CellLine.objects.using(DB).create(
            name="HeLa", genotype="KO", target_id=cls.target.pk,
            site_id=cls.site.pk, ko_validated=False,
            ko_validation_notes="Confirmed")

    def test_a_failed_check_is_not_a_confirmation(self):
        from pipeline.services import ko_validation

        self.assertFalse(ko_validation.confirmed(self.failed))
        self.assertTrue(ko_validation.disagrees(self.failed))
        self.assertTrue(ko_validation.confirmed(self.good))

    def test_an_unticked_row_stating_confirmed_is_a_question_not_a_verdict(self):
        """Forty-eight live rows are this way round. Silently promoting them
        would be inferring a verdict from free text; silently ignoring them is
        how the screen said *not* validated over a record saying it was."""
        from pipeline.services import ko_validation

        self.assertFalse(ko_validation.confirmed(self.unticked))
        self.assertTrue(ko_validation.disagrees(self.unticked))

    def test_free_text_that_states_no_verdict_leaves_the_tick_alone(self):
        """Finding a word in a string is not reading a verdict — the same rule
        `services/c_number.py` holds about finding a number in one."""
        from pipeline.services import ko_validation

        self.untold.ko_validation_notes = "blot run 12 Mar, see the folder"
        self.assertEqual(ko_validation.outcome(self.untold), ko_validation.UNCHECKED)
        self.untold.ko_validated = True
        self.assertEqual(ko_validation.outcome(self.untold), ko_validation.CONFIRMED)

    def test_no_row_is_ever_labelled_the_bare_word_validated(self):
        """Scanning the column, every row appeared to say "validated" — you had
        to notice the grey to know it meant the opposite. Colour is not a
        value."""
        from pipeline.services import cell_line_board

        for line in (self.failed, self.good, self.untold, self.unticked):
            label = cell_line_board.row_for(line)["ko_label"]
            self.assertNotEqual(label.lower().strip(), "validated", line.name)
        self.assertEqual(cell_line_board.row_for(self.untold)["ko_label"],
                         "not confirmed")

    def test_the_failed_row_is_not_green(self):
        from pipeline.services import cell_line_board

        self.assertNotEqual(cell_line_board.row_for(self.failed)["ko_tone"], "good")
        self.assertEqual(cell_line_board.row_for(self.good)["ko_tone"], "good")

    def test_the_filter_means_what_the_badge_means(self):
        """This is how it was found: filter to Validated, and the very first row
        is a failure. A filter that disagrees with the column it narrows reads
        as one of the two being broken, with no way to tell which."""
        from pipeline.services import cell_line_board

        confirmed = [r["name"] for r in cell_line_board.board_rows(ko_validated="yes")]
        self.assertEqual(confirmed, ["HAP1"])
        disputed = {r["name"] for r in cell_line_board.board_rows(ko_validated="disputed")}
        self.assertEqual(disputed, {"SKNFI", "HeLa"})
        not_confirmed = {r["name"] for r in cell_line_board.board_rows(ko_validated="no")}
        self.assertEqual(not_confirmed, {"SKNFI", "U2OS", "HeLa"})

    def test_the_gene_page_counts_the_same_way(self):
        """A count and a badge on two screens must not disagree about one row —
        the same family as `Target.ko_validated`, which said No beside a strip
        saying the knockout was confirmed."""
        client = _member_client(self, self.site)
        page = client.get(f"/pipeline/target/{self.target.pk}/").content.decode()
        self.assertIn("1 of 4", page,
                      "the gene page counted a failed knockout as confirmed")

    def test_the_progress_strip_counts_it_the_same_way_too(self):
        """The same page, a few centimetres down, and it was still reading the
        tick. `views/dashboard.py` learned this and `services/gene_progress.py`
        did not, so TARGET INFORMATION said **none confirmed yet** while the
        strip said **KO confirmed ✓ · 2 of 4** — two answers to one question in
        one viewport, which is the failure this whole family is about.

        The disputed rows are *named* rather than quietly dropped from the
        numerator: a line left out of a count is a line nobody knows to fix.
        """
        from pipeline.services import gene_progress

        steps = gene_progress.steps_for(self.target)
        ko = next(s for s in steps if s["key"] == "ko_validated")
        self.assertEqual(ko["detail"], "1 of 4 confirmed · 2 to check")

    def test_a_gene_whose_only_knockout_is_disputed_is_not_finished(self):
        """The third contradiction the field test read off one screen: with the
        strip counting the tick, the last unticked step went away and the page
        announced **Everything on this list is done.** over a knockout whose own
        record says the check failed."""
        from pipeline.services import gene_progress

        lone = Target.objects.using(DB).create(gene_name="SOD1")
        CellLine.objects.using(DB).create(
            name="SKNFI", genotype="KO", target_id=lone.pk, site_id=self.site.pk,
            ko_validated=True, ko_validation_notes="Failed-Truncated protein")

        steps = gene_progress.steps_for(lone)
        ko = next(s for s in steps if s["key"] == "ko_validated")
        self.assertFalse(ko["done"])
        self.assertEqual(ko["detail"], "0 of 1 confirmed · 1 to check")


class AResultRowIsNotAReadingTests(TestCase):
    """"It's headed RESULTS and says '22 results' when no result has been
    entered. It's counting antibody slots, not results."

    Planning a session writes one blank row per antibody it will test, so the
    column read "22 results" over two sessions nobody had run. The status was
    right and the count was the lie.

    What the count must NOT do is turn into a verdict. Gaps in a results section
    are normal and accepted, and what settles completion is a published report's
    DOI (owner, 5 Aug) — so this reports rows and readings and concludes
    nothing from the difference.
    """

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        from pipeline.models import (Antibody, Company, ExperimentSession, Member,
                                     WbResult)

        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)
        member = Member.objects.using(DB).get(site_id=self.site.pk)
        self.target = Target.objects.using(DB).create(gene_name="NR3C1")
        company = Company.objects.using(DB).create(name="Abcam")
        self.session = ExperimentSession.objects.using(DB).create(
            target_id=self.target.pk, procedure_type="WB", site_id=self.site.pk,
            experimenter_id=member.pk, date=date(2026, 4, 22),
            status=ExperimentSession.SessionStatus.PLANNED)
        # Three antibodies queued up and not one reading written on any of them
        # — the state the field test opened and the board called "3 results".
        for n in range(3):
            ab = Antibody.objects.using(DB).create(
                catalogue_number=f"ab{n}", company_id=company.pk,
                target_id=self.target.pk, site_id=self.site.pk)
            WbResult.objects.using(DB).create(session_id=self.session.pk,
                                              antibody_id=ab.pk)

    def test_a_planned_session_of_blank_rows_reports_no_readings(self):
        from pipeline.services import session_board

        row = session_board.row_for(
            self.session,
            session_board.result_counts([self.session.pk]),
            session_board.reading_counts([self.session.pk]))
        self.assertEqual(row["result_count"], 3)
        self.assertEqual(row["reading_count"], 0,
                         "blank result rows were counted as results")

    def test_a_written_row_counts(self):
        from pipeline.models import WbResult
        from pipeline.services import session_board

        first = WbResult.objects.using(DB).filter(session_id=self.session.pk).first()
        first.signal = "single band ~95 kDa"
        first.save(using=DB)
        self.assertEqual(
            session_board.reading_counts([self.session.pk])[self.session.pk], 1)

    def test_a_zero_is_a_reading_and_a_blank_is_not(self):
        from pipeline.services.session_board import is_reading

        self.assertTrue(is_reading(0))
        self.assertTrue(is_reading("0"))
        self.assertFalse(is_reading(""))
        self.assertFalse(is_reading(None))
        self.assertFalse(is_reading(False))
        self.assertFalse(is_reading("   "))

    def test_the_two_readers_of_this_rule_agree(self):
        """`report_generator._has_readings` walks prefetched objects and
        `reading_counts` asks the database. Two shapes of one question, and a
        report dropping a session the board says has readings — or the reverse
        — is the kind of disagreement nobody can see from either screen."""
        from pipeline.models import ExperimentSession, WbResult
        from pipeline.services import report_generator, session_board

        def report_says(session_pk):
            fresh = (ExperimentSession.objects.using(DB)
                     .prefetch_related("wb_results").get(pk=session_pk))
            return report_generator._has_readings(fresh)

        self.assertFalse(report_says(self.session.pk))
        self.assertEqual(
            session_board.reading_counts([self.session.pk]).get(self.session.pk, 0), 0)

        row = WbResult.objects.using(DB).filter(session_id=self.session.pk).first()
        row.rating = "good"
        row.save(using=DB)
        self.assertTrue(report_says(self.session.pk))
        self.assertEqual(
            session_board.reading_counts([self.session.pk])[self.session.pk], 1)

    def test_the_gene_page_does_not_call_a_planned_session_complete(self):
        """The same count, one page over, and the harmful direction: the gene
        page marks a procedure complete on `results > 0`, so three blank rows
        read as a finished western blot."""
        response = self.client.get(f"/pipeline/target/{self.target.pk}/")
        self.assertFalse(response.context["procedure_summary"]["WB"]["complete"],
                         "a planned session of blank rows marked WB complete")


class TargetsWithNoSiteAreReachableTests(TestCase):
    """"The Portfolio quietly leaves out 210 targets."

    The caveat naming them is on this branch already. What it did not have was
    anywhere to go: a number in a caveat with nothing to click is half a
    message, and this is the only page in the app that knows those targets
    exist.
    """

    databases = {"pipeline_db", "academy_db"}

    @classmethod
    def setUpTestData(cls):
        cls.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        cls.homed = Target.objects.using(DB).create(gene_name="TRPA1")
        TargetNomination.objects.using(DB).create(target_id=cls.homed.pk,
                                                  site_id=cls.site.pk)
        cls.stranded = Target.objects.using(DB).create(gene_name="STMN2")
        # A nomination with no site is the second way to be missing from the
        # table, and `filter(site__isnull=True)` finds this one while missing
        # the target with no nominations at all.
        cls.half = Target.objects.using(DB).create(gene_name="ELP3")
        TargetNomination.objects.using(DB).create(target_id=cls.half.pk)

    def test_site_none_finds_the_targets_with_no_site(self):
        from pipeline.services import target_board

        genes = {r["gene"] for r in target_board.board_rows(site="none")}
        self.assertEqual(genes, {"STMN2", "ELP3"})

    def test_the_form_can_carry_it(self):
        """board.js rebuilds its query from the form, so a filter the select
        cannot hold is dropped by the first rows fetch — the page looks filtered
        while the data is not."""
        from pipeline.services import sites as site_svc

        self.assertEqual(site_svc.form_value("none"), "none")
        client = _member_client(self, self.site)
        page = client.get("/pipeline/targets/board/?site=none").content.decode()
        self.assertIn('<option value="none" selected>No site recorded</option>',
                      page.replace(" \n", "").replace("  ", " "))
        self.assertNotIn("not a site on file", page,
                         "the sentinel was rendered as an unknown site as well")

    def test_the_portfolio_links_to_them(self):
        client = _member_client(self, self.site)
        page = client.get("/pipeline/targets/portfolio/").content.decode()
        self.assertIn("?site=none", page)


class RawDataIsReachableFromTheGenePageTests(TestCase):
    """The files a Zenodo deposit is made of, on the page the deposit is pressed.

    `services/deposit.py` packages every `FileAttachment` on a gene into
    `<GENE>_underlying_data.zip`, and the button that does it is at the bottom of
    the gene's own page — while the only surface that could *attach* one was the
    sessions board, inside a session's drawer. So the preview said "0 raw
    file(s)" and there was nothing on that page to do about it. A record whose
    underlying data is an empty archive is the shape of published work with
    nothing under it.

    Pins the three halves that would be wrong *silently*:

    * both surfaces mount **one** panel. Two copies writing raw lab data through
      the same endpoints drift on which categories they offer and on whether a
      failed read draws "no files yet", and from outside there is no telling
      which of the two is the broken one.
    * the mount carries its endpoints. A panel with no URLs is the mirror of an
      export with no importer: it draws, and nothing it does reaches anything.
    * the count beside the button is the files that are on file. A panel that
      counts must list what it counted, and this count is also what the deposit
      will find.
    """

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        from pipeline.models import (
            Antibody, Company, ExperimentSession, Member, WbResult)
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)
        self.member = Member.objects.using(DB).get(site_id=self.site.pk)
        self.target = Target.objects.using(DB).create(gene_name="STMN2")
        company = Company.objects.using(DB).create(name="Raw Supplier")
        self.ab = Antibody.objects.using(DB).create(
            catalogue_number="RAW-1", company=company, target=self.target)
        self.session = ExperimentSession.objects.using(DB).create(
            target=self.target, site=self.site, procedure_type="WB",
            date=date(2026, 8, 5), experimenter=self.member)
        WbResult.objects.using(DB).create(session=self.session, antibody=self.ab)

    def _page(self):
        page = self.client.get(f"/pipeline/target/{self.target.pk}/")
        self.assertEqual(page.status_code, 200)
        return page.content.decode()

    def _attach(self, filename):
        """Through the real write path, into a throwaway media root.

        `DEBUG` is off under the test runner, so without a persistent root the
        save is refused — correctly: a file the app knows it will lose is worse
        than no file.
        """
        from django.core.files.uploadedfile import SimpleUploadedFile
        from django.test import override_settings
        from pipeline.services import attachments
        media = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, media, True)
        with override_settings(MEDIA_ROOT=media, PERSISTENT_MEDIA_ROOTS=(media,)):
            result = attachments.save(
                self.session,
                SimpleUploadedFile(filename, b"gel", content_type="image/tiff"),
                category="wb_scan", member=self.member)
        self.assertTrue(result["ok"], result)

    def test_one_panel_serves_both_the_sessions_board_and_the_gene_page(self):
        """`OGABoard.filesPanel`, mounted twice — never written twice."""
        here = Path(settings.BASE_DIR) / "pipeline/templates/pipeline"
        for name in ("session_board.html", "target_detail.html"):
            source = (here / name).read_text()
            self.assertIn("OGABoard.filesPanel(", source,
                          f"{name} does not mount the shared raw-files panel")
        # And the panel itself is in the shared file, not in either page.
        board_js = (Path(settings.BASE_DIR)
                    / "pipeline/static/pipeline/board.js").read_text()
        self.assertIn("function filesPanel(", board_js)
        for name in ("session_board.html", "target_detail.html"):
            self.assertNotIn("function filesBodyHtml(",
                             (here / name).read_text(),
                             f"{name} still carries its own copy of the panel")

    def test_the_gene_page_mount_carries_its_endpoints(self):
        """A routed endpoint is not a reachable one. All three, or the panel
        draws and none of its buttons reaches anything."""
        body = self._page()
        for url in ("/pipeline/session/attachments/",
                    "/pipeline/session/attachments/upload/",
                    "/pipeline/session/attachments/delete/"):
            self.assertIn(url, body, url)

    def test_a_session_with_no_files_says_what_to_do_about_it(self):
        """The zero state is the one that matters — it is what the deposit will
        find — so it reads as an instruction, not as a fact about the record."""
        body = self._page()
        self.assertIn('class="gene-files-btn', body)
        self.assertIn(f'data-session="{self.session.pk}"', body)
        self.assertIn("+ Add raw data", body)

    def test_the_count_beside_the_button_is_the_files_on_file(self):
        self._attach("scan.tif")
        self._attach("ponceau.tif")
        body = self._page()
        self.assertIn("Raw data (2)", body)
        self.assertIn(f'data-session="{self.session.pk}" data-files="2"', body)
        page = self.client.get(f"/pipeline/target/{self.target.pk}/")
        session = page.context["sessions_by_type"]["WB"][0]
        self.assertEqual(session.attachment_count, 2)

    def test_the_result_counts_beside_it_are_not_multiplied_by_the_files(self):
        """The reason the count is its own query.

        A fifth `Count` on `views/dashboard.py`'s annotate joins a second
        multi-valued relation alongside the four result counts, and every one of
        them comes back multiplied: one result row and two files would read as
        two results — a number this page draws in a column headed Results and
        that PROCEDURE STATUS reads to decide whether a western blot was run.
        """
        self._attach("scan.tif")
        self._attach("ponceau.tif")
        page = self.client.get(f"/pipeline/target/{self.target.pk}/")
        session = page.context["sessions_by_type"]["WB"][0]
        self.assertEqual(session.result_count, 1)
        self.assertEqual(session.attachment_count, 2)

    def test_the_deposit_preview_says_where_the_raw_data_is_added(self):
        """A message that sends you somewhere must send you where the control
        is — and now it is on this page, a few sections up."""
        body = self._page()
        self.assertIn("Raw data</span> column of Experiment", body)


# ── Run 17, 17 Aug 2026 ──────────────────────────────────────────────────────


class ASearchReturnsOneRowPerLineNotOnePerBatchTests(TestCase):
    """A line's own C-number returned it once per freeze-down batch.

    `find.cell_line_q` reaches `vials`, so the `Q` it builds puts a LEFT JOIN on
    the batch table into the query. Searching a *batch* number matches one vial
    row and looks fine; searching the **line's own** number is true of every
    joined row at once, so a line with seventeen batches was drawn seventeen
    times — byte-identical rows, under a count that had been inflated to match.
    Seventeen copies of one line reads as seventeen tubes.

    `find.py::_cell_lines` has had `.distinct()` since C-numbers were first
    searchable and says in a comment exactly why, so `/pipeline/find/` answered
    *1 match* for the same string the board answered seventeen times. That is the
    one-reader rule failing at the second surface: the board reuses the module's
    `Q` builder and had to reproduce the `.distinct()` that goes with it.

    Both surfaces are asserted here, because a fix to one is what happened last
    time.
    """
    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)
        # HeLa's real shape on live: seventeen batches whose numbers are not
        # consecutive, with the line itself carrying the first.
        self.line = CellLine.objects.using(DB).create(
            name="HeLa", genotype="WT", site_id=self.site.pk, c_number=15)
        for n in (15, 16, 79):
            CellLineVial.objects.using(DB).create(
                cell_line_id=self.line.pk, c_number=n)

    def _rows(self, q):
        resp = self.client.get("/pipeline/cell-lines/board/rows/", {"q": q})
        self.assertEqual(resp.status_code, 200)
        return resp.json()

    def test_the_lines_own_number_draws_the_line_once(self):
        data = self._rows("C-15")
        self.assertEqual(len(data["rows"]), 1,
                         "one line, drawn once per batch it owns")
        self.assertEqual(data["rows"][0]["name"], "HeLa")

    def test_the_count_beside_the_grid_is_the_lines_not_the_batches(self):
        """The count and the list it totals come from one queryset, so an
        inflated count is the same defect seen from the other end — and it is
        the half that also invents pages."""
        self.assertEqual(self._rows("C-15")["count"], 1)

    def test_a_batch_number_still_finds_the_line_it_belongs_to(self):
        """The dedupe must not cost the reach. `C-79` is on no `CellLine` row —
        29% of live vial numbers are like that — and it is written on a tube."""
        data = self._rows("C-79")
        self.assertEqual(len(data["rows"]), 1)
        self.assertEqual(data["rows"][0]["name"], "HeLa")

    def test_the_board_and_the_nav_search_agree_about_how_many(self):
        board = self._rows("C-15")["count"]
        resp = self.client.get("/pipeline/find/", {"q": "C-15"})
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.context["results"]["total"], board,
                         "two search boxes disagreeing about one line reads as "
                         "one of them being broken, with no way to tell which")


class ASessionPickerNamesTheLinesOwnNumberTests(TestCase):
    """The dropdown labelled a line by its batches and not by itself.

    `_session_option` spliced in `_vial_numbers_all`, which reads
    `CellLineVial.c_number` alone — so a line showed the numbers of the batches
    frozen *from* it and never the number on its own record, and a line with no
    batches yet showed nothing at all however it was numbered.

    That is backwards for the case the picker is for: you walk to the freezer,
    pick up a tube, and come back to plan a session against it. The board finds
    that number; the picker could not say it. `batch_numbers` has always folded
    the two together for the board's own cell, which is the rendering a person
    recognises.
    """
    databases = {"pipeline_db"}

    def setUp(self):
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.target = Target.objects.using(DB).create(gene_name="STMN2")
        self.wt = CellLine.objects.using(DB).create(
            name="SK-N-SH", genotype="WT", site_id=self.site.pk, c_number=10000)
        self.ko = CellLine.objects.using(DB).create(
            name="SK-N-SH", genotype="KO", target_id=self.target.pk,
            site_id=self.site.pk, c_number=10002)
        for n in (10004, 10005):
            CellLineVial.objects.using(DB).create(
                cell_line_id=self.ko.pk, c_number=n)

    def _options(self):
        from pipeline.services import cell_lines
        return cell_lines.session_options(self.target, db=DB)

    def test_a_line_with_batches_is_labelled_with_its_own_number_too(self):
        label = self._options()["ko"][0]["label"]
        self.assertIn("C-10002", label, f"its own number is missing from {label!r}")
        self.assertIn("C-10004", label)
        self.assertIn("C-10005", label)

    def test_a_line_with_no_batches_still_shows_the_number_it_has(self):
        """The case that has no other way of being seen: nothing is frozen down
        yet, so the vial table knows nothing about this line at all."""
        label = self._options()["wt"][0]["label"]
        self.assertIn("C-10000", label, f"nothing numbers {label!r}")

    def test_the_numbers_come_before_the_site(self):
        """The site is the last thing said about a line everywhere else, so a
        number after it reads as belonging to the site."""
        label = self._options()["ko"][0]["label"]
        self.assertLess(label.index("C-10002"), label.index("Leicester"))


class FeasibilityReadsTheGeneItWasSentTests(TestCase):
    """`/pipeline/find/` offers a link this page ignored.

    A gene the pipeline has never heard of is the whole reason that link exists:
    the search box says *nothing matches "ELP3"*, and underneath, *check its
    feasibility and add it*. It carries `?gene=ELP3`. The page rendered its empty
    form and the reader had to type the gene again — the parameter arrived and
    nothing read it.

    A destination that does not read `?gene=` does not get one (CLAUDE.md), and
    Browse honours that by linking here bare. This link is the other half of the
    rule: it is offered, so it has to work.

    The lookup itself stays a browser-side call. The view makes no outbound
    request — it fills the box in and the page asks, exactly as pressing Search
    would.
    """
    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)

    def test_the_gene_lands_in_the_search_box(self):
        body = self.client.get("/pipeline/feasibility/", {"gene": "ELP3"}).content.decode()
        self.assertIn('id="gene-input"', body)
        self.assertIn('value="ELP3"', body)

    def test_the_page_looks_it_up_without_being_pressed(self):
        body = self.client.get("/pipeline/feasibility/", {"gene": "ELP3"}).content.decode()
        self.assertIn("doLookup()", body)
        self.assertIn('data-gene="ELP3"', body)

    def test_no_gene_leaves_the_box_empty_and_asks_nothing(self):
        """`None` prints as four characters in a template, and an autofocused
        box reading `None` is worse than an empty one."""
        body = self.client.get("/pipeline/feasibility/").content.decode()
        self.assertIn('value=""', body)
        self.assertNotIn("None", body[body.index('id="gene-input"'):][:400])
        self.assertNotIn('data-gene="', body)

    def test_the_link_that_sends_it_is_the_search_pages_own(self):
        """Pinned together: the parameter is only worth reading while something
        offers it, and only worth offering while something reads it."""
        Target.objects.using(DB).create(gene_name="STMN2")
        body = self.client.get("/pipeline/find/", {"q": "ELP3"}).content.decode()
        self.assertIn("/pipeline/feasibility/?gene=ELP3", body)


class TwoAddButtonsOnOnePageSayWhatEachAddsTests(TestCase):
    """Both read "Add to pipeline", and one of them writes a list.

    The gene card's button adds the one gene that was looked up; the paste
    panel's adds everything in the box. Matching on the label reached the wrong
    one — it was correctly disabled and nothing happened, so the app protected
    the press, but a page where two buttons of the same name do different things
    is one where being protected is the only thing standing in the way.

    Only the idle wording is asserted. Once a check has run, the paste panel's
    button carries its count (`OGABoard.targetAddSummary`), which is shared with
    the target board and already unambiguous.
    """
    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)

    def test_the_two_buttons_are_not_called_the_same_thing(self):
        body = self.client.get("/pipeline/feasibility/").content.decode()
        labels = re.findall(r"Add [^<\"']*?[Pp]ipeline", body)
        self.assertEqual(len(set(l.lower() for l in labels)), len(set(labels)) or 1,
                         "case is not a distinction a reader can act on")
        self.assertIn("Add the pasted genes", body)

    def test_the_reset_puts_back_the_label_it_starts_with(self):
        """`invalidateBulkCheck` rewrites the label when the list changes, so
        the two spellings have to be the same one."""
        here = Path(settings.BASE_DIR) / "pipeline/templates/pipeline/feasibility.html"
        source = here.read_text()
        self.assertEqual(source.count("Add the pasted genes"), 2,
                         "the markup and the reset must agree")


class AMergedHeaderCellIsNamedNotSilentlyZeroTests(TestCase):
    """"0 rows read" over a file holding four antibodies.

    Run 17's Part 3, case 1, reproduced from the tester's own file. Merging
    `A1:B1` in the header row keeps A1's text and blanks B1, so the `catalogue`
    heading disappears — and `metadata.parse_table` drops every row that has no
    catalogue. The preview said **0 rows read. 0 new, 0 already known**, with no
    error, no warning and no mention of a column. Somebody who tidied a sheet,
    uploaded it, read that and pressed Save would conclude their edits had
    already been applied.

    **Writer-independent, so it reaches Excel.** `<mergeCell ref="A1:B1"/>` is the
    only encoding for a merge and openpyxl and Excel both emit it — which is why
    this one case is worth testing with openpyxl at all, and it is checked here
    rather than asserted.

    Deliberately not a special case for merged cells: the refusal names the
    headings it read and which matching columns are missing, so a deleted column
    and a renamed one answer the same way.
    """
    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)

    def _sheet(self, *, merge=False, blank_a_row=False):
        import openpyxl
        from pipeline.services import board_columns
        cols = [c.heading for c in board_columns.registry("antibodies")]
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Antibodies"
        ws.append(cols)
        for i in (1, 2):
            row = [""] * len(cols)
            row[cols.index("catalogue")] = f"AB-RUN17-E{i}"
            row[cols.index("company")] = "abcam"
            row[cols.index("gene")] = "ELP3"
            ws.append(row)
        if merge:
            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=2)
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        buf.name = "antibodies.xlsx"
        return buf

    def _upload(self, f):
        return self.client.post("/pipeline/import/upload/antibodies/", {"file": f})

    def test_the_merge_is_what_openpyxl_and_excel_both_write(self):
        """The premise. If this stops being true the test below stops meaning
        anything, so it is asserted rather than trusted."""
        import zipfile
        with zipfile.ZipFile(self._sheet(merge=True)) as z:
            sheet = [n for n in z.namelist() if n.startswith("xl/worksheets/")][0]
            xml = z.read(sheet).decode()
        self.assertIn('<mergeCell ref="A1:B1"/>', xml)

    def test_the_control_reads_its_rows(self):
        """Any failure below is the merge, not the writer — the same control the
        tester ran first."""
        resp = self._upload(self._sheet())
        self.assertEqual(resp.status_code, 200, resp.content[:400])
        self.assertEqual(len(resp.json()["items"]), 2)

    def test_a_merged_header_is_refused_rather_than_previewed_as_nothing(self):
        resp = self._upload(self._sheet(merge=True))
        self.assertEqual(resp.status_code, 400,
                         "a sheet with rows in it must not preview as 0 rows")
        error = resp.json()["error"]
        self.assertIn("2 rows", error, error)
        self.assertIn("Column B", error, error)
        self.assertIn("Merging cells", error, error)
        self.assertIn("no catalogue column", error, error)
        # And the headings it did read, so the sheet can be compared against it.
        self.assertIn("(blank)", error, error)

    def test_the_refusal_says_how_many_rows_it_actually_read(self):
        """The number is the whole reason this is not "the file is empty": four
        rows went up and the screen said zero."""
        from pipeline.views.imports import _nothing_read_refusal
        text = "gene\t\tcompany\nELP3\tX\tabcam\n"
        self.assertIn("Read 1 row from that sheet",
                      _nothing_read_refusal("antibodies", text))


class ABlankHeadingIsNamedEvenWhenTheRowsSurviveTests(SimpleTestCase):
    """Merging two columns nobody matches on loses what was under the second
    one, silently — the same failure as the key column, one column smaller. The
    note rides on `sheet_note`, which `board.js::sheetLine` already draws in
    amber above every preview."""

    def test_a_blank_heading_is_named_by_its_column_letter(self):
        from pipeline.views.imports import _heading_note
        note = _heading_note("name\tgene\t\tsite\nHAP1\tELP3\tx\tLeicester\n")
        self.assertIn("Column C", note)
        self.assertIn("Merging cells in row 1", note)

    def test_a_clean_header_says_nothing(self):
        from pipeline.views.imports import _heading_note
        self.assertEqual(_heading_note("name\tgene\nHAP1\tELP3\n"), "")

    def test_two_blanks_are_both_named_and_read_as_plural(self):
        from pipeline.views.imports import _heading_note
        note = _heading_note("name\t\t\tsite\nHAP1\tx\ty\tLeicester\n")
        self.assertIn("Column B, C", note)
        self.assertIn("those columns", note)


class AFormulaCellIsReadAsItsValueTests(SimpleTestCase):
    """Run 17's Part 3 case 2 — **not a defect**, and the open half closes here.

    The tester wrote `=5*2` into a concentration cell with openpyxl, saw the
    preview say nothing about the concentration, and read that as "the preview
    did not name what it could not read". There was nothing it could not read:
    every reader in this app opens a workbook with `data_only=True`, so a
    formula cell is its **cached result**, and openpyxl writes no cached result —
    the cell arrives empty. A blank means "not written down", so fill-only-blank
    left the stored value alone and said nothing, which is correct.

    That also answers the half the report left open for real Excel. Excel writes
    `<v>10</v>` beside the formula, `data_only=True` hands back `10`, and the row
    is treated as though `10` had been typed — which is the value the spreadsheet
    itself says the cell holds, and the reason every reader here opens workbooks
    that way. No Excel needed to settle it.

    Pinned because the guarantee is `data_only=True` on **every** reader, and a
    new one written without it would silently start storing `=5*2` as text.
    """

    def test_every_workbook_reader_asks_for_values_not_formulas(self):
        import re
        root = Path(settings.BASE_DIR)
        readers = []
        for path in sorted(root.glob("pipeline/**/*.py")):
            if "/tests" in str(path) or "/.venv/" in str(path):
                continue
            for line in path.read_text().splitlines():
                if "load_workbook(" in line:
                    readers.append((path.name, line.strip()))
        self.assertTrue(readers, "no workbook readers found — has the API moved?")
        for name, line in readers:
            with self.subTest(reader=f"{name}: {line}"):
                self.assertIn("data_only=True", line,
                              "a reader without data_only stores the formula text")


class TheCellLinesSheetCarriesWhatIsOnTheTubeTests(TestCase):
    """Run 17 downloaded the board it had just used and the numbers were not in it.

    Two halves. The `c number` column wrote the bare integer `10003` where the
    board beside it draws `C-10003` — one field written two ways in two places a
    person reads in the same minute. And the **freeze-down batches were in no
    column at all**, so C-10004 and C-10005 — the numbers actually written on the
    vials — appeared nowhere in a file whose own panel calls itself a round trip.

    The batches column is read-only on purpose: a batch is one press of
    `+ batch`, and its number comes from `services/lab_numbers.py`.
    """
    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)
        self.target = Target.objects.using(DB).create(gene_name="STMN2")
        self.line = CellLine.objects.using(DB).create(
            name="SK-N-SH", genotype="KO", target_id=self.target.pk,
            site_id=self.site.pk, c_number=10002)
        for n in (10004, 10005):
            CellLineVial.objects.using(DB).create(
                cell_line_id=self.line.pk, c_number=n)

    def _sheet(self, **params):
        import openpyxl
        resp = self.client.get("/pipeline/cell-lines/export/", params)
        self.assertEqual(resp.status_code, 200)
        ws = openpyxl.load_workbook(io.BytesIO(resp.content)).active
        rows = list(ws.values)
        return rows[0], rows[1:]

    def test_the_c_number_is_written_the_way_the_tube_is(self):
        header, rows = self._sheet(q="SK-N-SH")
        cell = rows[0][header.index("c number")]
        self.assertEqual(cell, "C-10002")

    def test_the_freeze_down_batches_are_in_the_sheet(self):
        header, rows = self._sheet(q="SK-N-SH")
        self.assertIn("freeze-down batches (read-only)", header)
        cell = rows[0][header.index("freeze-down batches (read-only)")]
        for number in ("C-10002", "C-10004", "C-10005"):
            self.assertIn(number, cell)

    def test_the_batches_column_is_named_as_one_an_upload_ignores(self):
        from pipeline.services import board_columns
        self.assertIn("freeze-down batches (read-only)",
                      board_columns.read_only("cell-lines"))

    def test_a_prefixed_c_number_uploads_back_unchanged(self):
        """A download the importer cannot read back is not a round trip."""
        from pipeline.services import bulk_cell_lines
        rows = bulk_cell_lines.parse(
            "name\tgene\tgenotype\tc number\nSK-N-SH\tSTMN2\tKO\tC-10002\n")
        self.assertEqual(rows[0]["c_number"], "C-10002")
        items = bulk_cell_lines.plan(rows, member=None)
        self.assertFalse(items[0].get("c_number_dropped"),
                         "the sheet's own spelling must not be refused")

    def test_the_export_draws_one_row_per_line_not_one_per_batch(self):
        """Run 17's P3-1 — finding B-1 escaping the screen and getting into a
        file people edit and send each other. Two copies edited differently and
        uploaded means the second write wins, silently. Same `.distinct()`, since
        the export reads the board's own `apply_filters`."""
        _header, rows = self._sheet(q="C-10002")
        self.assertEqual(len(rows), 1, "one line, drawn once per batch it owns")


class ABlankTemplateStillPreviewsAsNothingToDoTests(TestCase):
    """The refusal must not fire on a sheet that legitimately holds nothing.

    A downloaded template uploaded untouched has one data row and it is the
    `e.g.` example, which every parser drops (`services/example_row.py`). Zero
    records out of it is the truth, not a failure, and turning "download the
    sheet and look at it" into an error message would be a worse bug than the one
    the refusal fixes. Caught by `tests_conventions` when the refusal was first
    written, and pinned here so it cannot come back the other way round.
    """
    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)

    def test_the_blank_template_is_previewed_not_refused(self):
        for kind in ("antibodies", "cell-lines"):
            with self.subTest(kind=kind):
                blank = self.client.get(f"/pipeline/import/template/{kind}/")
                self.assertEqual(blank.status_code, 200)
                f = io.BytesIO(blank.content)
                f.name = f"{kind}.xlsx"
                resp = self.client.post(f"/pipeline/import/upload/{kind}/", {"file": f})
                self.assertEqual(resp.status_code, 200, resp.content[:300])
                self.assertEqual(resp.json()["items"], [])


class AnUploadedSheetIsHeaderLedNotGuessedAtTests(TestCase):
    """Run 17's verification run, claim 8 — the half the first fix missed.

    The refusal shipped and then did not fire, because it was written against a
    fixture that differed from a real export in one way: **RRIDs**.
    `metadata.parse_table` falls back to the RRID-anchored parse — built for
    wrapped PDF text, where there is no header row to trust — as soon as it sees
    two or more RRIDs and fewer clean rows than that. A board export has an RRID
    on every row, so merging `A1:B1` sent it straight down the PDF path, which
    walks the flattened text by vendor name and **filled `catalogue` with the
    gene**: `ALKBH2` where `ARP54321_P050` belongs.

    So the preview did not say "0 rows read" any more — it said *4 rows read,
    4 skipped*, each row blaming *"no gene on this row"* while printing ALKBH2 in
    the first column. Both halves of that were the same fabricated row: the gene
    had been eaten by the catalogue field, so the gene really was empty.

    **And the skip was luck, not a guard.** The same sheet uploaded from a gene's
    page carries a `default_gene`, which fills the gene the mis-parse emptied —
    at which point the rows resolve, stop being skipped, and four antibodies are
    creatable with a catalogue number of `ALKBH2`. That is a silent wrong write
    off a spreadsheet whose only fault was a merged pair of header cells.

    The fix is not another special case: **a file that arrives with a header row
    is read by its header row.** `workbook.read` picks the sheet *by* recognising
    those headings, so an upload is header-led by construction; guessing from
    vendor names is for a paste, and the paste box keeps it.
    """
    databases = {"pipeline_db", "academy_db"}

    HEADERLESS_PDF = (
        "Aviva Systems Biology ARP54321_P050 AB_2224621 rabbit polyclonal\n"
        "Abcam ab138501 AB_2537855 rabbit monoclonal\n")

    def setUp(self):
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)
        Target.objects.using(DB).create(gene_name="ALKBH2")

    def _export_shaped(self, *, merge=False, rename_catalogue=None):
        """A sheet the way the antibodies board's own Download writes one —
        RRID on every row, which is the whole difference."""
        import openpyxl
        from pipeline.services import board_columns
        cols = [c.heading for c in board_columns.registry("antibodies")]
        if rename_catalogue:
            cols[cols.index("catalogue")] = rename_catalogue
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Antibodies"
        ws.append(cols)
        for i in (1, 2, 3, 4):
            row = [""] * len(cols)
            row[cols.index(rename_catalogue or "catalogue")] = f"ARP5432{i}_P050"
            row[cols.index("company")] = "Aviva Systems Biology"
            row[cols.index("gene")] = "ALKBH2"
            row[cols.index("rrid")] = f"AB_222462{i}"
            row[cols.index("site")] = "Leicester"
            ws.append(row)
        if merge:
            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=2)
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        buf.name = "antibodies.xlsx"
        return buf

    def _upload(self, f, **extra):
        return self.client.post("/pipeline/import/upload/antibodies/",
                                {"file": f, **extra})

    def test_the_control_still_reads_its_rows(self):
        resp = self._upload(self._export_shaped())
        self.assertEqual(resp.status_code, 200, resp.content[:300])
        items = resp.json()["items"]
        self.assertEqual(len(items), 4)

    def test_a_real_export_with_a_merged_header_is_refused(self):
        """The tester's own case, reproduced from an export-shaped sheet."""
        resp = self._upload(self._export_shaped(merge=True))
        self.assertEqual(resp.status_code, 400,
                         f"previewed instead of refusing: {resp.content[:400]}")
        error = resp.json()["error"]
        self.assertIn("Column B", error, error)
        self.assertIn("catalogue", error, error)

    def test_the_gene_is_never_fabricated_into_the_catalogue_column(self):
        """The sharp end. `catalogue` is half of what makes an antibody that
        antibody, and a value invented for it is not a bad message — it is a
        wrong record."""
        from pipeline.views.imports import _file_to_text
        from pipeline.services import bulk_antibodies
        text, _note = _file_to_text(self._export_shaped(merge=True), "antibodies")
        rows = bulk_antibodies.parse(text, "", header_led=True)
        self.assertNotIn("ALKBH2", [r.get("catalogue") for r in rows],
                         "the gene was stored as the catalogue number")

    def test_the_same_sheet_from_a_gene_page_cannot_create_anything(self):
        """`default_gene` is what a gene page's upload panel sends, and it is
        what turned four skipped rows into four creatable ones."""
        resp = self._upload(self._export_shaped(merge=True), default_gene="ALKBH2")
        self.assertEqual(resp.status_code, 400,
                         f"a gene page could create these: {resp.content[:400]}")

    def test_a_renamed_key_column_is_refused_with_no_blank_heading_to_spot(self):
        """No merge, nothing blank — just a heading this parser does not know.
        The anchored fallback made rows out of it all the same."""
        resp = self._upload(self._export_shaped(rename_catalogue="product code"))
        self.assertEqual(resp.status_code, 400,
                         f"previewed instead of refusing: {resp.content[:400]}")
        self.assertIn("catalogue", resp.json()["error"])

    def test_a_pasted_pdf_table_still_gets_the_anchored_parse(self):
        """The other half of the rule, and the reason this is a parameter rather
        than a deletion: a paste has no header row to be led by, which is the
        case the anchored walk was written for."""
        from pipeline.services import bulk_antibodies
        rows = bulk_antibodies.parse(self.HEADERLESS_PDF, "")
        self.assertEqual(len(rows), 2, rows)
        self.assertEqual(sorted(r["catalogue"] for r in rows),
                         ["ARP54321_P050", "ab138501"])


class AParentOfTheWrongBackgroundIsRefusedEverywhereTests(TestCase):
    """Run 19 pasted `HAP1 | STMN2 | KO | parent HeLa` and read *parent "HeLa"
    is your site's line — new*. The tester stopped because it re-read the brief,
    not because the app said anything; the row would have been written.

    `backfill_cell_line_parents` had refused exactly this since 4 Sep — a wild
    type is not enough, it has to be a wild type *of that background* — but the
    check lived in the command and nowhere else. So the paste door and the
    identity dialog, the two surfaces a person actually uses, would link a HAP1
    knockout to a HeLa parental, and every session planned against it afterwards
    would read the mismatched control as the matched one. One reader now
    (`cell_lines.wrong_background`), asked by all three.
    """

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)
        self.target = Target.objects.using(DB).create(gene_name="STMN2")
        self.hap1 = CellLine.objects.using(DB).create(
            name="HAP1", genotype="WT", site_id=self.site.pk)
        self.hela = CellLine.objects.using(DB).create(
            name="HeLa", genotype="WT", site_id=self.site.pk)
        # Leicester's parentals were imported without C-numbers and nothing
        # backfills one; `pre_save` numbers a fresh row, so take them off again.
        CellLine.objects.using(DB).filter(
            pk__in=[self.hap1.pk, self.hela.pk]).update(c_number=None)

    def _plan(self, name, parent):
        from pipeline.models import Member
        from pipeline.services import bulk_cell_lines
        member = Member.objects.using(DB).get(site_id=self.site.pk)
        rows = bulk_cell_lines.parse(
            "name\tgene\tgenotype\tparent\tclone\n"
            f"{name}\tSTMN2\tKO\t{parent}\tA1\n")
        return bulk_cell_lines.plan(rows, member=member)[0]

    def test_the_paste_door_blocks_a_hela_parent_on_a_hap1_knockout(self):
        item = self._plan("HAP1", "HeLa")
        self.assertEqual(item["status"], "blocked", item["note"])
        self.assertIn("Looks wrong", item["note"])
        self.assertIn("HAP1 knockout does not come from a HeLa", item["note"])
        # A blocked row is not about to be numbered either.
        self.assertFalse(item["will_be_numbered"])

    def test_the_refusal_names_the_wild_type_it_should_have_named(self):
        """Leicester's parentals carry no C-number, so the hint has to name the
        line by name — "no wild type on file" would be false here and would
        send somebody to add a second HAP1."""
        note = self._plan("HAP1", "HeLa")["note"]
        self.assertIn("name it as HAP1", note)

    def test_a_numbered_parental_is_offered_by_number(self):
        from pipeline.models import CellLineVial
        CellLineVial.objects.using(DB).create(
            cell_line=self.hap1, c_number=15, site_id=self.site.pk)
        note = self._plan("HAP1", "HeLa")["note"]
        self.assertIn("The HAP1 wild type on file is C-15", note)

    def test_the_right_parent_still_goes_through(self):
        item = self._plan("HAP1", "HAP1")
        self.assertEqual(item["status"], "create", item["note"])
        self.assertIn("your site's line", item["note"])

    def test_a_decorated_name_is_judged_by_its_background(self):
        """`HAP1 STMN2 KO` is a HAP1 — the label's decoration is not part of
        the cell line, so it must neither pass a HeLa nor refuse a HAP1."""
        self.assertEqual(self._plan("HAP1 STMN2 KO", "HAP1")["status"], "create")
        self.assertEqual(self._plan("HAP1 STMN2 KO", "HeLa")["status"], "blocked")

    def test_the_save_writes_nothing_for_the_blocked_row(self):
        from pipeline.models import Member
        from pipeline.services import bulk_cell_lines
        member = Member.objects.using(DB).get(site_id=self.site.pk)
        rows = bulk_cell_lines.parse(
            "name\tgene\tgenotype\tparent\tclone\nHAP1\tSTMN2\tKO\tHeLa\tA1\n")
        bulk_cell_lines.apply(rows, member=member)
        self.assertFalse(CellLine.objects.using(DB)
                         .filter(genotype="KO", target_id=self.target.pk).exists())

    def test_the_identity_dialog_refuses_the_same_link(self):
        ko = CellLine.objects.using(DB).create(
            name="HAP1", genotype="KO", target_id=self.target.pk,
            site_id=self.site.pk, parent_line_id=self.hap1.pk, clone="A1")
        resp = self.client.post(
            "/pipeline/cell-lines/board/identity/save/",
            {"cell_line_id": ko.pk, "name": "HAP1", "gene": "STMN2",
             "genotype": "KO", "clone": "A1", "parent": "HeLa"})
        self.assertEqual(resp.status_code, 400, resp.content[:300])
        self.assertIn("Looks wrong", resp.json()["error"])
        ko.refresh_from_db(using=DB)
        self.assertEqual(ko.parent_line_id, self.hap1.pk)

    def test_one_reader_for_all_three_surfaces(self):
        base = Path(settings.BASE_DIR)
        for rel in ("pipeline/services/bulk_cell_lines.py",
                    "pipeline/services/identity.py",
                    "pipeline/management/commands/backfill_cell_line_parents.py"):
            self.assertIn("wrong_background(", (base / rel).read_text(), rel)


class ANoteIsNotAReadingTests(TestCase):
    """Run 19 created a session from the sessions board's Add panel with a NOTES
    value on each antibody row and, before typing a single reading, the board
    read **"2 results — every row in this session has something recorded"**.
    The notes had landed in the result rows' `comments`, and `is_reading` took
    any result column as a measurement. A comment rides beside the readings; it
    does not make a row one — on the board, in the workbook importer and in the
    report, from one list."""

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        from pipeline.models import Antibody, Company, ExperimentSession, Member, WbResult
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)
        member = Member.objects.using(DB).get(site_id=self.site.pk)
        self.target = Target.objects.using(DB).create(gene_name="STMN2")
        company = Company.objects.using(DB).create(name="Abcam")
        self.ab = Antibody.objects.using(DB).create(
            target_id=self.target.pk, company_id=company.pk,
            catalogue_number="AB-RUN19-01", site_id=self.site.pk)
        self.session = ExperimentSession.objects.using(DB).create(
            target_id=self.target.pk, procedure_type="WB", date="2026-09-05",
            site_id=self.site.pk, experimenter_id=member.pk)
        self.row = WbResult.objects.using(DB).create(
            session_id=self.session.pk, antibody_id=self.ab.pk,
            comments="test at 1/1000 [COWORK RUN19]")

    def test_the_board_does_not_count_a_commented_row_as_read(self):
        from pipeline.services import session_board
        self.assertEqual(session_board.reading_counts([self.session.pk]), {})

    def test_a_reading_beside_the_comment_still_counts(self):
        from pipeline.services import session_board
        self.row.signal = "band in WT, absent in KO"
        self.row.save(using=DB)
        self.assertEqual(session_board.reading_counts([self.session.pk]),
                         {self.session.pk: 1})

    def test_the_report_agrees(self):
        from pipeline.services.report_generator import _readings_in
        self.assertEqual(_readings_in(self.session), 0)

    def test_the_workbook_importer_agrees(self):
        from pipeline.services import session_import
        row = {"comments": "only a note"}
        self.assertFalse(session_import._has_result(row, "WB"))
        self.assertTrue(session_import._has_result({"signal": "band"}, "WB"))

    def test_comments_are_still_a_result_column(self):
        """Not a reading, but still drawn, exported and editable."""
        from pipeline.services import session_board
        self.assertIn("comments", session_board.result_field_names("WB"))
        self.assertNotIn("comments", session_board.reading_fields("WB"))


class AZeroIsNotAboutToBeUpdatedTests(SimpleTestCase):
    """*"0 already known (will be updated)"* — run 19 read the suffix over a
    zero on both spreadsheet doors. Three surfaces print the line; all three
    now hedge only when there is something to hedge about."""

    def test_no_surface_prints_the_suffix_unconditionally(self):
        base = Path(settings.BASE_DIR)
        for rel in ("pipeline/static/pipeline/board.js",
                    "pipeline/templates/pipeline/cell_line_board.html",
                    "pipeline/templates/pipeline/antibody_board.html"):
            src = (base / rel).read_text()
            # The rendered line, not a comment quoting the old wording.
            self.assertNotIn("} already known (will be updated)", src, rel)
            self.assertIn("already known${s.update ? ' (will be updated)' : ''}", src, rel)


class TheStepFormSaysWhereAntibodiesGetOnTests(TestCase):
    """The three-step form asks for no antibodies, and the results drawer on
    the session it makes said "add the antibodies to the session first" — with
    nowhere to do that. Run 19 planned session 627 that way and had to read the
    guide to find the sessions board's Add panel. The place antibodies get onto
    a session that has none is its bench sheet, and both pages now say so."""

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)

    def test_the_form_says_antibodies_are_not_chosen_here(self):
        html = self.client.get("/pipeline/session/new/").content.decode()
        self.assertIn("Antibodies are not chosen here", html)
        self.assertIn("+ Record a session", html)
        self.assertIn("bench sheet", html)
        self.assertNotIn("session detail page", html)

    def test_the_drawer_points_at_the_bench_sheet_not_at_nothing(self):
        html = self.client.get("/pipeline/sessions/board/").content.decode()
        self.assertNotIn("add the antibodies to the session first", html)
        self.assertIn("picking list", html)


class ARecordedSheetSaysWhatItNumberedTests(SimpleTestCase):
    """Run 20 walked the picking-list bench sheet end to end, A-8 and A-9 were
    issued to exactly the two rows it wrote on — and it could not quote the
    receipt, because there wasn't one to quote. It filed that as its own
    driver's fault. It was ours, twice over.

    **The bench-sheet branch never called `saveNotes`.** Recording a sheet is
    one of the three doors that issues an A-number; the run-19 fix reached the
    workbook receipt and the gene page's panel and missed this one — the
    "seven call sites" failure the composed writer exists to prevent.

    **And the receipt was destroyed by the redraw.** Both branches ended
    `await openResults(id)`, which replaces the whole drawer body, so the
    message was written into an element that the very next line threw away.
    A save that wrote and announced nothing is this file's worst shape, and it
    is why the run reasonably concluded the message was merely transient.

    Ordering assertions rather than a browser test: what has to hold is that
    the receipt is written *after* the await, and that is a fact about the
    source's order that a browser would confirm more slowly and no response
    test can see at all.
    """

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.src = (Path(settings.BASE_DIR)
                   / "pipeline/templates/pipeline/session_board.html").read_text()
        cls.handler = cls.src[cls.src.index("const staged = benchRows.get(id);"):]
        cls.handler = cls.handler[:cls.handler.index("async function openResults")]

    def test_both_doors_name_the_numbers_they_issued(self):
        """The workbook branch and the bench-sheet branch, not one of them."""
        self.assertEqual(self.handler.count("OGABoard.saveNotes(d)"), 2,
                         "each branch of the commit handler must name what it numbered")

    def test_the_receipt_is_drawn_after_the_redraw_that_would_wipe_it(self):
        redraw = self.handler.index("await openResults(id)")
        drawn = self.handler.index("fresh.innerHTML = receipt")
        self.assertLess(redraw, drawn,
                        "the receipt is written before openResults replaces the body, "
                        "so the save announces itself to nobody")

    def test_the_element_is_re_queried_rather_than_reused(self):
        """The `out` the handler started with is gone after the redraw; writing
        to it puts the receipt into a detached node, which looks identical to
        not writing one at all."""
        tail = self.handler[self.handler.index("await openResults(id)"):]
        self.assertIn('drawerBody.querySelector(`.bench-out', tail)

    def test_the_receipt_scrolls_itself_into_view(self):
        self.assertIn("OGABoard.bringIntoView(fresh)", self.handler)


class ARefusalNamesOnlyControlsOnThisPanelTests(TestCase):
    """*"tick “Add the gene as a new target if it isn't one yet” above"* — on a
    panel whose only two tick boxes are Overwrite and Give new rows an A-number.
    Run 20 read that on the antibodies **Upload** panel, went looking for the
    control and found nothing.

    The first fix told the planners which surface was asking. The owner then
    settled the wider question underneath it: **a target is added on the targets
    doors and nowhere else** (5 Sep 2026), so both Add panels lost the box too
    and there is one sentence again, naming no control and pointing at the
    board. This pins that no surface offers or names it.
    """

    databases = {"pipeline_db", "academy_db"}

    def setUp(self):
        self.site = Site.objects.using(DB).create(name="Leicester", short_code="LEI")
        self.client = _member_client(self, self.site)

    def _upload(self, kind, body, filename="rows.csv"):
        f = io.BytesIO(body.encode())
        f.name = filename
        return self.client.post(f"/pipeline/import/upload/{kind}/", {"file": f})

    def test_the_antibodies_upload_does_not_name_a_missing_tick(self):
        resp = self._upload("antibodies",
                            "gene,catalogue,company\nZZZZZZ,AB-1,Abcam\n")
        self.assertEqual(resp.status_code, 200, resp.content[:300])
        note = " ".join(i["note"] for i in resp.json()["items"])
        self.assertIn("not in the pipeline yet", note)
        self.assertNotIn("tick", note, note)
        self.assertIn("target board", note)

    def test_the_cell_lines_upload_does_not_either(self):
        resp = self._upload("cell-lines",
                            "name,gene,genotype,parent\nHAP1,ZZZZZZ,KO,HAP1\n")
        self.assertEqual(resp.status_code, 200, resp.content[:300])
        note = " ".join(i["note"] for i in resp.json()["items"])
        self.assertIn("not in the pipeline yet", note)
        self.assertNotIn("tick", note, note)

    def test_the_add_panels_say_the_same_thing_as_the_uploads(self):
        """One sentence, because there is one way through."""
        from pipeline.services import bulk_antibodies, bulk_cell_lines
        ab = bulk_antibodies.plan(
            bulk_antibodies.parse("gene\tcatalogue\tcompany\nZZZZZZ\tAB-1\tAbcam\n", ""))[0]
        cl = bulk_cell_lines.plan(
            bulk_cell_lines.parse("name\tgene\tgenotype\nHAP1\tZZZZZZ\tKO\n"))[0]
        for note in (ab["note"], cl["note"]):
            self.assertIn("target board", note)
            self.assertNotIn("tick", note, note)

    def test_no_panel_offers_the_control_any_more(self):
        base = Path(settings.BASE_DIR) / "pipeline/templates/pipeline"
        for page in ("antibody_board.html", "cell_line_board.html", "data_io.html"):
            src = (base / page).read_text()
            self.assertNotIn("create-targets", src, page)
            self.assertNotIn("createtargets", src, page)
            self.assertNotIn("create_targets", src, page)


class APlannedSessionIsNotAnEmptyOneTests(SimpleTestCase):
    """The workbook preview warned *"1 of them still planned with no results"*
    about a session that had two results and one status. `_existing_sessions`
    selects on `status` alone, and the sentence claimed something about
    readings that the query never asked — the same planned-versus-recorded
    confusion the progress strip and the results column were fixed for, in a
    third place (run 20).

    The warning's job is *"you may have meant to fill that one in"*, which is
    just as true of a planned session that has readings. So it says what it
    knows and stops."""

    def test_the_guard_does_not_claim_a_session_is_empty(self):
        src = (Path(settings.BASE_DIR)
               / "pipeline/templates/pipeline/_bench_workbook_upload.html").read_text()
        self.assertIn("still", src)
        self.assertNotIn("planned</span> with no results", src)
